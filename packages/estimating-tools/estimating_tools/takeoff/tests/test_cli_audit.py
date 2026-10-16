"""CLI tests for est-takeoff audit (test_cli_ifc2dxf.py style)."""
import json
from pathlib import Path

import openpyxl
import pytest

from estimating_tools.takeoff.cli import main


def _takeoff_json(tmp_path, quantity=480.0):
    data = {"organization_id": 1, "project_ref": "B1", "rows": [{
        "description": "Structural Concrete", "kind": "quantity",
        "quantity": quantity, "unit": "SF", "cost_code": "03310",
        "item_number": "12", "schedule": "A", "structure_id": "HW-3",
        "elements": [{"name": "Deck-1", "global_id": "G1",
                      "section_type": "decks", "count": 1, "length_ft": 10.0,
                      "width_ft": 4.0, "height_ft": 0.67,
                      "measure": "GrossArea", "value_si": 44.59}]}]}
    p = tmp_path / "takeoff.json"
    p.write_text(json.dumps(data), encoding="utf-8")
    return p


def _bm_json(tmp_path, qty=480.0):
    data = {"estimate_ref": "B1", "items": [
        {"schedule": "A", "item_number": "12",
         "description": "Structural Concrete", "unit": "SF",
         "estimate_qty": qty,
         "structures": [{"structure_id": "HW-3", "sections": []}]}]}
    p = tmp_path / "bm.json"
    p.write_text(json.dumps(data), encoding="utf-8")
    return p


def test_audit_defaults_out_beside_takeoff(tmp_path, capsys):
    t = _takeoff_json(tmp_path)
    rc = main(["audit", "--takeoff", str(t), "--bm",
               str(_bm_json(tmp_path))])
    assert rc == 0
    out_path = tmp_path / "takeoff.audit.xlsx"
    assert out_path.exists()
    wb = openpyxl.load_workbook(str(out_path))
    assert wb.sheetnames == ["Summary", "Elements", "Exceptions"]
    msg = capsys.readouterr().out
    assert "1 item" in msg and "0 exception" in msg and "BM" in msg


def test_audit_strict_nonzero_on_exceptions(tmp_path):
    t = _takeoff_json(tmp_path)
    bm = _bm_json(tmp_path, qty=999.0)         # big variance
    assert main(["audit", "--takeoff", str(t), "--bm", str(bm)]) == 0
    assert main(["audit", "--takeoff", str(t), "--bm", str(bm),
                 "--strict"]) == 1


def test_audit_bad_snapshot_is_error_rc1(tmp_path, capsys):
    t = _takeoff_json(tmp_path)
    bad = tmp_path / "bm.json"
    bad.write_text('{"estimate_ref": "x"}', encoding="utf-8")
    rc = main(["audit", "--takeoff", str(t), "--bm", str(bad)])
    assert rc == 1
    assert "error:" in capsys.readouterr().err


def test_audit_takeoff_only_runs_offline(tmp_path, capsys):
    rc = main(["audit", "--takeoff", str(_takeoff_json(tmp_path)),
               "--out", str(tmp_path / "a.xlsx")])
    assert rc == 0
    assert (tmp_path / "a.xlsx").exists()
    assert "takeoff only" in capsys.readouterr().out


def test_audit_corrupt_ifc_is_error_rc1(tmp_path, capsys):
    # an EXISTING file that is not an IFC must be a clean `error:` + rc 1,
    # not a raw ifcopenshell traceback (a nonexistent path is OSError and
    # was already caught — this pins the unreadable-file case).
    corrupt = tmp_path / "corrupt.ifc"
    corrupt.write_text("not an ifc", encoding="utf-8")
    rc = main(["audit", "--takeoff", str(_takeoff_json(tmp_path)),
               "--ifc", str(corrupt)])
    assert rc == 1
    assert "error:" in capsys.readouterr().err


def _authored_ifc(tmp_path, lie_pct=0.0):
    """Emit a real 4x4x2 ft footing via the Phase 3a kernel; optionally
    author its volume Qto `lie_pct` percent off the true geometry. Returns
    (ifc_path, takeoff_json_path) with matching GlobalIds."""
    import uuid
    ifcopenshell = pytest.importorskip("ifcopenshell")
    from bimforge.author.emit_ifc import emit_ifc
    from bimforge.author.model import Element, StructureModel
    from bimforge.author.primitives import Prism
    from bimforge.author.project import ProjectContext

    ctx = ProjectContext(name="CLI Mesh Fixture", epsg=2226,
                         base_e=6_594_000.0, base_n=1_943_000.0, base_h=200.0,
                         guid_ns=uuid.UUID("7c1b9e3f-2a58-4d06-9f47-5e8d0a2c4b61"))
    m = StructureModel("cli-mesh", ctx)
    p = Prism("FTG-1", [(-2, -2), (2, -2), (2, 2), (-2, 2)],
              z_bot=0.0, z_top=2.0, ifc="IfcFooting")
    el = Element.from_geometry(p.name, p, ifc_class="IfcFooting")
    el.quantities["volume_cf"] *= (1.0 + lie_pct / 100.0)
    m.add(el)
    ifc_path = tmp_path / "model.ifc"
    emit_ifc(m, str(ifc_path))

    si = el.quantities["volume_cf"] * (1200.0 / 3937.0) ** 3
    data = {"organization_id": 1, "project_ref": "B1", "rows": [{
        "description": "Footing Concrete", "kind": "quantity",
        "quantity": el.quantities["volume_cf"], "unit": "CF",
        "cost_code": "03310", "item_number": "40", "schedule": "A",
        "structure_id": "HW-3",
        "elements": [{"name": "FTG-1",
                      "global_id": ctx.guid("element/FTG-1"),
                      "section_type": "footings", "count": 1,
                      "measure": "GrossVolume", "value_si": si}]}]}
    t = tmp_path / "takeoff.json"
    t.write_text(json.dumps(data), encoding="utf-8")
    return ifc_path, t


def test_audit_mesh_runs_by_default_with_ifc(tmp_path, capsys):
    ifc_path, t = _authored_ifc(tmp_path)
    out = tmp_path / "a.xlsx"
    rc = main(["audit", "--takeoff", str(t), "--ifc", str(ifc_path),
               "--out", str(out)])
    assert rc == 0
    wb = openpyxl.load_workbook(str(out))
    assert "Geometry" in wb.sheetnames
    msg = capsys.readouterr().out
    assert "mesh leg:" in msg and "1 meshed" in msg
    assert "MESH" in msg          # legs line includes the new leg


def test_audit_no_mesh_skips_the_leg(tmp_path, capsys):
    ifc_path, t = _authored_ifc(tmp_path)
    out = tmp_path / "b.xlsx"
    rc = main(["audit", "--takeoff", str(t), "--ifc", str(ifc_path),
               "--no-mesh", "--out", str(out)])
    assert rc == 0
    wb = openpyxl.load_workbook(str(out))
    assert "Geometry" not in wb.sheetnames
    assert "mesh leg:" not in capsys.readouterr().out


def test_audit_geom_tol_pct_threads_through(tmp_path):
    # a 5% lie passes at ±10% but fails --strict at the 0.5% default
    ifc_path, t = _authored_ifc(tmp_path, lie_pct=5.0)
    assert main(["audit", "--takeoff", str(t), "--ifc", str(ifc_path),
                 "--geom-tol-pct", "10.0", "--strict",
                 "--out", str(tmp_path / "c.xlsx")]) == 0
    assert main(["audit", "--takeoff", str(t), "--ifc", str(ifc_path),
                 "--strict", "--out", str(tmp_path / "d.xlsx")]) == 1


def test_acceptance_lying_qto_and_unclassified_solid(tmp_path, capsys):
    """Spec §5 acceptance: (1) a 5% lie between authored GrossVolume and the
    real solid is caught end-to-end (workbook verdict + strict exit code);
    (2) a solid with NO takeoff row and an IFC class the orphan check's
    measured-classes filter would skip still surfaces via
    unclassified-geometry."""
    import uuid
    pytest.importorskip("ifcopenshell")
    from bimforge.author.emit_ifc import emit_ifc
    from bimforge.author.model import Element, StructureModel
    from bimforge.author.primitives import Prism
    from bimforge.author.project import ProjectContext

    ctx = ProjectContext(name="Acceptance Fixture", epsg=2226,
                         base_e=6_594_000.0, base_n=1_943_000.0, base_h=200.0,
                         guid_ns=uuid.UUID("9d4f1a6b-3c27-4e85-b190-7f2e5c8a0d13"))
    m = StructureModel("acceptance", ctx)
    p = Prism("FTG-1", [(-2, -2), (2, -2), (2, 2), (-2, 2)],
              z_bot=0.0, z_top=2.0, ifc="IfcFooting")
    el = Element.from_geometry(p.name, p, ifc_class="IfcFooting")
    el.quantities["volume_cf"] *= 1.05                  # the lie: +5%
    m.add(el)
    # the forgotten solid: real geometry, PROXY class (outside the orphan
    # check's measured-classes filter), no takeoff row
    m.add(Element("forgotten pad", ifc_class="IfcBuildingElementProxy",
                  geometry={"kind": "prism", "units": "ft",
                            "base": [(10, 0), (12, 0), (12, 2), (10, 2)],
                            "z_bot": 0.0, "z_top": 1.0}))
    ifc_path = tmp_path / "acceptance.ifc"
    emit_ifc(m, str(ifc_path))

    si = el.quantities["volume_cf"] * (1200.0 / 3937.0) ** 3
    data = {"organization_id": 1, "project_ref": "B1", "rows": [{
        "description": "Footing Concrete", "kind": "quantity",
        "quantity": el.quantities["volume_cf"], "unit": "CF",
        "cost_code": "03310", "item_number": "40", "schedule": "A",
        "structure_id": "HW-3",
        "elements": [{"name": "FTG-1",
                      "global_id": ctx.guid("element/FTG-1"),
                      "section_type": "footings", "count": 1,
                      "measure": "GrossVolume", "value_si": si}]}]}
    t = tmp_path / "takeoff.json"
    t.write_text(json.dumps(data), encoding="utf-8")

    out = tmp_path / "acceptance.xlsx"
    rc = main(["audit", "--takeoff", str(t), "--ifc", str(ifc_path),
               "--strict", "--out", str(out)])
    assert rc == 1                                       # strict catches it
    err = capsys.readouterr().err
    assert "geometry-variance" in err and "FTG-1" in err
    assert "unclassified-geometry" in err and "forgotten pad" in err

    wb = openpyxl.load_workbook(str(out))
    geo = wb["Geometry"]
    flags = {row[2].value: row[7].value
             for row in geo.iter_rows(min_row=2)}
    assert flags["FTG-1"] == "GEOM VARIANCE"
    assert flags["forgotten pad"] == "UNCLASSIFIED"
    ws = wb["Summary"]
    hdr = [c.value for c in ws[7]]      # header row (rows 1-6 are preamble)
    geom_col = hdr.index("Geom") + 1
    assert ws.cell(8, geom_col).value == "GEOM VARIANCE"
