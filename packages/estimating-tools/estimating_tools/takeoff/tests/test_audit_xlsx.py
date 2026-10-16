"""Read-back tests for the audit workbook writer."""
import openpyxl
import pytest

from estimating_tools.takeoff.audit import (AuditFinding, AuditReport,
                                            ElementAudit, ItemAudit,
                                            build_audit)
from estimating_tools.takeoff.audit_xlsx import write_audit_xlsx
from estimating_tools.takeoff.model import CanonicalTakeoff, TakeoffRow


def _report():
    return AuditReport(
        items=[ItemAudit(item_key="A/12", schedule="A", item_number="12",
                         description="Structural Concrete", unit="SF",
                         takeoff_qty=480.0, ifc_qty=481.0, bm_qty=500.0,
                         delta_ifc_pct=0.21, delta_bm_pct=4.17,
                         status="VARIANCE", element_count=1,
                         structure_count=1)],
        elements=[ElementAudit(item_key="A/12", global_id="G1", name="Deck-1",
                               ifc_class="IfcSlab", section_type="decks",
                               structure_id="HW-3", count=1, length_ft=10.0,
                               width_ft=4.0, height_ft=0.67,
                               measure="GrossArea", takeoff_si=44.59,
                               ifc_si=44.68, delta_flag="OK", join="id")],
        exceptions=[AuditFinding("variance", "A/12", "BM 500.0 vs 480.0")],
        legs={"ifc": True, "bm": True}, tol_pct=1.0,
        provenance={"takeoff": "t.json", "ifc": "m.ifc", "bm": "bm.json",
                    "generated": "2026-08-11 12:00", "tool": "0.1.32"})


def test_workbook_has_three_tabs_and_provenance(tmp_path):
    out = tmp_path / "audit.xlsx"
    write_audit_xlsx(_report(), str(out))
    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames == ["Summary", "Elements", "Exceptions"]
    ws = wb["Summary"]
    assert ws["A1"].value == "Quantity Audit"
    head = [ws.cell(r, 1).value for r in range(2, 8)]
    assert any(v and "t.json" in str(v) for v in head)          # provenance
    assert any(v and "±1.0%" in str(v) for v in head)           # tolerance


def test_summary_row_values_and_status(tmp_path):
    out = tmp_path / "a.xlsx"
    write_audit_xlsx(_report(), str(out))
    ws = openpyxl.load_workbook(str(out))["Summary"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[rows.index(("Schedule", "Item", "Description", "Unit",
                           "Takeoff qty", "IFC qty", "BM qty", "Measured",
                           "IFC Δ%", "BM Δ%", "Geom Δ%", "Status", "Verdict",
                           "Geom", "Elements", "Structures"))]
    data = rows[rows.index(hdr) + 1]
    # openpyxl reads back an empty-string cell as None; _report()'s ItemAudit
    # is hand-built (not via build_audit) so verdict/verdict_tier default ""
    assert data == ("A", "12", "Structural Concrete", "SF", 480.0, 481.0,
                    500.0, None, 0.21, 4.17, None, "VARIANCE", None, None,
                    1, 1)
    # the VARIANCE status cell is yellow (1-based sheet row after the header)
    data_row = rows.index(hdr) + 2
    assert ws.cell(data_row, 12).fill.fgColor.rgb == "FFFFEB9C"


def test_elements_and_exceptions_rows(tmp_path):
    out = tmp_path / "b.xlsx"
    write_audit_xlsx(_report(), str(out))
    wb = openpyxl.load_workbook(str(out))
    el = list(wb["Elements"].iter_rows(values_only=True))[1]
    assert el[:4] == ("A/12", "G1", "Deck-1", "IfcSlab")
    ex = list(wb["Exceptions"].iter_rows(values_only=True))[1]
    assert ex == ("variance", "A/12", "BM 500.0 vs 480.0")


def test_missing_leg_renders_na(tmp_path):
    rep = _report()
    rep.legs = {"ifc": False, "bm": True}
    rep.items[0].ifc_qty = None
    rep.items[0].delta_ifc_pct = None
    out = tmp_path / "c.xlsx"
    write_audit_xlsx(rep, str(out))
    ws = openpyxl.load_workbook(str(out))["Summary"]
    data = [r for r in ws.iter_rows(values_only=True)
            if r and r[0] == "A"][0]
    assert data[5] == "N/A"


def test_status_fills_ok_green_and_missing_red(tmp_path):
    rep = _report()
    rep.items[0].status = "OK"
    rep.items.append(ItemAudit(
        item_key="A/13", schedule="A", item_number="13",
        description="Riprap", unit="TON", takeoff_qty=None, ifc_qty=None,
        bm_qty=90.0, delta_ifc_pct=None, delta_bm_pct=None,
        status="MISSING IN TAKEOFF", element_count=0, structure_count=0))
    out = tmp_path / "d.xlsx"
    write_audit_xlsx(rep, str(out))
    ws = openpyxl.load_workbook(str(out))["Summary"]
    fills = {ws.cell(r, 2).value: ws.cell(r, 12).fill.fgColor.rgb
             for r in range(1, ws.max_row + 1)}
    assert fills["12"] == "FFC6EFCE"        # OK -> green
    assert fills["13"] == "FFFFC7CE"        # MISSING -> red


def test_geometry_sheet_written_when_mesh_leg_ran(tmp_path):
    si = 64.0 * (1200.0 / 3937.0) ** 3
    row = TakeoffRow(description="Footing Concrete", kind="quantity",
                     quantity=64.0, unit="CF", cost_code="03310",
                     item_number="40", schedule="A", structure_id="HW-3",
                     elements=[{"name": "FTG-1", "global_id": "G1",
                                "section_type": "footings", "count": 1,
                                "measure": "GrossVolume", "value_si": si}])
    ifc = {"G1": {"name": "FTG-1", "ifc_class": "IfcFooting",
                  "qtys": {"GrossVolume": si}}}
    mesh = {"G1": {"name": "FTG-1", "ifc_class": "IfcFooting",
                   "mesh": {"volume_m3": si * 1.05, "plan_area_m2": 1.486,
                            "len_m": 1.2192, "wid_m": 1.2192,
                            "hgt_m": 0.6096, "z_min_m": 0.0,
                            "z_max_m": 0.6096}}}
    rep = build_audit(CanonicalTakeoff(organization_id=1, project_ref="B1",
                                       rows=[row]), ifc=ifc, mesh=mesh)
    out = tmp_path / "a.xlsx"
    write_audit_xlsx(rep, str(out))
    wb = openpyxl.load_workbook(str(out))
    assert "Geometry" in wb.sheetnames
    geo = wb["Geometry"]
    hdr = [c.value for c in geo[1]]
    assert hdr == ["Item", "GlobalId", "Name", "IFC class", "Authored (m3)",
                   "Mesh (m3)", "Δ%", "Flag", "Plan (SF)", "L (ft)",
                   "W (ft)", "H (ft)", "Z bot (ft, model)",
                   "Z top (ft, model)"]
    r = [c.value for c in geo[2]]
    assert r[0] == "A/40" and r[2] == "FTG-1" and r[7] == "GEOM VARIANCE"
    assert r[9] == pytest.approx(4.0, abs=1e-3)          # 1.2192 m -> 4 ft
    # Summary carries the measured column and the geometry verdict
    # (rows 1-5 are title/provenance, row 6 is blank, row 7 is the header)
    ws = wb["Summary"]
    hdr = [c.value for c in ws[7]]
    assert "Measured" in hdr and "Geom Δ%" in hdr and "Geom" in hdr


def test_no_geometry_sheet_without_mesh_leg(tmp_path):
    rep = build_audit(CanonicalTakeoff(organization_id=1, project_ref="B1",
                                       rows=[]))
    out = tmp_path / "b.xlsx"
    write_audit_xlsx(rep, str(out))
    wb = openpyxl.load_workbook(str(out))
    assert "Geometry" not in wb.sheetnames
    # geometry columns exist but render blank, not "N/A"
    # (rows 1-5 are title/provenance, row 6 is blank, row 7 is the header)
    ws = wb["Summary"]
    hdr = [c.value for c in ws[7]]
    assert "Measured" in hdr


def test_verdict_column_written(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from estimating_tools.takeoff.audit import AuditReport, ItemAudit
    from estimating_tools.takeoff.audit_xlsx import write_audit_xlsx
    it = ItemAudit(item_key="A/1", schedule="A", item_number="1",
                   description="X", unit="CY", takeoff_qty=100.0,
                   ifc_qty=None, bm_qty=150.0, delta_ifc_pct=None,
                   delta_bm_pct=50.0, status="VARIANCE",
                   verdict="large variance - do NOT use for bidding",
                   verdict_tier="red")
    rep = AuditReport(items=[it])
    p = tmp_path / "a.xlsx"
    write_audit_xlsx(rep, str(p))
    ws = openpyxl.load_workbook(p)["Summary"]
    header = [c.value for c in ws[7]]
    assert "Verdict" in header
    col = header.index("Verdict") + 1
    assert ws.cell(8, col).value == "large variance - do NOT use for bidding"
