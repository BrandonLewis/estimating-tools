"""AuditReport -> audit.xlsx. Formatting only — no reconciliation logic here.
House style: rebar/schedule.py (lazy openpyxl, bold 14pt title, grey header
fill); status colors green/yellow/red."""
from __future__ import annotations

_GREEN, _YELLOW, _RED = "FFC6EFCE", "FFFFEB9C", "FFFFC7CE"
_STATUS_FILL = {"OK": _GREEN, "VARIANCE": _YELLOW}   # else red

FT_PER_M = 3937.0 / 1200.0

# geometry flags -> fill color: variance/unclassified red, missing-info
# yellow, OK/REFERENCE unfilled
_GEOM_FILL = {"GEOM VARIANCE": _RED, "UNCLASSIFIED": _RED,
              "NO GEOMETRY": _YELLOW, "NO AUTHORED VOLUME": _YELLOW}

_TIER_FILL = {"green": _GREEN, "yellow": _YELLOW, "red": _RED}


def _fill(color: str):
    from openpyxl.styles import PatternFill
    return PatternFill("solid", fgColor=color)


def _fill_for(status: str):
    return _fill(_STATUS_FILL.get(status, _RED))


def _na(v):
    return "N/A" if v is None else v


def _blank(v):
    """Geometry cells: blank (not 'N/A') when the leg or value is absent."""
    return "" if v in (None, "n/a") else v


def write_audit_xlsx(report, path: str) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:                      # pragma: no cover
        raise RuntimeError(
            "writing .xlsx requires openpyxl: pip install "
            "'estimating_tools[excel]'") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Quantity Audit"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=14)
    legs = ", ".join(k.upper() for k, v in sorted(report.legs.items()) if v)
    prov = report.provenance
    ws.append([f"takeoff: {prov.get('takeoff', '?')}"])
    ws.append([f"ifc: {prov.get('ifc') or '(not provided)'}"])
    ws.append([f"bm: {prov.get('bm') or '(not provided)'}"])
    ws.append([f"generated {prov.get('generated', '?')} by estimating_tools "
               f"{prov.get('tool', '?')} — legs: {legs or 'takeoff only'} — "
               f"tolerance ±{report.tol_pct}%"])
    ws.append([])

    header_fill = PatternFill("solid", fgColor="FFDDDDDD")

    def _header(sheet, cols):
        sheet.append(cols)
        for c in sheet[sheet.max_row]:
            c.font = Font(bold=True)
            c.fill = header_fill

    _header(ws, ["Schedule", "Item", "Description", "Unit", "Takeoff qty",
                 "IFC qty", "BM qty", "Measured", "IFC Δ%", "BM Δ%",
                 "Geom Δ%", "Status", "Verdict", "Geom", "Elements",
                 "Structures"])
    for it in report.items:
        ws.append([it.schedule, it.item_number, it.description, it.unit,
                   _na(it.takeoff_qty), _na(it.ifc_qty), _na(it.bm_qty),
                   _blank(it.measured_qty), _na(it.delta_ifc_pct),
                   _na(it.delta_bm_pct), _blank(it.delta_geom_pct),
                   it.status, it.verdict, _blank(it.geom_status),
                   it.element_count, it.structure_count])
        ws.cell(ws.max_row, 12).fill = _fill_for(it.status)
        tier = _TIER_FILL.get(it.verdict_tier)
        if tier:
            ws.cell(ws.max_row, 13).fill = _fill(tier)
        color = _GEOM_FILL.get(it.geom_status)
        if color:
            ws.cell(ws.max_row, 14).fill = _fill(color)

    els = wb.create_sheet("Elements")
    _header(els, ["Item", "GlobalId", "Name", "IFC class", "Section",
                  "Structure", "Count", "L (ft)", "W (ft)", "H (ft)",
                  "Measure", "Takeoff (SI)", "IFC (SI)", "Flag", "Join"])
    for e in report.elements:
        els.append([e.item_key, _na(e.global_id), e.name, _na(e.ifc_class),
                    e.section_type, _na(e.structure_id), e.count,
                    _na(e.length_ft), _na(e.width_ft), _na(e.height_ft),
                    _na(e.measure), _na(e.takeoff_si), _na(e.ifc_si),
                    e.delta_flag, e.join])
        if e.delta_flag != "OK":
            els.cell(els.max_row, 14).fill = _fill_for(e.delta_flag)

    if report.legs.get("mesh"):
        geo = wb.create_sheet("Geometry")
        _header(geo, ["Item", "GlobalId", "Name", "IFC class",
                      "Authored (m3)", "Mesh (m3)", "Δ%", "Flag",
                      "Plan (SF)", "L (ft)", "W (ft)", "H (ft)",
                      "Z bot (ft, model)", "Z top (ft, model)"])
        for g in report.geometry:
            geo.append([g.item_key or "(unclassified)", _na(g.global_id),
                        g.name, _na(g.ifc_class),
                        _blank(g.authored_si), _blank(g.mesh_si),
                        _blank(g.delta_pct), g.flag,
                        round(g.plan_area_m2 * FT_PER_M ** 2, 2),
                        round(g.len_m * FT_PER_M, 3),
                        round(g.wid_m * FT_PER_M, 3),
                        round(g.hgt_m * FT_PER_M, 3),
                        round(g.z_min_m * FT_PER_M, 3),
                        round(g.z_max_m * FT_PER_M, 3)])
            color = _GEOM_FILL.get(g.flag)
            if color:
                geo.cell(geo.max_row, 8).fill = _fill(color)

    exs = wb.create_sheet("Exceptions")
    _header(exs, ["Kind", "Item", "Detail"])
    for f in report.exceptions:
        exs.append([f.kind, f.item_key, f.detail])

    wb.save(path)
