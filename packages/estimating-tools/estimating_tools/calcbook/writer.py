"""Styled, byte-reproducible calculator-workbook writing (#202).

The shared library the Work calculators (drilling, power_sizing, vibrators,
concrete_cooling, sheet-pile, …) each hand-rolled: the estimator-input yellow
cell, a styled header / key-value / section helper, and a RATES table with the
override-column pattern — so a rate can be flexed on bid day without re-running
Python. Pure styling/formula plumbing; no org data.

Determinism reuses the package's one xlsx normalizer (``estimating_tools.xlsx_determinism``):
openpyxl restamps ``docProps/modified`` and every zip member with the wall
clock on save, so two runs of the same inputs only match by luck unless the
member stamps AND that string are pinned. ``save`` pins both.
"""
from __future__ import annotations

from typing import Iterable, Mapping, Sequence

from openpyxl.styles import Alignment, Font, PatternFill

from estimating_tools.calcbook.formula import cell_ref, override

# ---------------------------------------------------------------------------
# the shared style vocabulary
# ---------------------------------------------------------------------------

#: Estimator-input cells are yellow. The one convention every calculator wants
#: and applied inconsistently before this library.
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=12)
BOLD_FONT = Font(bold=True)
NOTE_FONT = Font(italic=True, color="808080")


def new_workbook(*sheet_names: str):
    """A workbook with pinned doc-properties and the given named sheets.

    With names, the default sheet is removed and each name created in order;
    without, the single default sheet is kept.
    """
    from openpyxl import Workbook

    from estimating_tools.xlsx_determinism import BUILD_STAMP
    wb = Workbook()
    wb.properties.created = wb.properties.modified = BUILD_STAMP
    if sheet_names:
        wb.remove(wb.active)
        for name in sheet_names:
            wb.create_sheet(name)
    return wb


def save(wb, path: str) -> str:
    """Save and pin the zip so identical inputs give byte-identical output."""
    from estimating_tools.xlsx_determinism import normalize_xlsx
    wb.save(path)
    normalize_xlsx(path)
    return path


def mark_input(cell) -> None:
    """Flag a cell as an estimator input (yellow fill)."""
    cell.fill = INPUT_FILL


def header(ws, columns: Sequence[str], *, row: int = 1,
           freeze: bool = True) -> None:
    """Write a styled (bold, filled) header row and freeze the rows above."""
    for i, name in enumerate(columns, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if freeze:
        ws.freeze_panes = f"A{row + 1}"


def section(ws, title: str, *, row: int, span: int = 1) -> None:
    """A bold section-title cell, optionally merged across ``span`` columns."""
    c = ws.cell(row=row, column=1, value=title)
    c.font = TITLE_FONT
    if span > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row,
                       end_column=span)


def kv(ws, rows: Iterable[Sequence], *, start_row: int = 1,
       inputs: Iterable[str] = ()) -> dict[str, str]:
    """Write label / value / (optional note) rows into columns A / B / C.

    Value cells whose label is in ``inputs`` are marked yellow. Returns
    ``{label: value_cell_coord}`` so downstream formulas can reference them.
    """
    input_set = set(inputs)
    refs: dict[str, str] = {}
    for i, entry in enumerate(rows):
        r = start_row + i
        label, value = entry[0], entry[1]
        note = entry[2] if len(entry) > 2 else None
        ws.cell(row=r, column=1, value=label)
        vc = ws.cell(row=r, column=2, value=value)
        if note is not None:
            ws.cell(row=r, column=3, value=note)
        if label in input_set:
            mark_input(vc)
        refs[label] = vc.coordinate
    return refs


def set_widths(ws, widths: Mapping[str, float]) -> None:
    """Set column widths from ``{column_letter: width}``."""
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def rate_table(ws, *, headers: Sequence[str], rows: Sequence[Sequence],
               shipped_cols: Sequence[str], override_col: str,
               effective_cols: Sequence[str], start_row: int = 1,
               sheet: str | None = None) -> dict[str, list[str]]:
    """A RATES sheet with the override-column pattern.

    Each ``effective_cols[i]`` cell is ``=IF(override="", shipped[i], override)``
    off ``shipped_cols[i]`` and the single ``override_col`` (an estimator input,
    yellow). Returns ``{effective_header: [abs cell ref per data row]}`` for
    downstream formulas (e.g. a PERT off the effective low/likely/high).
    """
    if len(shipped_cols) != len(effective_cols):
        raise ValueError("shipped_cols and effective_cols must align")
    from openpyxl.utils import column_index_from_string, get_column_letter
    sheet = sheet or ws.title
    header(ws, headers, row=start_row)
    col_header = {get_column_letter(i + 1): h
                  for i, h in enumerate(headers)}
    override_idx = column_index_from_string(override_col)
    refs: dict[str, list[str]] = {}
    for r_off, values in enumerate(rows):
        r = start_row + 1 + r_off
        for c_off, val in enumerate(values, start=1):
            ws.cell(row=r, column=c_off, value=val)
        mark_input(ws.cell(row=r, column=override_idx))
        for eff_col, ship_col in zip(effective_cols, shipped_cols):
            ship_ref = f"{ship_col}{r}"              # relative
            over_ref = f"${override_col}{r}"         # abs col, rel row
            ws[f"{eff_col}{r}"] = override(ship_ref, over_ref)
            refs.setdefault(col_header.get(eff_col, eff_col), []).append(
                cell_ref(sheet, eff_col, r))
    return refs
