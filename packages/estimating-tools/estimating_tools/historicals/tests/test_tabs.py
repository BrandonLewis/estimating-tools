from pathlib import Path

from estimating_tools.historicals.tabs import cross_check, load_csv, load_excel, parse_table

FIX = Path(__file__).parent / "fixtures"


def test_parse_csv_fixture():
    tab = load_csv(FIX / "sr49_tab.csv")
    assert tab.agency == "caltrans"
    assert tab.project_ref == "sr-49-widening-2026"
    assert tab.bid_date == "2026-05-14"
    assert tab.contract == "03-1J2304"
    assert [b.name for b in tab.bidders] == ["Ghilotti", "Teichert"]
    assert tab.bidders[0].rank == 1 and tab.bidders[0].total_bid == 684800.0
    assert len(tab.items) == 2
    item = tab.items[1]
    assert item.bid_item_code == "510120" and item.quantity == 412.0
    assert item.unit_prices == {"Ghilotti": 1650.0, "Teichert": 1710.0}
    assert item.unit == "CY"


def test_leading_zeros_preserved():
    tab = load_csv(FIX / "sr49_tab.csv")
    assert tab.items[0].bid_item_code == "070030"   # string, not 70030


def test_bad_row_collected_not_fatal():
    rows = [
        ["Agency:", "caltrans"], ["Project:", "p1"],
        ["Item No", "Item Code", "Description", "Quantity", "Unit",
         "A Unit Price", "A Amount"],
        ["1", "070030", "MOB", "not-a-number", "LS", "5", "5"],
        ["2", "510120", "CONC", "10", "CY", "100", "1000"],
    ]
    tab = parse_table(rows)
    assert len(tab.items) == 1                       # bad row skipped
    assert any("row" in w for w in tab.warnings)     # ...and reported


def test_cross_check_flags_total_mismatch():
    tab = load_csv(FIX / "sr49_tab.csv")
    tab.bidders[0].total_bid = 999999.0
    warnings = cross_check(tab)
    assert any("Ghilotti" in w for w in warnings)
    tab2 = load_csv(FIX / "sr49_tab.csv")
    assert cross_check(tab2) == []                   # clean tab, no flags


def test_ragged_section_break_row_not_fatal():
    rows = [
        ["Agency:", "caltrans"], ["Project:", "p1"],
        ["Item No", "Item Code", "Description", "Quantity", "Unit",
         "A Unit Price", "A Amount"],
        ["1", "070030", "MOB", "1", "LS", "5", "5"],
        ["ROADWAY ITEMS"],                            # section break, 1 cell
        ["2", "510120", "CONC", "10", "CY", "100", "1000"],
    ]
    tab = parse_table(rows)                           # must not raise
    assert [it.bid_item_code for it in tab.items] == ["070030", "510120"]
    assert any("row 5" in w for w in tab.warnings)


def test_cross_check_flags_empty_extension_column():
    tab = load_csv(FIX / "sr49_tab.csv")
    for item in tab.items:                            # bidder column unparseable
        item.unit_prices["Teichert"] = None
        item.extensions["Teichert"] = None
    warnings = cross_check(tab)
    assert any("Teichert" in w and "no item extensions" in w for w in warnings)


def test_header_shape_mismatch_warns():
    rows = [
        ["Item No", "Item Code", "Description", "Quantity", "Unit",
         "A Unit Price", "A Amount", "Notes", "B Unit Price", "B Amount"],
        ["1", "070030", "MOB", "1", "LS", "5", "5", "x", "6", "6"],
    ]
    tab = parse_table(rows)
    assert any("misattributed" in w for w in tab.warnings)


def test_excel_trailing_empty_header_cells_no_shape_warning(tmp_path):
    # Excel's iter_rows commonly yields trailing None cells beyond the data
    # range (styled-but-empty cells widen max_column); they must not trip
    # the header-shape warning on a valid tab.
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.append(["Agency:", "caltrans"])
    ws.append(["Contract:", "03-1J2304"])
    ws.append(["Project:", "sr-49-widening-2026"])
    ws.append(["Bid Date:", "2026-05-14"])
    ws.append(["Region:", "sierra"])
    ws.append(["Item No", "Item Code", "Description", "Quantity", "Unit",
               "Ghilotti Unit Price", "Ghilotti Amount",
               "Teichert Unit Price", "Teichert Amount"])
    ws.append([1, "070030", "LEAD COMPLIANCE PLAN", 1, "LS",
               5000.00, 5000.00, 7500.00, 7500.00])
    ws.append(["TOTAL", "", "", "", "", "", 684800.00, "", 712020.00])
    # styled-but-empty cell beyond the data range -> iter_rows pads every
    # row (header included) with trailing None cells
    ws.cell(row=1, column=12).font = Font(bold=True)
    path = tmp_path / "sr49_tab_trailing.xlsx"
    wb.save(path)

    tab = load_excel(path)
    assert [b.name for b in tab.bidders] == ["Ghilotti", "Teichert"]
    assert not any("misattributed" in w for w in tab.warnings)
    assert tab.items[0].unit_prices == {"Ghilotti": 5000.0, "Teichert": 7500.0}


def test_total_row_variants_not_parsed_as_items():
    rows = [
        ["Item No", "Item Code", "Description", "Quantity", "Unit",
         "A Unit Price", "A Amount"],
        ["1", "070030", "MOB", "1", "LS", "5", "5"],
        ["Total Bid", "", "", "", "", "", "5.00"],
    ]
    tab = parse_table(rows)
    assert len(tab.items) == 1                        # not treated as an item
    assert tab.bidders[0].total_bid == 5.0


def test_region_lowercased():
    rows = [
        ["Agency:", "Caltrans"], ["Region:", "SIERRA"],
        ["Item No", "Item Code", "Description", "Quantity", "Unit",
         "A Unit Price", "A Amount"],
        ["1", "070030", "MOB", "1", "LS", "5", "5"],
    ]
    tab = parse_table(rows)
    assert tab.region == "sierra"


def test_load_pdf_concatenates_page_tables(monkeypatch):
    from estimating_tools.historicals import tabs

    pages_tables = [
        [[["Agency:", "caltrans"], ["Project:", "p1"],
          ["Item No", "Item Code", "Description", "Quantity", "Unit",
           "A Unit Price", "A Amount"]]],
        [[["1", "070030", "MOB", "1", "LS", "5", "5"]]],
    ]

    class _Page:
        def __init__(self, tables): self._t = tables
        def extract_tables(self): return self._t

    class _Pdf:
        pages = [_Page(t) for t in pages_tables]
        def __enter__(self): return self
        def __exit__(self, *a): return False

    monkeypatch.setattr(tabs, "_open_pdf", lambda path: _Pdf())
    tab = tabs.load_pdf("whatever.pdf")
    assert tab.agency == "caltrans"
    assert len(tab.items) == 1 and tab.items[0].bid_item_code == "070030"


def test_load_excel_matches_csv_fixture(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Agency:", "caltrans"])
    ws.append(["Contract:", "03-1J2304"])
    ws.append(["Project:", "sr-49-widening-2026"])
    ws.append(["Bid Date:", "2026-05-14"])
    ws.append(["Region:", "sierra"])
    ws.append(["Item No", "Item Code", "Description", "Quantity", "Unit",
               "Ghilotti Unit Price", "Ghilotti Amount",
               "Teichert Unit Price", "Teichert Amount"])
    ws.append([1, "070030", "LEAD COMPLIANCE PLAN", 1, "LS",
               5000.00, 5000.00, 7500.00, 7500.00])
    ws.append(["TOTAL", "", "", "", "", "", 684800.00, "", 712020.00])
    path = tmp_path / "sr49_tab.xlsx"
    wb.save(path)

    tab = load_excel(path)
    assert tab.agency == "caltrans"
    assert tab.project_ref == "sr-49-widening-2026"
    assert tab.contract == "03-1J2304"
    assert [b.name for b in tab.bidders] == ["Ghilotti", "Teichert"]
    assert tab.bidders[0].total_bid == 684800.0
    assert len(tab.items) == 1
    item = tab.items[0]
    assert item.bid_item_code == "070030"
    assert item.unit_prices == {"Ghilotti": 5000.0, "Teichert": 7500.0}


def test_load_dispatches_by_suffix_and_rejects_unsupported(tmp_path):
    from estimating_tools.historicals.tabs import load

    tab = load(FIX / "sr49_tab.csv")
    assert tab.agency == "caltrans"
    assert len(tab.items) == 2

    bad = tmp_path / "tab.docx"
    bad.write_text("nope")
    try:
        load(bad)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "docx" in str(exc)
