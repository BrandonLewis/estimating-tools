"""Bid-tab parsing: CSV/Excel/PDF-extracted tables -> CanonicalBidTab.

Layout contract (v1): optional metadata rows ("Agency:", value), a header
row starting Item No/Item Code/Description/Quantity/Unit followed by
per-bidder pairs "<Name> Unit Price","<Name> Amount", item rows, optional
TOTAL row. Per-row failures are collected in tab.warnings, never fatal.
"""
from __future__ import annotations

import csv
from pathlib import Path

from estimating_tools.historicals.model import Bidder, CanonicalBidTab, TabItem

_META_KEYS = {"agency": "agency", "contract": "contract", "project": "project_ref",
              "bid date": "bid_date", "region": "region"}
_FIXED_COLS = 5           # Item No, Item Code, Description, Quantity, Unit
_TOTAL_TOLERANCE = 0.005  # 0.5%


def _num(cell) -> float | None:
    if cell is None:
        return None
    s = str(cell).strip().replace("$", "").replace(",", "")
    if not s:
        return None
    return float(s)   # caller catches ValueError


def parse_table(rows: list[list]) -> CanonicalBidTab:
    meta: dict = {}
    header: list[str] | None = None
    items: list[TabItem] = []
    warnings: list[str] = []
    bidders: list[Bidder] = []
    totals_row = None

    for idx, raw in enumerate(rows):
        cells = ["" if c is None else str(c).strip() for c in raw]
        if not any(cells):
            continue
        if header is None:
            if cells[0].endswith(":"):
                key = cells[0].rstrip(":").strip().lower()
                if key in _META_KEYS and len(cells) > 1:
                    meta[_META_KEYS[key]] = cells[1]
                continue
            low = [c.lower() for c in cells]
            if "unit" in low and any("item" in c for c in low):
                header = cells
                names = []
                for h in header[_FIXED_COLS:]:
                    for suffix in (" Unit Price", " Amount"):
                        if h.endswith(suffix):
                            name = h[: -len(suffix)]
                            if name not in names:
                                names.append(name)
                bidders = [Bidder(name=n) for n in names]
                expected_tail = [f"{n}{suffix}" for n in names
                                 for suffix in (" Unit Price", " Amount")]
                tail = header[_FIXED_COLS:]
                while tail and not tail[-1]:
                    tail = tail[:-1]  # trailing empties (Excel range padding)
                if tail != expected_tail:
                    warnings.append(
                        "bidder columns not in expected Unit Price/Amount "
                        "pairs; prices may be misattributed")
                continue
            warnings.append(f"row {idx + 1}: unrecognized pre-header row")
            continue
        if cells[0].upper().startswith("TOTAL"):
            totals_row = cells
            continue
        try:
            item = TabItem(
                item_number=cells[0],
                bid_item_code=cells[1],
                description=cells[2],
                quantity=_num(cells[3]),
                unit=cells[4],
            )
            for b_i, bidder in enumerate(bidders):
                up_col = _FIXED_COLS + 2 * b_i
                ext_col = up_col + 1
                item.unit_prices[bidder.name] = (
                    _num(cells[up_col]) if up_col < len(cells) else None)
                item.extensions[bidder.name] = (
                    _num(cells[ext_col]) if ext_col < len(cells) else None)
            items.append(item)
        except (ValueError, IndexError) as exc:
            warnings.append(f"row {idx + 1}: {exc}")

    if totals_row is not None:
        for b_i, bidder in enumerate(bidders):
            ext_col = _FIXED_COLS + 2 * b_i + 1
            if ext_col < len(totals_row):
                try:
                    bidder.total_bid = _num(totals_row[ext_col])
                except ValueError:
                    warnings.append(f"TOTAL: unreadable amount for {bidder.name}")
    totaled = sorted((b for b in bidders if b.total_bid is not None),
                     key=lambda b: b.total_bid)
    for rank, bidder in enumerate(totaled, start=1):
        bidder.rank = rank

    if header is None:
        warnings.append("no header row found (need Item/Unit columns)")
    return CanonicalBidTab(
        agency=meta.get("agency", "").lower(),
        project_ref=meta.get("project_ref", ""),
        bid_date=meta.get("bid_date"),
        region=meta.get("region", "").lower(),
        contract=meta.get("contract", ""),
        bidders=bidders, items=items, warnings=warnings,
    )


def cross_check(tab: CanonicalBidTab) -> list[str]:
    """Per bidder: sum of extensions (or qty*price) vs stated total bid."""
    flags = []
    for bidder in tab.bidders:
        if bidder.total_bid is None:
            continue
        total = 0.0
        for item in tab.items:
            ext = item.extensions.get(bidder.name)
            if ext is None:
                price = item.unit_prices.get(bidder.name)
                ext = (price * item.quantity
                       if price is not None and item.quantity is not None else None)
            if ext is not None:
                total += ext
        if total == 0.0 and bidder.total_bid:
            flags.append(
                f"{bidder.name}: no item extensions to check against "
                f"stated total {bidder.total_bid:,.2f}")
        elif total and abs(total - bidder.total_bid) > _TOTAL_TOLERANCE * bidder.total_bid:
            flags.append(
                f"{bidder.name}: item extensions sum {total:,.2f} != "
                f"stated total {bidder.total_bid:,.2f}")
    return flags


def load_csv(path: str | Path) -> CanonicalBidTab:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return parse_table(list(csv.reader(fh)))


def load_excel(path: str | Path) -> CanonicalBidTab:
    """Load a bid tab from an .xlsx/.xlsm workbook.

    Only the workbook's ACTIVE sheet is read; other sheets are ignored.
    """
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    return parse_table([list(row) for row in ws.iter_rows(values_only=True)])


def _open_pdf(path):
    import pdfplumber
    return pdfplumber.open(path)


def load_pdf(path: str | Path) -> CanonicalBidTab:
    """Extract every table from every page, in order, into one row stream."""
    rows: list[list] = []
    with _open_pdf(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                rows.extend(table)
    return parse_table(rows)


def load(path: str | Path) -> CanonicalBidTab:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return load_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return load_excel(path)
    if suffix == ".pdf":
        return load_pdf(path)
    raise ValueError(f"unsupported bid-tab format: {suffix or path}")
