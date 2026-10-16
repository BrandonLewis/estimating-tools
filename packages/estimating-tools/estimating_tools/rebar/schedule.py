"""Bar bending schedule (BBS): cut lengths, weights, and LBS takeoff rows for
a set of rebar marks — an ESTIMATING AID (see rebar/detailing.py), not a
detailer's shop BBS. Hook allowances come from
``estimating_tools.rebar.detailing.hook_allowance_in``, a conservative
estimating allowance rather than a shop bend-deduction table; design/detailing
governs the actual bar list.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from estimating_tools.rebar.detailing import hook_allowance_in
from estimating_tools.rebar.model import BarSize
from estimating_tools.takeoff.core import WASTE_FACTORS

#: default rebar waste factor, imported from the shared takeoff table
#: (estimating_tools.takeoff.core.WASTE_FACTORS["rebar"])
DEFAULT_WASTE_FACTOR = WASTE_FACTORS["rebar"]


@dataclass
class BarMark:
    """One bar mark: a repeated bar shape (straight or hooked legs)."""

    mark: str
    size: BarSize
    count: int
    legs_in: list[float]
    hooks: list[str] = field(default_factory=list)
    note: str = ""

    def __post_init__(self) -> None:
        if self.count < 1:
            raise ValueError(
                f"mark {self.mark!r}: count must be >= 1, got {self.count}"
            )
        if not self.legs_in:
            raise ValueError(f"mark {self.mark!r}: legs_in must not be empty")
        for leg in self.legs_in:
            if leg <= 0:
                raise ValueError(
                    f"mark {self.mark!r}: leg lengths must be positive, got {leg}"
                )
        for h in self.hooks:
            try:
                hook_allowance_in(self.size, h)
            except ValueError as exc:
                raise ValueError(f"mark {self.mark!r}: {exc}") from exc

    @property
    def cut_length_in(self) -> float:
        """Straight legs plus each hook's allowance (bend arc + extension)."""
        hook_total = sum(hook_allowance_in(self.size, h) for h in self.hooks)
        return sum(self.legs_in) + hook_total

    @property
    def cut_length_ft(self) -> float:
        return self.cut_length_in / 12.0

    @property
    def weight_lb(self) -> float:
        return self.cut_length_ft * self.size.weight_lb_per_ft * self.count

    def legs_summary(self) -> str:
        return " + ".join(_fmt_num(leg) for leg in self.legs_in)

    def hooks_summary(self) -> str:
        return ", ".join(self.hooks) if self.hooks else "-"


@dataclass
class BarSchedule:
    """A bar bending schedule: a titled set of marks. Rejects duplicate marks
    rather than silently dropping/overwriting them."""

    title: str
    marks: list[BarMark]

    def __post_init__(self) -> None:
        seen: set[str] = set()
        for m in self.marks:
            if m.mark in seen:
                raise ValueError(
                    f"duplicate mark {m.mark!r} in schedule {self.title!r}"
                )
            seen.add(m.mark)

    @property
    def total_weight_lb(self) -> float:
        return sum(m.weight_lb for m in self.marks)

    def weight_by_size(self) -> dict[int, float]:
        """Total weight (lb), summed per bar-size number (e.g. 5 for #5)."""
        totals: dict[int, float] = {}
        for m in self.marks:
            totals[m.size.number] = totals.get(m.size.number, 0.0) + m.weight_lb
        return totals

    def to_rows(self) -> list[dict[str, Any]]:
        """One dict per mark: mark, size, count, legs, hooks, cut_ft, weight_lb."""
        rows = []
        for m in self.marks:
            rows.append({
                "mark": m.mark,
                "size": str(m.size),
                "count": m.count,
                "legs": m.legs_summary(),
                "hooks": m.hooks_summary(),
                "cut_ft": round(m.cut_length_ft, 2),
                "weight_lb": round(m.weight_lb, 1),
            })
        return rows

    def to_markdown(self) -> str:
        """GFM table of marks plus a bold total-weight row."""
        lines = [
            f"# {self.title}",
            "",
            "| Mark | Size | Count | Legs (in) | Hooks | Cut (ft) | Weight (lb) |",
            "|---|---|---|---|---|---|---|",
        ]
        for r in self.to_rows():
            lines.append(
                f"| {r['mark']} | {r['size']} | {r['count']} | {r['legs']} | "
                f"{r['hooks']} | {r['cut_ft']:.2f} | {r['weight_lb']:.1f} |"
            )
        lines.append(
            f"| **Total** | | | | | | **{self.total_weight_lb:.1f}** |"
        )
        return "\n".join(lines)

    def to_takeoff_rows(
        self,
        *,
        cost_code: str | None = None,
        item_number: str | None = None,
        schedule: str | None = None,
        waste_factor: float = DEFAULT_WASTE_FACTOR,
    ) -> list[dict[str, Any]]:
        """One quantity/LBS ``TakeoffRow``-shaped dict for the whole schedule
        (total weight x waste_factor); ``notes`` lists the marks included."""
        marks_note = ", ".join(m.mark for m in self.marks)
        return [{
            "description": self.title,
            "kind": "quantity",
            "quantity": round(self.total_weight_lb * waste_factor, 1),
            "unit": "LBS",
            "cost_code": cost_code,
            "item_number": item_number,
            "schedule": schedule,
            "waste_factor": waste_factor,
            "notes": f"marks: {marks_note}",
        }]

    def to_xlsx(self, path: str | Path) -> None:
        """Write a simple styled workbook: title row, header row (bold +
        fill), one row per mark, a bold totals row, and set column widths."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise ImportError(
                "writing .xlsx requires openpyxl: pip install 'estimating_tools[excel]'"
            ) from exc

        header = ["Mark", "Size", "Count", "Legs (in)", "Hooks", "Cut (ft)", "Weight (lb)"]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "BBS"

        ws.append([self.title])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=14)
        ws.append([])

        ws.append(header)
        header_row = ws.max_row
        header_fill = PatternFill("solid", fgColor="FFDDDDDD")
        for ci in range(1, len(header) + 1):
            cell = ws.cell(header_row, ci)
            cell.font = Font(bold=True)
            cell.fill = header_fill

        for r in self.to_rows():
            ws.append([r["mark"], r["size"], r["count"], r["legs"], r["hooks"],
                       r["cut_ft"], r["weight_lb"]])

        ws.append(["Total", "", "", "", "", "", round(self.total_weight_lb, 1)])
        total_row = ws.max_row
        for ci in range(1, len(header) + 1):
            ws.cell(total_row, ci).font = Font(bold=True)

        widths = [10, 8, 8, 20, 16, 10, 12]
        for ci, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = w

        wb.save(str(path))


def _fmt_num(value: float) -> str:
    """Compact numeric formatting for summaries: 24.0 -> "24", 24.5 -> "24.5"."""
    return f"{value:g}"
