"""CIP concrete-structure buildup workbook generator (`est-estimate cip-buildup`).

LEGACY GENERATOR, NOT A FORMAT PACK: unlike build.py (which renders a
format-pack recipe), `cip-buildup` is a standalone, hard-coded generator --
a port of the Work-tree structure_estimate.py. No format pack drives it and
no consuming org's estimate config runs through it any more; it stays for
whoever still wants the CIP buildup shape without adopting the recipe engine.

``DEFAULT_RATES`` and ``DEFAULT_LABOR`` below are ILLUSTRATIVE PLACEHOLDERS
ONLY -- round, made-up numbers in the same shape as a real rate table, not
any org's actual crew rates. A real user supplies real rates through their
own config's ``rates`` / ``labor_sections`` keys (see "Optional per-structure
keys" below); nothing here should be read as pricing guidance.

A config-driven generator for cast-in-place concrete-structure estimate
workbooks. The .xlsx carries a **Rates** sheet (single source of truth for
unit rates), one **buildup sheet per structure**, a **Summary** (Each x Qty =
Extended), and — when any structure carries subcontract lines — a **Sub
Spread** rollup sheet.

Each structure sheet follows this buildup template:

    concrete takeoff (closed box L/W/H  OR  component prisms -> CY w/ waste)
    LABOR     : lettered phases, crew composition x days x hrs/day
    EQUIPMENT : pickup, tools, skytrak, concrete pump, volume
    MATERIAL  : concrete, rebar (unless carried as a sub), short loads
    LUMBER / HARDWARE / FUEL / SPECIALS
    SELF-PERFORM SUBTOTAL [+ SUBCONTRACTORS - SUB SPREAD] -> STRUCTURE TOTAL

All inputs are editable yellow literals; every derived cell is a live Excel
formula (rates addressed via ``Rates!$B$n``), so the estimator can tune the
workbook without rerunning the tool.

Departures from the Work original:
  * config validation — unknown keys raise ValueError instead of being
    silently ignored (a typo like ``labour_sections`` is a silent-wrong-bid
    class, not a no-op);
  * rebar stays in MATERIAL unless actually subbed — the original dropped the
    per-$/CY rebar line whenever ANY ``subs`` existed, so an earthwork-only
    subs list silently un-priced the rebar; the port removes it only for a
    ``rebar_cost`` shorthand or an explicit ``kind: "rebar"`` sub;
  * deterministic output — docProps and zip member timestamps are pinned to
    CIP_BUILD_STAMP (mirrors bimforge's QTO_BUILD_STAMP) so rebuilding an
    unchanged config is byte-identical;
  * de-orged SAMPLE (no customer/vendor names).

Optional per-structure keys (backward compatible with the Work configs):
    components  : list of {name, count, L, W, H} rectangular prisms. When
                  present the takeoff uses these (CY = count*L*W*H/27, counts
                  may be negative for opening deductions) instead of the
                  closed-box L/W/H geometry.
    rebar_cost  : lump-sum rebar $ (e.g. a sub quote). Moves rebar OUT of
                  MATERIAL and into SUBCONTRACTORS; rebar_vendor names the sub.
    subs        : explicit list of {name, vendor, qty, unit,
                  unit_price|amount, kind} subcontract lines. Any sub also
                  rolls up to the Sub Spread sheet.
    days        : {phase-id: days} overrides of the per-phase day defaults.
    count       : structures of this type (default 1); Summary extends by it.

Pure file operations + openpyxl; the CLI surface lives in cli.py
(``est-estimate cip-buildup``).
"""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Font, PatternFill

from estimating_tools.xlsx_determinism import BUILD_STAMP, normalize_xlsx

#: Compatibility alias. This module owned the determinism helpers until they
#: were lifted out; three unrelated modules imported them from here. Remove
#: when cip.py becomes a format pack.
CIP_BUILD_STAMP = BUILD_STAMP
_normalize_xlsx = normalize_xlsx

# ---- defaults (illustrative placeholders -- see module docstring) ----
DEFAULT_RATES = [
    ('Labor FORE $/hr', 100), ('Labor CARP $/hr', 90), ('Labor LAB $/hr', 80), ('Labor MASON $/hr', 80),
    ('Hrs per day', 8), ('Pickup $/hr', 20), ('Tools $/hr', 20), ('Skytrak $/hr', 20), ('Concrete Pump $/hr', 250),
    ('Volume $/CY', 10), ('Concrete $/CY', 200), ('Rebar $/CY', 700), ('Short load $/EA', 400), ('Lumber $/CY', 100),
    ('Build panel $/hr', 200), ('Hardware $/CY', 50), ('Fuel $/hr', 20), ('Waterstop $/LF', 20), ('Waste fraction', 0.10),
]
# phase id, label, FORE, CARP, LAB, MASON, default days
DEFAULT_LABOR = [
    ('A', '[A] Footings/Invert - form, fine grade, strip', 1, 1, 1, 0, 2.0),
    ('B', '[B] Pour footings / invert', 1, 1, 2, 0, 1.0),
    ('C', '[C] Walls - form (incl. channel & bench)', 1, 2, 1, 0, 3.0),
    ('D', '[D] Pour walls', 1, 1, 1, 2, 1.0),
    ('E', '[E] Deck / top slab - form', 1, 2, 1, 0, 1.0),
    ('F', '[F] Pour deck', 1, 1, 2, 0, 1.0),
    ('G', '[G] Sandblast / clean / finish - clean-up', 1, 0, 1, 0, 1.0),
]
DEF = dict(wall=1.0, base=1.0, deck=1.0, pickup_hr=88, tools_hr=88, skytrak_hr=88, pump_hr=24,
           short_loads=3, panel_hrs=16, fuel_pickup_hr=88, fuel_equip_hr=88, waterstop_lf=0)

SAMPLE = {
    "title": "Sample Drainage Improvements - Junction Boxes A-B (CIP) Estimate",
    "basis_note": "basis = prior CIP outfall buildup. EXCLUDES excavation & dewatering. "
                  "Dims read from plan details - VERIFY.",
    "output": "Junction Boxes A-B - CIP Buildup.xlsx",
    "structures": [
        {"sheet": "JS-1 (Box A)", "title": "JUNCTION STRUCTURE 1 - Box A",
         "bid_item": "Item 13 - Junction Box A",
         "note": "24in, 36in (x2) SD mains", "L": 9.5, "W": 9.5, "H": 11.0},
        {"sheet": "JS-2 (Box B)", "title": "JUNCTION STRUCTURE 2 - Box B",
         "bid_item": "Item 14 - Junction Box B",
         "note": "24in, 36in, 54in FRP SD mains", "L": 7.5, "W": 7.5, "H": 10.1},
    ],
}

B = Font(bold=True); TITLE = Font(bold=True, size=13); WHITE = Font(bold=True, color='FFFFFF')
INP = PatternFill('solid', fgColor='FFF2CC'); SUB = PatternFill('solid', fgColor='D9E1F2')
HD = PatternFill('solid', fgColor='4472C4'); GT = PatternFill('solid', fgColor='C6E0B4')
MONEY = '#,##0'; NUM = '#,##0.00'

# =============================================================================
# Validation — the Work original silently ignored unknown keys; a typo'd
# override then rode the defaults into the bid. The port refuses instead.
# =============================================================================

_TOP_KEYS = {'title', 'basis_note', 'output', 'rates', 'labor_sections', 'structures'}
_STRUCT_KEYS = {'sheet', 'title', 'bid_item', 'note', 'count', 'components',
                'L', 'W', 'H', 'wall', 'base', 'deck', 'rebar_cost', 'rebar_vendor',
                'days', 'subs', 'pickup_hr', 'tools_hr', 'skytrak_hr', 'pump_hr',
                'short_loads', 'panel_hrs', 'fuel_pickup_hr', 'fuel_equip_hr',
                'waterstop_lf'}
_COMPONENT_KEYS = {'name', 'count', 'L', 'W', 'H'}
_SUB_KEYS = {'name', 'vendor', 'qty', 'unit', 'unit_price', 'amount', 'kind'}
_LABOR_KEYS = {'id', 'name', 'fore', 'carp', 'lab', 'mason', 'days'}
_RATE_LABELS = {lab for lab, _v in DEFAULT_RATES}


def _check_keys(mapping, allowed, where):
    unknown = sorted(set(mapping) - allowed)
    if unknown:
        raise ValueError(
            "unknown %s key(s) %s — known keys: %s"
            % (where, ", ".join(repr(k) for k in unknown), ", ".join(sorted(allowed))))


def validate_config(config):
    """Raise ValueError on unknown keys / missing geometry. Returns None."""
    if not isinstance(config, dict):
        raise ValueError("config must be a JSON object")
    _check_keys(config, _TOP_KEYS, "config")
    structures = config.get('structures')
    if not isinstance(structures, list) or not structures:
        raise ValueError("config needs a non-empty 'structures' list")

    unknown_rates = sorted(set(config.get('rates', {})) - _RATE_LABELS)
    if unknown_rates:
        raise ValueError(
            "unknown rate label(s) %s — known labels: %s"
            % (", ".join(repr(k) for k in unknown_rates),
               ", ".join(lab for lab, _v in DEFAULT_RATES)))

    sections = config.get('labor_sections')
    if sections is not None:
        if not isinstance(sections, list) or not sections:
            raise ValueError("'labor_sections' must be a non-empty list")
        for sec in sections:
            _check_keys(sec, _LABOR_KEYS, "labor_sections")
            missing = sorted(_LABOR_KEYS - set(sec))
            if missing:
                raise ValueError(
                    "labor_sections entry %r missing key(s): %s"
                    % (sec.get('id', sec.get('name', '?')), ", ".join(missing)))
        phase_ids = {sec['id'] for sec in sections}
    else:
        phase_ids = {s[0] for s in DEFAULT_LABOR}

    sheets = [s.get('sheet') for s in structures]
    dupes = sorted({n for n in sheets if n and sheets.count(n) > 1})
    if dupes:
        # openpyxl silently renames the second "S1" sheet to "S11" while both
        # Summary rows still reference 'S1' — one structure counted twice, the
        # other contributing $0. Refuse instead (silent-wrong-bid class).
        raise ValueError(
            "duplicate structure sheet name(s): %s — every structure needs a "
            "unique 'sheet'" % ", ".join(repr(d) for d in dupes))

    for s in structures:
        sheet = s.get('sheet')
        if not sheet:
            raise ValueError("every structure needs a 'sheet' name")
        _check_keys(s, _STRUCT_KEYS, "structure %r" % sheet)
        comps = s.get('components')
        if comps:
            for comp in comps:
                _check_keys(comp, _COMPONENT_KEYS, "structure %r component" % sheet)
                missing = sorted(_COMPONENT_KEYS - set(comp))
                if missing:
                    raise ValueError(
                        "structure %r component %r missing key(s): %s"
                        % (sheet, comp.get('name', '?'), ", ".join(missing)))
        elif not all(k in s for k in ('L', 'W', 'H')):
            raise ValueError(
                "structure %r has neither 'components' nor closed-box L/W/H "
                "geometry" % sheet)
        unknown_days = sorted(set(s.get('days', {})) - phase_ids)
        if unknown_days:
            raise ValueError(
                "structure %r 'days' key(s) %s do not match any labor phase id "
                "(%s)" % (sheet, ", ".join(repr(k) for k in unknown_days),
                          ", ".join(sorted(phase_ids))))
        for sub in s.get('subs', []):
            _check_keys(sub, _SUB_KEYS, "structure %r sub" % sheet)
            if not sub.get('name'):
                raise ValueError("structure %r has a sub without a 'name'" % sheet)
            if sub.get('unit_price') is None and sub.get('amount') is None:
                # a priceless sub prices the line at $0 — and a kind:"rebar"
                # sub ALSO removes the MATERIAL rebar line, silently
                # un-pricing the whole rebar scope. Refuse.
                raise ValueError(
                    "structure %r sub %r needs 'unit_price' or 'amount'"
                    % (sheet, sub['name']))


# =============================================================================
# Determinism — normalize the saved zip so byte content depends only on the
# config (openpyxl restamps docProps modified and every zip entry with
# wall-clock time on save).
# =============================================================================

# =============================================================================
# Generator
# =============================================================================

def build(config, out_path):
    """Render a validated CIP buildup config (dict) to out_path (.xlsx)."""
    validate_config(config)
    rate_vals = dict(DEFAULT_RATES)
    rate_vals.update(config.get('rates', {}))
    labor = config.get('labor_sections')  # optional full override
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = CIP_BUILD_STAMP

    rs = wb.active; rs.title = 'Rates'
    rs['A1'] = 'RATE BASIS - %s (edit here; all structures reference these)' % config.get('basis_note', '')
    rs['A1'].font = TITLE
    rr = {}
    for i, (lab, _default) in enumerate(DEFAULT_RATES):
        r = i + 2
        rs.cell(r, 1, lab).font = B
        c = rs.cell(r, 2, rate_vals[lab]); c.fill = INP
        c.number_format = NUM
        rr[lab] = r
    rs.column_dimensions['A'].width = 24; rs.column_dimensions['B'].width = 12

    def RT(lab):
        return "Rates!$B$%d" % rr[lab]

    sections = labor if labor else [dict(id=s[0], name=s[1], fore=s[2], carp=s[3], lab=s[4], mason=s[5], days=s[6]) for s in DEFAULT_LABOR]

    def build_struct(s):
        ws = wb.create_sheet(s['sheet'])
        ws.column_dimensions['A'].width = 42
        for col in 'BCDEFGH':
            ws.column_dimensions[col].width = 11
        ws['A1'] = s.get('title', s['sheet']); ws['A1'].font = TITLE
        ws['A2'] = 'Basis: Rates sheet. CIP. EXCLUDES excavation & dewatering.'
        ws['A2'].font = Font(italic=True, color='C00000')
        if s.get('note'):
            ws['A3'] = 'Mains/notes: ' + s['note']
        st = {'r': 5}

        def hdr(t):
            r = st['r']
            ws.cell(r, 1, t).font = WHITE
            for cc in range(1, 9):
                ws.cell(r, cc).fill = HD
            st['r'] = r + 1

        def io_row(label, val, unit=''):
            r = st['r']
            ws.cell(r, 1, label).font = B
            c = ws.cell(r, 2, val); c.fill = INP; c.number_format = NUM
            if unit:
                ws.cell(r, 3, unit)
            st['r'] = r + 1
            return "$B$%d" % r

        def calc(label, formula, unit='', bold=False, fill=None):
            r = st['r']
            lc = ws.cell(r, 1, label)
            if bold:
                lc.font = B
            c = ws.cell(r, 2, formula); c.number_format = NUM
            if fill:
                c.fill = fill
            if unit:
                ws.cell(r, 3, unit)
            st['r'] = r + 1
            return "$B$%d" % r

        hdr('CONCRETE TAKEOFF')
        if s.get('components'):
            # detailed take-off already done (e.g. wingwall/headwall outfall): each
            # component is a rectangular prism count*L*W*H (counts may be negative
            # for opening deductions). Feeds the neat CY directly.
            r = st['r']
            for i, t in enumerate(['Section (from take-off)', 'No.', 'L (ft)', 'W (ft)', 'H (ft)', 'CY']):
                ws.cell(r, i + 1, t).font = B
            st['r'] = r + 1
            comp_cells = []
            for comp in s['components']:
                r = st['r']
                ws.cell(r, 1, comp['name'])
                for col, key in ((2, 'count'), (3, 'L'), (4, 'W'), (5, 'H')):
                    c = ws.cell(r, col, comp.get(key, 0)); c.fill = INP; c.number_format = NUM
                ws.cell(r, 6, "=B%d*C%d*D%d*E%d/27" % (r, r, r, r)).number_format = NUM
                comp_cells.append("F%d" % r)
                st['r'] = r + 1
            r = st['r']; ws.cell(r, 1, 'Neat concrete subtotal').font = B
            sc = ws.cell(r, 6, "=SUM(%s:%s)" % (comp_cells[0], comp_cells[-1])); sc.number_format = NUM
            subCY = "$F$%d" % r; st['r'] = r + 1
            cyc = calc('TOTAL CONCRETE (w/ waste)', "=%s*(1+%s)" % (subCY, RT('Waste fraction')), 'CY', bold=True, fill=SUB)
            st['r'] += 1
        else:
            Lc = io_row('Ext. Length', s['L'], 'ft'); Wc = io_row('Ext. Width', s['W'], 'ft'); Hc = io_row('Ext. Height', s['H'], 'ft')
            twc = io_row('Wall thickness', s.get('wall', DEF['wall']), 'ft')
            tbc = io_row('Base thickness', s.get('base', DEF['base']), 'ft')
            tdc = io_row('Deck thickness', s.get('deck', DEF['deck']), 'ft')
            baseCF = calc('Base slab', "=%s*%s*%s" % (Lc, Wc, tbc), 'CF')
            deckCF = calc('Deck / top slab', "=%s*%s*%s" % (Lc, Wc, tdc), 'CF')
            wallCF = calc('Walls', "=2*((%s-%s)+(%s-%s))*%s*(%s-%s-%s)" % (Lc, twc, Wc, twc, twc, Hc, tbc, tdc), 'CF')
            subCF = calc('Subtotal', "=%s+%s+%s" % (baseCF, deckCF, wallCF), 'CF')
            subCY = calc('Subtotal', "=%s/27" % subCF, 'CY')
            cyc = calc('TOTAL CONCRETE (w/ waste)', "=%s*(1+%s)" % (subCY, RT('Waste fraction')), 'CY', bold=True, fill=SUB)
            st['r'] += 1

        hdr('LABOR  (crew rates from Rates; DAYS = input)')
        r = st['r']
        for i, t in enumerate(['Section', 'FORE', 'CARP', 'LAB', 'MASON', 'Crew $/hr', 'Days', 'Cost']):
            ws.cell(r, i + 1, t).font = B
        st['r'] = r + 1
        days_over = s.get('days', {})
        lab_cells = []
        for sec in sections:
            r = st['r']
            ws.cell(r, 1, sec['name'])
            for col, v in zip((2, 3, 4, 5), (sec['fore'], sec['carp'], sec['lab'], sec['mason'])):
                ws.cell(r, col, v).fill = INP
            crew = ws.cell(r, 6, "=B%d*%s+C%d*%s+D%d*%s+E%d*%s" % (
                r, RT('Labor FORE $/hr'), r, RT('Labor CARP $/hr'), r, RT('Labor LAB $/hr'), r, RT('Labor MASON $/hr')))
            crew.number_format = MONEY
            ws.cell(r, 7, days_over.get(sec['id'], sec['days'])).fill = INP
            ws.cell(r, 8, "=F%d*%s*G%d" % (r, RT('Hrs per day'), r)).number_format = MONEY
            lab_cells.append("H%d" % r)
            st['r'] = r + 1
        r = st['r']; ws.cell(r, 1, 'Labor subtotal').font = B
        c = ws.cell(r, 8, "=SUM(%s:%s)" % (lab_cells[0], lab_cells[-1])); c.number_format = MONEY; c.fill = SUB
        labsub = "$H$%d" % r; st['r'] = r + 2

        def block(title, items):
            hdr(title)
            r = st['r']
            for i, t in enumerate(['Item', 'Qty', 'Unit', 'Rate', '', '', '', 'Cost']):
                ws.cell(r, i + 1, t).font = B
            st['r'] = r + 1
            cells = []
            for name, qty, is_input, unit, rate_label in items:
                r = st['r']
                ws.cell(r, 1, name)
                q = ws.cell(r, 2, qty)
                if is_input:
                    q.fill = INP
                q.number_format = NUM
                ws.cell(r, 3, unit)
                ws.cell(r, 4, "=%s" % RT(rate_label)).number_format = MONEY
                ws.cell(r, 8, "=B%d*D%d" % (r, r)).number_format = MONEY
                cells.append("H%d" % r)
                st['r'] = r + 1
            return cells

        eq = block('EQUIPMENT', [
            ('Pickup', s.get('pickup_hr', DEF['pickup_hr']), True, 'hr', 'Pickup $/hr'),
            ('Tools', s.get('tools_hr', DEF['tools_hr']), True, 'hr', 'Tools $/hr'),
            ('Skytrak', s.get('skytrak_hr', DEF['skytrak_hr']), True, 'hr', 'Skytrak $/hr'),
            ('Concrete Pump', s.get('pump_hr', DEF['pump_hr']), True, 'hr', 'Concrete Pump $/hr'),
            ('Volume', "=%s" % cyc, False, 'CY', 'Volume $/CY'),
        ])
        r = st['r']; ws.cell(r, 1, 'Equipment subtotal').font = B
        c = ws.cell(r, 8, "=SUM(%s:%s)" % (eq[0], eq[-1])); c.number_format = MONEY; c.fill = SUB
        eqsub = "$H$%d" % r; st['r'] = r + 2

        # Subcontracts: explicit `subs` list, or the `rebar_cost` shorthand. When rebar
        # is a sub it moves out of MATERIAL into the SUB SPREAD block below.
        subs = list(s.get('subs', []))
        if s.get('rebar_cost') is not None and not any(x.get('kind') == 'rebar' for x in subs):
            subs.insert(0, {'name': 'Rebar - furnish & install', 'vendor': s.get('rebar_vendor', 'Rebar sub'),
                            'amount': s['rebar_cost'], 'kind': 'rebar'})
        # Rebar leaves MATERIAL only when it is ACTUALLY subbed (a rebar_cost
        # shorthand became a kind:"rebar" sub above, or an explicit one exists).
        # The Work original dropped the per-$/CY rebar line for ANY sub — an
        # earthwork-only subs list silently un-priced the rebar.
        rebar_is_sub = any(x.get('kind') == 'rebar' for x in subs)
        mat_items = [('Concrete', "=%s" % cyc, False, 'CY', 'Concrete $/CY')]
        if not rebar_is_sub:
            mat_items.append(('Rebar (per $/CY basis)', "=%s" % cyc, False, 'CY', 'Rebar $/CY'))
        mat_items.append(('Short loads', s.get('short_loads', DEF['short_loads']), True, 'EA', 'Short load $/EA'))
        mat = block('MATERIAL', mat_items)
        r = st['r']; ws.cell(r, 1, 'Material subtotal').font = B
        c = ws.cell(r, 8, "=" + "+".join(mat)); c.number_format = MONEY; c.fill = SUB
        matsub = "$H$%d" % r; st['r'] = r + 2

        oth = block('LUMBER / HARDWARE / FUEL / SPECIALS', [
            ('Lumber', "=%s" % cyc, False, 'CY', 'Lumber $/CY'),
            ('Build form panels', s.get('panel_hrs', DEF['panel_hrs']), True, 'hr', 'Build panel $/hr'),
            ('Hardware', "=%s" % cyc, False, 'CY', 'Hardware $/CY'),
            ('Fuel - pickup', s.get('fuel_pickup_hr', DEF['fuel_pickup_hr']), True, 'hr', 'Fuel $/hr'),
            ('Fuel - equipment', s.get('fuel_equip_hr', DEF['fuel_equip_hr']), True, 'hr', 'Fuel $/hr'),
            ('Specials - hydrophilic waterstop', s.get('waterstop_lf', DEF['waterstop_lf']), True, 'LF', 'Waterstop $/LF'),
        ])
        r = st['r']; ws.cell(r, 1, 'Lumber/HW/Fuel/Specials subtotal').font = B
        c = ws.cell(r, 8, "=" + "+".join(oth)); c.number_format = MONEY; c.fill = SUB
        othsub = "$H$%d" % r; st['r'] = r + 2

        # self-perform subtotal (labor + equipment + material + lumber/HW/fuel)
        r = st['r']; ws.cell(r, 1, 'SELF-PERFORM SUBTOTAL').font = B
        c = ws.cell(r, 8, "=%s+%s+%s+%s" % (labsub, eqsub, matsub, othsub)); c.number_format = MONEY; c.fill = SUB
        spsub = "$H$%d" % r; st['r'] = r + 2

        # SUBCONTRACTORS (sub spread) -- rebar etc. carried as quoted subcontracts
        sub_lines = []
        if subs:
            hdr('SUBCONTRACTORS - SUB SPREAD')
            r = st['r']
            for i, t in enumerate(['Subcontract item', 'Vendor', 'Qty', 'Unit', 'Unit Price', '', '', 'Amount']):
                ws.cell(r, i + 1, t).font = B
            st['r'] = r + 1
            cells = []
            for sub in subs:
                r = st['r']
                ws.cell(r, 1, sub['name'])
                ws.cell(r, 2, sub.get('vendor', ''))
                q = ws.cell(r, 3, sub.get('qty', 1)); q.fill = INP; q.number_format = NUM
                ws.cell(r, 4, sub.get('unit', 'LS'))
                p = ws.cell(r, 5, sub.get('unit_price', sub.get('amount', 0))); p.fill = INP; p.number_format = MONEY
                ws.cell(r, 8, "=C%d*E%d" % (r, r)).number_format = MONEY
                cells.append("H%d" % r)
                sub_lines.append({'name': sub['name'], 'vendor': sub.get('vendor', ''), 'kind': sub.get('kind', ''),
                                  'cell': "'%s'!$H$%d" % (s['sheet'], r)})
                st['r'] = r + 1
            r = st['r']; ws.cell(r, 1, 'Subcontract subtotal (sub spread)').font = B
            c = ws.cell(r, 8, "=SUM(%s:%s)" % (cells[0], cells[-1])); c.number_format = MONEY; c.fill = SUB
            subsub = "$H$%d" % r; st['r'] = r + 2
        else:
            subsub = None

        r = st['r']
        ws.cell(r, 1, 'STRUCTURE TOTAL (CIP, excl. excavation & dewatering)').font = Font(bold=True, size=12)
        parts = [spsub] + ([subsub] if subsub else [])
        tot = ws.cell(r, 8, "=" + "+".join(parts))
        tot.number_format = MONEY; tot.font = Font(bold=True, size=12); tot.fill = GT
        return {
            'total': "'%s'!$H$%d" % (s['sheet'], r),
            'selfperform': "'%s'!%s" % (s['sheet'], spsub),
            'sub': ("'%s'!%s" % (s['sheet'], subsub)) if subsub else None,
            'sub_lines': sub_lines,
        }

    totals = {}
    for s in config['structures']:
        totals[s['sheet']] = build_struct(s)

    sm = wb.create_sheet('Summary', 0)
    sm['A1'] = config.get('title', 'Structure Estimate'); sm['A1'].font = TITLE
    sm['A2'] = config.get('basis_note', ''); sm['A2'].font = Font(italic=True, color='C00000')
    cols = ['Bid Item', 'Structure', 'Self-Perform', 'Subcontract', 'Each Total', 'Qty', 'Extended Total']
    for i, t in enumerate(cols):
        c = sm.cell(4, i + 1, t); c.font = WHITE; c.fill = HD
    r = 5
    for s in config['structures']:
        ref = totals[s['sheet']]
        sm.cell(r, 1, s.get('bid_item', ''))
        sm.cell(r, 2, s['sheet'])
        sm.cell(r, 3, "=%s" % ref['selfperform']).number_format = MONEY
        sm.cell(r, 4, ("=%s" % ref['sub']) if ref['sub'] else 0).number_format = MONEY
        sm.cell(r, 5, "=C%d+D%d" % (r, r)).number_format = MONEY
        qc = sm.cell(r, 6, s.get('count', 1)); qc.fill = INP
        sm.cell(r, 7, "=E%d*F%d" % (r, r)).number_format = MONEY
        r += 1
    sm.cell(r, 1, 'TOTAL').font = Font(bold=True, size=12)
    for col in (3, 4, 7):
        L = openpyxl.utils.get_column_letter(col)
        g = sm.cell(r, col, "=SUM(%s5:%s%d)" % (L, L, r - 1)); g.number_format = MONEY
        g.font = Font(bold=True, size=12); g.fill = GT
    for col, w in zip('ABCDEFG', (24, 22, 13, 13, 13, 6, 15)):
        sm.column_dimensions[col].width = w

    # Sub Spread sheet -- every quoted subcontract line, mirrors the BM sub spread
    has_subs = any(totals[s['sheet']]['sub_lines'] for s in config['structures'])
    if has_subs:
        ss = wb.create_sheet('Sub Spread', 1)
        ss['A1'] = 'SUBCONTRACTOR SUB SPREAD'; ss['A1'].font = TITLE
        ss['A2'] = 'Quoted subcontracts spread across the structures (also entered as work-package quotes).'
        ss['A2'].font = Font(italic=True, color='C00000')
        for i, t in enumerate(['Vendor', 'Subcontract item', 'Bid Item', 'Structure', 'Qty', 'Each', 'Extended']):
            c = ss.cell(4, i + 1, t); c.font = WHITE; c.fill = HD
        r = 5
        for s in config['structures']:
            for ln in totals[s['sheet']]['sub_lines']:
                ss.cell(r, 1, ln['vendor']); ss.cell(r, 2, ln['name'])
                ss.cell(r, 3, s.get('bid_item', '')); ss.cell(r, 4, s['sheet'])
                qc = ss.cell(r, 5, s.get('count', 1)); qc.fill = INP
                ss.cell(r, 6, "=%s" % ln['cell']).number_format = MONEY
                ss.cell(r, 7, "=E%d*F%d" % (r, r)).number_format = MONEY
                r += 1
        ss.cell(r, 1, 'TOTAL SUBCONTRACTS').font = Font(bold=True, size=12)
        g = ss.cell(r, 7, "=SUM(G5:G%d)" % (r - 1)); g.number_format = MONEY
        g.font = Font(bold=True, size=12); g.fill = GT
        for col, w in zip('ABCDEFG', (20, 30, 22, 22, 6, 12, 13)):
            ss.column_dimensions[col].width = w

    wb.save(out_path)
    _normalize_xlsx(out_path)
    return out_path
