"""Generator: render an org-agnostic estimate workbook from the estimate_spec.

Consumes tokens.yaml + units.yaml + recipes.yaml + a validated estimate config
and writes an .xlsx via openpyxl. See README.md for the layer model and the
placeholder mini-language the Resolver implements.

Sheets are laid out first (collecting each sheet's published `exports:`) and
every cell is resolved afterwards, in `write()`. That ordering is what lets
`{ref:SHEET!NAME}` -- the opt-in cross-sheet form of `{ref:NAME}` -- name a
sheet that had not yet been laid out when the referring cell was placed.
"""
from __future__ import annotations

import argparse
import copy
import datetime
import hashlib
import json
import re
import sys
from collections.abc import Mapping
from importlib.metadata import entry_points
from pathlib import Path

import openpyxl
import yaml
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from estimating_tools.estimate_spec import layout_map as LM
from estimating_tools.estimate_spec import recipe as R

# Entry-point group a consumer package registers a pre-render config check under
# (#256). Each entry point resolves to `check(config: dict) -> Iterable[str]`.
CONFIG_CHECK_GROUP = "estimating_tools.estimate_spec.config_check"


def run_config_checks(config, quotes=None):
    """Run consumer-registered pre-render config checks (#256).

    A consumer whose policy lives in its own repo (e.g. "every sheet must carry
    a FUEL cost block, else it silently books $0") registers a check under the
    CONFIG_CHECK_GROUP entry-point group. Each check is `check(config)->iterable
    of finding strings`. Advisory ONLY: findings are returned prefixed with the
    check name; a check that raises is reported as a finding, never propagated —
    a broken or absent lint must not break the build. Returns [] when none are
    registered.

    `quotes` (the CanonicalQuotes the build loaded, or None when none were
    loaded) adds the built-in package lint: a package with no quote file. The
    function stays pure; the caller does the loading."""
    findings: list[str] = []
    if quotes is not None:
        from estimating_tools.estimate_spec.package_frames import _fold
        quoted = {_fold(q.work_package_name) for q in quotes}
        for pkg in config.get("work_packages") or []:
            if _fold(str(pkg.get("name", ""))) not in quoted:
                findings.append(f"package '{pkg.get('name')}' has no quote file")
    try:
        eps = entry_points(group=CONFIG_CHECK_GROUP)
    except Exception:
        return findings
    for ep in eps:
        try:
            check = ep.load()
            for finding in (check(config) or []):
                findings.append(f"{ep.name}: {finding}")
        except Exception as exc:  # a consumer's check must never break the build
            findings.append(f"{ep.name}: check errored: {exc}")
    return findings

# =============================================================================
# Resolver — substitutes the placeholder mini-language (units.yaml header).
# =============================================================================

_PLACEHOLDER = re.compile(r"\{([^{}]+)\}")
_LANDMARK = re.compile(r"^([A-Za-z_][A-Za-z0-9_]*)\s*([+-]\s*\d+)?$")
#: `{c}`, `{c+1}`, `{c-2}`: the current `repeat_cols` base column, offset.
_BAND_COL = re.compile(r"^c\s*([+-]\s*\d+)?$")
#: Column landmarks share the unit's landmark dict under this prefix.
COL_LANDMARK = "col:"


#: A literal brace: text marked with `literal()` carries these private-use
#: characters instead of `{`/`}`, so the fixed-point resolver never reads it
#: as a placeholder; the writer turns them back into braces in the cell's
#: value and comment. For text from outside the pack (a quote file's
#: conditions, say), which must render exactly as typed.
_LBRACE, _RBRACE = "\ue000", "\ue001"
_LITERAL_TEXT = "\ue002"


def literal(text) -> str:
    """Text whose braces stay literal and whose leading equals is not a formula."""
    text = str(text)
    prefix = _LITERAL_TEXT if text.startswith("=") else ""
    return prefix + text.replace("{", _LBRACE).replace("}", _RBRACE)


def _unliteral(value):
    if isinstance(value, str):
        return value.replace(_LBRACE, "{").replace(_RBRACE, "}").replace(_LITERAL_TEXT, "")
    return value


def col_offset(letter, offset):
    """`letter` moved `offset` columns right (negative: left)."""
    return get_column_letter(column_index_from_string(letter) + offset)
_QUOTED_SHEET = re.compile(r"^'(.*)'$", re.DOTALL)

# The separator that opts a {ref:} into another sheet's namespace. A ref key
# is `INSTANCE.EXPORT` (instance keys may carry [brackets]); none has ever
# contained this character, which is what keeps the syntax purely additive.
SHEET_SEP = "!"
# Excel quotes a sheet name that needs it, and escapes an apostrophe inside
# those quotes by doubling it. `sheet_qualified` doubles; `split_sheet_ref`
# undoubles.
QUOTE = "'"


def quoted_sheet(sheet):
    """A sheet name as a formula must quote it: `Bob's Sheet` -> `'Bob''s Sheet'`.

    Excel escapes an apostrophe inside those quotes by doubling it; anything
    that builds an address out of a sheet name goes through here, so nothing
    can emit the half-escaped form that only fails when the file is opened.
    """
    return f"{QUOTE}{sheet.replace(QUOTE, QUOTE * 2)}{QUOTE}"


def sheet_title(name: str, taken: set[str]) -> str:
    """`name` as a legal, unused Excel sheet title.

    Excel's refused characters ([ ] : * ? / \\) become `-`, the result is cut
    to Excel's 31 characters, and a title already in `taken` (compared
    case-insensitively, as Excel does) is suffixed ` (2)`, ` (3)` ... with
    the stem shortened so the suffix still fits in 31. `taken` is only read;
    the caller adds the returned title to it.
    """
    clean = "".join("-" if c in R._SHEET_NAME_BAD_CHARS else c for c in name)
    clean = clean[:R._SHEET_NAME_MAX]
    folded = {t.casefold() for t in taken}
    title, n = clean, 1
    while title.casefold() in folded:
        n += 1
        suffix = f" ({n})"
        title = clean[:R._SHEET_NAME_MAX - len(suffix)] + suffix
    return title


def sheet_qualified(sheet, addr):
    """An address as another sheet must spell it: `'Sheet'!$A$1`.

    One spelling for both cross-sheet forms (a declared global with
    `qualify: sheet`, and a `{ref:SHEET!NAME}`), so a pack never has to care
    which of the two produced the address it is reading.

    An apostrophe in the sheet name is doubled -- Excel's own escaping inside
    a quoted sheet reference, and the exact inverse of what `split_sheet_ref`
    unescapes. `Bob's Sheet` therefore round-trips, instead of rendering an
    address that fails only when someone opens the workbook.
    """
    return f"{quoted_sheet(sheet)}{SHEET_SEP}{addr}"


def split_sheet_ref(key):
    """`SHEET!NAME` -> (sheet, name); a key without a separator -> (None, key).

    The split is on the LAST separator: an export name never contains one,
    while a sheet name legally may. The sheet may be written bare or quoted
    (`'Crew Rates'!x`), the way the workbook itself spells it; inside quotes an
    apostrophe is doubled, the inverse of `sheet_qualified`.
    """
    if SHEET_SEP not in key:
        return None, key
    sheet, _, name = key.rpartition(SHEET_SEP)
    m = _QUOTED_SHEET.match(sheet)
    if m:
        sheet = m.group(1).replace(QUOTE * 2, QUOTE)
    return sheet, name


class CrossSheetRefError(R.RecipeError):
    """A `{ref:SHEET!NAME}` names an export no sheet in the workbook published.

    Its own class so the `optional:` escape hatch in `write()` cannot swallow
    it: an unresolvable cross-sheet reference is always an error, never a
    blank cell and never a silent zero.
    """


class Resolver:
    """Resolve {placeholders} in a cell value/formula for one render context.

    Refs always expand to full cell addresses; {r} and {row:NAME} expand to
    integers that the author prefixes with a literal column letter. {col:NAME}
    (and {col:NAME+N}) expands to a column letter recorded by a row's
    `mark_cols_first`/`mark_cols_last`; inside a `repeat_cols` band {c} (and
    {c+N}) is the band's current base column and {v}/{v.field} (or whatever
    the band's `as:` names) its bound element.

    `exports` is the referring sheet's own namespace and is consulted first,
    so every ref a pack can write today resolves exactly as it always has.
    `book_exports` (sheet -> that sheet's namespace) is the whole workbook's
    and backs the opt-in `{ref:SHEET!NAME}` form; `where` and `pack` only
    name the site in an error message.
    """

    def __init__(self, *, params=None, item=None, row=None, landmarks=None,
                 exports=None, globals=None, book_exports=None, where="", pack=None,
                 col=None, binds=None):
        self.params = params or {}
        self.item = item
        self.row = row
        self.col = col
        self.binds = binds or {}
        self.book_exports = book_exports
        self.where = where
        self.pack = pack
        self.landmarks = landmarks or {}
        self.exports = exports or {}
        self.globals = dict(globals or {})

    def resolve(self, text):
        """Resolve all placeholders. Non-strings pass through untouched.

        Runs to a fixed point so a parameter that itself contains placeholders
        (e.g. driver = "={ref:takeoff.ftg_cy}") is fully expanded.
        """
        if not isinstance(text, str):
            return text
        cur = text
        for _ in range(10):
            if "{" not in cur:
                break
            nxt = _PLACEHOLDER.sub(self._sub, cur)
            if nxt == cur:
                break
            cur = nxt
        return cur

    def _sub(self, match):
        token = match.group(1).strip()

        if token.startswith("p."):
            name = token[2:]
            if name not in self.params:
                raise KeyError(f"unknown parameter {name!r}")
            return str(self.params[name])

        if token == "it":
            return str(self.item)
        if token.startswith("it."):
            return str(self.item[token[3:]])

        if token == "r" or token.startswith("r-") or token.startswith("r+"):
            if self.row is None:
                raise ValueError("current row {r} used but not set")
            offset = token[1:]
            return str(self.row + (int(offset) if offset else 0))

        if token.startswith("row:"):
            return str(self._landmark(token[4:]))

        if token.startswith(COL_LANDMARK):
            return self._col_landmark(token[len(COL_LANDMARK):])

        if self.col is not None:
            m = _BAND_COL.match(token)
            if m:
                return col_offset(self.col, int(m.group(1).replace(" ", "")) if m.group(1) else 0)

        head, _, field = token.partition(".")
        if head in self.binds:
            value = self.binds[head]
            return str(value[field] if field else value)

        if token.startswith("ref:"):
            key = token[4:]
            if key in self.exports:          # this sheet's namespace always wins
                return str(self.exports[key])
            if SHEET_SEP in key:             # opt-in: another sheet's namespace
                return self._cross_sheet(key)
            raise KeyError(f"unknown ref {key!r}")

        if token in self.globals:
            value = self.globals[token]
            if value is None:
                raise ValueError(f"{{{token}}} used but not set")
            return value

        raise ValueError(f"unresolved placeholder {{{token}}}")

    def _cross_sheet(self, key):
        """`{ref:SHEET!NAME}` -> `'SHEET'!<that sheet's exported address>`.

        Resolved against the WHOLE workbook's exports, which is why the
        target sheet may be laid out before or after the referring one: every
        cell is resolved in one pass after all layout is done. Anything that
        does not resolve names itself, its site and its pack and stops the
        build -- see CrossSheetRefError.
        """
        sheet, name = split_sheet_ref(key)
        ref = f"{{ref:{key}}}"
        site = f"{self.where}: " if self.where else ""
        pack = f" in pack {self.pack!r}" if self.pack else ""
        if self.book_exports is None:
            raise CrossSheetRefError(
                f"{site}{ref} reaches across sheets, but no other sheet is in "
                "scope here. A cross-sheet ref belongs in a cell, where it is "
                "resolved against the whole workbook; a unit's `exports:` "
                "template publishes an address on its OWN sheet and cannot "
                "reach across")
        if sheet not in self.book_exports:
            raise CrossSheetRefError(
                f"{site}{ref} names sheet {sheet!r}, which this estimate does "
                f"not lay out{pack} (sheets: {sorted(self.book_exports)}). A "
                "cross-sheet ref names a sheet by the name the estimate config "
                "gives it; for one address published once and read from "
                "everywhere, declare it under `globals:` with scope: workbook "
                "instead of naming the sheet here"
            )
        published = self.book_exports[sheet]
        if name not in published:
            raise CrossSheetRefError(
                f"{site}{ref}: sheet {sheet!r} publishes no export {name!r}"
                f"{pack} (it offers: {sorted(published)}). An export is "
                "published as `INSTANCE.NAME`, where INSTANCE is the recipe "
                "entry's `key:` (the unit name when it has none) and NAME is a "
                "key under that unit's `exports:`"
            )
        return sheet_qualified(sheet, published[name])

    def _col_landmark(self, expr):
        m = _LANDMARK.match(expr)
        if not m:
            raise ValueError(f"malformed column landmark {expr!r}")
        name, offset = m.group(1), m.group(2)
        key = COL_LANDMARK + name
        if key not in self.landmarks:
            raise ValueError(f"unknown column landmark {name!r}")
        return col_offset(self.landmarks[key], int(offset.replace(" ", "")) if offset else 0)

    def _landmark(self, expr):
        m = _LANDMARK.match(expr)
        if not m:
            raise ValueError(f"malformed landmark {expr!r}")
        name, offset = m.group(1), m.group(2)
        if name not in self.landmarks:
            raise ValueError(f"unknown landmark {name!r}")
        val = self.landmarks[name]
        if offset:
            val += int(offset.replace(" ", ""))
        return val


# =============================================================================
# StyleFactory — turns a composite style name (+ overrides) into openpyxl style.
# =============================================================================

_RESERVED_CELL_KEYS = {"value", "formula", "style", "optional", "if_ref", "if_param", "if_it",
                        "validation", "ref_error", "comment", "hide_column"}

#: A cell key placed relative to a column landmark (`"{col:vl+2}"`) instead of
#: at a fixed letter -- for a cell that sits past a `repeat_cols` band.
_COL_KEY = re.compile(r"^\{col:[^{}]+\}$")


class StyleFactory:
    def __init__(self, tokens):
        self.base = tokens.get("meta", {}).get("base_font", {"name": "Arial", "size": 10})
        self.fills = tokens.get("fills", {})
        self.font_colors = tokens.get("font_colors", {})
        self.font_sizes = tokens.get("font_sizes", {})
        self.borders = tokens.get("borders", {})
        self.number_formats = tokens.get("number_formats", {})
        self.alignments = tokens.get("alignments", {})
        self.styles = tokens.get("styles", {})

    @staticmethod
    def _argb(hex6):
        h = hex6.lstrip("#")
        return h if len(h) == 8 else "FF" + h

    def _size(self, value):
        return self.font_sizes.get(value, value) if isinstance(value, str) else value

    def fill_for(self, token):
        return PatternFill("solid", fgColor=self._argb(self.fills[token]))

    def protection_for(self, style_name, overrides=None):
        """The `locked:` Protection for a style+overrides, or None if it sets no `locked`.

        Pulled out of `apply()` so `write()` can apply JUST the lock to a
        skipped optional cell (no value, no other styling) — a pack that
        never sets `locked` gets None here and that cell is left completely
        untouched, so its golden snapshot does not move.
        """
        style = dict(self.styles.get(style_name, {}))
        if overrides:
            style.update(overrides)
        if "locked" not in style:
            return None
        locked = style["locked"]
        if not isinstance(locked, bool):
            raise R.RecipeError(
                f"style {style_name!r}: locked must be true or false "
                f"(got {type(locked).__name__} {locked!r})")
        return Protection(locked=locked)

    def dxf(self, style_name):
        """(font, fill) for a conditional format -- the part of a composite
        style a rule may change (bold/italic/colour and fill only; a dxf
        carries no border/alignment/number_format). An unknown style name,
        or a style whose font_color/fill names a token this pack's
        tokens.yaml does not define, stops the build naming the style and
        the missing token."""
        if style_name not in self.styles:
            raise R.RecipeError(f"conditional: unknown style {style_name!r}")
        style = self.styles[style_name]
        color = None
        if "font_color" in style:
            token = style["font_color"]
            if token not in self.font_colors:
                raise R.RecipeError(
                    f"conditional: style {style_name!r} names font_color {token!r}, "
                    f"which is not in tokens.font_colors (has: {sorted(self.font_colors)})")
            color = self._argb(self.font_colors[token])
        font = Font(bold=bool(style.get("bold", False)),
                    italic=bool(style.get("italic", False)), color=color)
        fill = None
        if "fill" in style:
            token = style["fill"]
            if token not in self.fills:
                raise R.RecipeError(
                    f"conditional: style {style_name!r} names fill {token!r}, "
                    f"which is not in tokens.fills (has: {sorted(self.fills)})")
            c = self._argb(self.fills[token])
            fill = PatternFill(fill_type="solid", fgColor=c, bgColor=c)
        return font, fill

    def apply(self, cell, style_name, overrides=None):
        style = dict(self.styles.get(style_name, {}))
        if overrides:
            style.update(overrides)

        color = self._argb(self.font_colors[style["font_color"]]) if "font_color" in style else None
        cell.font = Font(
            name=self.base["name"],
            size=self._size(style.get("size", self.base["size"])),
            bold=bool(style.get("bold", False)),
            italic=bool(style.get("italic", False)),
            color=color,
        )
        if "fill" in style:
            cell.fill = self.fill_for(style["fill"])
        if "border" in style:
            b = self.borders[style["border"]]
            side = Side(style=b["style"])
            cell.border = Border(**{s: side for s in b["sides"]})
        if "align" in style:
            cell.alignment = Alignment(**self.alignments[style["align"]])
        if "number_format" in style:
            nf = style["number_format"]
            cell.number_format = self.number_formats.get(nf, nf)
        protection = self.protection_for(style_name, overrides)
        if protection is not None:
            cell.protection = protection


def coerce_cell_value(text):
    """Resolved string -> formula / date / number / text for openpyxl.

    Public: read-back compares a workbook cell with what this wrote for a spec
    value, so it must use the same conversion."""
    if not isinstance(text, str):
        return text
    if text.startswith("="):
        return text
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        from datetime import date
        return date(int(m[1]), int(m[2]), int(m[3]))
    if re.fullmatch(r"-?\d+", text):
        digits = text.lstrip("-")
        if len(digits) > 1 and digits[0] == "0":
            return text   # zero-padded (e.g. an item number "010") -- stays text
        return int(text)
    try:
        return float(text)
    except ValueError:
        return text


#: The old private name, kept so existing callers and tests keep working.
_coerce = coerce_cell_value


def _sheet_setup(ws, sdef: dict) -> None:
    """A recipe's or book sheet's `freeze:` and `print:` (validated at load
    by recipe.validate_recipes); a sheet that sets neither is untouched."""
    if "freeze" in sdef:
        ws.freeze_panes = sdef["freeze"]
    page = sdef.get("print")
    if not page:
        return
    if "orientation" in page:
        ws.page_setup.orientation = page["orientation"]
    if "fit_to_width" in page:
        # N pages wide, as many pages tall as it takes (0 = unlimited).
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = page["fit_to_width"]
        ws.page_setup.fitToHeight = 0
    if "repeat_rows" in page:
        ws.print_title_rows = page["repeat_rows"]


def _matches(element, wanted):
    """`if_it:` -- every listed field of the current repeat element equals the
    given value, compared as text (YAML may hand back 31 or "31").

    A field the element lacks reads as None, so it matches only null/"None";
    a misspelled field name matches no element. Numbers compare by their text
    form, so 31 and 31.0 are distinct (`"31" != "31.0"`); the comparison is
    case-sensitive.
    """
    if not isinstance(element, dict):
        return False
    return all(str(element.get(k)) == str(v) for k, v in wanted.items())


# ---- validation: dropdowns -------------------------------------------------

#: Excel's own cap on a List validation's literal `"a,b,c"` formula1 string
#: (quotes included) -- past this, Excel silently truncates or refuses the
#: validation, so a pack that writes a list this long needs an `=range`
#: instead.
_LITERAL_LIST_MAX = 255

# An optional sheet prefix in front of any of the range shapes below: either
# a bare identifier (`Codes!`) or a quoted name (`'Codes'!`), quoted because
# it has a space, a reserved character, or an apostrophe -- which Excel's
# own quoting escapes by doubling (`'Bob''s Codes'!`), the exact shape
# `(?:[^']|'')*` consumes.
_SHEET_PREFIX = r"(?:'(?:[^']|'')*'|[A-Za-z_][A-Za-z0-9_.]*)!"

# cell, or cell:cell -- `$A$2` or `$A$2:$B$9`.
_A1_CELL_RANGE_RE = re.compile(
    rf"^(?:{_SHEET_PREFIX})?"
    r"\$?([A-Za-z]{1,3})\$?(\d+)"
    r"(?::\$?([A-Za-z]{1,3})\$?(\d+))?$"
)
# whole-column range -- `$A:$B` (every row, those columns).
_A1_COLUMN_RANGE_RE = re.compile(
    rf"^(?:{_SHEET_PREFIX})?" r"\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})$")
# whole-row range -- `1:3` (every column, those rows).
_A1_ROW_RANGE_RE = re.compile(rf"^(?:{_SHEET_PREFIX})?" r"\$?(\d+):\$?(\d+)$")


def _range_spans_rows_and_columns(ref):
    """True if the RESOLVED `ref` (its leading `=` already stripped) is an
    A1 range spanning more than one row AND more than one column -- the one
    shape Excel's List validation refuses as a source (it accepts a single
    cell, a single column, or a single row, never a rectangle). An optional
    leading sheet prefix (`'Sheet'!` or `Sheet!`) is consumed first.

    Three shapes are recognised: a plain cell or cell:cell range; a
    whole-column range (`$A:$B`, implicitly every row -- a rectangle unless
    both ends name the same column); and a whole-row range (`1:3`,
    implicitly every column -- a rectangle unless both ends name the same
    row). A `ref` none of these match (a named range, a formula like
    `OFFSET(...)`) is not flagged -- there is nothing to check its shape
    against, so it fails OPEN (passes) rather than guessing.
    """
    m = _A1_CELL_RANGE_RE.match(ref)
    if m:
        if m.group(3) is None:             # a single cell, no second corner
            return False
        c1, r1, c2, r2 = m.groups()
        return column_index_from_string(c1) != column_index_from_string(c2) and r1 != r2
    m = _A1_COLUMN_RANGE_RE.match(ref)
    if m:
        c1, c2 = m.groups()
        return column_index_from_string(c1) != column_index_from_string(c2)
    m = _A1_ROW_RANGE_RE.match(ref)
    if m:
        r1, r2 = m.groups()
        return r1 != r2
    return False


#: What an estimator sees on typing a value a strict dropdown does not offer.
_DROPDOWN_ERROR_TITLE = "Not on the list"
_DROPDOWN_ERROR = "Pick a value from the list."


def _check_validation(v, where):
    """`validation: {list: [choices]}` or `{list: "=<range>"}`, optionally with
    `strict: false` -- nothing else.

    A dropdown is strict by default: Excel stops a typed value that is not on
    the list. `strict: false` lets the cell take other values too (the list
    is then only a suggestion).

    A literal list is checked as far as layout time can: no empty/comma'd/
    quote-containing choice (a comma or a double quote inside a choice would
    corrupt Excel's own `"a,b,c"` literal-list syntax), and the literal
    formula1 it will produce stays under Excel's 255-character cap. A range
    list must have something to resolve after its `=` -- a bare `"="` is
    never a usable range. What the range itself resolves TO (empty, or a
    rectangle spanning more than one row and column) can only be checked
    once placeholders are expanded, so that lives in `_dropdown` instead.
    """
    if not isinstance(v, dict) or set(v) not in ({"list"}, {"list", "strict"}):
        raise R.RecipeError(
            f"{where}: validation must be {{list: ...}} or {{list: ..., strict: ...}}, got {v!r}")
    if "strict" in v and not isinstance(v["strict"], bool):
        raise R.RecipeError(
            f"{where}: validation strict must be true or false, got {v['strict']!r}")
    lst = v["list"]
    if isinstance(lst, list):
        if not lst or not all(
                isinstance(c, str) and c and "," not in c and '"' not in c for c in lst):
            raise R.RecipeError(
                f"{where}: validation list must be non-empty text choices without commas "
                f"or double quotes, got {lst!r}")
        formula1 = f'"{",".join(lst)}"'
        if len(formula1) > _LITERAL_LIST_MAX:
            raise R.RecipeError(
                f"{where}: validation list {lst!r} is {len(formula1)} characters as a "
                f"literal list -- Excel caps a literal list at {_LITERAL_LIST_MAX} -- "
                "use an =range")
    elif not (isinstance(lst, str) and lst.startswith("=") and lst[1:].strip()):
        raise R.RecipeError(f"{where}: validation list must be choices or an =range, got {lst!r}")


# ---- conditional: validation ------------------------------------------------

_CONDITIONAL_RULE_KEYS = {"cells", "formula", "style"}
_CONDITIONAL_CELLS_RE = re.compile(r"^[A-Z]{1,3}(:[A-Z]{1,3})?$")
#: A row's `conditional:` key was absent (as opposed to present but falsy).
_NO_CONDITIONAL = object()


def _validate_conditional(conditional, unit, row):
    """`conditional:` rows fail loudly at layout time, before being recorded
    (and long before `write()` would otherwise surface a confusing KeyError
    or silently do nothing): a list of {cells, formula, style} mappings,
    `cells` a column or column range (`F`, `A:F`, first column before the
    second -- `F:A` is refused), `formula` a string starting with `=`. The
    `style` name itself is resolved -- and an unknown style or a style
    naming a missing font_color/fill token rejected -- by `StyleFactory.dxf`
    when the rule is emitted. The multi-column anchoring trap
    (`_check_conditional_anchoring`) is NOT checked here: it needs the
    RESOLVED formula (placeholders expanded), so it runs in `write()`.
    """
    where = f"unit {unit!r} row {row} conditional"
    if not isinstance(conditional, list):
        raise R.RecipeError(f"{where}: must be a list of rule mappings, got {conditional!r}")
    for rule in conditional:
        if not isinstance(rule, dict) or set(rule) != _CONDITIONAL_RULE_KEYS:
            raise R.RecipeError(
                f"{where}: each rule must be a mapping with exactly the keys "
                f"{sorted(_CONDITIONAL_RULE_KEYS)}, got {rule!r}")
        cells = rule["cells"]
        if not isinstance(cells, str) or not _CONDITIONAL_CELLS_RE.match(cells):
            raise R.RecipeError(
                f"{where}: cells must be a column or column range like 'F' "
                f"or 'A:F', got {cells!r}")
        for col in cells.split(":"):
            if R._column_index(col) > R._MAX_COLUMN_INDEX:
                raise R.RecipeError(
                    f"{where}: cells {cells!r} names column {col!r}, past "
                    "Excel's last column, XFD")
        if ":" in cells:
            c1, c2 = cells.split(":")
            if R._column_index(c1) >= R._column_index(c2):
                raise R.RecipeError(
                    f"{where}: cells {cells!r} must list its first column "
                    f"before its second (left to right) -- {c1!r} does not "
                    f"come before {c2!r}")
        formula = rule["formula"]
        if not isinstance(formula, str) or not formula.startswith("="):
            raise R.RecipeError(
                f"{where}: formula must be a string starting with '=', "
                f"got {formula!r}")


# A cell reference inside a RESOLVED conditional formula (every {r}/{row:}/
# {p.}/{ref:} placeholder already expanded to its final address): bounded on
# both sides -- (?<!...) refuses a preceding identifier/quote/$/! character,
# (?!...) refuses a following identifier/(/./! character -- so it matches
# only a free-standing reference, never a run of letters-then-digits inside
# a longer token. That bound is what lets LOG10(...), ATAN2(...),
# DAYS360(...), BIN2DEC(...), a scientific literal like 1E5, and a named
# range like RATE2024 or Q1_TOTAL all pass through untouched: none of them
# is preceded/followed the way a real reference is. An optional 'Sheet'! or
# Sheet! prefix is consumed so a cross-sheet reference anchors the same way;
# a quoted name may hold Excel's doubled apostrophe ('Bob''s'!), which
# `(?:[^']|'')*` consumes whole -- otherwise the reference behind it would be
# skipped unchecked.
# Letters are case-insensitive (`[A-Za-z]`, not `[A-Z]`) so a lowercase
# reference (`f10`) is still caught, not silently waved through.
# A false positive here is NOT harmless -- unlike most defensive checks, it
# would block a formula the pack author wrote correctly -- so this regex
# was deliberately tested against the function/literal/named-range cases
# above before being trusted; see test_a_conditional_... in
# test_engine_extensions.py for the fixture formulas it must accept.
_CELL_REF_RE = re.compile(
    r"(?<![A-Za-z0-9_.$'!])"
    r"(?:(?:'(?:[^']|'')*'|[A-Za-z_][A-Za-z0-9_.]*)!)?"
    r"(\$?)([A-Za-z]{1,3})\$?\d+"
    r"(?![A-Za-z0-9_(.!])"
)


def _check_conditional_anchoring(cells, formula, where):
    """A `cells:` range spanning more than one column needs every cell
    reference in its RESOLVED formula column-anchored (`$F`, not `F`):
    Excel evaluates a formula-based conditional format relative to the
    top-left cell of the range it is applied to, so `=F10=0` painted over
    `A10:F10` checks F10 for column A, G10 for column B, H10 for column C,
    and so on -- almost never what the author meant. A single-column range
    has nowhere to drift to, so it is exempt. Quoted string literals are
    blanked out first so a literal like `"F2"` in a text argument is never
    mistaken for a reference.
    """
    if ":" not in cells:
        return
    stripped = re.sub(r'"(?:[^"]|"")*"', '""', formula)
    for m in _CELL_REF_RE.finditer(stripped):
        anchor, col = m.group(1), m.group(2)
        if not anchor:
            raise R.RecipeError(
                f"{where}: resolved formula {formula!r} references column "
                f"{col} without a leading $, but cells {cells!r} spans more "
                "than one column -- Excel evaluates the rule relative to "
                f"the range's top-left cell, so anchor it as ${col}")


# =============================================================================
# EstimateBuilder — lays out every sheet (collecting exports), then resolves.
# =============================================================================

class EstimateBuilder:
    def __init__(self, config, pack=None, *, navigator=True):
        from estimating_tools.estimate_spec.format_pack import bundled_pack

        self.config = config
        self.pack = pack or bundled_pack()
        #: Lay out Estimate Navigator's `_templates` sheet (when the pack's
        #: layout_map declares templates). `generate` ties it to `stamp`: an
        #: unstamped workbook carries no navigator anchors at all.
        self.navigator = navigator
        self.tokens = self.pack.tokens
        self.units = self.pack.units
        self.recipes_doc = self.pack.recipes
        self.recipes = self.recipes_doc["recipes"]
        R.validate_recipes(self.recipes_doc, self.units)
        markup_by_type = (self.pack.vocabulary or {}).get("markup_by_type")
        markup = config.get("markup")
        if isinstance(markup, dict) and markup_by_type not in (True, "required"):
            raise R.RecipeError(
                "markup by cost type needs a pack that opts in (vocabulary "
                "markup_by_type: true) and lays out its own markup sheet")
        if markup_by_type == "required" and not isinstance(markup, dict):
            raise R.RecipeError(
                f"pack {self.pack.name!r} requires markup by cost type "
                "(vocabulary markup_by_type: required) -- markup must be an "
                f"object of percentages by cost type, got {type(markup).__name__}")
        self.globals_decl = self.recipes_doc.get("globals") or {}
        self.sheet_globals = {}       # sheet -> {name: addr}   (scope: sheet)
        self.book_globals = {}        # name -> addr            (scope: workbook; promoted once)
        self.book_global_origin = {}  # name -> (sheet, unit/key) that promoted it
        self.fill_span = {}           # sheet -> (first col, last col)
        self.styles = StyleFactory(self.tokens)

        self.wb = openpyxl.Workbook()
        self.placements = []          # deferred cell writes
        self.region_fills = []        # (sheet, r1, r2, c1, c2, fill_token)
        self.boxes = []               # (sheet, r1, r2, c1, c2) thick perimeter
        self.merges = []              # (sheet, "A1:F1")
        self.exports = {}             # sheet_name -> {ref_key: address}
        self.conditionals = []        # (row-scoped conditional-format rules, resolved in write)
        self.dropdowns = {}           # (sheet, formula1, strict) -> [cell addresses]
        #: sheet_name -> a description of what created it, recorded only for
        #: a sheet an anchored unit's `sheet:` override creates on first use
        #: (see `_target`). Lets a book-sheet collision (`_run_book`) name
        #: what it collides WITH, not just that it does.
        self.sheet_origin = {}
        self._book_plan = None        # _book_sheets(): computed once per build
        #: sheet_name -> {instance keys a unit WITH exports has already
        #: published on it}. A second instance publishing the same key would
        #: otherwise silently overwrite the first's exports (see
        #: `layout_unit`) -- whatever reads that key then reads the wrong
        #: block. A unit with no `exports:` (e.g. fieldbook's `code_rows`,
        #: laid out once per bid item under the shared unit-name key, with
        #: nothing to publish) has nothing to collide, so it is exempt.
        self._seen_instance_keys = {}
        #: {kind: [columns]} blanked on `_templates` at the end of `write()`
        #: (see `build_templates`); empty when the pack declares no templates.
        self.template_clear = {}
        #: (sheet, column) pairs a placed cell's `hide_column: true` hides.
        self.hidden_cols = set()

    # ---- layout helpers ----------------------------------------------------
    def _exp(self, key, addr):
        self.cur_exports[key] = addr

    def _place_row(self, cells, sheet, row, params, item, landmarks, merge, fill_row,
                   merge_down=0, unit=None, conditional=_NO_CONDITIONAL,
                   band_col=None, binds=None):
        for col, cell in cells.items():
            if _COL_KEY.match(col):
                # Placed past a band: the landmark is the unit's own, recorded
                # by an earlier row or by this row's band (placed first).
                try:
                    col = Resolver(landmarks=landmarks).resolve(col)
                except ValueError as e:
                    raise R.RecipeError(f"unit {unit!r} cell key {col!r}: {e}") from e
            raw = cell["formula"] if "formula" in cell else cell.get("value")
            overrides = {k: v for k, v in cell.items() if k not in _RESERVED_CELL_KEYS}
            if "if_it" in cell:
                wanted = cell["if_it"]
                if not isinstance(wanted, dict) or not wanted:
                    raise R.RecipeError(
                        f"unit {unit!r} cell {col}: if_it must be a non-empty mapping, "
                        f"got {wanted!r}"
                    )
                if not isinstance(item, dict):
                    raise R.RecipeError(
                        f"unit {unit!r} cell {col}: if_it needs a `repeat:` row over "
                        f"mappings, got item {item!r}"
                    )
            if "validation" in cell:
                _check_validation(cell["validation"], f"unit {unit!r} cell {col}")
                self._check_dropdown_unlocked(cell.get("style", "label"), overrides,
                                              f"unit {unit!r} cell {col}")
            if "ref_error" in cell:
                ref_error = cell["ref_error"]
                if not isinstance(ref_error, str) or not ref_error:
                    raise R.RecipeError(
                        f"unit {unit!r} cell {col}: ref_error must be a non-empty string, "
                        f"got {ref_error!r}")
                if cell.get("optional"):
                    raise R.RecipeError(
                        f"unit {unit!r} cell {col}: ref_error cannot be combined with "
                        "optional -- an optional cell already renders blank instead of "
                        "raising, so ref_error would never fire")
            if "hide_column" in cell:
                if not isinstance(cell["hide_column"], bool):
                    raise R.RecipeError(
                        f"unit {unit!r} cell {col}: hide_column must be true or false, "
                        f"got {cell['hide_column']!r}")
                if cell["hide_column"]:
                    self.hidden_cols.add((sheet, col))
            if "comment" in cell and not (isinstance(cell["comment"], str) and cell["comment"]):
                raise R.RecipeError(
                    f"unit {unit!r} cell {col}: comment must be a non-empty template, "
                    f"got {cell['comment']!r}")
            self.placements.append(dict(
                sheet=sheet, row=row, col=col, raw=raw, unit=unit, comment=cell.get("comment"),
                style=cell.get("style", "label"), overrides=overrides,
                optional=cell.get("optional", False), if_ref=cell.get("if_ref"),
                if_param=cell.get("if_param"), if_it=cell.get("if_it"),
                validation=cell.get("validation"), ref_error=cell.get("ref_error"),
                params=params, item=item, landmarks=landmarks,
                band_col=band_col, binds=binds,
                globals=dict(self.sheet_globals.get(sheet, {})),   # snapshot, as item_qty was
            ))
        if merge:
            c1, c2 = merge.split(":")
            self.merges.append((sheet, f"{c1}{row}:{c2}{row + merge_down}"))
        if fill_row:
            self.region_fills.append((sheet, row, row, *self._fill_cols(sheet), fill_row))
        if conditional is not _NO_CONDITIONAL:
            # Present at all -> validated, so `conditional: {}` or `""` is an
            # error rather than a silently-ignored falsy value.
            _validate_conditional(conditional, unit, row)
        if conditional is not _NO_CONDITIONAL and conditional:
            self.conditionals.append(dict(
                sheet=sheet, row=row, unit=unit, rules=conditional,
                params=params, item=item, landmarks=landmarks,
                globals=dict(self.sheet_globals.get(sheet, {}))))

    def _check_dropdown_unlocked(self, style, overrides, where):
        """On a pack that protects its sheets, a dropdown cell must be
        unlocked: Excel will not let anyone pick from a locked cell's list,
        so the dropdown would be dead on arrival. Its effective style (the
        composite plus the cell's own overrides) has to say `locked: false`.
        """
        if not self.pack.protect.get("sheets"):
            return
        protection = self.styles.protection_for(style, overrides)
        if protection is None or protection.locked:
            raise R.RecipeError(
                f"{where}: a dropdown on a locked cell can never be used on a "
                "protected sheet -- give its style locked: false")

    def _add_box(self, sheet, r1, r2, span):
        c1, c2 = span.split(":")
        self.boxes.append((sheet, r1, r2, c1, c2))

    def _fill_cols(self, sheet):
        # A `sheet:` target carries no recipe widths, so it keeps this A-H fallback.
        return self.fill_span.get(sheet, ("A", "H"))    # engine fallback for recipes without widths

    def _recipe(self, name, *, item=None):
        """Resolve a named recipe from the pack, or fail loudly naming what's
        missing: the item (when one is known), the recipe name, the pack, and
        what the pack does offer -- never a bare KeyError."""
        try:
            return self.recipes[name]
        except KeyError:
            where = f"item {item['item_no']} ({item['sheet']})" if item is not None else "recipes.cover"
            raise R.RecipeError(
                f"{where}: recipe {name!r} is not in pack {self.pack.name!r} "
                f"(has: {sorted(self.recipes)})"
            ) from None

    def _anchored(self, entry):
        return bool((self.units.get(entry["unit"]) or {}).get("anchor"))

    def _target(self, name, origin=None):
        """The sheet an entry writes to, created on first use.

        `origin` (a describing string) is recorded in `sheet_origin` only
        when this call is the one that creates the sheet -- a later call
        that merely reuses it (several items anchoring to the same `Rates`
        sheet, say) never overwrites that first description.
        """
        if name not in self.wb.sheetnames:
            self.wb.create_sheet(name)
            if origin is not None:
                self.sheet_origin[name] = origin
        self.exports.setdefault(name, {})
        return name

    def _promote_globals(self, namespace, sheet):
        """Publish declared globals whose export belongs to the instance just laid out."""
        for name, g in self.globals_decl.items():
            if g["export"].split(".", 1)[0] != namespace:
                continue
            addr = (self.cur_exports.get(g["export"], g["default"]) if "default" in g
                    else self.cur_exports[g["export"]])
            if g.get("qualify") == "sheet":
                addr = sheet_qualified(sheet, addr)
            if g.get("scope", "sheet") == "workbook":
                if name in self.book_global_origin:
                    prev_sheet, prev_ns = self.book_global_origin[name]
                    raise R.RecipeError(
                        f"globals.{name}: promoted twice at workbook scope -- "
                        f"first by {prev_ns!r} on sheet {prev_sheet!r}, again by "
                        f"{namespace!r} on sheet {sheet!r}, in pack {self.pack.name!r} "
                        "(a workbook-scope global may be promoted from only one sheet/unit)"
                    )
                self.book_global_origin[name] = (sheet, namespace)
                self.book_globals[name] = addr
            else:
                self.sheet_globals.setdefault(sheet, {})[name] = addr

    def _param_list(self, params, ref):
        return params.get(ref.split(".", 1)[1], []) or []

    @staticmethod
    def _apply_defaults(udef, params):
        merged = dict(params)
        for name, spec in (udef.get("params") or {}).items():
            if isinstance(spec, dict) and "default" in spec and name not in merged:
                merged[name] = spec["default"]
        return merged

    @staticmethod
    def _truthy(params, ref):
        if isinstance(ref, str) and ref.startswith("p."):
            return bool(params.get(ref[2:]))
        return True

    def _place_band(self, rowdef, sheet, row, params, item, landmarks, unit):
        """Lay out a row's `repeat_cols:` band, after its fixed cells.

        One block of `width` columns per element of the `over:` list, the
        i-th starting `i * width` columns right of `start`; each cell is
        `start + i*width + offset`, resolved with `{c}` = that block's base
        column and the element bound under `as:` (default `v`). The row's
        `mark_cols_first`/`mark_cols_last` record the first and last base
        columns as column landmarks -- or `start` when the list is empty, so
        a range written across the band stays well-formed with nothing in it.
        """
        band = rowdef.get("repeat_cols")
        if not band:
            return
        start, width, name = band["start"], band["width"], band.get("as", "v")
        bases = []
        for i, element in enumerate(self._param_list(params, band["over"])):
            base = col_offset(start, i * width)
            bases.append(base)
            cells = {col_offset(base, int(k)): cell for k, cell in band["cells"].items()}
            self._place_row(cells, sheet, row, params, item, landmarks, None, None,
                            unit=unit, band_col=base, binds={name: element})
        if "mark_cols_first" in rowdef:
            landmarks[COL_LANDMARK + rowdef["mark_cols_first"]] = bases[0] if bases else start
        if "mark_cols_last" in rowdef:
            landmarks[COL_LANDMARK + rowdef["mark_cols_last"]] = bases[-1] if bases else start

    def _place_cells(self, rowdef, sheet, r, params, item, landmarks, unit_name):
        """One row: its fixed cells, then its `repeat_cols` band, then any cell
        keyed by a column landmark (`"{col:NAME+N}"`), which may name a
        landmark this row's band has just recorded."""
        cells = rowdef.get("cells") or {}
        fixed = {c: v for c, v in cells.items() if not _COL_KEY.match(c)}
        after = {c: v for c, v in cells.items() if _COL_KEY.match(c)}
        self._place_row(fixed, sheet, r, params, item, landmarks,
                        rowdef.get("merge"), rowdef.get("fill_row"),
                        rowdef.get("merge_down", 0), unit=unit_name,
                        conditional=rowdef.get("conditional", _NO_CONDITIONAL))
        self._place_band(rowdef, sheet, r, params, item, landmarks, unit_name)
        if after:
            self._place_row(after, sheet, r, params, item, landmarks, None, None, unit=unit_name)

    def _publish_each(self, each_rows, ekey, landmarks, unit_name, sheet):
        """Publish a `repeat:` row's `export_each:` once per repeated row:
        `<instance>.<key>.<NAME>`, with `key` expanded from that row's `it`
        and each NAME's address template resolved at that row's `r`."""
        where = f"unit {unit_name!r} on sheet {sheet!r} (export_each:)"
        published = set()
        for each, it, r in each_rows:
            res = Resolver(item=it, row=r, landmarks=landmarks, pack=self.pack.name, where=where)
            try:
                key = res.resolve(each["key"])
            except (KeyError, ValueError, TypeError) as e:
                raise R.RecipeError(
                    f"{where}: export_each key {each['key']!r} does not resolve for "
                    f"element {it!r}: {e!r}") from e
            if key in published:
                raise R.RecipeError(
                    f"{where}: export_each key {key!r} is published twice by instance "
                    f"{ekey!r} -- two repeated rows share it, and the second would "
                    "silently overwrite the first's exports")
            published.add(key)
            for name, tmpl in each["cells"].items():
                try:
                    self.cur_exports[f"{ekey}.{key}.{name}"] = res.resolve(tmpl)
                except CrossSheetRefError:
                    raise
                except (KeyError, ValueError) as e:
                    raise R.RecipeError(
                        f"{where}: export_each cell {name!r} for {key!r}: {e}") from e

    def layout_unit(self, unit_name, sheet, cursor, params, instance_key=None):
        udef = self.units[unit_name]
        params = self._apply_defaults(udef, params)
        landmarks = {}
        each_rows = []        # (export_each, it, r) per repeated row, published below
        anchor = udef.get("anchor")
        local = [anchor["row"]] if anchor else cursor   # anchored units use their own cursor
        start = local[0]

        for rowdef in udef.get("rows", []):
            if "blank" in rowdef:
                local[0] += rowdef["blank"]
                continue
            if "repeat" in rowdef:
                items = self._param_list(params, rowdef["repeat"])
                first = last = None
                for it in items:
                    r = local[0]
                    first = r if first is None else first
                    last = r
                    self._place_cells(rowdef, sheet, r, params, it, landmarks, unit_name)
                    if "export_each" in rowdef:
                        each_rows.append((rowdef["export_each"], it, r))
                    local[0] += 1
                if first is not None:
                    if "mark_first" in rowdef:
                        landmarks[rowdef["mark_first"]] = first
                    if "mark_last" in rowdef:
                        landmarks[rowdef["mark_last"]] = last
                continue
            if "when" in rowdef and not self._truthy(params, rowdef["when"]):
                continue
            r = local[0]
            self._place_cells(rowdef, sheet, r, params, None, landmarks, unit_name)
            if "mark" in rowdef:
                landmarks[rowdef["mark"]] = r
            if rowdef.get("box"):
                self._add_box(sheet, r, r, rowdef["box"])
            local[0] += 1

        end = local[0] - 1
        if udef.get("fill_region"):
            self.region_fills.append((sheet, start, end, *self._fill_cols(sheet), udef["fill_region"]))
        if udef.get("box_region"):
            self._add_box(sheet, start, end, udef["box_region"])

        # publish exports under the instance key (singleton = unit name)
        ekey = instance_key or unit_name
        # Anchored units deliberately re-land on the same fixed spot every
        # time they're used (a shared "Rates" sheet reached from more than
        # one item, say) -- last write wins there by design, and a SECOND
        # promotion of a workbook-scope global already has its own, more
        # specific guard (`_promote_globals`). Only a non-anchored repeat
        # publishing the same key twice is the silent-overwrite bug this
        # catches.
        publishes = udef.get("exports") or any(
            "export_each" in row for row in udef.get("rows", []))
        if publishes and not anchor:
            seen = self._seen_instance_keys.setdefault(sheet, set())
            if ekey in seen:
                raise R.RecipeError(
                    f"sheet {sheet!r}: instance key {ekey!r} (unit {unit_name!r}) is "
                    "laid out twice on this sheet -- the second instance would silently "
                    "overwrite the first's exports, pricing the wrong block. Two recipe "
                    "entries (or two repeated elements) are publishing the same key; give "
                    "each instance its own `key:`"
                )
            seen.add(ekey)
        exp_resolver = Resolver(landmarks=landmarks, pack=self.pack.name,
                                where=f"unit {unit_name!r} on sheet {sheet!r} (exports:)")
        for name, tmpl in (udef.get("exports") or {}).items():
            try:
                self.cur_exports[f"{ekey}.{name}"] = exp_resolver.resolve(tmpl)
            except CrossSheetRefError:
                raise      # the one other broad catch: never swallow this one either
            except (KeyError, ValueError):
                pass  # landmark omitted by a `when:` row -> export simply not published
        self._publish_each(each_rows, ekey, landmarks, unit_name, sheet)
        return landmarks, (start, end)

    # ---- per-sheet build ---------------------------------------------------
    def build_item_sheet(self, item):
        sheet = item["sheet"]
        # Excel (and openpyxl) treat sheet names case-insensitively, so
        # 'Front' and 'FRONT' collide too -- casefold() both sides.
        if sheet.casefold() in {s.casefold() for s in self.wb.sheetnames}:
            # Without this, openpyxl would silently create "Front1" (or
            # whatever) instead -- a broken-but-quiet sheet nobody's {ref:}
            # can ever reach. Catches a duplicate item sheet name and an
            # item named like the cover sheet (legacy path); an item named
            # like a BOOK sheet is caught earlier, in `_run_book`, against
            # every book sheet_name -- not just the ones already created.
            raise R.RecipeError(
                f"item {item['item_no']}: a sheet named {sheet!r} already "
                "exists in the workbook -- rename the item or whatever "
                "already claimed that name"
            )
        ws = self.wb.create_sheet(sheet)
        self.cur_exports = self.exports.setdefault(sheet, {})
        recipe = self._recipe(item["recipe"], item=item)
        for col, width in (recipe.get("column_widths") or {}).items():
            ws.column_dimensions[col].width = width
        for col in recipe.get("hidden_columns") or []:
            ws.column_dimensions[col].hidden = True
        _sheet_setup(ws, recipe)

        cols = sorted(recipe.get("column_widths") or {}, key=column_index_from_string)
        if cols:
            self.fill_span[sheet] = (cols[0], cols[-1])

        cursor = [1]
        spans = {}
        scopes = {"item": item, "config": self.config, "span": spans,
                  "exports": self.cur_exports}
        for entry in recipe["units"]:
            if R.is_active(entry, scopes):
                unit = R.concrete_unit(entry, scopes, self.recipes_doc)
                target = self._target(
                    entry.get("sheet", sheet),
                    origin=f"item {item['item_no']} ({sheet})'s unit {unit!r}")
                # A unit placed on another sheet publishes into that sheet's
                # namespace (and globals promote from it), never the host's.
                self.cur_exports = self.exports[target]
                try:
                    for key, params in R.instances(entry, scopes):
                        _, span = self.layout_unit(unit, target, cursor, params, instance_key=key)
                        first = spans.get(entry["unit"], (None, None))[0]
                        spans[entry["unit"]] = (span[0] if first is None else first, span[1])
                        self._promote_globals(key or unit, target)
                finally:
                    self.cur_exports = self.exports[sheet]
            if not self._anchored(entry):
                cursor[0] += 1             # one blank row between sections (visual spacing)

    def build_cover(self):
        cover = self._recipe("cover")
        sheet = cover["sheet_name"]
        self.cur_exports = self.exports.setdefault(sheet, {})
        if len(cover["units"]) != 1:
            raise R.RecipeError("recipes.cover: exactly one unit is supported")
        entry = cover["units"][0]
        items = []
        # _book_items skips a shared-cost sheet (`in_summary: false`): the
        # items that pay for it reach its TOTAL with a cross-sheet ref, so
        # listing it here too would carry every dollar on it into the bid
        # twice. Still rendered and still publishes its total -- only the
        # summary skips it.
        for it in self._book_items():
            isheet = it["sheet"]
            grand = self.exports[isheet][cover["item_total_export"]]
            items.append({"sheet_ref": quoted_sheet(isheet),
                          "total_export": "=" + sheet_qualified(isheet, grand)})
        scopes = {"item": {}, "config": self.config, "span": {}, "exports": self.cur_exports}
        (_, params), = R.instances(entry, scopes)
        params["items"] = items                   # engine-supplied: cross-sheet totals
        self.layout_unit(entry["unit"], sheet, [1], params)

    def _book_items(self):
        """The bid items a book sheet (or the legacy cover) lists: all but a
        shared-cost sheet (`in_summary: false`), whose dollars reach the bid
        through the items that pay for it."""
        return [it for it in self.config["items"] if it.get("in_summary") is not False]

    def _book_sheets(self):
        """[(sheet def, title, scopes)] for every book sheet to lay out, in order.

        A `for_each` entry yields one sheet per element of its list (none for
        an empty or absent list), the element in scope as `sheet`; its
        `sheet_name` template is expanded per element. An entry's own `when:`
        is checked against those scopes -- per element under `for_each`.
        Every title goes through `sheet_title`: the static names are reserved
        first, so a repeated sheet whose title collides with one is the one
        suffixed. Computed once; `_run_book` and `build_book` share it.
        """
        if self._book_plan is not None:
            return self._book_plan
        base = {"item": {}, "config": self.config, "span": {}, "exports": {},
                "book": {"items": self._book_items()}}
        laid = []
        for sdef in self.recipes_doc["book"]:
            if "for_each" not in sdef:
                laid.append((sdef, [base]))
                continue
            seq = R.get_path(base, sdef["for_each"])
            seq = [] if seq is R.MISSING or seq is None else seq
            if not isinstance(seq, list):
                raise R.RecipeError(f"book sheet {sdef['sheet_name']!r}: for_each "
                                    f"{sdef['for_each']} is not a list")
            laid.append((sdef, [{**base, "sheet": el} for el in seq]))
        laid = [(sdef, [s for s in per if R.is_active(sdef, s)]) for sdef, per in laid]
        taken = {sdef["sheet_name"] for sdef, per in laid if per and "for_each" not in sdef}
        plan = []
        for sdef, per in laid:
            for scopes in per:
                if "for_each" in sdef:
                    title = sheet_title(str(R.expand(sdef["sheet_name"], scopes)), taken)
                    taken.add(title)
                else:
                    title = sdef["sheet_name"]
                plan.append((sdef, title, scopes))
        self._book_plan = plan
        return plan

    def build_book(self):
        """Lay out the pack's `book:` sheets; return (before_items, after_items) names.

        A unit entry's `when:` on a book sheet is checked against THIS sheet's
        own `scopes` (`exports` scoped to just this one book sheet) -- it
        cannot see an item sheet's or another book sheet's exports.
        """
        before, after = [], []
        for sdef, title, sheet_scopes in self._book_sheets():
            sheet = self._target(title)
            widths = sdef.get("column_widths") or {}
            for col, width in widths.items():
                self.wb[sheet].column_dimensions[col].width = width
            for col in sdef.get("hidden_columns") or []:
                self.wb[sheet].column_dimensions[col].hidden = True
            if sdef.get("hidden"):
                self.wb[sheet].sheet_state = "hidden"
            _sheet_setup(self.wb[sheet], sdef)
            cols = sorted(widths, key=column_index_from_string)
            if cols:
                self.fill_span[sheet] = (cols[0], cols[-1])
            self.cur_exports = self.exports[sheet]
            scopes = {**sheet_scopes, "span": {}, "exports": self.cur_exports}
            cursor = [1]
            for entry in sdef.get("units") or []:
                if R.is_active(entry, scopes):
                    unit = R.concrete_unit(entry, scopes, self.recipes_doc)
                    for key, params in R.instances(entry, scopes):
                        self.layout_unit(unit, sheet, cursor, params, instance_key=key)
                        self._promote_globals(key or unit, sheet)
                if not self._anchored(entry):
                    cursor[0] += 1
            (before if sdef["position"] == "before_items" else after).append(sheet)
        return before, after

    def build_templates(self):
        """The hidden `_templates` item sheet Estimate Navigator copies rows from.

        Laid out by the pack's own recipe from `layout_map.templates.item`, so
        its formulas are exactly an item sheet's. Excel's copy then adjusts
        them as it would for an estimator's copy-paste. The inputs listed under
        `templates.clear` are blanked afterwards: a prototype row carries no
        trade, quantity or price. Returns the sheet name, or None when the pack
        declares no templates or the builder is not laying out navigator anchors.

        The template recipe must lay out only on its own sheet: no anchored
        unit writing to a shared sheet and no global promotion, or the
        prototype item would leak into the estimate (fieldbook's `lump_sum`
        satisfies this).
        """
        tpl = (self.pack.layout_map or {}).get("templates") or {}
        if not tpl or not self.navigator:
            return None
        item = {**copy.deepcopy(tpl["item"]), "sheet": LM.TEMPLATES_SHEET}
        self.build_item_sheet(item)
        self.wb[LM.TEMPLATES_SHEET].sheet_state = "hidden"
        self.template_clear = tpl["clear"]
        return LM.TEMPLATES_SHEET

    # ---- resolve + write ---------------------------------------------------
    def _ref_published(self, key, sheet):
        """Did `key` get published? Backs `if_ref:`, which takes both forms.

        A bare key asks the referring sheet's own namespace, exactly as it
        always has. A `SHEET!NAME` key asks that sheet's -- the one way a pack
        may deliberately skip a cell whose cross-sheet target is absent,
        instead of having the error swallowed by `optional:`.
        """
        if key in self.exports[sheet]:
            return True
        target, name = split_sheet_ref(key)
        return target is not None and name in self.exports.get(target, {})

    def _lock_blank_cell(self, p):
        """An optional placement skipped blank gets no value and no other
        styling, but its style's `locked:` key (if any) still has to govern
        whether the estimator can type into it once the sheet is protected
        -- so apply just that, leaving every untouched pack's cells (and
        goldens) exactly as before.
        """
        protection = self.styles.protection_for(p["style"], p["overrides"])
        if protection is not None:
            self.wb[p["sheet"]][f"{p['col']}{p['row']}"].protection = protection

    def _dropdown(self, p, res):
        """Record `p`'s `validation:` (if any) for the data validation emitted
        after the placements loop. Called BEFORE value resolution so a blank
        optional cell -- skipped by `write()` before it ever sets a value --
        still gets its dropdown; an optional cell does not excuse a bad list,
        so a range that resolves empty or to a rectangle (more than one row
        AND more than one column) still stops the build even though the
        cell's own value would have been allowed to go blank. Placements are
        grouped by (sheet, formula1, strict) so every cell sharing the same
        list/range and strictness gets exactly one DataValidation object.

        Raises plain `KeyError`/`ValueError` (never wrapped here) -- `write()`
        wraps whatever escapes with the placement's own location, the one
        exception being `CrossSheetRefError`, which already names its site
        and must reach the caller unchanged.
        """
        v = p.get("validation")
        if not v:
            return
        lst = v["list"]
        if isinstance(lst, list):
            formula = f'"{",".join(lst)}"'
        else:
            formula = res.resolve(lst).lstrip("=")
            if not formula:
                raise ValueError(f"validation list {lst!r} resolved to an empty range")
            if _range_spans_rows_and_columns(formula):
                raise ValueError(
                    f"validation list {lst!r} resolved to {formula!r}, which spans more "
                    "than one row and more than one column -- Excel's List validation "
                    "needs a single cell, a single column, or a single row")
        strict = v.get("strict", True)
        self.dropdowns.setdefault((p["sheet"], formula, strict), []).append(
            f"{p['col']}{p['row']}")

    def _resolver(self, p, where):
        return Resolver(params=p["params"], item=p["item"], row=p["row"],
                        col=p.get("band_col"), binds=p.get("binds"),
                        landmarks=p["landmarks"], exports=self.exports[p["sheet"]],
                        globals={**{n: None for n in self.globals_decl},
                                 **self.book_globals, **p["globals"]},
                        book_exports=self.exports, pack=self.pack.name, where=where)

    def write(self):
        # region fills first (so cells with no fill keep the band colour)
        for sheet, r1, r2, c1, c2, token in self.region_fills:
            ws = self.wb[sheet]
            fill = self.styles.fill_for(token)
            for r in range(r1, r2 + 1):
                for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
                    ws.cell(r, c).fill = fill

        # Every placement is resolved here, after the LAST sheet has been laid
        # out, so `self.exports` is the finished workbook's namespace: a
        # {ref:SHEET!NAME} reads the same whether its target sheet came before
        # or after the referring one. One pass, no ordering rule for a pack
        # author to remember.
        for p in self.placements:
            where = f"sheet {p['sheet']!r} cell {p['col']}{p['row']} (unit {p['unit']!r})"
            res = self._resolver(p, where)
            if "{" in p["style"]:
                # A style chosen per element (`style: "vendor_{v.role}"`).
                try:
                    p["style"] = res.resolve(p["style"])
                except CrossSheetRefError:
                    raise
                except (KeyError, ValueError) as e:
                    raise R.RecipeError(f"{where}: style {p['style']!r}: {e}") from e
                if p["style"] not in self.styles.styles:
                    raise R.RecipeError(f"{where}: style {p['style']!r} is not in tokens.styles")
            if p["if_ref"] and not self._ref_published(p["if_ref"], p["sheet"]):
                continue
            if p["if_param"] and p["params"].get(p["if_param"]) is None:
                continue   # render when the param is present (incl. 0); skip only when absent
            if p["if_it"] and not _matches(p["item"], p["if_it"]):
                continue
            try:
                self._dropdown(p, res)
            except CrossSheetRefError:
                raise      # already names its own site; never re-wrapped
            except (KeyError, ValueError) as e:
                raise R.RecipeError(f"{where}: validation list: {e}") from e
            try:
                value = res.resolve(p["raw"])
            except CrossSheetRefError:
                raise      # never optional: an unresolvable cross-sheet ref stops the build
            except (KeyError, ValueError) as e:
                if p["optional"]:
                    self._lock_blank_cell(p)
                    continue
                if p["ref_error"]:
                    # A pack-authored message (may itself use {it}/{p.x}/etc,
                    # which still resolve fine -- only the failing {ref:...}
                    # in the cell's own formula could not). If the message's
                    # OWN template fails to resolve too (a typo in a {p.x} it
                    # names, say), that must still surface as a named
                    # RecipeError, not the bare KeyError/ValueError leaking
                    # out of a cell that was trying to fail more gently.
                    try:
                        message = res.resolve(p["ref_error"])
                    except (KeyError, ValueError) as e2:
                        raise R.RecipeError(
                            f"{where}: ref_error itself failed to resolve: {e2}") from e2
                    raise R.RecipeError(message) from e
                raise
            literal_text = isinstance(value, str) and value.startswith(_LITERAL_TEXT)
            value = _unliteral(_coerce(value))
            if value == "" or value is None:
                if p["optional"]:
                    self._lock_blank_cell(p)
                    continue
            ws = self.wb[p["sheet"]]
            cell = ws[f"{p['col']}{p['row']}"]
            cell.value = value
            if literal_text:
                cell.data_type = "s"
            self.styles.apply(cell, p["style"], p["overrides"])
            if p.get("comment"):
                try:
                    text = res.resolve(p["comment"])
                except CrossSheetRefError:
                    raise
                except (KeyError, ValueError) as e:
                    raise R.RecipeError(f"{where}: comment: {e}") from e
                if text:
                    cell.comment = Comment(_unliteral(text), "estimate")

        for sheet, col in sorted(self.hidden_cols):
            self.wb[sheet].column_dimensions[col].hidden = True

        # Data validations are grouped by (sheet, formula1, strict): every
        # placement recorded against the same list/range and strictness gets
        # one DataValidation, not one per cell. A strict dropdown (the
        # default) stops a typed value that is not on the list; `strict:
        # false` shows no error, so the list is only a suggestion.
        for (sheet, formula, strict), cells in self.dropdowns.items():
            if strict:
                dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                                    showErrorMessage=True, errorStyle="stop",
                                    errorTitle=_DROPDOWN_ERROR_TITLE, error=_DROPDOWN_ERROR)
            else:
                dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                                    showErrorMessage=False)
            for addr in cells:
                dv.add(addr)
            self.wb[sheet].add_data_validation(dv)

        # A `conditional:` rule belongs to the ROW, not to any one cell on it:
        # it is emitted once per placed row regardless of which cells on that
        # row resolved blank or were skipped by `if_it`/`if_ref`/`if_param`/
        # `optional` -- a pack that wants the rule itself to go quiet for a
        # particular row writes that into the formula as a guard (e.g.
        # `=AND($G{r}<>"",$G{r}>10000)`), rather than relying on the engine to
        # infer it from a sibling cell's placement skip. Rules are added in
        # placement order, then in each row's own list order; on a
        # conflicting property the earlier-listed rule wins, while formats
        # that don't conflict (e.g. one rule's fill and another's font
        # colour) combine. A merged row (see `merge_down`) gets its rule on
        # the merge's FIRST row only, same as every other per-row placement.
        # `StyleFactory.dxf` lifts only font bold/italic/colour and fill from
        # the named style -- a dxf cannot carry border, alignment or
        # number_format, so those keys on the style are silently inert here
        # (never applied, never an error).
        for c in self.conditionals:
            res = self._resolver(c, f"sheet {c['sheet']!r} row {c['row']} "
                                    f"(unit {c['unit']!r} conditional)")
            ws = self.wb[c["sheet"]]
            for rule in c["rules"]:
                cols = rule["cells"].split(":")
                rng = f"{cols[0]}{c['row']}" + (f":{cols[1]}{c['row']}" if len(cols) == 2 else "")
                font, fill = self.styles.dxf(rule["style"])
                where = (f"sheet {c['sheet']!r} row {c['row']} (unit {c['unit']!r} "
                         f"conditional {rule['formula']!r})")
                try:
                    formula = res.resolve(rule["formula"]).lstrip("=")
                except CrossSheetRefError:
                    raise      # already names its own site; never re-wrapped
                except (KeyError, ValueError) as e:
                    raise R.RecipeError(f"{where}: conditional: {e}") from e
                _check_conditional_anchoring(rule["cells"], formula, where)
                ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], font=font, fill=fill))

        # thick perimeter boxes, merged with each cell's existing borders
        thick = Side(style="thick")
        for sheet, r1, r2, c1, c2 in self.boxes:
            ws = self.wb[sheet]
            ci1, ci2 = column_index_from_string(c1), column_index_from_string(c2)
            for r in range(r1, r2 + 1):
                for ci in range(ci1, ci2 + 1):
                    b = ws.cell(r, ci).border
                    ws.cell(r, ci).border = Border(
                        left=thick if ci == ci1 else b.left,
                        right=thick if ci == ci2 else b.right,
                        top=thick if r == r1 else b.top,
                        bottom=thick if r == r2 else b.bottom,
                    )

        for sheet, rng in self.merges:
            try:
                self.wb[sheet].merge_cells(rng)
            except Exception:
                pass

        # The prototype rows on `_templates` keep their formulas but lose
        # their inputs: an inserted crew/line/task starts blank.
        if LM.TEMPLATES_SHEET in self.wb.sheetnames and self.template_clear:
            ws = self.wb[LM.TEMPLATES_SHEET]
            kind_col = self.pack.layout_map["tags"]["kind"]
            for r in range(1, ws.max_row + 1):
                for col in self.template_clear.get(ws[f"{kind_col}{r}"].value, ()):
                    ws[f"{col}{r}"].value = None

        for ws in self.wb.worksheets:
            self.protect_sheet(ws)

    def protect_sheet(self, ws):
        """Protect `ws` as the pack asks (a no-op when it protects nothing)."""
        protect = self.pack.protect
        if protect.get("sheets"):
            ws.protection.sheet = True
            if protect.get("password"):
                ws.protection.password = protect["password"]
            # A protected sheet still has to let the estimator resize
            # columns/rows to fit what they typed -- only these two
            # are explicitly allowed. Everything else openpyxl protects
            # by default (insert/delete rows, sort, autoFilter,
            # formatCells) stays blocked.
            ws.protection.formatColumns = False
            ws.protection.formatRows = False

    def run(self):
        if "book" in self.recipes_doc:
            return self._run_book()
        self.wb.active.title = self._recipe("cover")["sheet_name"]   # placeholder; cover laid out last
        for item in self.config["items"]:
            self.build_item_sheet(item)
        self.build_cover()
        self.write()
        return self.wb

    def _run_book(self):
        # Every book sheet as laid out -- a `for_each` entry's expanded titles,
        # not its template. Excel (and openpyxl) treat sheet names
        # case-insensitively -- every comparison below casefold()s both sides
        # so 'Front' and 'FRONT' collide exactly as they would in the workbook.
        titles = [title for _, title, _ in self._book_sheets()]
        book_names = {t.casefold(): t for t in titles}
        # openpyxl's own default sheet is just a blank "Sheet" nothing has
        # claimed yet; remove it so that name is ordinary again (an item or
        # a book sheet may use it) and so build_book can create book[0] via
        # `_target` exactly like every other book sheet -- no placeholder
        # rename, and no exemption for book[0] in the collision check below.
        self.wb.remove(self.wb.active)
        # Estimate Navigator's hidden sheet is reserved whenever this build
        # writes navigator anchors (`_write_layout` creates it even for a
        # layout_map without templates). Checked before anything is laid out.
        if self.navigator and self.pack.layout_map:
            for item in self.config["items"]:
                if item["sheet"].casefold() == LM.TEMPLATES_SHEET.casefold():
                    raise R.RecipeError(
                        f"item {item['item_no']} ({item['sheet']}): its sheet name "
                        "collides with Estimate Navigator's hidden sheet "
                        f"{LM.TEMPLATES_SHEET!r} -- rename the item's sheet")
        for item in self.config["items"]:
            if item["sheet"].casefold() in book_names:
                raise R.RecipeError(
                    f"item {item['item_no']} ({item['sheet']}): its sheet name "
                    f"collides with book sheet_name "
                    f"{book_names[item['sheet'].casefold()]!r} -- rename one of them"
                )
            self.build_item_sheet(item)
        # Any book sheet_name an item sheet's anchored unit already created
        # (e.g. a `sheet: Rates` target) is a collision the layout below
        # would otherwise paper over -- `_target` would just reuse that
        # sheet instead of laying the book sheet out on its own.
        existing_by_fold = {s.casefold(): s for s in self.wb.sheetnames}
        for name in titles:
            actual = existing_by_fold.get(name.casefold())
            if actual is not None:
                origin = self.sheet_origin.get(actual, "an item sheet")
                raise R.RecipeError(
                    f"book sheet_name {name!r} already exists in the "
                    f"workbook (as {actual!r}), created by {origin} -- "
                    "rename one of them"
                )
        before, after = self.build_book()
        tpl = self.build_templates()
        middle = [s for s in self.wb.sheetnames if s not in before and s not in after and s != tpl]
        # openpyxl offers move_sheet(offset) only; set the whole order at once.
        self.wb._sheets = [self.wb[name] for name in before + middle + after + ([tpl] if tpl else [])]
        # book[0] -- the natural "index 0" -- may itself be `hidden:`; Excel
        # refuses to open a workbook whose active sheet is hidden, so the
        # active sheet is the first VISIBLE one in the final order, not
        # blindly the first. A workbook with NO visible sheet (no items, and
        # every book sheet hidden) cannot be opened at all -- openpyxl only
        # fails on save, with a bare IndexError -- so refuse it here, by name.
        visible = next(
            (i for i, ws in enumerate(self.wb._sheets) if ws.sheet_state == "visible"), None)
        if visible is None:
            raise R.RecipeError(
                "at least one sheet must be visible -- every book sheet is "
                "hidden and the estimate lays out no item sheets")
        self.wb.active = visible
        self.write()
        return self.wb


def _canonical(v):
    """`v` as plain JSON data: mappings of str keys, lists, and scalars --
    whatever mapping/sequence/scalar subclasses the loader produced (PyYAML's
    dicts, ruamel.yaml's round-trip types) and dates as ISO text."""
    if isinstance(v, Mapping):
        return {str(k): _canonical(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [_canonical(x) for x in v]
    if v is None or isinstance(v, bool):
        return v
    if isinstance(v, int):
        return int(v)
    if isinstance(v, float):
        return float(v)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.isoformat()
    return str(v)


def spec_fingerprint(config) -> str:
    """sha256 of the spec's data (a canonical JSON dump, keys sorted) -- the
    `estimate_spec_sha` stamp, so read-back can tell a workbook built from
    this very spec from one built before the spec last changed."""
    text = json.dumps(_canonical(config), sort_keys=True, ensure_ascii=False,
                      separators=(",", ":"))
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _template_names(wb, lm) -> dict:
    """est.tpl.task|crew|line -> whole-row refs on `_templates` (none without templates)."""
    if LM.TEMPLATES_SHEET not in wb.sheetnames:
        return {}
    ws, kind_col = wb[LM.TEMPLATES_SHEET], lm["tags"]["kind"]
    kinds = [(r, ws[f"{kind_col}{r}"].value) for r in range(1, ws.max_row + 1)]

    def first(kind):
        rows = [r for r, k in kinds if k == kind]
        if not rows:
            raise R.RecipeError(f"{LM.TEMPLATES_SHEET}: the template item laid out no {kind} row")
        return rows[0]

    head, sub = first(lm["blocks"]["head"]), first(lm["blocks"]["sub"])
    crew, line = first(lm["templates"]["rows"]["crew"]), first(lm["templates"]["rows"]["line"])
    return {"est.tpl.task": (LM.TEMPLATES_SHEET, f"${head}:${sub}"),
            "est.tpl.crew": (LM.TEMPLATES_SHEET, f"${crew}:${crew}"),
            "est.tpl.line": (LM.TEMPLATES_SHEET, f"${line}:${line}")}


def _is_value(v) -> bool:
    return v is not None and v != ""


def _last_value_row(ws) -> int:
    """The last row holding a value in any column (0 for an empty sheet).
    Not `ws.max_row`: a styled but empty trailing cell would stretch an
    `after:` table down to it."""
    last = 0
    for row in ws.iter_rows():
        if any(_is_value(c.value) for c in row):
            last = row[0].row
    return last


def _last_value_col(ws) -> int:
    """The last column holding a value in any row (0 for an empty sheet)."""
    return max((c.column for row in ws.iter_rows() for c in row if _is_value(c.value)), default=0)


def _write_layout(wb, builder, *, spec_name=None, brand=None) -> None:
    """Estimate Navigator's anchors: workbook-scoped defined names, the layout
    JSON that says how to read them, and the `estimate_layout` marker.

    Office.js truncates a custom document property past 255 characters on
    desktop Excel, so the property carries only the layout version. The JSON
    lives in one cell on the hidden `_templates` sheet -- row 1, two columns
    past the sheet's last used column -- that the name `est.layout` points at.
    A layout_map with no templates still gets an (otherwise empty) hidden
    `_templates` sheet, placed last, to hold it.
    """
    from openpyxl.packaging.custom import StringProperty
    from openpyxl.workbook.defined_name import DefinedName

    lm = builder.pack.layout_map
    max_rows = {ws.title: _last_value_row(ws) for ws in wb.worksheets}
    names, tables = LM.resolve(lm, builder.exports, max_rows)
    tpl_names = _template_names(wb, lm)
    names.update(tpl_names)
    if LM.TEMPLATES_SHEET in wb.sheetnames:
        ws = wb[LM.TEMPLATES_SHEET]
    else:
        ws = wb.create_sheet(LM.TEMPLATES_SHEET)     # created last
        ws.sheet_state = "hidden"
        builder.protect_sheet(ws)
    layout = {"pack": builder.pack.name, "pack_version": builder.pack.version,
              "layout_version": LM.LAYOUT_VERSION, "tables": tables,
              "tags": dict(lm["tags"]), "blocks": dict(lm["blocks"]),
              "templates": bool(tpl_names)}
    if builder.pack.derived_from:
        # Which layout contract the grid keeps: Estimate Navigator offers
        # read-back for a fork of a pack read-back reads.
        layout["derived_from"] = builder.pack.derived_from
    if spec_name:
        layout["spec"] = spec_name
    if brand:
        layout["brand"] = brand
    col = get_column_letter(_last_value_col(ws) + 2)
    ws[f"{col}1"] = json.dumps(layout, sort_keys=True, separators=(",", ":"))
    names["est.layout"] = (LM.TEMPLATES_SHEET, f"${col}$1")
    for name, (sheet, ref) in sorted(names.items()):
        wb.defined_names[name] = DefinedName(name, attr_text=sheet_qualified(sheet, ref))
    wb.custom_doc_props.append(StringProperty(name="estimate_layout", value=str(LM.LAYOUT_VERSION)))


def generate(config, pack=None, *, stamp=True, spec_name=None, brand=None, quotes=None):
    """Render an estimate workbook (openpyxl Workbook) from a config in a format pack.

    `pack` defaults to the bundled pack. With `stamp`, the workbook carries the
    pack lineage stamp (`FormatPack.stamp`, e.g. "arr 1 / fieldbook 2") as the
    custom property `estimate_format`, so any
    .xlsx traces to the format that produced it, and the spec's own
    fingerprint as `estimate_spec_sha` (`spec_fingerprint`, taken before
    anything reads the config), so read-back can refuse a workbook built
    from an older version of the spec.

    A pack that sets `strict: true` in its manifest also gets its composed
    schema checked here, before anything is laid out -- the same jsonschema
    validation `scaffold.validate_spec` runs, so a spec that pack's own
    vocabulary refuses (an item_rule, a `false`-schema'd field such as
    fieldbook's `cost_details`/`spreads`) is refused by `build`/the CLI too,
    not only by callers that happen to validate before building.

    A spec with `work_packages` is levelled against `quotes` (CanonicalQuote
    list; `build` reads them from the spec's quotes directory) into package
    frames, and the config is replaced by `package_frames.apply_packages`'s
    transformed copy before anything is laid out -- after validation, so the
    schema only ever sees the spec as written.
    """
    from openpyxl.packaging.custom import StringProperty

    from estimating_tools.estimate_spec.format_pack import bundled_pack
    fingerprint = spec_fingerprint(config) if stamp else None
    effective_pack = pack or bundled_pack()
    if effective_pack.strict:
        import jsonschema
        try:
            jsonschema.validate(config, effective_pack.schema)
        except jsonschema.ValidationError as exc:
            path = "$" + "".join(f"[{p!r}]" for p in exc.absolute_path)
            raise ValueError(f"schema violation: {path}: {exc.message}") from exc

    from estimating_tools.estimate_spec.validate import validate_cross_references
    validate_cross_references(config)
    if config.get("work_packages"):
        from estimating_tools.estimate_spec.package_frames import (
            apply_packages,
            build_package_frames,
        )
        book = (effective_pack.recipes or {}).get("book") or []
        taken = {sdef["sheet_name"] for sdef in book if "for_each" not in sdef}
        frames = build_package_frames(
            config, quotes or [], taken=taken,
            spare_columns=config.get("work_packages_spare_vendor_columns", 4))
        config = apply_packages(config, frames, shift_hours=config.get("shift_hours", 8))
    builder = EstimateBuilder(config, pack, navigator=stamp)
    try:
        wb = builder.run()
    except (ValueError, KeyError) as exc:
        _blame_quote_text(exc, quotes if config.get("_packages") else None)
        raise
    if stamp:
        wb.custom_doc_props.append(StringProperty(
            name="estimate_format", value=builder.pack.stamp))
        wb.custom_doc_props.append(StringProperty(name="estimate_spec_sha", value=fingerprint))
        if builder.pack.layout_map:
            _write_layout(wb, builder, spec_name=spec_name, brand=brand)
    return wb


def _quote_texts(q):
    """Every free text a quote file carries into a package sheet."""
    yield from (q.quote_ref, q.notes, q.contact_info)
    for k, v in (q.conditions or {}).items():
        yield k
        yield v
    for ln in q.lines or []:
        yield from (ln.line_name, ln.exclusion_notes, ln.folded_into)


def _blame_quote_text(exc, quotes) -> None:
    """Re-raise a resolver failure whose placeholder came from a quote file's
    text as a PackageError naming that file (quote text is bound literally,
    so this is a guard: no quote text should ever reach the resolver)."""
    if not quotes:
        return
    from estimating_tools.estimate_spec.package_frames import PackageError

    tokens = re.findall(r"\{[^{}]+\}", str(exc))
    for q in quotes:
        for text in _quote_texts(q):
            for token in tokens:
                if isinstance(text, str) and token in text:
                    where = getattr(q, "source_file", None) or f"{q.vendor_name}'s quote"
                    raise PackageError(f"{where}: text {token} cannot be rendered") from exc


def build(config_path, out_path=None, pack=None, *, stamp=True, lint=True, brand=None):
    config = yaml.safe_load(Path(config_path).read_text(encoding="utf-8"))
    quotes = None
    if config.get("work_packages"):
        from estimating_tools.quotes.files import load_quotes
        quotes = load_quotes(Path(config_path).parent
                             / config.get("work_packages_quotes_dir", "quotes"))
    if lint:
        for finding in (run_config_checks(config) if quotes is None
                        else run_config_checks(config, quotes)):
            print(f"lint: {finding}", file=sys.stderr)
    wb = generate(config, pack, stamp=stamp, spec_name=Path(config_path).name, brand=brand,
                  quotes=quotes)
    out = out_path or config.get("output") or (Path(config_path).with_suffix(".xlsx"))
    wb.save(out)
    return out


def main(argv=None):
    from estimating_tools.estimate_spec.format_pack import load_pack

    ap = argparse.ArgumentParser(description="Render an estimate workbook from an estimate_spec config.")
    ap.add_argument("config", help="estimate config YAML")
    ap.add_argument("-o", "--output", help="output .xlsx path")
    ap.add_argument("--format", help="format pack directory (default: the bundled pack)")
    ap.add_argument("--no-stamp", action="store_true", help="omit the estimate_format property")
    ap.add_argument("--no-lint", action="store_true",
                    help="skip consumer-registered pre-render config checks (#256)")
    args = ap.parse_args(argv)
    pack = load_pack(args.format) if args.format else None
    print("WROTE:", build(args.config, args.output, pack, stamp=not args.no_stamp, lint=not args.no_lint))
    return 0


if __name__ == "__main__":
    sys.exit(main())
