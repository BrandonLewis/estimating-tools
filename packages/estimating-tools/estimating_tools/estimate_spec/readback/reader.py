"""Read an edited fieldbook workbook back into its spec: refuse, read, merge, check."""
from __future__ import annotations

import copy
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path

from estimating_tools.estimate_spec import format_pack as fp
from estimating_tools.estimate_spec.readback import layout as L
from estimating_tools.estimate_spec.readback import sheets as S
from estimating_tools.estimate_spec.readback.merge import (
    Findings, driven_cells, find_needs_rate, merge_bid_date, merge_item, merge_markup,
    merge_packages, merge_rates)
from estimating_tools.estimate_spec.readback.values import blank, fmt, same_text

FIELDBOOK_DIR = fp.ENGINE_DIR / "packs" / "fieldbook"


class ReadBackRefused(ValueError):
    """The workbook cannot be read back against this spec at all."""


@dataclass
class ReadBack:
    """`spec` is the updated spec only when read-back succeeded: when nothing
    changed, or when `blocked`, it is an untouched copy of the input -- the
    merge works on its own copy and commits only on success, so a blocked
    result never hands back a half-applied spec.

    `quote_writes` are the vendor quote files the workbook's package sheets
    change or add (`merge.QuoteWrite`), returned on a dry run too -- the
    caller decides whether to write them; empty when `blocked`, like the
    spec."""
    spec: dict
    findings: Findings
    quote_writes: list = field(default_factory=list)

    @property
    def changed(self) -> bool:
        return bool(self.findings.changes)

    @property
    def blocked(self) -> bool:
        return bool(self.findings.problems)


def fieldbook_pack():
    return fp.load_pack(FIELDBOOK_DIR)


def load_workbook(path):
    """The workbook at `path` with its formulas (not cached values), or
    ReadBackRefused naming why it cannot be opened."""
    from openpyxl import load_workbook as _load
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        return _load(Path(path))
    except FileNotFoundError:
        raise ReadBackRefused(f"no such file: {path}") from None
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ReadBackRefused(f"cannot read {path} as an .xlsx workbook: {exc}") from None


def check_stamp(wb, pack) -> None:
    """Refuse another pack, version or base-pack lineage before reading cells."""
    props = {p.name: p.value for p in wb.custom_doc_props.props}
    stamp = props.get("estimate_format")
    if not stamp:
        raise ReadBackRefused(
            "not an est-estimate workbook: it carries no estimate_format property "
            "(est-estimate build stamps every workbook it writes)")
    chain = [tuple(part.strip().partition(" ")[::2]) for part in str(stamp).split("/")]
    name = chain[0][0]
    if name != pack.name:
        fix = (f"pass --format {name}" if name in fp.DERIVABLE else
               f"pass --format with the {name!r} pack's directory, or set the org "
               "profile's estimating.format to it")
        raise ReadBackRefused(
            f"not a {pack.name} workbook: this workbook was built in the {name!r} "
            f"format, but read-back resolved the {pack.name!r} pack -- {fix}")
    family = fp.readback_family(pack)
    if family != L.FAMILY:
        raise ReadBackRefused(
            f"pack {pack.name!r} is in the {family!r} family, and read-back reads only "
            f"{L.FAMILY!r} workbooks (fieldbook, or a pack whose format.yaml declares "
            f"derived_from: {L.FAMILY} or extends: {L.FAMILY})")
    want = [(n, str(v)) for n, v in (pack.chain or ((pack.name, pack.version),))]
    if chain != want:
        raise ReadBackRefused(
            f"built with {str(stamp).strip()}; this read-back "
            f"needs {pack.stamp} -- rebuild from the spec")


def _item_sheets(wb) -> list:
    return [ws.title for ws in wb.worksheets
            if ws.title != L.TEMPLATES_SHEET
            and any(ws[f"{L.KIND_COL}{r}"].value == "ITEM" for r in range(1, min(ws.max_row, 5) + 1))]


def _package_sheets(wb) -> list:
    """Titles of the `Sub <package>` sheets (the PKG tag in row 1 -- built in
    Q1, but a column inserted before it moves it), in workbook order."""
    return [ws.title for ws in wb.worksheets
            if any(isinstance(c.value, str) and c.value.strip() == L.PKG_TAG for c in ws[1])]


def check_items(spec: dict, wb) -> None:
    """Refuse a workbook whose item sheets are not the spec's items."""
    items = spec.get("items") or []
    want = sorted(str(it.get("sheet")) for it in items)
    have = sorted(_item_sheets(wb))
    if want != have:
        raise ReadBackRefused(
            f"the workbook's item sheets {have} are not the spec's {want} -- "
            "was it built from this spec?")
    for it in items:
        ws = wb[it["sheet"]]
        (row,) = [r for r in range(1, 6) if ws[f"{L.KIND_COL}{r}"].value == "ITEM"][:1]
        if not same_text(ws[f"A{row}"].value, it.get("item_no")):
            raise ReadBackRefused(
                f"sheet {it['sheet']!r} is item {ws[f'A{row}'].value}, but the spec "
                f"puts item {it.get('item_no')} there -- was it built from this spec?")
    for name in (L.RATES_SHEET, L.MARKUP_SHEET):
        if name not in wb.sheetnames:
            raise ReadBackRefused(f"the workbook has no {name!r} sheet")


STALE = ("this workbook was built from a different version of the spec (it has changed "
         "since -- another read-back, a hand edit or a bm-sync fetch). Rebuild the workbook "
         "from the current spec and redo the edits, or pass --force to apply the workbook "
         "over the current spec.")
UNSTAMPED = ("this workbook does not record which version of the spec it was built from "
             "(no estimate_spec_sha stamp: it was built before read-back stamped one), so it "
             "may undo newer spec edits. Rebuild the workbook from the current spec and redo "
             "the edits, or pass --force to apply the workbook over the current spec.")


def check_fingerprint(spec: dict, wb, f: Findings, *, force: bool) -> None:
    """A workbook built from an older version of the spec would silently put
    every spec edit made since back the way it was: a blocking problem, or
    with `force` a notice."""
    from estimating_tools.estimate_spec.build import spec_fingerprint

    stamp = {p.name: p.value for p in wb.custom_doc_props.props}.get("estimate_spec_sha")
    if stamp == spec_fingerprint(spec):
        return
    why = UNSTAMPED if not stamp else STALE
    if force:
        f.notices.append(f"--force: {why.split(' Rebuild')[0].rstrip('.')} -- applied over the current spec "
                         "anyway; the changes below include every newer spec edit it undoes")
    else:
        f.problems.append(why)


def _describe_path(spec, path: list) -> str:
    """`["items", 0, "labor", 0, "crew"]` as "item 1 · 01-100 · crew"."""
    parts, node, i = [], spec, 0
    while i < len(path):
        key = path[i]
        if key in ("items", "labor") and i + 1 < len(path) and isinstance(node.get(key), list):
            node = node[key][path[i + 1]]
            parts.append(f"item {node.get('item_no')}" if key == "items"
                         else str(node.get("cost_code")))
            i += 2
            continue
        parts.append(f"row {key + 1}" if isinstance(key, int) else str(key))
        node = node[key]
        i += 1
    return " · ".join(parts) or "the spec"


def unshare(spec) -> list:
    """Give every path through `spec` its own mapping/list, in place.

    A YAML anchor and its aliases (`crew: &mobcrew` ... `crew: *mobcrew`)
    load as ONE object reachable by several paths, so an edit read back from
    one task's rows would land in every task sharing it, and the merge of the
    other task would then report (and write) its unchanged rows as a change.
    Each later occurrence is replaced by a deep copy of its own. Returns
    `[(first path, later path)]`, described for the report."""
    first: dict = {}
    shared = []

    def walk(node, path):
        if isinstance(node, Mapping):
            children = list(node.items())
        elif isinstance(node, list):
            children = list(enumerate(node))
        else:
            return
        for key, child in children:
            if not isinstance(child, (Mapping, list)):
                continue
            here = path + [key]
            if id(child) in first:
                shared.append((_describe_path(spec, first[id(child)]), _describe_path(spec, here)))
                child = node[key] = copy.deepcopy(child)     # a fresh copy: no memo, no sharing
                _drop_anchor(child)
            first.setdefault(id(child), here)
            walk(child, here)

    walk(spec, [])
    return shared


def _drop_anchor(node) -> None:
    """A ruamel.yaml node copied from an anchored one carries the anchor's
    name; clear it, so the copy is never written out as a second `&name`."""
    anchor = getattr(node, "anchor", None)
    if anchor is not None and getattr(anchor, "value", None) is not None:
        anchor.value = None
        anchor.always_dump = False


def explain_build_error(spec, exc) -> str:
    """The test build's refusal, in the estimator's terms where it is a
    known one (a task with no crew rows left, an item with no tasks or no
    takeoff rows), with the raw message kept as a second line."""
    raw = f"the updated spec would not build: {exc}"
    cause = exc.__cause__
    path = list(getattr(cause, "absolute_path", None) or [])
    instance = getattr(cause, "instance", None)
    friendly = None
    if len(path) >= 2 and path[0] == "items" and isinstance(path[1], int):
        item = spec["items"][path[1]]
        at = f"item {item.get('item_no')}"
        rest = path[2:]
        if rest == ["labor"] and instance == []:
            friendly = f"{at} has no tasks left -- every item needs at least one task"
        elif len(rest) == 3 and rest[0] == "labor" and rest[2] == "crew" and instance == []:
            code = item["labor"][rest[1]].get("cost_code")
            friendly = (f"{at} · task {code} has no crew rows left -- every task needs at "
                        "least one crew row")
        elif rest == ["takeoff"] and isinstance(instance, Mapping) and any(
                v == [] for v in instance.values()):
            friendly = (f"{at} has no takeoff rows left -- a {item.get('recipe')} item needs "
                        "at least one")
    return raw if friendly is None else f"{friendly}\n  ({raw})"


def _quote_files(directory) -> tuple[list, list]:
    """(quotes, file names) of the quote files in `directory`, in
    `load_quotes`'s order -- the names, so an edited quote is written back
    to the file it came from."""
    from estimating_tools.quotes.files import QuoteFileError, load_quote_file

    d = Path(directory)
    if not d.is_dir():
        return [], []
    try:
        pairs = [(load_quote_file(p), p.name) for p in d.glob("*.json")
                 if not p.name.endswith(".example.json")]
    except QuoteFileError as exc:
        raise ReadBackRefused(f"cannot read the quote files: {exc}") from None
    pairs.sort(key=lambda qn: (qn[0].quote_date or "", qn[1]))
    return [q for q, _ in pairs], [n for _, n in pairs]


def _with_writes(quotes: list, names: list, writes: list) -> list:
    """`quotes` as they will be once `writes` are written: a rewritten file's
    quote replaced, a new file's added."""
    out = list(quotes)
    for w in writes:
        at = next((i for i, (q, n) in enumerate(zip(quotes, names))
                   if (n or "") == w.file_name
                   or (n is None and q.vendor_name.strip().casefold()
                       == w.quote.vendor_name.strip().casefold()
                       and q.work_package_name.strip().casefold()
                       == w.quote.work_package_name.strip().casefold())), None)
        if at is None:
            out.append(w.quote)
        else:
            out[at] = w.quote
    return out


def read_back(spec: dict, wb, pack=None, *, force: bool = False, quotes=None,
              quotes_dir=None) -> ReadBack:
    """What the edited fieldbook workbook `wb` changes in `spec`.

    `pack` is the pack that built `wb` -- fieldbook, or a pack derived from
    it -- and the one the fresh baseline build uses; None is the bundled
    fieldbook.

    Never mutates `spec`. Raises ReadBackRefused when `wb` cannot be read back
    against `spec` at all; everything else -- including problems that must
    block writing -- comes back in the result's findings. A workbook built
    from another version of `spec` (its `estimate_spec_sha` stamp) is a
    problem unless `force`.

    A spec with `work_packages` is read against its vendor quotes: `quotes`
    (CanonicalQuote list) when given, else the files in `quotes_dir`, else
    none. They must be the quotes the workbook was built with -- the baseline
    build levels them into the same package sheets."""
    from estimating_tools.estimate_spec.build import generate
    from estimating_tools.estimate_spec.package_frames import build_package_frames

    names = None
    if quotes is None and quotes_dir is not None:
        quotes, names = _quote_files(quotes_dir)
    quotes = list(quotes or [])
    names = names or [None] * len(quotes)
    pack = pack or fieldbook_pack()
    check_stamp(wb, pack)
    # The baseline build validates the spec BEFORE anything reads its item
    # records: `items: [null]` or a non-mapping item is the build's refusal,
    # never an AttributeError from check_items.
    try:
        baseline = generate(copy.deepcopy(spec), pack, quotes=quotes)
        frames = []
        if spec.get("work_packages"):
            book = (pack.recipes or {}).get("book") or []
            frames = build_package_frames(
                copy.deepcopy(spec), quotes,
                taken={sdef["sheet_name"] for sdef in book if "for_each" not in sdef},
                spare_columns=spec.get("work_packages_spare_vendor_columns", 4))
    except Exception as exc:
        # Broad on purpose: the spec itself not building is always a
        # refusal, never an uncaught traceback, whatever shape the pack's
        # own build-time error takes.
        raise ReadBackRefused(f"the spec itself does not build: {exc}") from None
    check_items(spec, wb)

    new = copy.deepcopy(spec)
    f = Findings()
    check_fingerprint(spec, wb, f, force=force)
    for a, b in unshare(new):
        f.notices.append(f"the YAML anchor shared by {a} and {b} was expanded so each can be "
                         "edited separately")
    merge_rates(new, S.read_rates(wb[L.RATES_SHEET]), f)
    if L.BID_FORM_SHEET in wb.sheetnames:
        merge_bid_date(new, *S.read_bid_date(wb[L.BID_FORM_SHEET]), f)
    merge_markup(new, S.read_markup(wb[L.MARKUP_SHEET]), f)
    pkg_sheets = {t: S.read_package_sheet(wb[t]) for t in _package_sheets(wb)}
    pkg_base = {fr.title: S.read_package_sheet(baseline[fr.title])
                for fr in frames if fr.title in baseline.sheetnames}
    writes = merge_packages(new, list(pkg_sheets.values()), quotes, frames, f,
                            file_names=names, baseline=pkg_base)
    # Package sheets are read by header, so columns may have been inserted
    # there: {title: {baseline column: edited column}} for the driven cells.
    pkg_cols = {t: S.column_map(pkg_base[t], baseline[t], pkg_sheets[t], wb[t])
                for t in pkg_base if t in pkg_sheets}
    driven = driven_cells(spec, frames)
    rate_names = [r.get("trade") for r in new.get("rate_table") or []]
    trades = list((pack.vocabulary or {}).get("trades") or [])
    for item in new.get("items") or []:
        takeoff_kind = (item.get("takeoff") or {}).get("kind")
        sheet = S.read_item_sheet(wb[item["sheet"]], takeoff_kind=takeoff_kind)
        # The same sheet read from the unedited spec's own fresh build --
        # the only way a placeholder formula field (`={ref:...}`, written
        # resolved) can tell "unchanged" from "a different formula was
        # typed over it". Its own `noticed` is never absorbed: a pristine
        # build has nothing to notice.
        baseline_sheet = S.read_item_sheet(baseline[item["sheet"]], takeoff_kind=takeoff_kind)
        # ...with its formulas moved onto the edited sheet's rows: a row the
        # estimator inserted or deleted shifts every formula below it.
        S.remap_formulas(baseline_sheet, S.row_map(baseline[item["sheet"]], wb[item["sheet"]]))
        for title, cols in pkg_cols.items():
            S.remap_columns(baseline_sheet, title, cols)
        merge_item(item, sheet, f, rate_names=rate_names, trades=trades, baseline=baseline_sheet,
                   driven=driven.get(str(item.get("item_no"))))
    if L.BID_FORM_SHEET in wb.sheetnames:
        was = baseline[L.BID_FORM_SHEET][L.BID_FORM_PROJECT].value
        now = wb[L.BID_FORM_SHEET][L.BID_FORM_PROJECT].value
        if not (blank(was) and blank(now)) and was != now:
            f.notices.append(
                f"{L.BID_FORM_SHEET}!{L.BID_FORM_PROJECT}: the project name {fmt(was)!r} → "
                f"{fmt(now)!r} is not read back -- change it in the spec")
    noticed = S.Noticed()
    for name in L.BASELINE_SHEETS:
        if name in wb.sheetnames and name in baseline.sheetnames:
            S.scan_against_baseline(
                wb[name], baseline[name], noticed,
                free_text=(L.BID_FORM_PROJECT,) if name == L.BID_FORM_SHEET else ())
    # The package sheets are read by header: their columns are compared
    # through the header map, so an inserted column moves nothing.
    for title, cols in pkg_cols.items():
        S.scan_against_baseline(wb[title], baseline[title], noticed, cols=cols)
    f.absorb(noticed)
    find_needs_rate(new, f)
    if (f.changes or writes) and not f.problems:
        try:
            generate(copy.deepcopy(new), pack, quotes=_with_writes(quotes, names, writes))
        except Exception as exc:          # broad on purpose -- see the baseline build above
            f.problems.append(explain_build_error(new, exc))
    return ReadBack(new if f.changes and not f.problems else copy.deepcopy(spec), f,
                    [] if f.problems else writes)
