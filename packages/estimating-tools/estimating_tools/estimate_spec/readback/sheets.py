"""Read an edited fieldbook workbook's sheets into plain records.

Nothing here knows the spec: each reader walks one sheet and returns what is
typed in its input cells, row by row, plus what it noticed on the way (a
formula typed over, a row with no kind tag). Item sheets are walked by the
hidden kind tag in column Q -- a TASK row opens a block, every row until the
next TASK belongs to it, END closes it -- so nothing depends on the cost-code
formulas in column R having been recalculated.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from difflib import SequenceMatcher

from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from estimating_tools.estimate_spec.readback import layout as L
from estimating_tools.estimate_spec.readback.values import blank as _blank
from estimating_tools.estimate_spec.readback.values import excel_date, remap_rows


@dataclass
class Row:
    """One sheet row: its number and the cells the reader takes from it."""
    row: int
    cells: dict

    def __getitem__(self, col: str):
        return self.cells.get(col)


@dataclass
class Task:
    row: int
    code: object
    description: object
    notes: list = field(default_factory=list)     # [Row] (column A)
    crew: list = field(default_factory=list)      # [Row]
    lines: list = field(default_factory=list)     # [Row]
    closed: bool = False                          # its END (or SUBTOTAL) row has been seen


@dataclass
class Noticed:
    """What a reader noticed besides the inputs; merged into the report."""
    overwrites: list = field(default_factory=list)   # formula cells typed over
    untagged: list = field(default_factory=list)     # rows with no (known) kind tag
    notices: list = field(default_factory=list)      # read, but worth a look
    problems: list = field(default_factory=list)     # cannot be read back as is


@dataclass
class ItemSheet:
    sheet: str
    item_no: object = None                        # column A of the ITEM row, as written
    qty: object = None
    item_row: int | None = None
    description: object = None                    # shown on the ITEM row, never read back
    uom: object = None                            # likewise
    notes: list = field(default_factory=list)     # [Row]: NOTE rows before the first task
    takeoff: list = field(default_factory=list)   # [Row]: TAKEOFF rows
    tasks: list = field(default_factory=list)     # [Task], top to bottom
    noticed: Noticed = field(default_factory=Noticed)


def _cells(ws, r: int, cols: str) -> dict:
    return {c: ws[f"{c}{r}"].value for c in cols}


def _is_formula(v) -> bool:
    return isinstance(v, (ArrayFormula, DataTableFormula)) or (isinstance(v, str) and v.startswith("="))


def _visible(ws, r: int) -> str:
    """'A: text, B: 12' for the row's non-blank, non-formula visible cells --
    a stray formula left over from a paste is not content anyone typed."""
    return ", ".join(f"{c}: {ws[f'{c}{r}'].value}" for c in L.VISIBLE_COLS
                     if not _blank(ws[f"{c}{r}"].value) and not _is_formula(ws[f"{c}{r}"].value))


def check_formulas(ws, r: int, cols, what: str, noticed: Noticed) -> None:
    """Report every cell of `cols` (column letters: a string of one-letter
    columns, or a sequence) on row `r` that no longer holds a formula."""
    for c in cols:
        v = ws[f"{c}{r}"].value
        if _is_formula(v):
            continue
        typed = "cleared" if _blank(v) else f"typed {v!r} over"
        noticed.overwrites.append(f"{ws.title}!{c}{r} ({what}): {typed} the formula")


#: The takeoff total's sum (`=SUM(B6:B7)`, anchors allowed): its first and last row.
_TAKEOFF_SUM = re.compile(r"=SUM\(\$?B\$?(\d+):\$?B\$?(\d+)\)", re.IGNORECASE)


def _check_takeoff_span(ws, out: ItemSheet, total_row: int | None) -> None:
    """A TAKEOFF row outside the rows the takeoff total sums (Excel does not
    grow `SUM(B6:B7)` for a row inserted at the total row) is still read as a
    segment -- so the rebuilt total differs from what the workbook showed.
    Say so. A total that no longer holds its SUM is reported as typed over."""
    if total_row is None:
        return
    m = _TAKEOFF_SUM.fullmatch(str(ws[f"B{total_row}"].value or "").replace(" ", ""))
    if m is None:
        return
    first, last = sorted((int(m.group(1)), int(m.group(2))))
    for row in out.takeoff:
        if not first <= row.row <= last:
            label = row["A"] if not _blank(row["A"]) else "(no label)"
            out.noticed.notices.append(
                f"{ws.title}: takeoff row {row.row} ({label}) is not in the workbook's takeoff "
                "total -- it is read back as a segment, so the total will change after the "
                "rebuild (insert segments above the total row to see them priced now)")


def read_item_sheet(ws, *, takeoff_kind: str | None) -> ItemSheet:
    """Walk an item sheet by its kind tags."""
    out = ItemSheet(sheet=ws.title)
    seen = out.noticed
    task = None
    takeoff_total = None
    takeoff_cols = "".join(col for _, col, _ in L.TAKEOFF_FIELDS.get(takeoff_kind, ()))
    for r in range(1, ws.max_row + 1):
        raw = ws[f"{L.KIND_COL}{r}"].value
        kind = raw.strip() if isinstance(raw, str) else raw
        if _blank(kind) or kind not in L.KINDS:
            text = _visible(ws, r)
            if _blank(kind) and not text:
                continue                                  # an empty row
            tag = "" if _blank(kind) else f" (unknown tag {raw!r})"
            seen.untagged.append(f"{ws.title} row {r}{tag}: {text or '(no visible text)'}")
            continue
        label = ws[f"A{r}"].value
        check_formulas(ws, r, L.formula_columns(
            kind, label=label, in_task=task is not None,
            takeoff_kind=takeoff_kind, has_takeoff=takeoff_kind is not None),
            kind, seen)
        if kind == "TOTAL" and label == "Takeoff total":
            takeoff_total = r
        if kind == "ITEM":
            out.item_no, out.qty, out.item_row = label, ws[f"{L.ITEM_QTY}{r}"].value, r
            out.description, out.uom = ws[f"{L.ITEM_DESC}{r}"].value, ws[f"{L.ITEM_UOM}{r}"].value
        elif kind == "TAKEOFF":
            out.takeoff.append(Row(r, _cells(ws, r, takeoff_cols)))
        elif kind == "TASK":
            task = Task(r, ws[f"{L.TASK_CODE}{r}"].value, ws[f"{L.TASK_DESC}{r}"].value)
            out.tasks.append(task)
        elif kind == "NOTE":
            note = Row(r, _cells(ws, r, L.NOTE_TEXT))
            (task.notes if task is not None else out.notes).append(note)
        elif kind in ("CREW", "LINE"):
            if task is None:
                seen.problems.append(
                    f"{ws.title} row {r}: a {kind} row above the first task block "
                    f"belongs to no task -- move it into a block (above an END row)")
                continue
            cols = (L.CREW_TRADE + L.CREW_COUNT + L.CREW_DAYS if kind == "CREW" else
                    L.LINE_NAME + L.LINE_QTY + L.LINE_UOM + L.LINE_PRICE
                    + L.LINE_TYPE + L.LINE_SOURCE + L.LINE_VENDOR)
            cells = _cells(ws, r, cols)
            if all(_blank(v) for v in cells.values()):
                # Cleared, not deleted -- the estimator's way of removing a
                # row without moving every other row's formulas.
                seen.notices.append(f"{ws.title} row {r}: {kind} row was cleared -- read as removed")
                continue
            if task.closed:
                seen.notices.append(
                    f"{ws.title} row {r}: {kind} row below task {task.code}'s END row "
                    f"(outside its subtotal in this workbook) -- read into {task.code}")
            (task.crew if kind == "CREW" else task.lines).append(Row(r, cells))
        elif kind in ("END", "SUBTOTAL"):
            if task is not None:
                task.closed = True
    _check_takeoff_span(ws, out, takeoff_total)
    return out


def _tagged_rows(ws) -> list:
    """[(row, kind, column-A literal or "")] for every kind-tagged row. A
    formula in column A (a RECAP's `=$A$18`, a SUBTOTAL's label) is "" --
    its text moves with the rows, so it cannot identify one."""
    out = []
    for r in range(1, ws.max_row + 1):
        kind = ws[f"{L.KIND_COL}{r}"].value
        if isinstance(kind, str) and kind.strip() in L.KINDS:
            a = ws[f"A{r}"].value
            out.append((r, kind.strip(), "" if _blank(a) or _is_formula(a) else str(a)))
    return out


def _pair_rows(base: list, edited: list, out: dict, *, by_text: bool = True) -> None:
    """Pair tagged rows in order: by (kind, column-A text) first; inside a run
    that differs (a label retyped, rows inserted or deleted), by kind alone."""
    key = (lambda t: (t[1], t[2])) if by_text else (lambda t: t[1])
    matcher = SequenceMatcher(None, [key(t) for t in base], [key(t) for t in edited],
                              autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            out.update((base[i1 + k][0], edited[j1 + k][0]) for k in range(i2 - i1))
        elif tag == "replace" and by_text:
            _pair_rows(base[i1:i2], edited[j1:j2], out, by_text=False)


def row_map(base_ws, ws) -> dict:
    """{baseline row: edited row} for an item sheet: where each row of a
    fresh build of the spec now sits in the edited sheet, matched by the
    hidden kind tags (and column-A text, where it is literal). Rows the
    estimator inserted have no baseline row; rows deleted have no entry."""
    out: dict = {}
    _pair_rows(_tagged_rows(base_ws), _tagged_rows(ws), out)
    return out


def remap_formulas(sheet: ItemSheet, rows: dict) -> None:
    """Move every formula read from a fresh build's item sheet onto the
    edited sheet's rows (`row_map`), in place -- so a placeholder input
    (`={ref:quantity.total}`, built as `=$B$8`) still reads as unchanged
    after Excel shifted it to `=$B$9` for an inserted takeoff row."""
    sheet.qty = remap_rows(sheet.qty, rows)
    groups = [sheet.notes, sheet.takeoff]
    for task in sheet.tasks:
        groups += [task.notes, task.crew, task.lines]
    for group in groups:
        for row in group:
            row.cells = {c: remap_rows(v, rows) for c, v in row.cells.items()}


@dataclass
class Rates:
    shift_hours: object = None
    shift_found: bool = False                     # the label row exists (its value may be blank)
    rows: list = field(default_factory=list)      # [Row]: columns A..G
    noticed: Noticed = field(default_factory=Noticed)


def read_rates(ws) -> Rates:
    """The Rates table: every non-empty row between its head row and the
    insert-above marker (or the sheet's end, if the marker was deleted).
    Rows below the marker are not in the table at all -- noticed, not
    silently dropped, since a rate pasted there misses both the crew
    VLOOKUP table and the resource dropdown."""
    out = Rates()
    head = None
    for r in range(1, ws.max_row + 1):
        a = ws[f"A{r}"].value
        if a == L.SHIFT_LABEL:
            out.shift_hours, out.shift_found = ws[f"B{r}"].value, True
        elif a == L.RATES_HEAD:
            head = r
            break
    if head is None:
        out.noticed.problems.append(f"{ws.title}: no '{L.RATES_HEAD}' head row -- the rate table cannot be read")
        return out
    marker = None
    for r in range(head + 1, ws.max_row + 1):
        cells = _cells(ws, r, "ABCDEFG")
        if cells["A"] == L.RATES_MARKER:
            marker = r
            break
        if all(_blank(v) for v in cells.values()):
            continue
        out.rows.append(Row(r, cells))
        # H ($/hr) is never read back into the spec -- the rate is derived
        # from A-G -- so a blank H on a plainly inserted row is not "cleared"
        # (there was nothing to clear); a literal value typed there is.
        h = ws[f"{L.RATE_HOURLY}{r}"].value
        if not _blank(h) and not _is_formula(h):
            out.noticed.overwrites.append(
                f"{ws.title}!{L.RATE_HOURLY}{r} (rate $/hr): typed {h!r} over the formula")
    if marker is not None:
        for r in range(marker + 1, ws.max_row + 1):
            cells = _cells(ws, r, "ABCDEFG")
            if all(_blank(v) for v in cells.values()):
                continue
            name = cells["A"] if not _blank(cells["A"]) else "(no name)"
            out.noticed.notices.append(
                f"{ws.title} row {r} ({name}): below the Rates marker -- not in the rate "
                "table; insert rates above the marker")
    return out


@dataclass
class Markup:
    rows: dict = field(default_factory=dict)       # lower-case row name -> Row (B..E)
    bond: object = None
    bond_found: bool = False                       # the label row exists (its value may be blank)
    unbalance: list = field(default_factory=list)  # [Row]: columns A, C
    noticed: Noticed = field(default_factory=Noticed)


def read_markup(ws) -> Markup:
    """The Markup grid, bond and per-item unbalances. Bond and the unbalance
    header/rows are checked for FIRST on every row, before the grid-row
    capture -- so a grid whose Total row was deleted or relabelled (and so
    never closes on its own) cannot swallow them. The grid also closes on
    the first blank column-A row under it, as a second line of defence."""
    out = Markup()
    grid = unbal = False
    saw_total = False
    for r in range(1, ws.max_row + 1):
        a, c = ws[f"A{r}"].value, ws[f"C{r}"].value
        if a == L.BOND_LABEL:
            out.bond, out.bond_found = ws[f"B{r}"].value, True
            grid = False
            continue
        if c == L.UNBALANCE_HEAD:
            unbal, grid = True, False
            continue
        if unbal:
            if a == L.UNBALANCE_TOTAL:
                unbal = False
            elif not _blank(a):
                out.unbalance.append(Row(r, _cells(ws, r, "AC")))
            continue
        if a == L.MARKUP_HEAD:
            grid = True
            continue
        if grid:
            if a == L.MARKUP_TOTAL:
                grid, saw_total = False, True
            elif _blank(a):
                grid = False       # no Total row to close it -- stop here instead
            elif isinstance(a, str):
                name = a.strip().lower()
                if name in out.rows:
                    # Excel sums both rows; keeping either one alone would
                    # silently drop the other's markup.
                    out.noticed.problems.append(
                        f"{L.MARKUP_SHEET} rows {out.rows[name].row} and {r} are both "
                        f"{a.strip()!r} -- combine them into one row")
                    continue
                out.rows[name] = Row(r, _cells(ws, r, "ABCDE"))
    if not saw_total:
        out.noticed.notices.append(
            f"{L.MARKUP_SHEET}: no '{L.MARKUP_TOTAL}' row found under the grid -- "
            "closed it at the first blank row")
    return out


def read_bid_date(ws):
    """(value, True) when the 'Bid date' label row exists -- `value` may be
    blank, a date, or whatever free text was typed there. (None, False)
    when the label row itself is missing, never confused with a blank one."""
    for r in range(1, ws.max_row + 1):
        if ws[f"A{r}"].value == L.BID_DATE_LABEL:
            v = ws[f"B{r}"].value
            day = excel_date(v)
            return (v if day is None else day), True
    return None, False


def scan_against_baseline(ws, base_ws, noticed: Noticed, *, free_text=(), cols=None) -> None:
    """A fixed-shape book sheet: every cell that holds a formula in a fresh
    build of the spec must still hold one. Skipped, with a notice, when the
    sheet's column A no longer matches the fresh build's (rows were moved).
    `free_text` names column-A cells whose text may differ without that
    meaning anything moved (the Bid Form's project name).

    `cols` ({baseline column: edited column}, `column_map`), when given,
    is where each baseline column now sits -- a package sheet read by
    header, where columns may have been inserted; a baseline column with no
    counterpart (deleted) is not checked."""
    rows = max(ws.max_row, base_ws.max_row)
    if any(ws[f"A{r}"].value != base_ws[f"A{r}"].value for r in range(1, rows + 1)
           if f"A{r}" not in free_text):
        noticed.notices.append(
            f"{ws.title}: rows were added, removed or moved -- not checked for typed-over formulas")
        return
    for row in base_ws.iter_rows(min_row=1, max_row=base_ws.max_row):
        for base in row:
            if _is_formula(base.value):
                col = base.column_letter if cols is None else cols.get(base.column_letter)
                if col is None:
                    continue
                # A one-column tuple: `cols` is iterated, and "AA" is not "A" twice.
                check_formulas(ws, base.row, (col,), "formula", noticed)


@dataclass
class VendorCol:
    """One column of a package sheet's vendor band. A spare is a blank header
    (`name` is ""); a vendor typed into it is a vendor like any other."""
    name: str
    col: str
    quote_ref: object = None
    surcharge: object = None
    is_spare: bool = False


@dataclass
class PkgRow:
    """One package line. `cells` holds each vendor column's price cell, keyed
    by vendor name -- by column letter for a spare, so a typed-in price under a
    spare header can be told apart."""
    row: int
    item_no: object
    line_name: object
    pick: object = None                           # H: only an explicit per-line pick
    carry: object = None
    cells: dict = field(default_factory=dict)


@dataclass
class PackageSheet:
    name: str                                     # A1: the package's name
    sheet: str = ""                               # the sheet's title
    selector: object = None
    vendors: list = field(default_factory=list)   # [VendorCol], left to right
    rows: list = field(default_factory=list)      # [PkgRow]
    conditions: dict = field(default_factory=dict)  # (condition text, vendor key) -> cell
    cols: dict = field(default_factory=dict)      # PKG_HEADERS key -> column letter found
    header_row: int | None = None
    noticed: Noticed = field(default_factory=Noticed)


def _norm(v) -> str:
    return str(v).strip().casefold() if not _blank(v) else ""


def read_package_sheet(ws) -> PackageSheet:
    """A `Sub <package>` sheet, located by header text: nothing depends on the
    column a header sits in, except that the vendor band is whatever lies
    between the `ARR ext` and `CARRY` headers."""
    from openpyxl.utils import get_column_letter as letter

    H = L.PKG_HEADERS
    a1 = ws["A1"].value
    out = PackageSheet(name=str(a1).strip() if not _blank(a1) else ws.title, sheet=ws.title)
    problems = out.noticed.problems

    def missing(what: str) -> None:
        problems.append(f"{ws.title}: no '{what}' found -- the package sheet cannot be read as is")

    head = next((r for r in range(1, ws.max_row + 1)
                 if _norm(ws[f"A{r}"].value) == _norm(H["item"])
                 and _norm(ws[f"B{r}"].value) != _norm(L.PKG_ITEMS_HEAD)), None)
    if head is None:
        missing(H["item"])
        return out
    out.header_row = head
    for r in range(1, head):
        if _norm(ws[f"E{r}"].value) == _norm(L.PKG_SELECTOR_LABEL):
            out.selector = ws[f"F{r}"].value
            break
    else:
        missing(L.PKG_SELECTOR_LABEL)

    # Header text -> column index (first occurrence), for the fixed columns.
    texts = {c: _norm(ws.cell(head, c).value) for c in range(1, ws.max_column + 1)}
    found = {}
    for key in ("line", "qty", "uom", "rate", "pick", "arr", "arr_ext", "carry", "source", "state"):
        found[key] = next((c for c in sorted(texts) if texts[c] == _norm(H[key])), None)
        if found[key] is not None:
            out.cols[key] = letter(found[key])
    for key in ("carry", "source", "arr_ext", "pick"):
        if found[key] is None:
            missing(H[key])
    lo, hi = found["arr_ext"], found["carry"]
    if lo is not None and hi is not None and hi > lo:
        seen: dict = {}
        for c in range(lo + 1, hi):
            raw = ws.cell(head, c).value
            if _norm(raw) == _norm(H["ext"]):
                continue                              # a vendor's leveled-ext partner column
            name = "" if _blank(raw) else str(raw).strip()
            if name and name.casefold() in seen:
                problems.append(
                    f"{ws.title}: vendor {name!r} appears in columns {seen[name.casefold()]} and "
                    f"{letter(c)} -- vendor names must be unique")
                continue
            if name:
                seen[name.casefold()] = letter(c)
            out.vendors.append(VendorCol(name, letter(c), is_spare=not name))
    elif lo is not None and hi is not None:
        problems.append(f"{ws.title}: the '{H['carry']}' column sits before '{H['arr_ext']}'")

    qrow = srow = None
    for r in range(head + 1, head + 4):
        a = _norm(ws[f"A{r}"].value)
        if a == _norm(L.PKG_QUOTE_REF_LABEL) and qrow is None:
            qrow = r
        elif a == _norm(L.PKG_SURCHARGE_LABEL) and srow is None:
            srow = r
    if qrow is None:
        missing(L.PKG_QUOTE_REF_LABEL)
    if srow is None:
        missing(L.PKG_SURCHARGE_LABEL)
    for v in out.vendors:
        if qrow is not None:
            v.quote_ref = None if _blank(ws[f"{v.col}{qrow}"].value) else ws[f"{v.col}{qrow}"].value
        if srow is not None:
            v.surcharge = None if _blank(ws[f"{v.col}{srow}"].value) else ws[f"{v.col}{srow}"].value

    def stop(r: int) -> bool:
        b = _norm(ws[f"B{r}"].value)
        return (b in (_norm(L.PKG_ITEM_ROW_LABEL), _norm(L.PKG_ITEMS_HEAD))
                or _norm(ws[f"A{r}"].value) == _norm(L.PKG_TOTAL_LABEL))

    first = max(head, qrow or 0, srow or 0) + 1
    end = first
    while end <= ws.max_row and not stop(end):
        end += 1
    pick, carry = out.cols.get("pick"), out.cols.get("carry")
    for r in range(first, end):
        a, b = ws[f"A{r}"].value, ws[f"B{r}"].value
        cells = {(v.name or v.col): (None if _blank(ws[f"{v.col}{r}"].value) else ws[f"{v.col}{r}"].value)
                 for v in out.vendors}
        if _blank(a) and _blank(b) and all(_blank(x) for x in cells.values()):
            continue
        out.rows.append(PkgRow(
            r, a, b,
            pick=None if not pick or _blank(ws[f"{pick}{r}"].value) else ws[f"{pick}{r}"].value,
            carry=None if not carry or _blank(ws[f"{carry}{r}"].value) else ws[f"{carry}{r}"].value,
            cells=cells))

    ch = next((r for r in range(end, ws.max_row + 1)
               if _norm(ws[f"B{r}"].value) == _norm(L.PKG_CONDITIONS_LABEL)), None)
    if ch is not None:
        r = ch + 1
        while r <= ws.max_row and not _blank(ws[f"B{r}"].value):
            text = str(ws[f"B{r}"].value).strip()
            for v in out.vendors:
                cell = ws[f"{v.col}{r}"].value
                out.conditions[(text, v.name or v.col)] = None if _blank(cell) else cell
            r += 1
    return out


def column_map(base: PackageSheet, base_ws, edited: PackageSheet, ws) -> dict:
    """{baseline column letter: edited column letter} for a package sheet,
    matched by the header row's text (a column inserted in Excel has no
    baseline column; a vendor typed into a spare header pairs by position).
    Empty when either sheet has no header row."""
    from openpyxl.utils import get_column_letter as letter

    if base.header_row is None or edited.header_row is None:
        return {}

    def heads(sheet, w):
        return [(letter(c), _norm(w.cell(sheet.header_row, c).value))
                for c in range(1, w.max_column + 1)]

    a, b = heads(base, base_ws), heads(edited, ws)
    out: dict = {}
    matcher = SequenceMatcher(None, [t for _, t in a], [t for _, t in b], autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal" or (tag == "replace" and i2 - i1 == j2 - j1):
            out.update((a[i1 + k][0], b[j1 + k][0]) for k in range(i2 - i1))
    return out


def _sheet_cells_ref(title: str):
    quoted = re.escape("'" + title.replace("'", "''") + "'")
    bare = re.escape(title)
    return re.compile(rf"((?:{quoted}|(?<![\w']){bare})!\$?)([A-Z]{{1,3}})(\$?\d+)")


def remap_columns(sheet: ItemSheet, title: str, cols: dict) -> None:
    """Move every reference into the package sheet `title` in a fresh
    build's item-sheet formulas onto the edited package sheet's columns
    (`column_map`), in place -- what Excel did to them when a column was
    inserted there. A column with no counterpart is left as it was."""
    if not cols:
        return
    ref = _sheet_cells_ref(title)

    def move(v):
        if not (isinstance(v, str) and v.startswith("=")):
            return v
        return ref.sub(lambda m: m.group(1) + cols.get(m.group(2), m.group(2)) + m.group(3), v)

    groups = [sheet.notes, sheet.takeoff]
    for task in sheet.tasks:
        groups += [task.notes, task.crew, task.lines]
    for group in groups:
        for row in group:
            row.cells = {c: move(v) for c, v in row.cells.items()}
