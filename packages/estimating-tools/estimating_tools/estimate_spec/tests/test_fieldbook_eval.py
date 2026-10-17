"""The fieldbook workbook computes the bid it describes -- evaluated with pycel.

Every number here is what the built workbook's own formulas compute for the
pack's example spec; none is re-derived in Python except the independent
recap-vs-detail sum, which adds the CREW and LINE rows' totals by their tags.
"""
from __future__ import annotations

import copy
import sys

import pytest
from openpyxl.formula.translate import Translator

from estimating_tools.estimate_spec.tests._fieldbook import (
    CODE, KIND, build, example_spec, rows_of_kind, task_block)
from estimating_tools.estimate_spec.tests._formula_eval import evaluate

ITEMS = ("1 - MOBILIZATION", "2 - 8IN MAIN", "3 - VALVES")
ITEM2 = "2 - 8IN MAIN"


def _check_row(wb, rule: int, item) -> int:
    ws = wb["Checks"]
    for r in range(1, ws.max_row + 1):
        if ws[f"A{r}"].value == rule and ws[f"C{r}"].value == item:
            return r
    raise AssertionError(f"no check row for rule {rule}, item {item!r}")


def _statuses(wb, ev) -> dict:
    """{(rule, item): status} for every check row on the Checks sheet."""
    ws = wb["Checks"]
    return {(ws[f"A{r}"].value, ws[f"C{r}"].value): ev[f"Checks!F{r}"]
            for r in range(1, ws.max_row + 1) if isinstance(ws[f"A{r}"].value, int)}


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    wb = build()
    return wb, evaluate(wb, tmp_path_factory.mktemp("fieldbook"))


def test_missing_pycel_fails_under_ci_and_skips_locally(monkeypatch):
    from estimating_tools.estimate_spec.tests import _formula_eval
    monkeypatch.setitem(sys.modules, "pycel", None)          # `import pycel` now raises
    monkeypatch.setenv("CI", "true")
    with pytest.raises(pytest.fail.Exception, match="pycel is required under CI"):
        _formula_eval._require_pycel()
    monkeypatch.setenv("CI", "")
    with pytest.raises(pytest.skip.Exception):
        _formula_eval._require_pycel()


def test_the_example_prices_to_the_mockup_figures(built):
    wb, ev = built
    total = {c: ev[f"Cost Summary!{c}8"] for c in "CDEFGHIJKL"}
    assert total["C"] == 552                       # MH
    assert (total["D"], total["E"], total["F"], total["G"]) == (40320, 14795, 33390, 10880)
    assert total["H"] == 99385                     # direct
    assert total["I"] == 14790                     # plugs
    assert total["J"] == pytest.approx(18024.65)   # markup by cost type
    assert total["K"] == 0                         # unbalance nets out
    assert total["L"] == pytest.approx(118835.67813765182)
    assert ev["Bid Form!F9"] == pytest.approx(118837.20)
    # plugs moved from the Bid Form strip to the Cost Summary (Brandon,
    # option A: the Bid Form is agency-facing and carries no internal cost).
    assert ev["Cost Summary!L10"] == 14790
    assert ev["Cost Summary!L11"] == pytest.approx(0.148815213563415)


def test_each_item_bid_solves_markup_unbalance_and_bond(built):
    _, ev = built
    assert [ev[f"Cost Summary!L{r}"] for r in (4, 5, 6)] == pytest.approx(
        [5440.0, 106315.27327935223, 7080.404858299596])
    # item 1: 3,472 labor x 26% markup = 902.72; (3,472 + 902.72 + 1,000) / (1 - 1.2%) = 5,440.00
    assert ev["Cost Summary!J4"] == pytest.approx(902.72)


def test_the_bid_form_extends_rounded_unit_prices(built):
    _, ev = built
    assert [ev[f"Bid Form!E{r}"] for r in (5, 6, 7)] == pytest.approx([5440.00, 166.12, 1770.10])
    assert [ev[f"Bid Form!F{r}"] for r in (5, 6, 7)] == pytest.approx([5440.00, 106316.80, 7080.40])
    gap = abs(ev["Bid Form!F9"] - ev["Cost Summary!L8"])
    allowance = sum(ev[f"Bid Form!C{r}"] for r in (5, 6, 7)) * 0.005 + 0.005
    assert gap <= allowance                          # within the unit-price rounding allowance


def test_every_item_recap_equals_its_detail_rows(built):
    wb, ev = built
    for name in ITEMS:
        ws = wb[name]
        detail = sum(ev[f"{name}!P{r}"] for r in rows_of_kind(ws, "CREW") + rows_of_kind(ws, "LINE"))
        (total,) = [r for r in rows_of_kind(ws, "TOTAL") if ws[f"A{r}"].value == "ITEM TOTAL"]
        assert ev[f"{name}!P{total}"] == pytest.approx(detail), name
    assert ev[f"{ITEM2}!P13"] == 89969
    (sub,) = task_block(wb[ITEM2], "02-100")["SUBTOTAL"]
    assert [ev[f"{ITEM2}!{c}{sub}"] for c in "IKLMNOP"] == [6510, 384, 28608, 11315, 33390, 0, 73313]


def test_every_check_passes(built):
    # 22, not 20: check 6 (Bid Form = Cost Summary) moved from one pooled
    # row to one row per item (review fix), so the denominator grew by 2.
    wb, ev = built
    statuses = _statuses(wb, ev)
    assert len(statuses) == 22
    assert set(statuses.values()) == {"✓"}
    assert (ev["Checks!B2"], ev["Checks!D2"]) == (22, 22)
    assert (ev["Bid Form!B11"], ev["Bid Form!D11"]) == (22, 22)


# ---------------------------------------------------------------- tampering

def _failing(spec, tmp_path) -> set:
    wb = build(spec)
    ev = evaluate(wb, tmp_path)
    return {key for key, status in _statuses(wb, ev).items() if status != "✓"}


def test_a_takeoff_that_disagrees_with_bid_qty_fails_check_1(tmp_path):
    spec = example_spec()
    spec["items"][1]["qty"] = 600                    # the takeoff still says 640
    assert _failing(spec, tmp_path) == {(1, 2)}


def test_an_erroring_takeoff_total_fails_check_1(tmp_path):
    # ISNUMBER(error) is FALSE, same as ISNUMBER("n/a") -- a check that
    # treated "not a number" as the lump-sum sentinel would wrongly call an
    # erroring takeoff total the explicit "n/a" case and report ✓.
    wb = build()
    ws = wb[ITEM2]
    row = rows_of_kind(ws, "TAKEOFF")[0]
    ws[f"B{row}"] = "=1/0"
    ev = evaluate(wb, tmp_path)
    assert {k for k, s in _statuses(wb, ev).items() if s != "✓"} == {(1, 2)}


def test_a_crew_resource_without_a_rate_fails_check_2(tmp_path):
    spec = example_spec()
    spec["rate_table"] = [r for r in spec["rate_table"] if r["trade"] != "Plate compactor"]
    assert _failing(spec, tmp_path) == {(2, 2)}


def test_a_crew_resource_with_no_recognized_kind_fails_check_2(tmp_path):
    # Pipelayer's Kind cleared directly on the Rates sheet: its $/hr ($/hr
    # column, VLOOKUP col 8) still resolves numerically (the rate survives),
    # so a check that only asked ISNUMBER(rate) would stay green while the
    # crew's own labor/equip lanes price at $0 (neither IF(Kind="Labor",...)
    # nor IF(Kind="Equipment",...) matches a blank Kind).
    wb = build()
    wb["Rates"]["B6"] = None
    ev = evaluate(wb, tmp_path)
    assert ev[f"{ITEM2}!L21"] == 0
    assert ev[f"{ITEM2}!M21"] == 0
    assert {k for k, s in _statuses(wb, ev).items() if s != "✓"} == {(2, 2), (2, 3)}


def test_a_crew_resource_picked_as_the_rates_insert_marker_fails_check_2(tmp_path):
    # The insert-above marker row's own name sits inside the crew-resource
    # dropdown range (by design -- a pasted rate must land inside it), so an
    # estimator editing the built workbook in Excel can pick it as if it
    # were a real resource (the spec-level route is schema-refused outright
    # -- `trade` is one of the pack's named trades -- so this is tampered
    # directly on the built workbook, the way a dropdown pick would land).
    # Its $/hr cell is blank, which VLOOKUP returns as 0 -- a NUMBER -- so a
    # check that only asked ISNUMBER(rate) would wrongly call this fine.
    wb = build()
    ws = wb[ITEM2]
    row = task_block(ws, "02-100")["CREW"][1]        # "Pipelayer"
    ws[f"A{row}"] = "insert new rates above this row"
    ev = evaluate(wb, tmp_path)
    assert {k for k, s in _statuses(wb, ev).items() if s != "✓"} == {(2, 2)}


def test_a_cost_line_without_a_price_fails_check_3(tmp_path):
    spec = example_spec()
    spec["items"][2]["labor"][0]["lines"][0]["unit_price"] = 0
    assert _failing(spec, tmp_path) == {(3, 3)}


def test_a_cost_line_with_a_blank_type_fails_check_3(tmp_path):
    # Clearing H (Type) directly on the built workbook: B (qty) and G (price)
    # both stay numeric and non-zero, so a check that only asked
    # ISNUMBER(qty)/ISNUMBER(price) would wrongly call this priced.
    wb = build()
    ws = wb[ITEM2]
    line_row = rows_of_kind(ws, "LINE")[0]            # "Trench box (rental)": qty 11, price 185
    ws[f"H{line_row}"] = None
    ev = evaluate(wb, tmp_path)
    assert ev[f"{ITEM2}!B{line_row}"] not in (0, None)
    assert ev[f"{ITEM2}!G{line_row}"] not in (0, None)
    assert {k for k, s in _statuses(wb, ev).items() if s != "✓"} == {(3, 2)}


def test_an_unbalance_that_does_not_net_out_fails_check_4(tmp_path):
    spec = example_spec()
    spec["items"][1]["unbalance"] = -900
    assert _failing(spec, tmp_path) == {(4, "all")}


def test_a_unit_price_typed_over_on_the_bid_form_fails_check_6(tmp_path):
    # E6 is item 2's row (Task 10's layout): check 6 is now per item (review
    # fix), so only item 2's row trips, not a pooled "all".
    wb = build()
    wb["Bid Form"]["E6"] = 170                       # typed over the rounded formula
    ev = evaluate(wb, tmp_path)
    assert {k for k, s in _statuses(wb, ev).items() if s != "✓"} == {(6, 2)}


def test_an_operator_included_machine_beside_an_operator_fails_check_8(tmp_path):
    spec = example_spec()
    excavator = next(r for r in spec["rate_table"] if r["trade"] == "Excavator, 320 class")
    excavator["includes_operator"] = True            # 02-100 also carries an Equipment Operator
    assert _failing(spec, tmp_path) == {(8, 2)}


# ---------------------------------------------------------------- copy and paste

def _copy_rows(ws, first: int, last: int, to: int) -> None:
    """Copy rows first..last to start at row `to`, as Excel's copy/paste does:
    values and styles travel, relative references shift with the paste."""
    for offset, src_row in enumerate(range(first, last + 1)):
        for src in ws[src_row]:
            dst = ws.cell(to + offset, src.column)
            value = src.value
            if isinstance(value, str) and value.startswith("="):
                value = Translator(value, origin=src.coordinate).translate_formula(dst.coordinate)
            dst.value = value
            if src.has_style:
                dst._style = copy.copy(src._style)


def test_a_pasted_task_block_follows_its_new_header(tmp_path):
    wb = build()
    ws = wb[ITEM2]
    block = task_block(ws, "02-200")
    (first,), (sub,) = block["TASK"], block["SUBTOTAL"]
    to = sub + 2
    _copy_rows(ws, first, sub, to)
    ws[f"A{to}"] = "02-300"                          # pick a new code in the pasted header
    last = to + (sub - first)
    assert ws[f"{CODE}{to}"].value == f'=A{to}&""'
    assert [ws[f"{CODE}{r}"].value for r in range(to + 1, last + 1)] == [
        f"={CODE}{r - 1}" for r in range(to + 1, last + 1)]
    assert [ws[f"{KIND}{r}"].value for r in range(to, last + 1)] == [
        ws[f"{KIND}{r}"].value for r in range(first, sub + 1)]
    ev = evaluate(wb, tmp_path)
    assert {ev[f"{ITEM2}!{CODE}{r}"] for r in range(to, last + 1)} == {"02-300"}
    assert ev[f"{ITEM2}!{CODE}{sub}"] == "02-200"    # the original block keeps its code
    # until the estimate is read back and rebuilt, the recap misses the pasted rows
    assert ev[f"Checks!F{_check_row(wb, 5, 2)}"] == "✗"


def test_a_task_header_left_without_a_code_fails_check_7(tmp_path):
    wb = build()
    ws = wb[ITEM2]
    (head,) = task_block(ws, "02-200")["TASK"]
    ws[f"A{head}"] = None                            # the header's code, cleared in Excel
    ev = evaluate(wb, tmp_path)
    assert ev[f"Checks!F{_check_row(wb, 7, 2)}"] == "✗"
    assert ev[f"Checks!F{_check_row(wb, 7, 3)}"] == "✓"
