"""`est-estimate read-back`: dry run, --write with a backup, --report, exit codes."""
from __future__ import annotations

import io
import sys

import pytest
import yaml
from openpyxl import load_workbook

from estimating_tools.estimate_spec import cli
from estimating_tools.estimate_spec.tests._fieldbook import EXAMPLE, FIELDBOOK_DIR
from estimating_tools.estimate_spec.tests.test_readback_spec_io import commented_spec


@pytest.fixture(autouse=True)
def _no_profile(monkeypatch):
    """read-back resolves its pack the way build does; no org profile on the
    test machine may pick it (these workbooks are built with fieldbook)."""
    monkeypatch.delenv("ESTIMATING_PROFILES", raising=False)


@pytest.fixture
def bid(tmp_path):
    """A spec (the example, comments and all) and the workbook built from it."""
    spec = tmp_path / "bid.estimate.yaml"
    spec.write_text(EXAMPLE.read_text(encoding="utf-8"), encoding="utf-8")
    xlsx = tmp_path / "bid.xlsx"
    assert cli.main(["build", str(spec), "-o", str(xlsx), "--format", str(FIELDBOOK_DIR)]) == 0
    return spec, xlsx


def _edit(xlsx, **cells):
    """Set {"Sheet!A1": value} cells in the saved workbook, as Excel would."""
    wb = load_workbook(xlsx)
    for address, value in cells.items():
        sheet, cell = address.split("!")
        wb[sheet][cell] = value
    wb.save(xlsx)


def _backups(spec):
    return sorted(spec.parent.glob(spec.name + ".bak-*"))


def test_a_dry_run_reports_and_writes_nothing(bid, capsys):
    spec, xlsx = bid
    before = spec.read_bytes()
    _edit(xlsx, **{"Markup!B5": 0.10})
    assert cli.main(["read-back", str(spec), str(xlsx)]) == 0
    out = capsys.readouterr().out
    assert "# Read-back: bid.xlsx → bid.estimate.yaml" in out
    assert "- markup · profit labor 0.12 → 0.1" in out
    assert out.rstrip().endswith("dry run: 1 change(s) -- run again with --write to apply them")
    assert spec.read_bytes() == before
    assert _backups(spec) == []


def test_write_saves_a_backup_then_the_updated_spec(bid, tmp_path, capsys):
    spec, xlsx = bid
    before = spec.read_bytes()
    _edit(xlsx, **{"Markup!B5": 0.10})
    report = tmp_path / "readback.md"
    assert cli.main(["read-back", str(spec), str(xlsx), "--write", "--report", str(report)]) == 0
    (backup,) = _backups(spec)
    assert backup.read_bytes() == before
    written = yaml.safe_load(spec.read_text(encoding="utf-8"))
    expected = yaml.safe_load(before.decode("utf-8"))
    expected["markup"]["profit"]["labor"] = 0.10
    assert written == expected
    comments = [line for line in before.decode("utf-8").splitlines() if line.startswith("#")]
    assert [line for line in spec.read_text(encoding="utf-8").splitlines()
            if line.startswith("#")] == comments                # the example's header survives
    assert "- markup · profit labor 0.12 → 0.1" in report.read_text(encoding="utf-8")
    assert capsys.readouterr().out.rstrip().endswith(
        f"wrote 1 change(s) to {spec} (backup: {backup.name})")


@pytest.mark.parametrize("newline", ["\n", "\r\n"])
def test_write_keeps_the_specs_comments_key_order_layout_and_line_endings(tmp_path, newline):
    text = commented_spec()
    spec = tmp_path / "bid.estimate.yaml"
    spec.write_text(text, encoding="utf-8", newline=newline)
    xlsx = tmp_path / "bid.xlsx"
    assert cli.main(["build", str(spec), "-o", str(xlsx), "--format", str(FIELDBOOK_DIR)]) == 0
    _edit(xlsx, **{"Markup!B5": 0.10})
    assert cli.main(["read-back", str(spec), str(xlsx), "--write"]) == 0
    after = spec.read_bytes().decode("utf-8")
    assert ("\r\n" in after) == (newline == "\r\n")
    before_lines, after_lines = text.splitlines(), after.replace("\r\n", "\n").splitlines()
    assert len(after_lines) == len(before_lines)
    assert [(a, b) for a, b in zip(before_lines, after_lines) if a != b] == [
        ("    labor: 0.12", "    labor: 0.1")]


def test_nothing_changed_writes_nothing_even_with_write(bid, capsys):
    spec, xlsx = bid
    before = spec.read_bytes()
    assert cli.main(["read-back", str(spec), str(xlsx), "--write"]) == 0
    assert "No changes: the workbook matches the spec." in capsys.readouterr().out
    assert spec.read_bytes() == before                  # byte for byte: comments survive
    assert _backups(spec) == []


def test_blocking_problems_exit_1_and_write_nothing(bid, capsys):
    spec, xlsx = bid
    before = spec.read_bytes()
    _edit(xlsx, **{"Markup!B5": "twelve"})
    assert cli.main(["read-back", str(spec), str(xlsx), "--write"]) == 1
    captured = capsys.readouterr()
    assert "## Problems — nothing is written until these are fixed" in captured.out
    assert "blocked: 1 problem(s) -- nothing written" in captured.err
    assert spec.read_bytes() == before
    assert _backups(spec) == []


def test_refusals_exit_2(bid, tmp_path, capsys):
    spec, xlsx = bid
    other = tmp_path / "other.xlsx"
    assert cli.main(["build", str(spec), "-o", str(other), "--format", str(FIELDBOOK_DIR),
                     "--no-stamp"]) == 0
    assert cli.main(["read-back", str(spec), str(other)]) == 2
    assert "read-back refused: not an est-estimate workbook" in capsys.readouterr().err
    assert cli.main(["read-back", str(spec), str(tmp_path / "missing.xlsx")]) == 2
    assert "read-back refused: no such file" in capsys.readouterr().err
    assert cli.main(["read-back", str(tmp_path / "missing.yaml"), str(xlsx)]) == 2
    assert "no such file" in capsys.readouterr().err


def test_a_console_that_cannot_show_the_report_gets_escapes_not_a_crash(bid, monkeypatch):
    spec, xlsx = bid
    _edit(xlsx, **{"Markup!B5": 0.10})
    raw = io.BytesIO()
    monkeypatch.setattr(sys, "stdout", io.TextIOWrapper(raw, encoding="cp1252"))
    assert cli.main(["read-back", str(spec), str(xlsx)]) == 0
    sys.stdout.flush()
    assert b"# Read-back: bid.xlsx \\u2192 bid.estimate.yaml" in raw.getvalue()


def _rebuilds(spec, tmp_path):
    out = tmp_path / "rebuilt.xlsx"
    return cli.main(["build", str(spec), "-o", str(out), "--format", str(FIELDBOOK_DIR)]) == 0


def test_text_yaml_1_1_would_misread_is_written_so_build_reads_it_as_typed(bid, tmp_path):
    """`build` reads the spec with PyYAML (YAML 1.1): a plain `No`, `Off` or
    `1:30` would come back a bool or a base-60 int and fail the schema."""
    spec, xlsx = bid
    from estimating_tools.estimate_spec.tests._fieldbook import (
        simulate_insert_copied_row, task_block)
    wb = load_workbook(xlsx)
    ws = wb["2 - 8IN MAIN"]
    ws["C6"] = "1:30"                       # a takeoff ref
    (bedding,) = [r for r in range(1, ws.max_row + 1)
                  if ws[f"Q{r}"].value == "LINE" and ws[f"A{r}"].value == "Bedding rock"]
    ws[f"J{bedding}"], ws[f"C{bedding}"] = "Yes", "No"     # vendor, uom
    end = task_block(ws, "02-100")["END"][0]
    simulate_insert_copied_row(ws, end, end - 1)
    ws[f"A{end}"], ws[f"B{end}"], ws[f"C{end}"], ws[f"G{end}"] = "Off", 1, "On", 10
    ws[f"I{end}"], ws[f"J{end}"] = "PLUG", "Yes"
    wb.save(xlsx)
    assert cli.main(["read-back", str(spec), str(xlsx), "--write"]) == 0
    written = yaml.safe_load(spec.read_text(encoding="utf-8"))
    main = written["items"][1]
    assert main["takeoff"]["segments"][0]["ref"] == "1:30"
    lines = main["labor"][0]["lines"]
    assert (lines[-2]["uom"], lines[-2]["vendor"]) == ("No", "Yes")
    assert (lines[-1]["name"], lines[-1]["uom"], lines[-1]["vendor"]) == ("Off", "On", "Yes")
    assert _rebuilds(spec, tmp_path)


@pytest.mark.parametrize("edit", [{}, {"2 - 8IN MAIN!G26": 44}])
def test_a_spec_using_yaml_1_1_booleans_reads_back_as_build_reads_it(tmp_path, capsys, edit):
    text = EXAMPLE.read_text(encoding="utf-8").replace(
        'rate: 145.00, kind: equipment, includes_operator: false',
        'rate: 145.00, kind: equipment, includes_operator: no')
    assert "includes_operator: no" in text
    spec = tmp_path / "bid.estimate.yaml"
    spec.write_text(text, encoding="utf-8")
    xlsx = tmp_path / "bid.xlsx"
    assert cli.main(["build", str(spec), "-o", str(xlsx), "--format", str(FIELDBOOK_DIR)]) == 0
    _edit(xlsx, **edit)
    assert cli.main(["read-back", str(spec), str(xlsx), "--write"]) == 0
    out = capsys.readouterr().out
    assert "includes_operator" not in out           # `no` is false, as build reads it
    assert ("No changes" in out) == (not edit)
    assert yaml.safe_load(spec.read_text(encoding="utf-8"))["rate_table"][4]["includes_operator"] is False
    assert _rebuilds(spec, tmp_path)


def test_a_spec_build_reads_differently_is_refused(tmp_path, capsys):
    """A mapping key YAML 1.1 reads as a bool cannot be patched: exit 2."""
    spec = tmp_path / "bid.estimate.yaml"
    spec.write_text(EXAMPLE.read_text(encoding="utf-8") + "x-notes: {yes: 1, 'True': 2}\n",
                    encoding="utf-8")
    assert cli.main(["read-back", str(spec), str(tmp_path / "missing.xlsx")]) == 2
    assert "cannot be read the way build reads it" in capsys.readouterr().err


def test_a_dump_that_would_read_back_differently_writes_nothing(bid, monkeypatch, capsys):
    from estimating_tools.estimate_spec.readback import spec_io
    spec, xlsx = bid
    before = spec.read_bytes()
    _edit(xlsx, **{"Markup!B5": 0.10})
    real = spec_io.SpecYaml._dump_text
    monkeypatch.setattr(spec_io.SpecYaml, "_dump_text",
                        lambda self, s: real(self, s).replace("labor: 0.1", "labor: 0.2"))
    assert cli.main(["read-back", str(spec), str(xlsx), "--write"]) == 1
    assert "would read back differently" in capsys.readouterr().err
    assert spec.read_bytes() == before and _backups(spec) == []


def test_a_second_workbook_from_the_same_spec_is_refused_after_the_first_is_written(
        bid, tmp_path, capsys):
    """Two estimators, two workbooks built from one spec: once A is written,
    B is stale -- applying it would silently revert A's edits."""
    spec, a = bid
    b = tmp_path / "b.xlsx"
    assert cli.main(["build", str(spec), "-o", str(b), "--format", str(FIELDBOOK_DIR)]) == 0
    _edit(a, **{"Markup!B5": 0.10})
    _edit(b, **{"Markup!B7": 0.015})
    assert cli.main(["read-back", str(spec), str(a), "--write"]) == 0
    capsys.readouterr()
    written = spec.read_bytes()
    assert cli.main(["read-back", str(spec), str(b), "--write"]) == 1
    out = capsys.readouterr().out
    assert "built from a different version of the spec" in out
    assert "- markup · profit labor 0.1 → 0.12" in out           # what it would revert
    assert spec.read_bytes() == written
    assert cli.main(["read-back", str(spec), str(b), "--write", "--force"]) == 0
    assert "different version of the spec" in capsys.readouterr().out    # now a notice
    assert yaml.safe_load(spec.read_text(encoding="utf-8"))["markup"]["profit"]["labor"] == 0.12


def test_read_back_without_ruamel_says_how_to_install_it(bid, monkeypatch, capsys):
    spec, xlsx = bid
    for name in ("ruamel", "ruamel.yaml", "estimating_tools.estimate_spec.readback.spec_io"):
        monkeypatch.setitem(sys.modules, name, None)        # import -> ImportError
    assert cli.main(["read-back", str(spec), str(xlsx)]) == 2
    assert ("read-back needs ruamel.yaml: pip install 'estimating_tools[estimate_spec]'"
            in capsys.readouterr().err)


def _aliases(path):
    """The same file named another way: different case (on a case-folding
    file system), relative to the working directory, and via `..`."""
    import os
    yield str(path)
    if os.path.normcase("A") == os.path.normcase("a"):
        yield str(path.parent / path.name.upper())
    yield os.path.relpath(path)
    yield str(path.parent / ".." / path.parent.name / path.name)


@pytest.mark.parametrize("target", ["spec", "workbook"])
def test_a_report_path_naming_the_spec_or_the_workbook_is_refused(bid, monkeypatch, capsys, target):
    spec, xlsx = bid
    monkeypatch.chdir(spec.parent)
    _edit(xlsx, **{"Markup!B5": 0.10})
    victim = spec if target == "spec" else xlsx
    for alias in _aliases(victim):
        spec_bytes, xlsx_bytes = spec.read_bytes(), xlsx.read_bytes()
        assert cli.main(["read-back", str(spec), str(xlsx), "--write", "--report", alias]) == 2, alias
        err = capsys.readouterr().err
        assert f"--report {alias} is the {target} itself" in err
        assert spec.read_bytes() == spec_bytes and xlsx.read_bytes() == xlsx_bytes
        assert _backups(spec) == []


def test_a_report_path_that_is_a_symlink_to_the_spec_is_refused(bid, tmp_path, capsys):
    spec, xlsx = bid
    link = tmp_path / "report.md"
    try:
        link.symlink_to(spec)
    except (OSError, NotImplementedError):
        pytest.skip("this platform/account cannot create symlinks")
    before = spec.read_bytes()
    assert cli.main(["read-back", str(spec), str(xlsx), "--report", str(link)]) == 2
    assert "is the spec itself" in capsys.readouterr().err
    assert spec.read_bytes() == before


def test_a_spec_with_a_null_item_is_refused_with_exit_2(bid, capsys):
    spec, xlsx = bid
    text = spec.read_text(encoding="utf-8")
    spec.write_text(text.replace("\nitems:\n", "\nitems:\n  - null\n", 1), encoding="utf-8")
    assert "  - null" in spec.read_text(encoding="utf-8")
    assert cli.main(["read-back", str(spec), str(xlsx)]) == 2
    assert "read-back refused: the spec itself does not build" in capsys.readouterr().err
