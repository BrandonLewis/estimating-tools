"""Fieldbook's work-package sheets: layout only (Task 12 evaluates the values).

`packaged_spec()` adds two packages to the example: Hauling (item 2's haul
line; Dirt Movers priced, Bulk Haul Co a lump sum) and Valves (item 3's
valve-box line, item 3 task 01 self-perform; Valve Pros priced). Each package
gets a `Sub <name>` sheet; `Sub Spreads` summarises them; the item sheets read
the carried price and the self-perform switch from those sheets.
"""
from __future__ import annotations

import copy

import pytest
from openpyxl.utils import column_index_from_string, get_column_letter

from estimating_tools.estimate_spec import build as B
from estimating_tools.estimate_spec.package_frames import build_package_frames
from estimating_tools.estimate_spec.tests._fieldbook import (
    QUOTES, example_spec, fieldbook_pack, packaged_spec, rows_of_kind, task_block)

ITEM2, ITEM3 = "2 - 8IN MAIN", "3 - VALVES"


@pytest.fixture(scope="module")
def wb():
    return B.generate(packaged_spec(), fieldbook_pack(), quotes=copy.deepcopy(QUOTES))


def _row(ws, col, value, start=1):
    return next(r for r in range(start, ws.max_row + 1) if ws[f"{col}{r}"].value == value)


def _blank(v):
    return v is None or v == ""


def _off(col, n):
    return get_column_letter(column_index_from_string(col) + n)


def _heads(ws):
    """(header row, {label: column}) of a package sheet's line table."""
    hdr = next(r for r in range(1, ws.max_row + 1)
               if ws[f"A{r}"].value == "Item" and ws[f"B{r}"].value == "Line")
    cols = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(hdr, c).value
        if isinstance(v, str) and v and v not in cols:
            cols[v] = get_column_letter(c)
    return hdr, cols


def _line_row(ws, name):
    hdr, _ = _heads(ws)
    return _row(ws, "B", name, start=hdr + 1)


def _item_row(ws, item_no):
    """The item subtotal row (B "item ARR cost") of `item_no`."""
    return next(r for r in range(1, ws.max_row + 1)
                if ws[f"B{r}"].value == "item ARR cost" and str(ws[f"A{r}"].value) == item_no)


# ---------------------------------------------------------------- the sheets

def test_one_sheet_per_package_after_the_items_behind_the_summary(wb):
    names = wb.sheetnames
    assert names[-4:] == ["Sub Spreads", "Sub Hauling", "Sub Valves", "_templates"]
    assert names.index(ITEM3) < names.index("Sub Spreads")


def test_an_unpackaged_spec_lays_out_no_package_sheet_and_no_check_10():
    wb = B.generate(example_spec(), fieldbook_pack())
    assert not [n for n in wb.sheetnames if n.startswith("Sub ")]
    ws = wb["Checks"]
    assert 10 not in {ws[f"A{r}"].value for r in range(1, ws.max_row + 1)}
    for name in (ITEM2, ITEM3):
        sheet = wb[name]
        assert all(_blank(sheet[f"U{r}"].value) for r in range(1, sheet.max_row + 1))


# ---------------------------------------------------------------- package sheet head

def test_the_head_names_the_package_and_tags_it(wb):
    ws = wb["Sub Hauling"]
    assert ws["A1"].value == "Hauling"
    assert ws["Q1"].value == "PKG"
    assert (ws["A2"].value, ws["B2"].value) == ("Kind", "subcontract")
    assert (ws["E3"].value, ws["F3"].value) == ("CARRYING", "Dirt Movers")


def test_the_selector_is_a_dropdown_of_arr_carry_and_the_vendors(wb):
    ws = wb["Sub Hauling"]
    (dv,) = [d for d in ws.data_validations.dataValidation if "F3" in str(d.sqref)]
    assert dv.formula1 == "$H$3:$O$3"                       # ARR, CARRY, six vendor slots
    assert (ws["H3"].value, ws["I3"].value) == ("ARR", "CARRY")
    hdr, _ = _heads(ws)
    assert ws["J3"].value == f'=""&INDEX($K${hdr}:$W${hdr},1,0*2+1)'
    assert ws["O3"].value == f'=""&INDEX($K${hdr}:$W${hdr},1,5*2+1)'


# ---------------------------------------------------------------- the vendor band

def test_the_vendors_row_holds_the_quoted_vendors_the_spares_then_carry(wb):
    ws = wb["Sub Hauling"]
    hdr, cols = _heads(ws)
    assert [ws[f"{c}{hdr}"].value for c in "ABCDEFGHIJ"] == [
        "Item", "Line", "Code", "Qty", "U/M", "$ / unit", "Extension", "Carry from", "ARR", "ARR ext"]
    band = [ws[f"{c}{hdr}"].value for c in ("K", "M", "O", "Q", "S", "U")]
    assert band[:2] == ["Dirt Movers", "Bulk Haul Co"]
    assert all(_blank(v) for v in band[2:])                 # four spares
    assert [ws[f"{c}{hdr}"].value for c in ("L", "N", "V")] == ["ext", "ext", "ext"]
    assert [ws[f"{c}{hdr}"].value for c in ("W", "X", "Y", "Z")] == [
        "CARRY", "CARRY ext", "Source", "State"]
    assert ws.column_dimensions["AA"].hidden and ws.column_dimensions["AB"].hidden


def test_the_spare_vendor_heads_are_inputs_and_the_quoted_ones_are_not(wb):
    ws = wb["Sub Hauling"]
    hdr, _ = _heads(ws)
    assert ws[f"K{hdr}"].fill.fgColor.rgb != ws[f"O{hdr}"].fill.fgColor.rgb
    assert ws[f"O{hdr}"].font.color.rgb.endswith("1747B5")     # input blue


def test_the_quote_ref_and_surcharge_rows(wb):
    ws = wb["Sub Hauling"]
    hdr, _ = _heads(ws)
    assert ws[f"A{hdr + 1}"].value == "Quote ref"
    assert ws[f"N{hdr + 1}"].value == "LUMP SUM"              # Bulk Haul Co's amount is prorated
    assert _blank(ws[f"L{hdr + 1}"].value)
    # A quoted vendor's ref is never blank: its quote_ref, else its quote date,
    # else its quote file's name; a spare's is blank.
    assert str(ws[f"K{hdr + 1}"].value)[:10] == "2026-10-20"   # Dirt Movers: its date
    assert ws[f"M{hdr + 1}"].value == "bulk-haul-co--hauling.json"
    assert _blank(ws[f"O{hdr + 1}"].value)
    assert ws[f"A{hdr + 2}"].value == "Surcharge % of invoice"
    assert ws[f"K{hdr + 2}"].value == 0.04 and ws[f"M{hdr + 2}"].value == 0


# ---------------------------------------------------------------- line rows

def test_a_line_row_reads_its_qty_from_the_item_sheet(wb):
    ws = wb["Sub Hauling"]
    r = _line_row(ws, "Haul, import + export")
    item = wb[ITEM2]
    ir = _row(item, "A", "Haul, import + export")
    assert [ws[f"{c}{r}"].value for c in "ABCE"] == [2, "Haul, import + export", "02-200", "CY"]
    assert ws[f"D{r}"].value == f"='{ITEM2}'!$B${ir}"
    assert ws[f"G{r}"].value == f"=F{r}*D{r}"
    assert ws[f"I{r}"].value == 18
    assert ws[f"J{r}"].value == f"=I{r}*D{r}"
    assert _blank(ws[f"H{r}"].value)                          # no per-item pick: the selector rules


def test_the_vendor_cells_hold_prices_and_a_prorated_share_says_so(wb):
    ws = wb["Sub Hauling"]
    r = _line_row(ws, "Haul, import + export")
    assert ws[f"K{r}"].value == 16.5
    assert ws[f"M{r}"].value == pytest.approx(19.0)
    assert ws[f"K{r}"].comment is None
    assert ws[f"M{r}"].comment.text == "prorated from $8,740.00 lump sum"
    assert ws[f"M{r}"].font.i and not ws[f"K{r}"].font.i     # derived_money vs input_money
    assert ws[f"M{r}"].fill.fgColor.rgb.endswith("F6FAF7")    # derived_money's shaded `note` fill
    assert _blank(ws[f"O{r}"].value)                         # a spare


def test_source_state_and_the_carried_rate(wb):
    ws = wb["Sub Hauling"]
    hdr, _ = _heads(ws)
    r = _line_row(ws, "Haul, import + export")
    s, sur = hdr + 2, hdr + 2
    # The vendor ranges end at the CARRY column (W), so a column inserted at
    # CARRY -- or a vendor typed into the last spare -- is inside them.
    assert ws[f"W{hdr}"].value == "CARRY"
    head = f"$K${hdr}:$W${hdr}"
    match = f"MATCH($Y{r},{head},0)"
    vendor = f"INDEX($K{r}:$W{r},{match})"
    assert ws[f"Y{r}"].value == f'=IF($H{r}="",$F$3,$H{r})'
    assert ws[f"Z{r}"].value == (
        f'=IF($Y{r}="ARR","ARR",IF($Y{r}="CARRY",IF(ISNUMBER($W{r}),"price",NA()),'
        f'IF(ISNUMBER({vendor}),"price",IF(OR({vendor}="incl",{vendor}="excl"),{vendor},'
        f'IF({vendor}="","excl",NA())))))')
    assert ws[f"F{r}"].value == (
        f'=IF(OR($Z{r}="ARR",$Z{r}="excl"),$I{r},IF($Z{r}="incl",0,'
        f'IF($Y{r}="CARRY",$W{r},{vendor})*(1+IFERROR(INDEX($K${sur}:$W${sur},{match}),0))))')
    assert ws[f"AA{r}"].value == f'=IF($Z{r}="price","QUOTE","PLUG")'
    assert "IFERROR" not in ws[f"Z{r}"].value               # a vendor that is not there stays #N/A
    assert s == sur


def test_a_vendor_extension_levels_an_exclusion_at_the_arr_price(wb):
    ws = wb["Sub Hauling"]
    hdr, _ = _heads(ws)
    r = _line_row(ws, "Haul, import + export")
    assert ws[f"L{r}"].value == (
        f'=IF(K${hdr}="","",IF(ISNUMBER(K{r}),K{r}*$D{r}*(1+K${hdr + 2}),'
        f'IF(K{r}="incl",0,$J{r})))')


# ---------------------------------------------------------------- item rows, totals

def test_an_item_row_switches_its_self_perform_tasks(wb):
    ws = wb["Sub Valves"]
    hdr, cols = _heads(ws)
    assert cols["State"] == "X"                               # 1 vendor + 4 spares: K..S
    r = _item_row(ws, "3")
    first, last = hdr, _line_row(ws, "Valve boxes (subcontract)")   # ranges start at the head row
    assert ws[f"W{r}"].value == (f"=INDEX($W${first}:$W${last},"
                                 f"MATCH($A{r},$A${first}:$A${last},0))")
    assert ws[f"X{r}"].value == f'=IF($W{r}="ARR",1,0)'
    task = task_block(wb[ITEM3], "02-510")
    assert ws[f"J{r}"].value == (f"=SUMIF($A${first}:$A${last},$A{r},$J${first}:$J${last})"
                                 f"+('{ITEM3}'!$V${task['SUBTOTAL'][0]})")


def test_the_totals_row_publishes_the_arr_and_carried_totals(wb):
    ws = wb["Sub Hauling"]
    t = _row(ws, "A", "TOTAL")
    i = _item_row(ws, "2")
    assert ws[f"J{t}"].value == f"=SUM(J{i - 1}:J{i})"         # from the items' head row
    assert ws[f"G{t}"].value == f"=SUM(G{i - 1}:G{i})"
    summary = wb["Sub Spreads"]
    r = _row(summary, "A", "Hauling")
    assert summary[f"D{r}"].value == f"='Sub Hauling'!$J${t}"
    assert summary[f"H{r}"].value == f"='Sub Hauling'!$G${t}"


# ---------------------------------------------------------------- conditions, detail

def test_the_conditions_grid_holds_each_vendors_text(wb):
    ws = wb["Sub Hauling"]
    r = _row(ws, "B", "Fuel surcharge")
    assert ws[f"K{r}"].value == "4% of invoice"
    assert _blank(ws[f"M{r}"].value)
    assert ws[f"K{_row(ws, 'B', 'Dump fees')}"].value == "included"


def test_the_detail_block_lists_the_takeoff_when_asked():
    spec = packaged_spec()
    spec["work_packages"][1]["detail"] = "precast"
    wb = B.generate(spec, fieldbook_pack(), quotes=copy.deepcopy(QUOTES))
    ws = wb["Sub Valves"]
    r = _row(ws, "A", "Gate valve")
    assert [ws[f"{c}{r}"].value for c in "BCD"] == [4, "EA", "takeoff"]
    assert "Gate valve" not in {c.value for c in wb["Sub Hauling"]["A"]}


# ---------------------------------------------------------------- item sheets

def test_a_package_line_reads_its_rate_source_and_vendor_from_the_package_sheet(wb):
    item, pkg = wb[ITEM2], wb["Sub Hauling"]
    ir = _row(item, "A", "Haul, import + export")
    pr = _line_row(pkg, "Haul, import + export")
    assert item[f"G{ir}"].value == f"='Sub Hauling'!$F${pr}"
    assert item[f"I{ir}"].value == f"='Sub Hauling'!$AA${pr}"
    assert item[f"J{ir}"].value == f"='Sub Hauling'!$Y${pr}"
    assert (item[f"U{ir}"].value, item[f"V{ir}"].value) == ("PKG", "=(460)*(18)")


def test_a_self_perform_crew_reads_the_switch_and_keeps_its_arr_basis(wb):
    item, pkg = wb[ITEM3], wb["Sub Valves"]
    block = task_block(item, "02-510")
    (pipe,) = [r for r in block["CREW"] if item[f"A{r}"].value == "Pipelayer"]
    assert item[f"E{pipe}"].value == f"=2*'Sub Valves'!$X${_item_row(pkg, '3')}"
    assert item[f"U{pipe}"].value == "SP"
    assert item[f"V{pipe}"].value == f"=D{pipe}*(2)*'Rates'!$B$2*G{pipe}"
    sub = block["SUBTOTAL"][0]
    head, end = block["TASK"][0], block["END"][0]
    assert item[f"V{sub}"].value == f'=SUMIF(U{head}:U{end},"SP",V{head}:V{end})'
    assert item[f"K{head}"].value == "(self-perform: off when a vendor is carried)"


def test_a_supply_package_never_claims_the_self_perform_switch():
    # A supply package listed FIRST on item 3 (valve-box castings, furnished):
    # the crew still reads the subcontract package's factor (Sub Valves).
    spec = packaged_spec()
    (item3,) = [i for i in spec["items"] if i["item_no"] == "3"]
    item3["labor"][0]["lines"].append(
        {"cost_type": "material", "name": "Valve box castings", "qty": 4, "uom": "EA",
         "unit_price": 120, "package": "Castings"})
    spec["work_packages"].insert(0, {"name": "Castings", "kind": "supply", "items": ["3"],
                                     "selected": "ARR"})
    wb = B.generate(spec, fieldbook_pack(), quotes=QUOTES)
    item, pkg = wb[ITEM3], wb["Sub Valves"]
    block = task_block(item, "02-510")
    (pipe,) = [r for r in block["CREW"] if item[f"A{r}"].value == "Pipelayer"]
    assert item[f"E{pipe}"].value == f"=2*'Sub Valves'!$X${_item_row(pkg, '3')}"
    assert "Sub Castings" not in str(item[f"E{pipe}"].value)
    frames = {f.name: f for f in build_package_frames(spec, QUOTES, spare_columns=4)}
    assert frames["Castings"].sp_tasks == {}
    assert frames["Valves"].sp_tasks == {"3": ["task[02-510]"]}


def test_a_support_task_and_its_lines_render_as_before(wb):
    item = wb[ITEM2]
    block = task_block(item, "02-100")
    for r in block["CREW"] + block["LINE"]:
        assert _blank(item[f"U{r}"].value)
    (crew,) = [r for r in block["CREW"] if item[f"A{r}"].value == "Pipelayer"]
    assert item[f"E{crew}"].value == 8
    assert item[f"V{crew}"].value == f"=D{crew}*E{crew}*'Rates'!$B$2*G{crew}"
    (rock,) = [r for r in block["LINE"] if item[f"A{r}"].value == "Bedding rock"]
    assert item[f"V{rock}"].value == f"=B{rock}*G{rock}"
    assert item[f"I{rock}"].value == "PLUG"
    assert item.column_dimensions["U"].hidden and item.column_dimensions["V"].hidden


# ---------------------------------------------------------------- summary, checks

def test_the_summary_has_a_row_per_package(wb):
    ws = wb["Sub Spreads"]
    hauling, valves = _row(ws, "A", "Hauling"), _row(ws, "A", "Valves")
    assert valves == hauling + 1
    assert [ws[f"{c}{hauling}"].value for c in "BC"] == ["subcontract", "sub_out"]
    assert ws[f"G{hauling}"].value == "='Sub Hauling'!$F$3"
    assert ws[f"I{hauling}"].value == f"=H{hauling}-D{hauling}"
    assert ws[f"J{hauling}"].value == "LUMP SUM"
    assert _blank(ws[f"J{valves}"].value)
    t = _row(wb["Sub Hauling"], "A", "TOTAL")
    hdr, _ = _heads(wb["Sub Hauling"])
    exts, heads = f"'Sub Hauling'!$K${t}:$W${t}", f"'Sub Hauling'!$K${hdr}:$W${hdr}"
    assert ws[f"E{hauling}"].value == f'=IFERROR(INDEX({heads},1,MATCH(MIN({exts}),{exts},0)-1),"")'
    assert ws[f"F{hauling}"].value == f'=IF(E{hauling}="","",MIN({exts}))'


def test_check_10_has_a_row_per_package(wb):
    ws = wb["Checks"]
    rows = [r for r in range(1, ws.max_row + 1) if ws[f"A{r}"].value == 10]
    assert [ws[f"C{r}"].value for r in rows] == ["Hauling", "Valves"]
    r = rows[0]
    pkg = wb["Sub Hauling"]
    line = _line_row(pkg, "Haul, import + export")
    hdr, _ = _heads(pkg)
    ok = f"'Sub Hauling'!$AB${hdr}:$AB${line}"
    assert ws[f"D{r}"].value == f"=COUNT({ok})"
    assert ws[f"E{r}"].value == f"=SUM({ok})"
    assert pkg[f"AB{line}"].value == f'=IF(ISNA($Z{line}),0,IF(ISNUMBER($F{line}),1,0))'


def test_every_package_line_is_tagged_and_every_item_row_has_a_switch(wb):
    for title in ("Sub Hauling", "Sub Valves"):
        ws = wb[title]
        assert ws["Q1"].value == "PKG"
    assert rows_of_kind(wb[ITEM3], "LINE")


def test_carry_with_no_figure_is_loud_not_excluded(wb):
    """Source CARRY and a blank CARRY cell is #N/A (check 10 fails), never `excl`
    (which would quietly carry the ARR plug)."""
    ws = wb["Sub Hauling"]
    r = _line_row(ws, "Haul, import + export")
    state = ws[f"Z{r}"].value
    carry_branch = state[state.index('IF($Y{0}="CARRY"'.format(r)):state.index("IF(ISNUMBER(INDEX")]
    assert carry_branch == f'IF($Y{r}="CARRY",IF(ISNUMBER($W{r}),"price",NA()),'
    assert "excl" not in carry_branch


def _two_line_hauling():
    """Hauling with a second package line on item 2 task 02 (Dump fees, 1 LS at 500)."""
    spec = packaged_spec()
    (item2,) = [i for i in spec["items"] if i["item_no"] == "2"]
    (block,) = [b for b in item2["labor"] if b["slot"] == "02"]
    block["lines"].append({"cost_type": "subcontract", "name": "Dump fees", "qty": 1, "uom": "LS",
                           "unit_price": 500, "package": "Hauling"})
    return spec


def test_a_lump_sum_note_names_the_quotes_total_not_the_lines_share():
    wb = B.generate(_two_line_hauling(), fieldbook_pack(), quotes=copy.deepcopy(QUOTES))
    ws = wb["Sub Hauling"]
    haul, dump = _line_row(ws, "Haul, import + export"), _line_row(ws, "Dump fees")
    # 8,740 shared over plug extensions 8,280 and 500: neither share is 8,740
    assert ws[f"M{haul}"].value * 460 != pytest.approx(8740)
    for r in (haul, dump):
        assert ws[f"M{r}"].comment.text == "prorated from $8,740.00 lump sum"


def test_a_carry_figure_is_the_lines_dollars_per_unit():
    spec = packaged_spec()
    spec["work_packages"][0]["carry"] = {"2": 17.5}
    ws = B.generate(spec, fieldbook_pack(), quotes=copy.deepcopy(QUOTES))["Sub Hauling"]
    assert ws[f"W{_line_row(ws, 'Haul, import + export')}"].value == 17.5

    spec = _two_line_hauling()
    spec["work_packages"][0]["carry"] = {"2/Haul, import + export": 17.5, "2/Dump fees": 400}
    ws = B.generate(spec, fieldbook_pack(), quotes=copy.deepcopy(QUOTES))["Sub Hauling"]
    assert ws[f"W{_line_row(ws, 'Haul, import + export')}"].value == 17.5
    assert ws[f"W{_line_row(ws, 'Dump fees')}"].value == 400


def test_inserted_vendor_can_be_typed_and_is_in_summary_ranges(wb):
    ws = wb["Sub Hauling"]
    hdr, cols = _heads(ws)
    for address in ("F3", f"H{_line_row(ws, 'Haul, import + export')}"):
        (validation,) = [dv for dv in ws.data_validations.dataValidation
                         if address in dv.sqref]
        assert validation.showErrorMessage is False
    carry = cols["CARRY"]
    assert ws["J3"].value == f'=\"\"&INDEX($K${hdr}:${carry}${hdr},1,0*2+1)'
    total = _row(ws, "A", "TOTAL")
    assert any(f"$K${total}:${carry}${total}" in str(cell.value)
               for row in wb["Sub Spreads"].iter_rows() for cell in row)


def test_quote_text_starting_with_equals_remains_text_after_saving(tmp_path):
    from openpyxl import load_workbook
    spec = packaged_spec()
    quotes = copy.deepcopy(QUOTES)
    quotes[0].vendor_name = "=1+1"
    quotes[0].quote_ref = "=2+2"
    quotes[0].conditions["Payment"] = "=3+3"
    spec["work_packages"][0]["selected"] = "=1+1"
    built = B.generate(spec, fieldbook_pack(), quotes=quotes)
    path = tmp_path / "quotes.xlsx"
    built.save(path)
    sheet = load_workbook(path)["Sub Hauling"]
    for value in ("=1+1", "=2+2", "=3+3"):
        cells = [cell for row in sheet.iter_rows() for cell in row if cell.value == value]
        assert cells and all(cell.data_type == "s" for cell in cells)
