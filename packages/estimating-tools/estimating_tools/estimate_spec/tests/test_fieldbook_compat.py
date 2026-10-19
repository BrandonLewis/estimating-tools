"""The verbs that read a pack's seed -- scaffold, apply-takeoff, compact and
bm-sync's planner -- all work against the fieldbook pack and its new fields."""
from __future__ import annotations

import copy

import pytest
import yaml

from estimating_tools.estimate_spec.apply import apply_takeoff
from estimating_tools.estimate_spec.build import generate
from estimating_tools.estimate_spec.compact import parse_compact, render_compact
from estimating_tools.estimate_spec.scaffold import scaffold_spec, validate_spec
from estimating_tools.estimate_spec.sync import plan_sync
from estimating_tools.estimate_spec.tests._fieldbook import (
    FIELDBOOK_DIR, example_spec, fieldbook_pack, rows_of_kind)
from estimating_tools.estimate_spec.tests.test_apply import sample_bridge_takeoff
from estimating_tools.estimate_spec.tests.test_scaffold import (
    sample_bridge_manifest, sample_bridge_structures)


def _scaffolded(pack):
    return scaffold_spec(sample_bridge_manifest(), sample_bridge_structures(),
                         generated_at="2026-08-02T00:00:00+00:00", pack=pack)


def test_a_scaffolded_fieldbook_spec_validates_and_builds():
    pack = fieldbook_pack()
    spec = _scaffolded(pack).spec
    validate_spec(spec, pack=pack)
    assert {b["cost_code"] for item in spec["items"] for b in item["labor"]} == {"01-100"}
    kinds = [it["takeoff"]["kind"] for it in spec["items"] if it["recipe"] == "line_item"]
    assert kinds == ["measured"] * 3                    # the seed's placeholder takeoff
    wb = generate(spec, pack)
    assert wb.sheetnames[:6] == ["Bid Form", "Cost Summary", "Markup", "Checks", "Rates", "Codes"]


def test_a_scaffold_without_a_mobilization_item_injects_a_setup_sheet():
    pack = fieldbook_pack()
    manifest = sample_bridge_manifest()
    manifest.bid_frame = [e for e in manifest.bid_frame if e.unit != "LS"]
    result = scaffold_spec(manifest, sample_bridge_structures(), pack=pack)
    assert result.injected_rate_host
    validate_spec(result.spec, pack=pack)
    assert "Setup" in generate(result.spec, pack).sheetnames


def test_apply_takeoff_writes_measured_geometry_that_fieldbook_lays_out():
    pack = fieldbook_pack()
    result = apply_takeoff(_scaffolded(pack).spec, sample_bridge_takeoff(), pack=pack,
                           today="2026-08-02")
    assert [(m["item_no"], m["new_qty"]) for m in result.matched] == [
        ("2217-38", 29.63), ("2217-39", 104.44), ("2217-40", 75.0)]
    validate_spec(result.spec, pack=pack)
    wb = generate(result.spec, pack)
    ws = wb["38c0366 2217-39"]
    assert [ws[f"A{r}"].value for r in rows_of_kind(ws, "TAKEOFF")] == [
        "S-ABUT-1", "S-ABUT-2", "S-WING-LONG x2", "S-WING-SHORT x2"]


def test_compact_round_trips_the_fieldbook_example_unchanged():
    pack = fieldbook_pack()
    spec = example_spec()
    assert parse_compact(render_compact(spec), base_spec=copy.deepcopy(spec), pack=pack) == spec


def test_bm_sync_plans_a_new_equipment_rate_by_its_full_name():
    spec = example_spec()
    edited = copy.deepcopy(spec)
    edited["rate_table"].append(
        {"trade": "Backhoe, 416 class", "rate": 95.0, "kind": "equipment", "includes_operator": False})
    plan = plan_sync(spec, edited)
    assert [op.trade for op in plan.rate_ops] == ["Backhoe, 416 class"]
    assert plan.item_ops == [] and plan.crew_ops == []


# ---------------------------------------------------------------- strict: true
#
# fieldbook declares strict: true in its manifest (Task 15), so its own
# refusals -- item-level cost_details, spreads, both `false`-schema'd by its
# vocabulary -- now stop `generate`/`build` itself, not only the
# scaffold.validate_spec callers that happened to check first.

def test_build_refuses_a_priced_item_level_cost_details_block():
    pack = fieldbook_pack()
    assert pack.strict is True
    spec = example_spec()
    spec["items"][2]["cost_details"] = [
        {"name": "Material", "lines": [{"name": "x", "qty": 1, "unit_price": 1}]}]
    with pytest.raises(ValueError, match="schema violation"):
        generate(spec, pack)


def test_build_refuses_spreads():
    pack = fieldbook_pack()
    spec = example_spec()
    spec["items"][2]["spreads"] = {"haul": {"import_cy": 1}}
    with pytest.raises(ValueError, match="schema violation"):
        generate(spec, pack)


def test_cli_build_refuses_the_same_specs_with_a_nonzero_exit(tmp_path, capsys):
    from estimating_tools.estimate_spec.cli import main

    for key, value in (
        ("cost_details", [{"name": "Material",
                           "lines": [{"name": "x", "qty": 1, "unit_price": 1}]}]),
        ("spreads", {"haul": {"import_cy": 1}}),
    ):
        spec = example_spec()
        spec["items"][2][key] = value
        bad = tmp_path / f"bad-{key}.estimate.yaml"
        bad.write_text(yaml.dump(spec), encoding="utf-8")
        rc = main(["build", str(bad), "-o", str(tmp_path / f"{key}.xlsx"),
                   "--format", str(FIELDBOOK_DIR)])
        assert rc != 0
        assert "schema violation" in capsys.readouterr().err
        assert not (tmp_path / f"{key}.xlsx").exists()


def test_a_scaffolded_fieldbook_rate_table_carries_the_seeds_equipment_kinds():
    pack = fieldbook_pack()
    spec = _scaffolded(pack).spec
    assert spec["rate_table"] == [
        {"trade": "Superintendent", "rate": 0},
        {"trade": "Equipment Operator", "rate": 0},
        {"trade": "Laborer", "rate": 0},
        {"trade": "Pipelayer", "rate": 0},
        {"trade": "Excavator, 320 class", "rate": 0, "kind": "equipment"},
        {"trade": "Loader, 950 class", "rate": 0, "kind": "equipment"},
        {"trade": "Plate compactor", "rate": 0, "kind": "equipment"}]
    rates = generate(spec, pack)["Rates"]
    shown = {rates[f"A{r}"].value: rates[f"B{r}"].value for r in range(1, rates.max_row + 1)}
    assert shown["Laborer"] == "Labor"
    assert shown["Excavator, 320 class"] == shown["Plate compactor"] == "Equipment"


def test_a_compact_new_row_stubs_an_item_the_fieldbook_schema_accepts():
    """compact's new-row stub carries the pack's skeleton task when the pack
    requires one (fieldbook: every item needs a task with a cost code); an
    empty `labor: []` would fail fieldbook's own schema."""
    pack = fieldbook_pack()
    spec = example_spec()
    new_row = "| - | - | - | - | 4 | Traffic control | - | - | - | LS | - |"
    parsed = parse_compact(render_compact(spec) + "\n" + new_row,
                           base_spec=copy.deepcopy(spec), pack=pack)
    stub = parsed["items"][-1]
    assert (stub["item_no"], stub["recipe"], stub["sheet"]) == ("4", "lump_sum", "4")
    assert stub["labor"] == [{"slot": "01", "description": "(fill in)",
                              "crew": [{"trade": "Laborer", "count": 2, "days": 0}],
                              "cost_code": "01-100"}]
    validate_spec(parsed, pack=pack)
    assert "4" in generate(parsed, pack).sheetnames


def test_a_scaffolded_fieldbook_spec_carries_the_seeds_code_list():
    pack = fieldbook_pack()
    spec = _scaffolded(pack).spec
    assert [c["code"] for c in spec["cost_codes"]] == [
        "01-100", "02-100", "02-200", "02-300", "02-510"]


# ---------------------------------------------------------------- the bundled default

def test_with_no_pack_named_build_lays_out_fieldbook(tmp_path, monkeypatch):
    """What an estimator gets with no org pack: no --format, no profile ->
    the bundled pack, which is fieldbook (phase 3b)."""
    import openpyxl

    from estimating_tools.estimate_spec.cli import main
    from estimating_tools.estimate_spec.tests._fieldbook import EXAMPLE

    monkeypatch.delenv("ESTIMATING_PROFILES", raising=False)
    out = tmp_path / "bid.xlsx"
    assert main(["build", str(EXAMPLE), "-o", str(out)]) == 0
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[:6] == ["Bid Form", "Cost Summary", "Markup", "Checks", "Rates", "Codes"]
    assert {p.name: p.value for p in wb.custom_doc_props.props}["estimate_format"] == "fieldbook 2"
