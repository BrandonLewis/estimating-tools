"""Merging what the workbook says into the spec, one difference at a time."""
from __future__ import annotations

import copy

import pytest

from estimating_tools.estimate_spec.readback import sheets as S
from estimating_tools.estimate_spec.readback.merge import (
    Findings, _merge_takeoff, _merge_tasks, align, find_needs_rate, merge_item, merge_markup,
    merge_rates)
from estimating_tools.estimate_spec.tests._fieldbook import (
    build, copy_rows, example_spec, fieldbook_pack, rows_of_kind, simulate_delete_row,
    simulate_insert_copied_row, task_block)


def test_align_pairs_equal_rows_and_edits_in_place():
    # unchanged, retyped in place, deleted, inserted, replaced by more rows
    assert align(["a", "b", "c"], ["a", "b", "c"]) == [(0, 0), (1, 1), (2, 2)]
    assert align(["a", "b", "c"], ["a", "x", "c"]) == [(0, 0), (1, 1), (2, 2)]
    assert align(["a", "b", "c"], ["a", "c"]) == [(0, 0), (1, None), (2, 1)]
    assert align(["a", "b"], ["a", "b", "z"]) == [(0, 0), (1, 1), (None, 2)]
    assert align(["a"], ["x", "y"]) == [(0, None), (None, 0), (None, 1)]


def _rates(wb, spec=None):
    spec = copy.deepcopy(spec or example_spec())
    f = Findings()
    merge_rates(spec, S.read_rates(wb["Rates"]), f)
    return spec, f


def _markup(wb, spec=None):
    spec = copy.deepcopy(spec or example_spec())
    f = Findings()
    merge_markup(spec, S.read_markup(wb["Markup"]), f)
    return spec, f


def test_unedited_rates_and_markup_merge_to_themselves():
    wb = build()
    assert _rates(wb) == (example_spec(), Findings())
    assert _markup(wb) == (example_spec(), Findings())


def test_rate_edits_new_rates_and_deleted_rates():
    wb = build()
    ws = wb["Rates"]
    ws["B2"] = 10                                   # hours per crew-day
    ws["G5"] = 98                                   # Superintendent's flat rate
    ws["C6"] = 46                                   # Pipelayer's base: 46 x 1.30 + 16.80
    ws["F10"] = "Yes"                               # the loader now includes its operator
    simulate_delete_row(ws, 11)                     # Plate compactor
    simulate_insert_copied_row(ws, 11, 10)          # two rates pasted above the marker
    simulate_insert_copied_row(ws, 12, 8)
    ws["A11"], ws["F11"], ws["G11"] = "Water truck", "No", 95
    ws["A12"], ws["C12"], ws["D12"], ws["E12"] = "Truck Driver", 40, 0.25, 12
    spec, f = _rates(wb)
    assert f.changes == [
        "rates · shift_hours 8 → 10",
        "rates · Superintendent: rate 95 → 98",
        "rates · Pipelayer: base 44 → 46, rate 74 → 76.6",
        "rates · Loader, 950 class: includes_operator False → True",
        "rates · − Plate compactor (equipment, 35/hr) (removed)",
        "rates · + Water truck (equipment, 95/hr)",
        "rates · + Truck Driver (labor, 62/hr)"]
    assert f.problems == []
    assert spec["shift_hours"] == 10
    assert spec["rate_table"][1] == {
        "trade": "Pipelayer", "rate": 76.6, "base": 46, "burden_pct": 0.3, "fringe": 16.8}
    assert spec["rate_table"][-2:] == [
        {"trade": "Water truck", "rate": 95, "kind": "equipment", "includes_operator": False},
        {"trade": "Truck Driver", "rate": 62.0, "base": 40, "burden_pct": 0.25, "fringe": 12}]


def test_a_rate_within_a_cent_of_its_breakdown_is_not_rereported_unedited():
    # review fix (Tasks 5-7): rounding differences baked into a historical
    # spec must not look like a change on an otherwise unedited book --
    # only re-derive/report the rate when a breakdown cell actually changed.
    spec = example_spec()
    spec["rate_table"][1]["rate"] = 74.005              # a penny off its 44 x 1.3 + 16.8 breakdown
    spec, f = _rates(build(), spec)
    assert f.changes == [] and f.problems == []
    assert spec["rate_table"][1]["rate"] == 74.005       # left exactly as the spec had it


def test_a_flat_rate_typed_on_a_breakdown_row_is_a_notice_not_silently_ignored():
    wb = build()
    ws = wb["Rates"]
    ws["G6"] = 999                                      # Pipelayer already has a breakdown
    _, f = _rates(wb)
    assert f.notices == [
        "Rates row 6 (rate Pipelayer): flat rate (G) is ignored on a breakdown row"]
    assert f.changes == [] and f.problems == []


def test_includes_operator_yes_on_a_labor_row_is_a_notice_not_silently_ignored():
    wb = build()
    ws = wb["Rates"]
    ws["F6"] = "Yes"                                     # Pipelayer is labor, not equipment
    _, f = _rates(wb)
    assert f.notices == [
        "Rates row 6 (rate Pipelayer): includes operator (F) is ignored on a labor rate"]
    assert f.changes == [] and f.problems == []


def test_a_rate_row_that_cannot_be_read_is_a_problem():
    wb = build()
    ws = wb["Rates"]
    ws["B5"] = None                                 # Superintendent: no Kind
    ws["C9"] = 100                                  # an equipment row given a base wage
    ws["G7"] = "=G5*1.1"                            # a rate that refers to another cell
    ws["A8"] = None                                 # Laborer's name cleared: not the table's end
    spec, f = _rates(wb)
    assert f.problems == [
        "Rates row 5 (rate Superintendent): kind (blank) is not one of Labor, Equipment",
        "Rates row 7 (rate Equipment Operator): rate holds the formula '=G5*1.1' -- type a number",
        "Rates row 8 (rate (no name)): trade is blank",
        "Rates row 9 (rate Excavator, 320 class): an equipment rate is a flat rate "
        "-- clear Base, Burden % and Fringe"]
    assert len(spec["rate_table"]) == 7                 # nothing read as deleted


def test_markup_bond_and_unbalance_edits():
    wb = build()
    ws = wb["Markup"]
    ws["B5"] = 0.10                                 # profit on labor
    ws["E4"] = 0.005                                # insurance on subcontract
    ws["B7"] = 0.015                                # bond
    ws["C12"], ws["C14"] = 500, 500                 # item 1 down, item 3 up: still nets to 0
    spec, f = _markup(wb)
    assert f.changes == [
        "markup · insurance subcontract 0 → 0.005",
        "markup · profit labor 0.12 → 0.1",
        "markup · bond_pct 0.012 → 0.015",
        "item 1 · unbalance 1,000 → 500",
        "item 3 · unbalance 0 → 500"]
    assert spec["markup"]["profit"]["labor"] == 0.10
    assert spec["markup"]["insurance"]["subcontract"] == 0.005
    assert spec["markup"]["bond_pct"] == 0.015
    assert [it.get("unbalance") for it in spec["items"]] == [500, -1000, 500]


def test_a_markup_cell_that_is_not_a_number_is_a_problem():
    wb = build()
    wb["Markup"]["C3"] = "ten percent"
    _, f = _markup(wb)
    assert f.problems == [
        "Markup row 3 (markup overhead): equipment is 'ten percent', not a number"]


def test_a_blank_markup_bond_or_unbalance_cell_means_zero():
    wb = build()
    ws = wb["Markup"]
    ws["C3"] = None                                 # overhead on equipment
    ws["B7"] = None                                 # bond
    ws["C13"] = None                                # item 2's unbalance (-1,000)
    ws["C14"] = None                                # item 3's unbalance, already 0
    spec, f = _markup(wb)
    assert f.problems == []
    assert f.changes == [
        "markup · overhead equipment 0.1 → 0",
        "markup · bond_pct 0.012 → 0",
        "item 2 · unbalance -1,000 → 0"]
    assert spec["markup"]["overhead"]["equipment"] == 0 and spec["markup"]["bond_pct"] == 0
    assert [it.get("unbalance") for it in spec["items"]] == [1000, 0, None]


def test_a_markup_row_outside_overhead_insurance_profit_is_a_blocking_problem():
    # review fix (Tasks 5-7): Excel prices ANY row inside the grid -- a row
    # the estimator inserts (e.g. "Contingency") must not be silently
    # dropped; it has to block until it is folded in or removed.
    wb = build()
    ws = wb["Markup"]
    simulate_insert_copied_row(ws, 5, 4)                       # a row pasted above Profit
    ws["A5"] = "Contingency"
    _, f = _markup(wb)
    assert f.problems == [
        "Markup row 5: 'Contingency' is not a fieldbook markup row -- fold it into "
        "Overhead/Profit or delete it"]


def test_a_duplicate_markup_row_is_a_blocking_problem_naming_both_rows():
    # Copilot review (PR #355): Excel sums BOTH Profit rows; collapsing them
    # (last one wins) silently dropped the other row's markup.
    wb = build()
    ws = wb["Markup"]
    simulate_insert_copied_row(ws, 6, 5)                       # Profit pasted below itself
    spec, f = _markup(wb)
    assert f.problems == [
        "Markup rows 5 and 6 are both 'Profit' -- combine them into one row"]
    assert spec["markup"] == example_spec()["markup"]


def test_a_relabelled_bond_row_is_a_problem_not_a_cleared_bond():
    # Copilot review (PR #355): Excel still prices B7 as the bond; reading the
    # missing label as "blank" set bond_pct to 0.
    wb = build()
    wb["Markup"]["A7"] = "Bond %"
    spec, f = _markup(wb)
    assert f.problems == [
        "Markup: no 'Bond (% of bid)' row -- the Bond row's label was changed; "
        "restore 'Bond (% of bid)'"]
    assert spec["markup"]["bond_pct"] == 0.012 and f.changes == []


def test_a_relabelled_shift_hours_row_is_a_problem_naming_the_label():
    wb = build()
    wb["Rates"]["A2"] = "Hours/day"
    spec, f = _rates(wb)
    assert f.problems == [
        "Rates: no 'Hours per crew-day' row -- its label was changed; "
        "restore 'Hours per crew-day'"]
    assert spec.get("shift_hours") == example_spec().get("shift_hours") and f.changes == []


# ---------------------------------------------------------------- items

ITEM2 = "2 - 8IN MAIN"


def _item(wb, index=1, spec=None):
    """Merge item `index`'s sheet into a copy of the spec: (item, findings)."""
    spec = copy.deepcopy(spec or example_spec())
    item = spec["items"][index]
    f = Findings()
    sheet = S.read_item_sheet(wb[item["sheet"]],
                              takeoff_kind=(item.get("takeoff") or {}).get("kind"))
    merge_item(item, sheet, f, rate_names=[r["trade"] for r in spec["rate_table"]],
               trades=fieldbook_pack().vocabulary["trades"])
    return item, f


def _row(ws, kind, name):
    (r,) = [r for r in rows_of_kind(ws, kind) if ws[f"A{r}"].value == name]
    return r


def test_an_unedited_item_merges_to_itself():
    wb = build()
    for index in range(3):
        assert _item(wb, index) == (example_spec()["items"][index], Findings())


def test_every_input_on_an_item_sheet_reads_back():
    wb = build()
    ws = wb[ITEM2]
    ws["G1"] = 650                                            # bid qty
    ws["A3"] = "Trench to 5 ft max."                          # the item's scope note
    ws["B7"], ws["C7"] = 330, "C-4A"                          # a takeoff segment
    (head,) = task_block(ws, "02-200")["TASK"]
    ws[f"B{head}"] = "Backfill, compaction & testing"          # task description
    ws[f"A{head + 1}"] = "95% relative compaction everywhere."  # task note
    crew = _row(ws, "CREW", "Pipelayer")
    ws[f"D{crew}"], ws[f"E{crew}"] = 3, 9
    line = _row(ws, "LINE", "Bedding rock")
    ws[f"B{line}"], ws[f"C{line}"], ws[f"G{line}"] = 220, "TN", 33.5
    ws[f"H{line}"], ws[f"I{line}"], ws[f"J{line}"] = "Other", "QUOTE", "Quarry D"
    item, f = _item(wb)
    assert f.changes == [
        "item 2 · qty 640 → 650",
        "item 2 · note 'Trench to 6 ft max; native backfill above the pipe zone.' → "
        "'Trench to 5 ft max.'",
        "item 2 · takeoff · STA 13+20 to 16+40: length_lf 320 → 330, ref C-4 → C-4A",
        "item 2 · 02-100 · Pipelayer: count 2 → 3, days 8 → 9",
        "item 2 · 02-100 · Bedding rock: cost_type material → other, qty 210 → 220, "
        "uom TON → TN, unit_price 31 → 33.5, source plug → quote, vendor (blank) → Quarry D",
        "item 2 · 02-200 · description Backfill & compaction → Backfill, compaction & testing",
        "item 2 · 02-200 · note 'Native backfill above the pipe zone; 95% relative compaction "
        "under pavement.' → '95% relative compaction everywhere.'"]
    assert f.problems == []
    assert item["qty"] == 650 and item["notes"] == ["Trench to 5 ft max."]
    assert item["takeoff"]["segments"][1] == {
        "label": "STA 13+20 to 16+40", "length_lf": 330, "ref": "C-4A"}
    assert item["labor"][0]["crew"][1] == {"trade": "Pipelayer", "count": 3, "days": 9}
    assert item["labor"][0]["lines"][2] == {
        "cost_type": "other", "name": "Bedding rock", "qty": 220, "uom": "TN",
        "unit_price": 33.5, "source": "quote", "vendor": "Quarry D"}
    assert item["labor"][1]["description"] == "Backfill, compaction & testing"
    assert item["labor"][1]["notes"] == ["95% relative compaction everywhere."]
    assert item["labor"][1]["slot"] == "02"                    # not shown, carried over


def test_rows_added_and_deleted_in_excel():
    wb = build()
    ws = wb[ITEM2]
    simulate_insert_copied_row(ws, 8, 7)                      # a takeoff segment
    ws["A8"], ws["B8"], ws["C8"] = "STA 16+40 to 17+00", 60, "C-5"
    block = task_block(ws, "02-100")
    simulate_insert_copied_row(ws, block["END"][0], block["CREW"][3])     # Laborer, copied
    simulate_delete_row(ws, _row(ws, "LINE", "Bedding rock"))
    block = task_block(ws, "02-200")
    simulate_insert_copied_row(ws, block["END"][0], block["LINE"][0])     # the haul line, copied
    end = block["END"][0]
    ws[f"A{end}"], ws[f"B{end}"], ws[f"C{end}"], ws[f"G{end}"] = "Saw cutting", 120, "LF", 3.25
    item, f = _item(wb)
    assert f.changes == [
        "item 2 · takeoff · + takeoff row label STA 16+40 to 17+00, length_lf 60, ref C-5",
        "item 2 · 02-100 · + crew Laborer × 2, 8 days",
        "item 2 · 02-100 · − line Bedding rock 210 TON @ 31 material (removed)",
        "item 2 · 02-200 · + line Saw cutting 120 LF @ 3.25 subcontract"]
    assert item["labor"][0]["crew"][-1] == {"trade": "Laborer", "count": 2, "days": 8}
    assert [ln["name"] for ln in item["labor"][0]["lines"]] == [
        "Trench box (rental)", "8-inch ductile iron pipe"]
    assert item["labor"][1]["lines"][-1] == {
        "cost_type": "subcontract", "name": "Saw cutting", "qty": 120, "uom": "LF",
        "unit_price": 3.25, "source": "plug"}


def test_a_pasted_block_with_a_new_code_is_a_new_task():
    wb = build()
    ws = wb[ITEM2]
    block = task_block(ws, "02-200")
    (first,), (sub,) = block["TASK"], block["SUBTOTAL"]
    copy_rows(ws, first, sub, sub + 1)
    ws[f"A{sub + 1}"], ws[f"B{sub + 1}"] = "02-300", "Pavement restoration"
    item, f = _item(wb)
    assert f.changes[0] == ("item 2 · + new task 02-300 Pavement restoration "
                            "(4 crew rows, 1 cost line, 1 note)")
    assert len(f.changes) == 7                                # + 1 note, 4 crew, 1 line
    new = item["labor"][2]
    assert (new["slot"], new["cost_code"], new["description"]) == (
        "03", "02-300", "Pavement restoration")
    assert new["crew"] == item["labor"][1]["crew"]
    assert new["lines"] == item["labor"][1]["lines"]


def test_a_deleted_block_removes_its_task():
    wb = build()
    ws = wb[ITEM2]
    block = task_block(ws, "02-200")
    for _ in range(block["SUBTOTAL"][0] - block["TASK"][0] + 1):
        simulate_delete_row(ws, block["TASK"][0])
    item, f = _item(wb)
    assert f.changes == ["item 2 · − task 02-200 Backfill & compaction (removed)"]
    assert [t["cost_code"] for t in item["labor"]] == ["02-100"]


def test_a_pasted_block_that_keeps_an_existing_code_is_a_conflict():
    wb = build()
    ws = wb[ITEM2]
    block = task_block(ws, "02-200")
    (first,), (sub,) = block["TASK"], block["SUBTOTAL"]
    copy_rows(ws, first, sub, sub + 1)                        # the header still says 02-200
    _, f = _item(wb)
    assert f.problems == [
        f"2 - 8IN MAIN rows {first} and {sub + 1} (item 2): two task blocks show cost code "
        "02-200 -- give the pasted block its own code, or delete one"]


def _rate_names_and_trades():
    spec = example_spec()
    return [r["trade"] for r in spec["rate_table"]], fieldbook_pack().vocabulary["trades"]


def test_a_pure_task_reorder_is_reported_and_applied():
    # review fix (Tasks 5-7): a whole-block reorder, with nothing else
    # edited, must not be silently discarded by `_store_list` (it used to
    # duplicate the moved block) nor silently applied with no report.
    wb = build()
    sheet = S.read_item_sheet(wb[ITEM2], takeoff_kind="linear")
    sheet.tasks = [sheet.tasks[1], sheet.tasks[0]]            # 02-200 now before 02-100
    item = copy.deepcopy(example_spec()["items"][1])
    f = Findings()
    rate_names, trades = _rate_names_and_trades()
    _merge_tasks(item, sheet, "item 2", f, rate_names=rate_names, trades=trades)
    assert f.changes == ["item 2 · tasks reordered: 02-200 now before 02-100"]
    assert [t["cost_code"] for t in item["labor"]] == ["02-200", "02-100"]
    assert len(item["labor"]) == 2                            # nothing duplicated


def test_a_task_reorder_plus_an_edit_both_report():
    wb = build()
    sheet = S.read_item_sheet(wb[ITEM2], takeoff_kind="linear")
    sheet.tasks = [sheet.tasks[1], sheet.tasks[0]]
    sheet.tasks[1].description = "Trench excavation & re-laying"
    item = copy.deepcopy(example_spec()["items"][1])
    f = Findings()
    rate_names, trades = _rate_names_and_trades()
    _merge_tasks(item, sheet, "item 2", f, rate_names=rate_names, trades=trades)
    assert f.changes == [
        "item 2 · 02-100 · description Trench excavation & pipe laying → "
        "Trench excavation & re-laying",
        "item 2 · tasks reordered: 02-200 now before 02-100"]
    assert [t["cost_code"] for t in item["labor"]] == ["02-200", "02-100"]


def test_blank_or_unreadable_inputs_in_a_new_row_are_problems():
    wb = build()
    ws = wb[ITEM2]
    block = task_block(ws, "02-200")
    crew = block["END"][0]
    simulate_insert_copied_row(ws, crew, block["CREW"][1])
    ws[f"A{crew}"], ws[f"D{crew}"], ws[f"E{crew}"] = "Truck Driver", None, "=E20*2"
    line = crew + 1
    simulate_insert_copied_row(ws, line, block["LINE"][0])
    ws[f"A{line}"], ws[f"G{line}"], ws[f"H{line}"] = "Asphalt", None, "Asphalt"
    _, f = _item(wb)
    assert f.problems == [
        f"2 - 8IN MAIN row {crew} (item 2 · 02-200): count is blank",
        f"2 - 8IN MAIN row {crew} (item 2 · 02-200): days holds the formula '=E20*2' "
        "-- type a number",
        f"2 - 8IN MAIN row {line} (item 2 · 02-200): cost_type Asphalt is not one of "
        "Material, Subcontract, Equipment, Other",
        f"2 - 8IN MAIN row {line} (item 2 · 02-200): unit_price is blank"]


def test_a_resource_typed_in_another_case_reads_as_the_rate_rows_name():
    wb = build()
    ws = wb[ITEM2]
    crew = _row(ws, "CREW", "Pipelayer")
    ws[f"A{crew}"] = "pipelayer"                              # Excel's VLOOKUP still finds the rate
    item, f = _item(wb)
    assert f.changes == [] and f.problems == []
    assert f.notices == [f"2 - 8IN MAIN row {crew} (item 2 · 02-100): resource 'pipelayer' "
                         "read as 'Pipelayer' (same name, different case)"]
    assert item == example_spec()["items"][1]


def test_a_resource_the_pack_does_not_know_is_a_problem():
    wb = build()
    ws = wb[ITEM2]
    crew = _row(ws, "CREW", "Pipelayer")
    ws[f"A{crew}"] = "Asphalt Raker"
    _, f = _item(wb)
    assert f.problems == [
        f"2 - 8IN MAIN row {crew} (item 2 · 02-100): resource 'Asphalt Raker' is not one "
        "of this pack's resources (pick one from the dropdown)"]


def test_needs_rate_lists_every_crew_resource_with_no_rate_row():
    spec = example_spec()
    spec["items"][1]["labor"][0]["crew"].append({"trade": "Truck Driver", "count": 1, "days": 8})
    spec["items"][2]["labor"][0]["crew"].append({"trade": "Truck Driver", "count": 1, "days": 2})
    f = Findings()
    find_needs_rate(spec, f)
    assert f.needs_rate == {"Truck Driver": ["item 2 · 02-100", "item 3 · 02-510"]}


def test_needs_rate_matches_case_insensitively():
    # review fix (Tasks 5-7): consistent with decision 8 -- Excel's VLOOKUP
    # ignores case, so a resource spelled differently than its rate row
    # already has a rate and must not show up as needing one.
    spec = example_spec()
    spec["items"][1]["labor"][0]["crew"][1]["trade"] = "PIPELAYER"   # rate row says "Pipelayer"
    f = Findings()
    find_needs_rate(spec, f)
    assert f.needs_rate == {}


# ------------------------------------------------------- placeholder formulas

def _placeholder_merge(edited_days):
    """Pipelayer's `days` is a placeholder the builder resolves to the real
    formula `=$B$8` (stood in for a baseline the reader would have built);
    the edited workbook's cell holds `edited_days` instead."""
    wb = build()
    ws = wb[ITEM2]
    crew = _row(ws, "CREW", "Pipelayer")
    ws[f"E{crew}"] = edited_days
    sheet = S.read_item_sheet(ws, takeoff_kind="linear")
    baseline_sheet = S.read_item_sheet(build()[ITEM2], takeoff_kind="linear")
    baseline_sheet.tasks[0].crew[1].cells["E"] = "=$B$8"      # the placeholder, resolved
    spec = example_spec()
    item = spec["items"][1]
    item["labor"][0]["crew"][1]["days"] = "={ref:quantity.total}"
    f = Findings()
    merge_item(item, sheet, f, rate_names=[r["trade"] for r in spec["rate_table"]],
               trades=fieldbook_pack().vocabulary["trades"], baseline=baseline_sheet)
    return item, f


def test_the_placeholders_own_resolved_formula_is_unchanged():
    item, f = _placeholder_merge("=$B$8")                     # matches the baseline exactly
    assert f.changes == [] and f.problems == []
    assert item["labor"][0]["crew"][1]["days"] == "={ref:quantity.total}"


def test_a_pure_arithmetic_formula_over_a_placeholder_is_reported_as_a_change():
    # review fix (Tasks 1-4): any formula used to count as "unchanged" merely
    # because the spec's old value looked like a placeholder.
    item, f = _placeholder_merge("=700")
    assert f.changes == ["item 2 · 02-100 · Pipelayer: days ={ref:quantity.total} → =700"]
    assert f.problems == []
    assert item["labor"][0]["crew"][1]["days"] == "=700"


def test_a_cell_referencing_formula_over_a_placeholder_is_a_problem():
    item, f = _placeholder_merge("=B12*2")
    assert f.changes == []
    assert f.problems == [
        f"2 - 8IN MAIN row {_row(build()[ITEM2], 'CREW', 'Pipelayer')} (item 2 · 02-100): "
        "days holds the formula '=B12*2' -- type a number"]


# -------------------------------------------------- matching tolerates coercion

def test_notes_matching_tolerates_the_builders_own_number_coercion():
    # review fix (Tasks 5-7): a note "1.50" survives the builder's own
    # text -> cell conversion as the number 1.5; matching the OLD side by
    # raw text (not the same conversion) reported a false change.
    entry = {"notes": ["1.50"]}
    rows = [S.Row(5, {"A": 1.5})]              # what the cell actually holds
    f = Findings()
    from estimating_tools.estimate_spec.readback.merge import _merge_notes
    _merge_notes(entry, "notes", rows, "item 2", f)
    assert f.changes == []
    assert entry["notes"] == ["1.50"]           # untouched


def test_task_matching_tolerates_the_builders_own_number_coercion():
    # review fix (Tasks 5-7): a cost code "31.20" becomes the cell 31.2;
    # matching by raw text read it as one task removed and a new one added.
    wb = build()
    sheet = S.read_item_sheet(wb[ITEM2], takeoff_kind="linear")
    sheet.tasks[0].code = 31.2                  # what "31.20" becomes in the cell
    item = copy.deepcopy(example_spec()["items"][1])
    item["labor"][0]["cost_code"] = "31.20"
    f = Findings()
    rate_names, trades = _rate_names_and_trades()
    _merge_tasks(item, sheet, "item 2", f, rate_names=rate_names, trades=trades)
    assert f.changes == []
    assert item["labor"][0]["cost_code"] == "31.20"   # untouched, not read as removed+new
    assert len(item["labor"]) == 2


def test_unbalance_matching_tolerates_the_builders_own_number_coercion():
    # review fix (Tasks 5-7): an item_no "2.10" becomes the cell 2.1.
    wb = build()
    spec = copy.deepcopy(example_spec())
    spec["items"][0]["item_no"] = "2.10"
    m = S.read_markup(wb["Markup"])
    m.unbalance[0].cells["A"] = 2.1              # what "2.10" would read back as
    f = Findings()
    merge_markup(spec, m, f)
    assert f.problems == [] and f.changes == []
    assert spec["items"][0]["item_no"] == "2.10"  # untouched


# ------------------------------------------------------- hidden derived field

def test_a_changed_footing_cy_recomputes_the_hidden_footing_cf():
    # review fix (Tasks 5-7): footing_cf is never shown in the workbook, so
    # it is never read -- but it must not go stale when footing_cy changes.
    takeoff = {"kind": "measured", "lines": [
        {"label": "Wingwall A", "length": 20, "height": 8, "footing_cy": 10, "footing_cf": 270,
         "wall_cy": 5, "wall_cf": 135}]}
    sheet = S.ItemSheet(sheet="3 - VALVES")
    sheet.takeoff = [S.Row(10, {"A": "Wingwall A", "E": 20, "G": 8, "H": 15, "I": 5})]
    f = Findings()
    _merge_takeoff(takeoff, sheet, "item 3", f)
    assert takeoff["lines"][0]["footing_cy"] == 15
    assert takeoff["lines"][0]["footing_cf"] == 405        # 15 x 27
    assert takeoff["lines"][0]["wall_cy"] == 5              # unchanged
    assert takeoff["lines"][0]["wall_cf"] == 135             # unchanged -- not re-derived


def test_a_measured_line_reads_unit_count_width_weight_and_adjust_back():
    takeoff = {"kind": "measured", "lines": [
        {"label": "Footing A", "unit": "CY", "count": 1, "length": 10, "width": 4, "height": 2,
         "footing_cy": 2.96, "adjust": 0.04}]}
    sheet = S.ItemSheet(sheet="3 - VALVES")
    sheet.takeoff = [S.Row(10, {"A": "Footing A", "C": "CY", "D": 1, "E": 10, "F": 5, "G": 2,
                                "H": 2.96, "L": 120, "N": 0})]
    f = Findings()
    _merge_takeoff(takeoff, sheet, "item 3", f)
    assert f.problems == []
    line = takeoff["lines"][0]
    assert (line["unit"], line["count"], line["length"], line["width"], line["height"]) == ("CY", 1, 10, 5, 2)
    assert (line["footing_cy"], line["weight_lb"], line["adjust"]) == (2.96, 120, 0)


def test_a_formula_in_a_measured_cy_cell_is_a_problem_not_a_crash():
    # Copilot review (PR #355): `=10+5` was kept as a string, then
    # round("=10+5" * 27) raised TypeError. Takeoff fields are schema numbers.
    takeoff = {"kind": "measured", "lines": [
        {"label": "Wingwall A", "length": 20, "height": 8, "footing_cy": 10, "footing_cf": 270}]}
    sheet = S.ItemSheet(sheet="3 - VALVES")
    sheet.takeoff = [S.Row(10, {"A": "Wingwall A", "E": 20, "G": 8, "H": "=10+5"})]
    f = Findings()
    _merge_takeoff(takeoff, sheet, "item 3", f)
    assert f.problems == [
        "3 - VALVES row 10 (item 3 takeoff): footing_cy holds the formula '=10+5' -- type a number"]
    assert takeoff["lines"][0]["footing_cy"] == 10 and takeoff["lines"][0]["footing_cf"] == 270


def test_a_formula_in_a_linear_takeoff_cell_is_a_problem_not_a_schema_dump():
    wb = build()
    ws = wb[ITEM2]
    (first, _) = rows_of_kind(ws, "TAKEOFF")
    ws[f"B{first}"] = "=300+20"
    item, f = _item(wb)
    assert f.problems == [
        f"2 - 8IN MAIN row {first} (item 2 takeoff): length_lf holds the formula '=300+20' "
        "-- type a number"]
    assert item["takeoff"] == example_spec()["items"][1]["takeoff"]


def test_a_blank_source_on_a_cost_line_is_a_problem_not_a_silent_plug():
    wb = build()
    ws = wb[ITEM2]
    pipe = _row(ws, "LINE", "8-inch ductile iron pipe")           # a QUOTE line
    ws[f"I{pipe}"] = None
    item, f = _item(wb)
    assert f.problems == [
        f"2 - 8IN MAIN row {pipe} (item 2 · 02-100): Source is blank -- choose QUOTE or PLUG"]
    assert item["labor"][0]["lines"][1]["source"] == "quote"      # kept, not turned into plug


def test_a_retyped_cost_code_is_a_rename_that_keeps_the_task():
    """One task gone and one new at the same place in the same item is the
    estimator retyping the code: the task keeps its slot and every field the
    workbook does not show -- not a removal plus a brand-new task."""
    spec = example_spec()
    spec["items"][1]["labor"][1]["crew_factor"] = 1.15     # not shown in the workbook
    wb = build(spec)
    ws = wb[ITEM2]
    (head,) = task_block(ws, "02-200")["TASK"]
    ws[f"A{head}"] = "02-250"
    ws[f"D{task_block(ws, '02-250')['CREW'][0]}"] = 2   # and an edit inside it
    sheet = S.read_item_sheet(ws, takeoff_kind="linear")
    item = copy.deepcopy(spec["items"][1])
    task = item["labor"][1]
    f = Findings()
    rate_names, trades = _rate_names_and_trades()
    _merge_tasks(item, sheet, "item 2", f, rate_names=rate_names, trades=trades)
    assert f.changes == ["item 2 · task 02-200 renamed to 02-250",
                         "item 2 · 02-250 · Equipment Operator: count 1 → 2"]
    assert item["labor"][1] is task
    assert (task["slot"], task["cost_code"], task["crew_factor"]) == ("02", "02-250", 1.15)
    assert [t["cost_code"] for t in item["labor"]] == ["02-100", "02-250"]


def test_a_new_code_elsewhere_is_still_a_removal_and_a_new_task():
    wb = build()
    ws = wb[ITEM2]
    block = task_block(ws, "02-100")
    for _ in range(block["SUBTOTAL"][0] - block["TASK"][0] + 1):
        simulate_delete_row(ws, block["TASK"][0])               # 02-100 gone (position 0)
    block = task_block(ws, "02-200")
    (first,), (sub,) = block["TASK"], block["SUBTOTAL"]
    copy_rows(ws, first, sub, sub + 1)
    ws[f"A{sub + 1}"] = "02-300"                                 # new at position 1
    item, f = _item(wb)
    assert f.changes[0].startswith("item 2 · + new task 02-300")
    assert f.changes[-1] == "item 2 · − task 02-100 Trench excavation & pipe laying (removed)"


def _swap_rows(ws, a, b, cols):
    va = [ws[f"{c}{a}"].value for c in cols]
    vb = [ws[f"{c}{b}"].value for c in cols]
    for c, v in zip(cols, vb):
        ws[f"{c}{a}"] = v
    for c, v in zip(cols, va):
        ws[f"{c}{b}"] = v


def test_crew_rows_or_lines_swapped_in_place_are_a_reorder_not_remove_and_add():
    wb = build()
    ws = wb[ITEM2]
    block = task_block(ws, "02-100")
    _swap_rows(ws, block["CREW"][0], block["CREW"][1], "ADE")      # Superintendent <-> Pipelayer
    _swap_rows(ws, block["LINE"][0], block["LINE"][2], "ABCGHIJ")  # Trench box <-> Bedding rock
    ws[f"E{block['CREW'][0]}"] = 9                                 # and Pipelayer's days edited
    spec = example_spec()
    before = spec["items"][1]["labor"][0]
    crew, lines = list(before["crew"]), list(before["lines"])
    item, f = _item(wb)
    assert f.changes == ["item 2 · 02-100 · crew reordered",
                         "item 2 · 02-100 · Pipelayer: days 8 → 9",
                         "item 2 · 02-100 · lines reordered"]
    task = item["labor"][0]
    assert [c["trade"] for c in task["crew"]][:2] == ["Pipelayer", "Superintendent"]
    assert task["crew"][0] == dict(crew[1], days=9) and task["crew"][1] == crew[0]
    assert [ln["name"] for ln in task["lines"]] == [
        lines[2]["name"], lines[1]["name"], lines[0]["name"]]
    assert task["lines"][0] == lines[2]


# ------------------------------------------------------------------ work packages

def _packaged():
    from estimating_tools.estimate_spec.build import generate
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    return generate(packaged_spec(), fieldbook_pack(), quotes=QUOTES)


def _read_packaged(wb):
    from estimating_tools.estimate_spec.readback import read_back
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    before = copy.deepcopy(QUOTES)
    rb = read_back(packaged_spec(), wb, fieldbook_pack(), quotes=QUOTES)
    assert QUOTES == before                      # the caller's quotes are never edited
    return rb


def _hauling(spec):
    (pkg,) = [p for p in spec["work_packages"] if p["name"] == "Hauling"]
    return pkg


def test_an_unedited_packaged_workbook_reads_back_to_itself():
    from estimating_tools.estimate_spec.tests._fieldbook import packaged_spec
    rb = _read_packaged(_packaged())
    assert rb.findings == Findings()
    assert rb.quote_writes == []
    assert rb.spec == packaged_spec()


def test_a_selector_change_is_the_packages_selection():
    wb = _packaged()
    wb["Sub Hauling"]["F3"] = "arr"
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert rb.findings.changes == ["packages · Hauling · selected Dirt Movers → ARR"]
    assert _hauling(rb.spec)["selected"] == "ARR"
    assert rb.quote_writes == []


def test_a_pick_names_the_vendor_for_the_item():
    wb = _packaged()
    wb["Sub Hauling"]["H8"] = " bulk haul co "
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert rb.findings.changes == ["packages · Hauling · pick item 2 → Bulk Haul Co"]
    assert _hauling(rb.spec)["picks"] == {"2": "Bulk Haul Co"}


def test_a_carry_figure_typed_in_and_selected():
    wb = _packaged()
    ws = wb["Sub Hauling"]
    ws["F3"], ws["W8"] = "CARRY", 17.5
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert rb.findings.changes == ["packages · Hauling · selected Dirt Movers → CARRY",
                                   "packages · Hauling · carry 2 → 17.5"]
    assert _hauling(rb.spec)["selected"] == "CARRY"
    assert _hauling(rb.spec)["carry"] == {"2": 17.5}


def test_a_surcharge_typed_under_a_vendor():
    wb = _packaged()
    wb["Sub Hauling"]["K7"] = 0.05
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert rb.findings.changes == ["packages · Hauling · surcharge Dirt Movers 0.04 → 0.05"]
    assert _hauling(rb.spec)["surcharge"] == {"Dirt Movers": 0.05}


def test_a_vendor_typed_into_a_spare_column_is_a_new_quote_file():
    from datetime import date
    wb = _packaged()
    ws = wb["Sub Hauling"]
    ws["O5"], ws["O8"], ws["O16"] = "Coastal Trucking", 15.25, "3% of invoice"
    ws["F3"] = "Coastal Trucking"
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert rb.findings.changes == [
        "packages · Hauling · selected Dirt Movers → Coastal Trucking",
        "packages · Hauling · new quote Coastal Trucking → coastal-trucking--hauling.json"]
    (w,) = rb.quote_writes
    assert w.file_name == "coastal-trucking--hauling.json"
    assert w.reason == "typed into Sub Hauling"
    q = w.quote
    assert (q.vendor_name, q.work_package_name) == ("Coastal Trucking", "Hauling")
    assert q.quote_date == date.today().isoformat()
    assert q.notes == f"entered in workbook on {date.today().isoformat()}"
    assert [(ln.item_number, ln.line_name, ln.unit_price, ln.is_included) for ln in q.lines] == [
        ("2", "Haul, import + export", 15.25, True)]
    assert q.conditions == {"Fuel surcharge": "3% of invoice"}
    assert _hauling(rb.spec)["selected"] == "Coastal Trucking"


def test_typed_vendors_with_incl_and_with_no_price():
    wb = _packaged()
    ws = wb["Sub Hauling"]
    ws["O5"], ws["O8"] = "Coastal Trucking", "Incl"
    ws["Q5"], ws["Q8"] = "Shore Haul", 14
    ws["S5"], ws["S8"] = "Nobody Yet", "excl"
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    by = {w.quote.vendor_name: w.quote for w in rb.quote_writes}
    assert sorted(by) == ["Coastal Trucking", "Shore Haul"]
    (ln,) = by["Coastal Trucking"].lines
    assert (ln.is_included, ln.folded_into, ln.unit_price) == (True, "Haul, import + export", None)
    assert by["Shore Haul"].lines[0].unit_price == 14
    assert rb.findings.notices == ["Sub Hauling: vendor Nobody Yet (column S) has no prices "
                                   "typed under it -- no quote file written"]


def test_an_existing_vendor_excluding_a_line():
    wb = _packaged()
    wb["Sub Hauling"]["K8"] = None
    rb = _read_packaged(wb)
    (w,) = rb.quote_writes
    (ln,) = w.quote.lines
    assert (ln.is_included, ln.unit_price) == (False, None)


def test_a_typed_vendor_that_already_has_a_quote_file_is_a_problem():
    from estimating_tools.estimate_spec.readback import read_back
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    from estimating_tools.quotes.model import CanonicalQuote
    wb = _packaged()
    ws = wb["Sub Hauling"]
    ws["O5"], ws["O8"] = "Late Hauler", 15
    # Its quote file arrived after the workbook was built, so the sheet has no
    # column for it: typing it into a spare does not make it one.
    late = CanonicalQuote(vendor_name="late hauler", work_package_name="Hauling", amount=1)
    rb = read_back(packaged_spec(), wb, fieldbook_pack(), quotes=QUOTES + [late])
    assert ("Sub Hauling: Hauling's quote files changed since the workbook was built (7 vendor "
            "columns now, 6 on the sheet; a vendor typed in may already have a quote file) -- "
            "rebuild the workbook and redo the edits") in rb.findings.problems, rb.findings.problems
    assert rb.quote_writes == []


def test_an_edited_vendor_price_rewrites_that_quote_file():
    wb = _packaged()
    wb["Sub Hauling"]["K8"] = 15.9
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert rb.findings.overwrites == []
    assert rb.findings.quote_edits == [
        "Sub Hauling!K8 (Dirt Movers, item 2 Haul, import + export): 16.5 → 15.9 "
        "written to dirt-movers--hauling.json"]
    (w,) = rb.quote_writes
    assert (w.file_name, w.reason) == ("dirt-movers--hauling.json", "edited in Sub Hauling")
    assert [ln.unit_price for ln in w.quote.lines] == [15.9]
    assert w.quote.conditions == {"Fuel surcharge": "4% of invoice", "Dump fees": "included"}


def test_edited_prices_and_conditions_make_one_write_per_vendor():
    wb = _packaged()
    ws = wb["Sub Hauling"]
    ws["K8"], ws["K17"] = "incl", "extra $5/load"
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    (w,) = rb.quote_writes
    (ln,) = w.quote.lines
    assert (ln.is_included, ln.folded_into, ln.unit_price) == (
        True, "Haul, import + export", None)
    assert w.quote.conditions["Dump fees"] == "extra $5/load"


def test_a_prorated_lump_sum_cell_typed_over_is_a_real_price():
    wb = _packaged()
    wb["Sub Hauling"]["M8"] = 18.25
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    (w,) = rb.quote_writes
    assert w.quote.vendor_name == "Bulk Haul Co" and w.quote.amount == 8740
    assert [(ln.item_number, ln.line_name, ln.unit_price) for ln in w.quote.lines] == [
        ("2", "Haul, import + export", 18.25)]


def test_a_number_typed_over_a_package_lines_price_is_a_problem():
    wb = _packaged()
    ws = wb["2 - 8IN MAIN"]
    (row,) = [r for r in task_block(ws, "02-200")["LINE"] if ws[f"A{r}"].value.startswith("Haul")]
    ws[f"G{row}"] = 20
    rb = _read_packaged(wb)
    assert rb.findings.problems == [
        f"2 - 8IN MAIN row {row} (item 2 · 02-200): unit price is driven by the Sub Hauling "
        "sheet -- change the selector or the vendor price there"]
    assert rb.blocked


def _pipelayer(ws):
    (row,) = [r for r in task_block(ws, "02-510")["CREW"] if ws[f"A{r}"].value == "Pipelayer"]
    return row


def test_a_number_typed_over_self_perform_days_is_the_new_days():
    wb = _packaged()
    wb["Sub Valves"]["F3"] = "Valve Pros"
    ws = wb["3 - VALVES"]
    row = _pipelayer(ws)
    ws[f"E{row}"] = 3
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert ("item 3 · 02-510 · Pipelayer: days 2 → 3" in rb.findings.changes), \
        rb.findings.changes
    assert (f"3 - VALVES row {row} (item 3 · 02-510): days 3 typed over the self-perform "
            "formula -- the self-perform factor (Sub Valves) is re-applied on rebuild"
            ) in rb.findings.notices
    (item3,) = [i for i in rb.spec["items"] if i["item_no"] == "3"]
    (pipe,) = [m for m in item3["labor"][0]["crew"] if m["trade"] == "Pipelayer"]
    assert pipe["days"] == 3


def test_the_built_shape_with_a_new_number_is_the_new_days():
    wb = _packaged()
    ws = wb["3 - VALVES"]
    row = _pipelayer(ws)
    built = ws[f"E{row}"].value
    assert built.startswith("=2*'Sub Valves'!")
    ws[f"E{row}"] = "=3*" + built[len("=2*"):]
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert any(c.endswith("Pipelayer: days 2 → 3") for c in rb.findings.changes)
    assert any("re-applied on rebuild" in n for n in rb.findings.notices)


def test_another_formula_over_self_perform_days_is_a_problem():
    wb = _packaged()
    ws = wb["3 - VALVES"]
    row = _pipelayer(ws)
    ws[f"E{row}"] = "=3*1"
    rb = _read_packaged(wb)
    assert rb.findings.problems == [
        f"3 - VALVES row {row} (item 3 · 02-510): days is driven by the Sub Valves sheet -- type "
        "a number (the new days; the self-perform factor is re-applied on rebuild) or change "
        "the selector there"]
    assert rb.blocked


def test_a_number_typed_over_a_self_perform_lines_qty_is_the_new_qty():
    from estimating_tools.estimate_spec.build import generate
    from estimating_tools.estimate_spec.readback import read_back
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    spec = packaged_spec()
    (item3,) = [i for i in spec["items"] if i["item_no"] == "3"]
    item3["labor"][0]["lines"].append({"cost_type": "material", "name": "Gaskets", "qty": 4,
                                       "uom": "EA", "unit_price": 10, "source": "plug"})
    wb = generate(copy.deepcopy(spec), fieldbook_pack(), quotes=QUOTES)
    ws = wb["3 - VALVES"]
    (row,) = [r for r in task_block(ws, "02-510")["LINE"] if ws[f"A{r}"].value == "Gaskets"]
    assert ws[f"B{row}"].value.startswith("=4*'Sub Valves'!")
    ws[f"B{row}"] = 6
    rb = read_back(spec, wb, fieldbook_pack(), quotes=QUOTES)
    assert rb.findings.problems == []
    (got,) = [i for i in rb.spec["items"] if i["item_no"] == "3"]
    (gaskets,) = [ln for ln in got["labor"][0]["lines"] if ln["name"] == "Gaskets"]
    assert gaskets["qty"] == 6


def test_a_selector_naming_no_vendor_lists_the_vendors():
    wb = _packaged()
    wb["Sub Hauling"]["F3"] = "Dirt Mover"
    rb = _read_packaged(wb)
    assert rb.findings.problems == [
        "Sub Hauling: selector 'Dirt Mover' is not ARR, CARRY or a quoted vendor "
        "(vendors: Dirt Movers, Bulk Haul Co)"]


def test_a_pick_naming_no_vendor_names_the_item():
    wb = _packaged()
    wb["Sub Hauling"]["H8"] = "Nobody"
    rb = _read_packaged(wb)
    assert rb.findings.problems == [
        "Sub Hauling row 8: pick 'Nobody' for item 2 is not ARR, CARRY or a quoted vendor "
        "(vendors: Dirt Movers, Bulk Haul Co)"]


def test_a_missing_package_sheet_is_a_problem():
    wb = _packaged()
    del wb["Sub Valves"]
    rb = _read_packaged(wb)
    assert "Sub Valves: the workbook has no such sheet for package Valves -- rebuild the workbook" in (
        rb.findings.problems)


def test_quotes_read_from_a_directory_are_written_back_to_their_own_files(tmp_path):
    from estimating_tools.estimate_spec.readback import read_back
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    from estimating_tools.quotes.files import write_quote_file
    for i, q in enumerate(QUOTES):
        write_quote_file(tmp_path / f"q{i}.json", q)
    wb = _packaged()
    wb["Sub Hauling"]["K8"] = 15.9
    rb = read_back(packaged_spec(), wb, fieldbook_pack(), quotes_dir=tmp_path)
    assert rb.findings.problems == []
    assert [w.file_name for w in rb.quote_writes] == ["q0.json"]
    assert rb.findings.quote_edits[0].endswith("written to q0.json")


def test_a_typed_vendor_whose_file_name_is_taken_is_a_problem():
    from estimating_tools.estimate_spec.build import generate
    from estimating_tools.estimate_spec.readback import read_back
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    from estimating_tools.quotes.model import CanonicalQuote, QuoteLine
    other = CanonicalQuote(vendor_name="Coastal-Trucking", work_package_name="Hauling",
                           lines=[QuoteLine(schedule_name="Base", item_number="2", unit_price=17,
                                            line_name="Haul, import + export")])
    quotes = QUOTES + [other]
    wb = generate(packaged_spec(), fieldbook_pack(), quotes=quotes)
    ws = wb["Sub Hauling"]
    (spare,) = [c.column_letter for c in ws[5] if c.value == ""][:1]
    ws[f"{spare}5"], ws[f"{spare}8"] = "Coastal Trucking", 15
    rb = read_back(packaged_spec(), wb, fieldbook_pack(), quotes=quotes)
    assert rb.findings.problems == [
        "Sub Hauling: typed vendor 'Coastal Trucking' would write coastal-trucking--hauling.json, "
        "which already belongs to 'Coastal-Trucking' -- rename the vendor"]
    assert rb.quote_writes == []


def test_a_condition_typed_under_a_lump_sum_keeps_it_a_lump_sum():
    wb = _packaged()
    wb["Sub Hauling"]["M17"] = "extra"
    rb = _read_packaged(wb)
    assert rb.findings.problems == [] and rb.findings.notices == []
    (w,) = rb.quote_writes
    assert w.quote.vendor_name == "Bulk Haul Co"
    assert (w.quote.amount, w.quote.lines) == (8740, [])
    assert w.quote.conditions == {"Dump fees": "extra"}


def test_prices_under_a_blank_spare_header_are_a_problem():
    wb = _packaged()
    wb["Sub Hauling"]["O8"] = 15
    rb = _read_packaged(wb)
    assert rb.findings.problems == [
        "Sub Hauling: column O has prices but no vendor name in its header -- type the "
        "vendor's name in row 5"]
    assert rb.quote_writes == []


def test_a_vendor_whose_quote_file_was_removed_is_a_problem():
    from estimating_tools.estimate_spec.readback import read_back
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    # Built WITH Bulk Haul Co (no quote_ref of its own: its column shows the
    # default ref), read back after its quote file is gone.
    wb = _packaged()
    rb = read_back(packaged_spec(), wb, fieldbook_pack(),
                   quotes=[q for q in QUOTES if q.vendor_name != "Bulk Haul Co"])
    assert ("Sub Hauling: vendor 'Bulk Haul Co' has a quote ref but no quote file -- a quote "
            "file was removed since the build; rebuild the workbook (or, for a vendor you just "
            "typed in, clear its Quote ref cell)") in rb.findings.problems
    assert not any("new quote" in c for c in rb.findings.changes)
    assert rb.quote_writes == []


def _insert_cols(wb, sheet: str, at: int, amount: int) -> None:
    """Excel's column insert on `sheet` at column index `at`: every reference to
    a column >= `at` of that sheet -- its own unqualified ones, and qualified ones
    on every sheet -- moves right by `amount`."""
    import re
    from openpyxl.utils import column_index_from_string as idx
    from openpyxl.utils import get_column_letter as letter
    ref = re.compile(r"((?:'[^']*(?:''[^']*)*'|[A-Za-z0-9_ ]+)!)?(\$?)([A-Z]{1,3})(\$?\d+)")
    quoted = "'" + sheet.replace("'", "''") + "'"

    def shift(own):
        def move(m):
            q = m.group(1)
            if (q is None and not own) or (q is not None and q[:-1] not in (sheet, quoted)):
                return m.group(0)
            c = idx(m.group(3))
            return (q or "") + m.group(2) + letter(c + amount if c >= at else c) + m.group(4)
        return move

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    parts = re.split(r'("[^"]*")', v)
                    cell.value = "".join(p if i % 2 else ref.sub(shift(ws.title == sheet), p)
                                         for i, p in enumerate(parts))
    wb[sheet].insert_cols(at, amount)


def test_a_vendor_typed_into_an_inserted_column_pair():
    from openpyxl.utils import column_index_from_string as idx
    wb = _packaged()
    ws = wb["Sub Hauling"]
    assert ws["W5"].value == "CARRY"
    _insert_cols(wb, "Sub Hauling", idx("W"), 2)
    assert wb["2 - 8IN MAIN"]["I36"].value == "='Sub Hauling'!$AC$8"     # Label moved
    ws["W5"], ws["X5"], ws["W8"] = "Coastal Trucking", "ext", 15
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert rb.findings.overwrites == []
    (w,) = rb.quote_writes
    assert (w.file_name, w.quote.lines[0].unit_price) == ("coastal-trucking--hauling.json", 15)


def test_a_vendor_typed_into_a_column_inserted_at_carry_is_priced(tmp_path):
    # One pair inserted AT the CARRY column (W): the F/State vendor ranges end
    # at CARRY, so they grow over the new pair, and carrying the typed-in
    # vendor prices the line (no #N/A); read-back writes its one quote.
    from openpyxl.utils import column_index_from_string as idx
    from estimating_tools.estimate_spec.tests._formula_eval import evaluate
    wb = _packaged()
    ws = wb["Sub Hauling"]
    assert ws["W5"].value == "CARRY"
    _insert_cols(wb, "Sub Hauling", idx("W"), 2)
    assert ws["Y5"].value == "CARRY"
    assert "$K8:$Y8" in ws["F8"].value and "$K$5:$Y$5" in ws["F8"].value
    ws["W5"], ws["X5"], ws["W8"] = "Coastal Trucking", "ext", 15
    ws["F3"] = "Coastal Trucking"
    ev = evaluate(wb, tmp_path)
    assert ev["Sub Hauling!AB8"] == "price"                    # State, built at Z
    assert ev["Sub Hauling!F8"] == 15
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    (w,) = rb.quote_writes
    assert (w.file_name, w.quote.lines[0].unit_price) == ("coastal-trucking--hauling.json", 15)
    assert "packages · Hauling · selected Dirt Movers → Coastal Trucking" in rb.findings.changes


def test_a_formula_typed_over_on_a_package_sheet_is_reported_not_dropped():
    wb = _packaged()
    wb["Sub Hauling"]["F8"] = 99
    rb = _read_packaged(wb)
    assert rb.findings.problems == []
    assert rb.findings.overwrites == ["Sub Hauling!F8 (formula): typed 99 over the formula"]
    assert rb.quote_writes == []


def test_a_formula_typed_over_after_a_column_insert_is_found_where_it_moved():
    from openpyxl.utils import column_index_from_string as idx
    wb = _packaged()
    _insert_cols(wb, "Sub Hauling", idx("W"), 2)
    wb["Sub Hauling"]["Z8"] = 5                     # CARRY ext, built at X
    rb = _read_packaged(wb)
    assert rb.findings.overwrites == ["Sub Hauling!Z8 (formula): typed 5 over the formula"]


def test_an_alternate_row_is_shown_not_read_back():
    from estimating_tools.estimate_spec.build import generate
    from estimating_tools.estimate_spec.readback import read_back
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    from estimating_tools.quotes.model import QuoteLine
    alt = copy.deepcopy(QUOTES)
    alt[0].lines.append(QuoteLine(schedule_name="Base", item_number="2", unit_price=12.0,
                                  line_name="Haul, import + export", is_alternate=True))
    wb = generate(packaged_spec(), fieldbook_pack(), quotes=alt)
    ws = wb["Sub Hauling"]
    (r,) = [r for r in range(1, ws.max_row + 1)
            if ws[f"B{r}"].value == "Alternate: Haul, import + export"]
    assert ws[f"K{r}"].value == "$12.00 / CY"
    assert read_back(packaged_spec(), wb, fieldbook_pack(), quotes=alt).quote_writes == []
    ws[f"K{r}"] = "$11.00 / CY"
    rb = read_back(packaged_spec(), wb, fieldbook_pack(), quotes=alt)
    assert rb.findings.problems == [] and rb.quote_writes == []
    assert ("Sub Hauling (Dirt Movers, 'Alternate: Haul, import + export'): an alternate row "
            "shows the quote file's alternate line -- edit the quote file to change it"
            ) in rb.findings.notices


def _two_line_valves():
    from estimating_tools.estimate_spec.build import generate
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    spec = packaged_spec()
    (item3,) = [i for i in spec["items"] if i["item_no"] == "3"]
    item3["labor"][0]["lines"].append({"cost_type": "subcontract", "name": "Box risers",
                                       "qty": 4, "uom": "EA", "unit_price": 80,
                                       "package": "Valves"})
    wb = generate(copy.deepcopy(spec), fieldbook_pack(), quotes=QUOTES)
    ws = wb["Sub Valves"]
    rows = {ws[f"B{r}"].value: r for r in range(1, ws.max_row + 1)}
    return spec, wb, rows["Valve boxes (subcontract)"], rows["Box risers"]


def _read(spec, wb):
    from estimating_tools.estimate_spec.readback import read_back
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES
    return read_back(spec, wb, fieldbook_pack(), quotes=QUOTES)


def test_a_pick_on_one_line_of_a_two_line_item_is_a_problem():
    spec, wb, boxes, risers = _two_line_valves()
    wb["Sub Valves"][f"H{boxes}"] = "Valve Pros"
    rb = _read(spec, wb)
    assert rb.findings.problems == [
        f"Sub Valves: item 3's lines carry different picks (row {boxes} Valve Pros, row {risers} "
        "(blank)) -- set the same pick on every line of the item or clear them"]
    assert rb.blocked


def test_two_different_picks_on_one_item_are_a_problem():
    spec, wb, boxes, risers = _two_line_valves()
    wb["Sub Valves"][f"H{boxes}"] = "Valve Pros"
    wb["Sub Valves"][f"H{risers}"] = "carry"
    rb = _read(spec, wb)
    assert rb.findings.problems == [
        f"Sub Valves: item 3's lines carry different picks (row {boxes} Valve Pros, row {risers} "
        "CARRY) -- set the same pick on every line of the item or clear them"]


def test_the_same_pick_on_every_line_of_an_item_is_read():
    spec, wb, boxes, risers = _two_line_valves()
    wb["Sub Valves"][f"H{boxes}"] = "Valve Pros"
    wb["Sub Valves"][f"H{risers}"] = "valve pros"
    rb = _read(spec, wb)
    assert rb.findings.problems == []
    (valves,) = [p for p in rb.spec["work_packages"] if p["name"] == "Valves"]
    assert valves["picks"] == {"3": "Valve Pros"}


def test_a_price_edit_preserves_an_alternate_before_the_base_line():
    from estimating_tools.estimate_spec.build import generate
    from estimating_tools.estimate_spec.readback import read_back
    from estimating_tools.estimate_spec.tests._fieldbook import QUOTES, packaged_spec
    from estimating_tools.quotes.model import QuoteLine
    quotes = copy.deepcopy(QUOTES)
    alternate = QuoteLine(schedule_name="Base", item_number="2", unit_price=12.0,
                          line_name="Haul, import + export", is_alternate=True)
    quotes[0].lines.insert(0, alternate)
    spec = packaged_spec()
    wb = generate(spec, fieldbook_pack(), quotes=quotes)
    wb["Sub Hauling"]["K8"] = 15
    result = read_back(spec, wb, fieldbook_pack(), quotes=quotes)
    assert result.findings.problems == []
    (write,) = result.quote_writes
    assert write.quote.lines[0] == alternate
    (base,) = [line for line in write.quote.lines
               if line.item_number == "2" and not line.is_alternate]
    assert base.unit_price == 15
    rebuilt = generate(result.spec, fieldbook_pack(), quotes=[write.quote, *quotes[1:]])
    assert rebuilt["Sub Hauling"]["K8"].value == 15


@pytest.mark.parametrize("value", ["NaN", "inf", "1e999"])
def test_nonfinite_quote_prices_are_blocked_before_writing(value):
    wb = _packaged()
    wb["Sub Hauling"]["K8"] = value
    result = _read_packaged(wb)
    assert result.blocked and result.quote_writes == []
    assert any("not a price" in problem for problem in result.findings.problems)


@pytest.mark.parametrize("value", ["fee", "NaN"])
def test_a_typed_vendors_invalid_surcharge_is_blocked_instead_of_dropped(value):
    wb = _packaged()
    ws = wb["Sub Hauling"]
    ws["O5"], ws["O8"], ws["O7"] = "Late Vendor", 15, value
    result = _read_packaged(wb)
    assert result.blocked and result.quote_writes == []
    assert any("surcharge" in problem for problem in result.findings.problems)
