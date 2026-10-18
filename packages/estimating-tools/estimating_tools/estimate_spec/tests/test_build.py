"""Tests for the estimate_spec generator (build.py).

The Resolver is the fiddly pure logic: it substitutes the placeholder
mini-language documented in units.yaml. These tests pin every placeholder form
— pure, no pack involved.

The integration tests render the alt pack's own example
(tests/fixtures/packs/alt/examples/sample.estimate.yaml, authored for this
test suite) with `pack=alt` passed explicitly, and assert against cell
coordinates/values read back from the actual generated workbook (derived by
running the render, not hand-guessed) — see the fix-round-1 report for the
derivation.
"""
from __future__ import annotations

from pathlib import Path

import pytest
import yaml

from estimating_tools.estimate_spec.build import Resolver
from estimating_tools.estimate_spec.format_pack import load_pack

SPEC_DIR = Path(__file__).resolve().parents[1]
ALT_DIR = Path(__file__).resolve().parent / "fixtures" / "packs" / "alt"
ALT_SAMPLE = ALT_DIR / "examples" / "sample.estimate.yaml"
LEGACY_DIR = SPEC_DIR / "tests" / "fixtures" / "packs" / "legacy_demo"


@pytest.fixture(scope="module")
def alt():
    return load_pack(ALT_DIR)


def _cfg():
    return yaml.safe_load(ALT_SAMPLE.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def wb(alt):
    from estimating_tools.estimate_spec.build import generate
    return generate(_cfg(), alt)


# --- {p.NAME} unit parameters ------------------------------------------------
def test_param_substitution():
    r = Resolver(params={"slot": "AA", "description": "Footings"})
    assert r.resolve("[ {p.slot} ]") == "[ AA ]"
    assert r.resolve("{p.description}") == "Footings"


# --- {it...} repeat-item fields ----------------------------------------------
def test_repeat_item_field():
    r = Resolver(item={"name": "Fine Grade", "factor": 0.25})
    assert r.resolve("{it.name}") == "Fine Grade"


def test_repeat_item_scalar():
    r = Resolver(item="Mobilization is Critical")
    assert r.resolve("{it}") == "Mobilization is Critical"


# --- {r} current-row arithmetic ----------------------------------------------
def test_current_row():
    r = Resolver(row=44)
    assert r.resolve("=E{r}*D{r}*8") == "=E44*D44*8"


def test_current_row_arithmetic():
    r = Resolver(row=44)
    assert r.resolve("=SUM(F{r-1})") == "=SUM(F43)"
    assert r.resolve("=D{r+2}") == "=D46"


# --- {row:NAME} landmarks ----------------------------------------------------
def test_landmark_rows():
    r = Resolver(landmarks={"crew_first": 44, "crew_last": 46})
    assert r.resolve("=SUM(F{row:crew_first}:F{row:crew_last})") == "=SUM(F44:F46)"


# --- {ref:UNIT.EXPORT} cross-unit references (always full addresses) --------
def test_ref_returns_full_address():
    r = Resolver(exports={"takeoff.total_cy": "$D$19",
                          "labor_rollup.foreman_days": "$G$120"})
    assert r.resolve("={ref:takeoff.total_cy}") == "=$D$19"
    assert r.resolve("={ref:labor_rollup.foreman_days}*8") == "=$G$120*8"


def test_ref_keyed():
    r = Resolver(exports={"labor[A].driver": "$D$41"})
    assert r.resolve("={ref:labor[A].driver}") == "=$D$41"


# --- {item_qty} / {item_uom} -------------------------------------------------
def test_item_qty_uom():
    r = Resolver(row=15, landmarks={"total": 15}, globals={"item_qty": "$G$2", "item_uom": "$H$2"})
    assert r.resolve("=G{row:total}/{item_qty}") == "=G15/$G$2"
    assert r.resolve("={item_uom}") == "=$H$2"


# --- {rate_table} ------------------------------------------------------------
def test_rate_table():
    r = Resolver(row=44, globals={"rate_table": "'Crew Rates'!$A$2:$B$6"})
    out = r.resolve("=VLOOKUP($B{r},{rate_table},2,FALSE)*C{r}")
    assert out == "=VLOOKUP($B44,'Crew Rates'!$A$2:$B$6,2,FALSE)*C44"


# --- non-string passthrough --------------------------------------------------
def test_non_string_passthrough():
    r = Resolver(row=5)
    assert r.resolve(0.25) == 0.25
    assert r.resolve(10) == 10
    assert r.resolve(None) is None


# --- unresolved placeholder is a clear error ---------------------------------
def test_unknown_placeholder_raises():
    r = Resolver(params={"slot": "AA"})
    with pytest.raises(KeyError):
        r.resolve("{p.missing}")
    with pytest.raises(ValueError):
        r.resolve("{row:nope}")


# =============================================================================
# Integration: generate the alt pack's own example config.
# =============================================================================
def test_sheets_created_in_order(wb):
    assert wb.sheetnames == ["Recap", "PRE", "Crew Rates", "2 - DUCT RUN", "3 - HANGER SET"]


def test_rate_table_placed_on_its_own_sheet(wb):
    ws = wb["Crew Rates"]
    assert ws["A2"].value == "FITTER"
    assert ws["B2"].value == 88
    assert ws["A3"].value == "RIGGER"


def test_item_identity_row(wb):
    ws = wb["2 - DUCT RUN"]
    assert ws["B2"].value == "24-inch duct run, level 2"
    assert ws["H2"].value == "CY"


def test_crew_rate_vlookups_target_the_rate_sheet(wb):
    ws = wb["2 - DUCT RUN"]
    vlookups = [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=VLOOKUP")]
    assert vlookups, "expected crew-rate VLOOKUP formulas"
    assert all("'Crew Rates'!$A$2:$B$6" in f for f in vlookups)


def test_a_trade_added_to_the_rate_table_grows_the_vlookup_range(alt):
    # The exported range is row-derived ("$A$2:$B${row:rate_last}"), so a new
    # trade must grow the VLOOKUP range rather than landing outside it (where
    # every crew row's VLOOKUP would return #N/A, silently: openpyxl reads
    # formulas, so only an Excel recalc would show it).
    from estimating_tools.estimate_spec.build import generate
    cfg = _cfg()
    before_last_row = 1 + len(cfg["rate_table"])       # header row + N trades
    cfg["rate_table"].append({"trade": "IRONWORKER", "rate": 95})
    after_last_row = before_last_row + 1
    wb10 = generate(cfg, alt)
    assert wb10["Crew Rates"][f"A{after_last_row}"].value == "IRONWORKER"
    vlookups = [c.value for ws in wb10 for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=VLOOKUP")]
    assert vlookups, "expected crew-rate VLOOKUP formulas"
    assert all(f"'Crew Rates'!$A$2:$B${after_last_row}" in f for f in vlookups)


def test_rate_cells_carry_input_yellow_fill(wb):
    ws = wb["2 - DUCT RUN"]
    yellow = [c.coordinate for row in ws.iter_rows() for c in row
              if c.fill and c.fill.patternType and c.fill.fgColor.rgb in ("FFFFE699", "00FFE699")]
    assert yellow, "expected at least one input_yellow cell"


def test_total_block_sums_components(wb):
    ws = wb["2 - DUCT RUN"]
    total_rows = [c.row for row in ws.iter_rows() for c in row if c.value == "TOTAL"]
    assert total_rows
    g = ws.cell(total_rows[0], 7).value
    assert isinstance(g, str) and g.startswith("=SUM(")


def test_unit_cost_divides_total_by_qty(wb):
    ws = wb["2 - DUCT RUN"]
    rows = [c.row for row in ws.iter_rows() for c in row if c.value == "UNIT COSTS"]
    assert rows
    g = ws.cell(rows[0], 7).value
    assert isinstance(g, str) and g.startswith("=") and "/$G$2" in g


def test_recap_pulls_item_totals_and_has_markup(wb):
    ws = wb["Recap"]
    refs = [c.value for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and "'2 - DUCT RUN'" in c.value]
    assert refs, "Recap should reference the duct-run sheet total"
    pct = [c.coordinate for row in ws.iter_rows() for c in row if c.number_format == "0.00%"]
    assert pct, "expected a percent-formatted markup cell"


def test_hanger_set_total_includes_vendor_spread(wb):
    ws = wb["3 - HANGER SET"]
    ss = [c.row for row in ws.iter_rows() for c in row if c.value == "Vendor Spread"]
    assert ss, "expected a Vendor Spread row"
    total = [c.row for row in ws.iter_rows() for c in row if c.value == "TOTAL"]
    tf = ws.cell(total[0], 7).value
    assert f"$G${ss[0]}" in tf, f"vendor-spread $G${ss[0]} not in TOTAL {tf!r}"


def test_costdetail_unit_header_box_spans_A_to_F(wb):
    ws = wb["2 - DUCT RUN"]
    rows = [c.row for row in ws.iter_rows() for c in row
            if c.column == 1 and c.value == "Unit"]
    assert rows, "expected cost-detail 'Unit' header rows in column A"
    for r in rows:
        assert ws.cell(r, 1).border.left.style == "thick", f"A{r} left edge"
        assert ws.cell(r, 1).border.top.style == "thick", f"A{r} top edge"
        assert ws.cell(r, 6).border.right.style == "thick", f"F{r} right edge"


def test_crew_unit_driver_row_boxed_B_to_G(wb):
    ws = wb["2 - DUCT RUN"]
    rows = [c.row for row in ws.iter_rows() for c in row
            if c.column == 2 and c.value == "Unit"]
    assert rows, "expected crew-stack 'Unit' driver rows in column B"
    for r in rows:
        assert ws.cell(r, 2).border.left.style == "thick", f"B{r} left edge"
        assert ws.cell(r, 2).border.top.style == "thick", f"B{r} top edge"
        assert ws.cell(r, 7).border.right.style == "thick", f"G{r} right edge"
        # the driver-qty cell (D) keeps its pale fill inside the box
        assert ws.cell(r, 4).fill.patternType and \
            ws.cell(r, 4).fill.fgColor.rgb in ("FFFFF2CC", "00FFF2CC")


def test_production_summary_bar_spans_D_to_G(wb):
    ws = wb["2 - DUCT RUN"]
    target = None
    for row in ws.iter_rows():
        e = row[4]  # column E
        if (isinstance(e.value, str) and e.value.startswith("=SUM(E")
                and e.fill.patternType and e.fill.fgColor.rgb in ("FFFFE699", "00FFE699")):
            target = e.row
            break
    assert target, "no production factor-sum bar found"
    for col in (4, 5, 6, 7):  # D E F G
        c = ws.cell(target, col)
        assert c.fill.patternType and c.fill.fgColor.rgb in ("FFFFE699", "00FFE699"), \
            f"col {col} on the bar is not the input-yellow fill"


def test_crew_rollup_title_is_2x2_merge(wb):
    from openpyxl.utils import range_boundaries
    ws = wb["2 - DUCT RUN"]
    title = next(row[4] for row in ws.iter_rows()
                 if row[4].fill.patternType and row[4].fill.fgColor.rgb in ("FFE2EFDA", "00E2EFDA")
                 and isinstance(row[4].value, str) and row[4].value.startswith("="))
    found = None
    for mr in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(mr))
        if min_row == title.row and min_col == 5:  # anchored at E{title row}
            found = (min_col, min_row, max_col, max_row)
            break
    assert found, "no merge anchored at the rollup title (E) cell"
    min_col, min_row, max_col, max_row = found
    assert (max_col - min_col, max_row - min_row) == (1, 1), \
        f"expected 2x2 (E:F over 2 rows), got cols {min_col}-{max_col} rows {min_row}-{max_row}"


def test_crew_rollup_block_is_boxed(wb):
    ws = wb["2 - DUCT RUN"]
    title = next((row[4] for row in ws.iter_rows()
                  if row[4].fill.patternType and row[4].fill.fgColor.rgb in ("FFE2EFDA", "00E2EFDA")
                  and isinstance(row[4].value, str) and row[4].value.startswith("=")), None)
    assert title is not None, "no crew-rollup title cell (fill, formula) found"
    assert title.border.top.style == "thick", "rollup needs a top border"
    assert ws.cell(title.row, 5).border.left.style == "thick", "rollup left edge"


def test_parts_lease_allowance_in_header(wb):
    """PARTS & LEASE carry their $/unit allowance in the Unit header row
    (=E*B), this pack's own convention (see recipes.yaml's
    instance_defaults) — never as a line below it."""
    ws = wb["2 - DUCT RUN"]
    for block in ("PARTS", "LEASE"):
        label_row = next(c.row for row in ws.iter_rows() for c in row
                         if c.column == 1 and c.value == block)
        unit_row = label_row + 1
        assert ws.cell(unit_row, 1).value == "Unit", f"{block}: no Unit header"
        f = ws.cell(unit_row, 6).value  # F = extension on the header row
        assert isinstance(f, str) and f.startswith("=E") and "*B" in f, \
            f"{block}: allowance not in header (F{unit_row}={f!r})"


def test_roundtrip_save_reload(tmp_path, wb):
    import openpyxl
    p = tmp_path / "out.xlsx"
    wb.save(p)
    wb2 = openpyxl.load_workbook(p)
    assert wb2.sheetnames == wb.sheetnames


# --- Phase 4.5: cross-reference validation on generate() ---------------------
def test_generate_raises_on_missing_crew_reference(alt):
    from estimating_tools.estimate_spec.build import generate

    cfg = {
        "project": "Sample",
        "markup": 100,
        "rate_table": [{"trade": "FITTER", "rate": 150}],
        "items": [
            {"item_no": "Item 1", "recipe": "preamble", "sheet": "PRE",
             "description": "Mob", "uom": "LS", "labor": [],
             "crew_id": "crew_nonexistent"},
        ],
    }
    with pytest.raises(ValueError, match=r"Item 1.*crew_nonexistent"):
        generate(cfg, alt)


def test_generate_still_works_when_no_new_fields(wb):
    # Regression: the alt pack's own example (which exercises crews[],
    # item.crew_id, item.wbs, item.cost_code) still builds successfully.
    # The `wb` fixture from the module scope already exercises generate().
    assert wb is not None


# --- markup: number or formula (R1 item 8, ported from the Work repo) --------
def _markup_cfg(markup):
    return {
        "project": "Test",
        "markup": markup,
        "rate_table": [{"trade": "FITTER", "rate": 175}],
        "items": [{"recipe": "preamble", "sheet": "PRE", "item_no": "1",
                   "description": "Test", "uom": "EA", "labor": []}],
    }


def _schema(alt):
    return alt.schema


def test_markup_accepts_a_formula_string(alt):
    """A percentage markup has to be a formula so it reflows when items change.

    markup is a DOLLAR figure whose percentage is derived (units.yaml:
    D = "=E{r}/E{row:subtotal}"). Hand-entering the dollar strands it the
    moment an item is added, so allow a formula in its place.
    """
    import jsonschema
    jsonschema.validate(_markup_cfg("=E{row:subtotal}*0.10"), _schema(alt))


def test_markup_still_accepts_a_number(alt):
    import jsonschema
    jsonschema.validate(_markup_cfg(10000), _schema(alt))


def test_markup_rejects_a_bare_string(alt):
    import jsonschema
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(_markup_cfg("ten percent"), _schema(alt))


# --- #257: x-* namespaced passthrough (consumer generation metadata) ---------
def test_item_accepts_x_prefixed_passthrough_key(alt):
    """A consumer must be able to persist its own generation metadata on an item
    (e.g. which work-type profile produced its cost blocks) without the closed
    item object rejecting it. Namespaced x-* keys are the escape hatch, so the
    tag survives instead of being dropped after generation (#257)."""
    import jsonschema
    cfg = _markup_cfg(100)
    cfg["items"][0]["x-work-type"] = "formliner"
    jsonschema.validate(cfg, _schema(alt))


def test_top_level_accepts_x_prefixed_passthrough_key(alt):
    import jsonschema
    cfg = _markup_cfg(100)
    cfg["x-basis"] = "some-basis-2026Q1"
    jsonschema.validate(cfg, _schema(alt))


def test_non_namespaced_unknown_key_still_rejected(alt):
    """The passthrough is namespaced on purpose: a bare (typo'd) key must still
    fail so additionalProperties keeps catching mistakes."""
    import jsonschema
    cfg = _markup_cfg(100)
    cfg["items"][0]["worktype"] = "formliner"   # missing x- prefix
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(cfg, _schema(alt))


def test_resolver_declared_globals():
    res = Resolver(globals={"crew_rate_sheet": "'Rates'!$A$1:$B$9"})
    assert res.resolve("=VLOOKUP(X,{crew_rate_sheet},2,0)") == "=VLOOKUP(X,'Rates'!$A$1:$B$9,2,0)"
    with pytest.raises(ValueError, match="used but not set"):
        Resolver(globals={"g": None}).resolve("{g}")


def test_resolver_unset_global():
    assert Resolver(globals={"item_qty": "$G$2"}).resolve("{item_qty}") == "$G$2"
    with pytest.raises(ValueError, match="used but not set"):
        Resolver(globals={"rate_table": None}).resolve("{rate_table}")


# --- an item's recipe must resolve in the chosen pack, or fail loudly --------
# (fix-round-1: a bare KeyError named neither the item nor the pack.)

def test_unknown_item_recipe_names_the_item_recipe_and_pack():
    """A synthetic spec (never a real one) whose item names a recipe the
    resolved pack does not define must raise RecipeError, not a bare
    KeyError -- and the message must carry the item identifier, the recipe
    name, and the pack name, so a misconfigured estimator can actually act
    on it."""
    from estimating_tools.estimate_spec.build import generate
    from estimating_tools.estimate_spec.recipe import RecipeError

    demo = load_pack(LEGACY_DIR)
    config = {
        "project": "Synthetic",
        "bid_date": "2026-01-01",
        "markup": 0,
        "rate_table": [],
        "items": [
            {"recipe": "not_a_real_recipe", "sheet": "SHEET1",
             "item_no": "SPEC-ITEM-99", "description": "Synthetic item", "uom": "LS"},
        ],
    }
    with pytest.raises(RecipeError) as exc_info:
        generate(config, demo)
    msg = str(exc_info.value)
    assert "SPEC-ITEM-99" in msg
    assert "SHEET1" in msg
    assert "not_a_real_recipe" in msg
    assert demo.name in msg


# --- a workbook-scope global promoted twice must fail loudly ----------------
# (the recorded "last-sheet {rate_table} wins" quirk: promoting the same
# workbook-scope global from two sheets used to silently keep only the last
# promotion, with no error at all.)

def _rate_host_config(*item_nos):
    """A synthetic config (demo pack) whose items all use the `setup` recipe,
    which hosts the `rates` unit that promotes the workbook-scope
    `rate_lookup` global. One item promotes it once; more than one promotes
    it more than once."""
    return {
        "project": "Synthetic",
        "bid_date": "2026-01-01",
        "markup": 0,
        "rate_table": [{"trade": "FITTER", "rate": 50}],
        "items": [
            {"recipe": "setup", "sheet": f"SHEET{i}", "item_no": no,
             "description": f"item {no}", "uom": "LS"}
            for i, no in enumerate(item_nos, start=1)
        ],
    }


def test_a_global_promoted_once_still_works():
    """A pack that promotes a workbook-scope global exactly once (the normal
    case: one rate-host item) renders without incident."""
    from estimating_tools.estimate_spec.build import generate

    demo = load_pack(LEGACY_DIR)
    wb = generate(_rate_host_config("ITEM-1"), demo)
    assert "Rates" in wb.sheetnames


def test_a_global_promoted_twice_names_both_origins_and_the_pack():
    """_promote_globals must refuse a second promotion of the same
    workbook-scope global, naming the global, both (sheet, unit/key) origins
    involved, and the pack -- the same error style build.py already uses for
    an unknown recipe (naming item, recipe, pack, what the pack offers)."""
    from estimating_tools.estimate_spec.build import EstimateBuilder
    from estimating_tools.estimate_spec.recipe import RecipeError

    demo = load_pack(LEGACY_DIR)
    builder = EstimateBuilder({"project": "Synthetic", "items": []}, demo)

    builder.cur_exports = {"rates.table": "$A$2:$B$2"}
    builder._promote_globals("rates", "SheetA")
    assert builder.book_globals["rate_lookup"] == "'SheetA'!$A$2:$B$2"

    builder.cur_exports = {"rates.table": "$A$5:$B$5"}
    with pytest.raises(RecipeError) as exc_info:
        builder._promote_globals("rates", "SheetB")
    msg = str(exc_info.value)
    assert "rate_lookup" in msg
    assert "SheetA" in msg
    assert "SheetB" in msg
    assert demo.name in msg


def test_a_global_promoted_twice_via_two_items_fails_end_to_end():
    """The same guard, exercised through a full render: two items both
    using the rate-hosting recipe must fail generate(), not silently render
    with the second item's rate table winning."""
    from estimating_tools.estimate_spec.build import generate
    from estimating_tools.estimate_spec.recipe import RecipeError

    demo = load_pack(LEGACY_DIR)
    with pytest.raises(RecipeError) as exc_info:
        generate(_rate_host_config("ITEM-1", "ITEM-2"), demo)
    msg = str(exc_info.value)
    assert "rate_lookup" in msg
    assert demo.name in msg


# =============================================================================
# {ref:SHEET!NAME} — the opt-in cross-sheet reference.
#
# Resolution is a single deferred pass: every sheet is laid out (publishing
# its `exports:`) before any cell is resolved, so a ref may name a sheet that
# comes earlier OR later in the workbook. These first tests are pure — the
# Resolver alone, no pack — and pin the syntax, the address it produces, and
# the fact that a sheet-local ref is unaffected by any of it.
# =============================================================================
def test_cross_sheet_ref_is_a_sheet_qualified_address():
    r = Resolver(exports={}, book_exports={"Crew Rates": {"rate_sheet.table": "$A$2:$B$6"}})
    assert r.resolve("={ref:Crew Rates!rate_sheet.table}") == "='Crew Rates'!$A$2:$B$6"


def test_cross_sheet_ref_accepts_the_quoted_spelling():
    """A pack author who writes the sheet the way Excel does must get the
    same address as one who writes it bare."""
    book = {"Crew Rates": {"rate_sheet.table": "$A$2:$B$6"}}
    bare = Resolver(exports={}, book_exports=book).resolve("={ref:Crew Rates!rate_sheet.table}")
    quoted = Resolver(exports={}, book_exports=book).resolve("={ref:'Crew Rates'!rate_sheet.table}")
    assert bare == quoted == "='Crew Rates'!$A$2:$B$6"


def test_the_sheets_own_namespace_wins():
    """The local namespace is consulted first, so no key a pack can write
    today ever reaches the cross-sheet path."""
    r = Resolver(exports={"takeoff.total_cy": "$D$19"},
                 book_exports={"Other": {"takeoff.total_cy": "$Z$99"}})
    assert r.resolve("={ref:takeoff.total_cy}") == "=$D$19"


def test_a_bare_unknown_ref_still_raises_the_same_keyerror():
    """A key with no sheet separator behaves exactly as it always has."""
    r = Resolver(exports={}, book_exports={"Crew Rates": {}})
    with pytest.raises(KeyError, match="unknown ref"):
        r.resolve("={ref:takeoff.total_cy}")


def test_cross_sheet_ref_to_an_unknown_sheet_names_ref_site_and_pack():
    from estimating_tools.estimate_spec.build import CrossSheetRefError

    r = Resolver(exports={}, book_exports={"Crew Rates": {"rate_sheet.table": "$A$2:$B$6"}},
                 where="sheet 'PROBE' cell B7 (unit 'crosslink')", pack="alt")
    with pytest.raises(CrossSheetRefError) as exc_info:
        r.resolve("={ref:NO SUCH SHEET!item_total.grand_total}")
    msg = str(exc_info.value)
    assert "{ref:NO SUCH SHEET!item_total.grand_total}" in msg   # the reference
    assert "sheet 'PROBE' cell B7 (unit 'crosslink')" in msg     # the referring site
    assert "'alt'" in msg                                        # the pack
    assert "Crew Rates" in msg                                   # what there IS
    assert "globals:" in msg                                     # and the other tool


def test_cross_sheet_ref_to_an_unknown_export_lists_what_the_sheet_offers():
    from estimating_tools.estimate_spec.build import CrossSheetRefError

    r = Resolver(exports={}, book_exports={"Crew Rates": {"rate_sheet.table": "$A$2:$B$6"}},
                 where="sheet 'PROBE' cell B7 (unit 'crosslink')", pack="alt")
    with pytest.raises(CrossSheetRefError) as exc_info:
        r.resolve("={ref:Crew Rates!nothing_here}")
    msg = str(exc_info.value)
    assert "{ref:Crew Rates!nothing_here}" in msg
    assert "'nothing_here'" in msg
    assert "rate_sheet.table" in msg
    assert "'alt'" in msg


def test_a_cross_sheet_error_is_not_swallowed_as_an_optional_cell():
    """CrossSheetRefError must not be catchable as the ordinary KeyError /
    ValueError the `optional:` escape hatch swallows -- but it stays a
    RecipeError, so a caller already handling pack errors still sees it."""
    from estimating_tools.estimate_spec.build import CrossSheetRefError
    from estimating_tools.estimate_spec.recipe import RecipeError

    assert issubclass(CrossSheetRefError, RecipeError)


# --- the same thing through a real render of the alt pack -------------------
# The alt pack carries a `crosslink_probe` recipe, authored for these tests:
# one sheet that reads the rate sheet (laid out BEFORE it) and a later bid
# item's TOTAL (laid out AFTER it).

def _probe_config(later_sheet="3 - HANGER SET"):
    cfg = _cfg()
    cfg["items"].insert(1, {
        "recipe": "crosslink_probe", "sheet": "PROBE", "item_no": "Item 1a",
        "description": "cross-sheet probe", "qty": 1, "uom": "LS",
        "x-later-sheet": later_sheet,
    })
    return cfg


@pytest.fixture(scope="module")
def probe_wb(alt):
    from estimating_tools.estimate_spec.build import generate
    return generate(_probe_config(), alt)


def test_probe_sheet_lands_before_the_sheet_it_references(probe_wb):
    assert probe_wb.sheetnames == ["Recap", "PRE", "Crew Rates", "PROBE",
                                   "2 - DUCT RUN", "3 - HANGER SET"]


def test_cross_sheet_ref_backward_to_an_earlier_sheet(probe_wb):
    ws = probe_wb["PROBE"]
    assert ws["B5"].value == "=ROWS('Crew Rates'!$A$2:$B$6)"    # sheet named by a param
    assert ws["B6"].value == "=ROWS('Crew Rates'!$A$2:$B$6)"    # sheet named inline, quoted
    # and that really is where the rate table landed
    assert probe_wb["Crew Rates"]["A2"].value == "FITTER"
    assert probe_wb["Crew Rates"]["A6"].value == "HELPER"


def test_cross_sheet_ref_forward_to_a_sheet_laid_out_later(probe_wb):
    """PROBE is laid out third, '3 - HANGER SET' fifth: the reference still
    resolves, because every cell is resolved after ALL layout is done."""
    ws = probe_wb["PROBE"]
    assert ws["B7"].value == "='3 - HANGER SET'!$G$40"
    later = probe_wb["3 - HANGER SET"]
    assert later["E40"].value == "TOTAL"        # $G$40 is that sheet's TOTAL cell
    assert str(later["G40"].value).startswith("=SUM(")


def test_an_if_ref_guard_reads_the_other_sheets_namespace_too(probe_wb):
    ws = probe_wb["PROBE"]
    assert ws["B8"].value == "=ROWS('Crew Rates'!$A$2:$B$6)"   # target published -> rendered
    assert ws["B9"].value is None                              # target absent -> skipped


def test_an_unresolvable_cross_sheet_ref_stops_the_render():
    from estimating_tools.estimate_spec.build import CrossSheetRefError, generate

    pack = load_pack(ALT_DIR)
    with pytest.raises(CrossSheetRefError) as exc_info:
        generate(_probe_config(later_sheet="NO SUCH SHEET"), pack)
    msg = str(exc_info.value)
    assert "{ref:NO SUCH SHEET!item_total.grand_total}" in msg
    assert "PROBE" in msg and "crosslink" in msg
    assert pack.name in msg


def test_a_cross_sheet_ref_to_a_real_sheet_that_lacks_the_export_stops_it_too():
    from estimating_tools.estimate_spec.build import CrossSheetRefError, generate

    pack = load_pack(ALT_DIR)
    # 'Crew Rates' IS laid out, but it publishes rate_sheet.table, not a TOTAL.
    with pytest.raises(CrossSheetRefError) as exc_info:
        generate(_probe_config(later_sheet="Crew Rates"), pack)
    msg = str(exc_info.value)
    assert "'item_total.grand_total'" in msg
    assert "rate_sheet.table" in msg


def test_optional_does_not_turn_an_unresolvable_cross_sheet_ref_into_a_blank():
    """`optional: true` swallows an unresolved placeholder and leaves the cell
    empty. It must not do that to a cross-sheet ref -- a blank cell there is a
    silent zero in someone's bid."""
    from estimating_tools.estimate_spec.build import CrossSheetRefError, generate

    pack = load_pack(ALT_DIR)          # load_pack is uncached: this copy is ours
    pack.units["crosslink"]["rows"][3]["cells"]["B"]["optional"] = True
    with pytest.raises(CrossSheetRefError):
        generate(_probe_config(later_sheet="NO SUCH SHEET"), pack)


def test_the_declared_globals_path_is_untouched(wb, probe_wb):
    """The workbook-scope global (`crew_rate_lookup`, qualify: sheet) still
    produces its address the way it always did, in a workbook with cross-sheet
    refs and in one without."""
    for book in (wb, probe_wb):
        vlookups = [c.value for row in book["2 - DUCT RUN"].iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=VLOOKUP")]
        assert vlookups
        assert all("'Crew Rates'!$A$2:$B$6" in f for f in vlookups)


# --- fix round 1: an apostrophe in a sheet name ------------------------------
# `sheet_qualified` used to emit the half-escaped `'Bob's Sheet'!$A$1`, an
# invalid address that fails only when the file is opened, while
# `split_sheet_ref` already undoubled on the way in. Escaping (not refusing)
# restores the symmetry and makes the case work.

def test_a_quoted_sheet_name_doubles_an_apostrophe():
    from estimating_tools.estimate_spec.build import quoted_sheet, sheet_qualified

    assert quoted_sheet("Crew Rates") == "'Crew Rates'"
    assert quoted_sheet("Bob's Sheet") == "'Bob''s Sheet'"
    assert sheet_qualified("Bob's Sheet", "$A$1") == "'Bob''s Sheet'!$A$1"


def test_an_apostrophe_sheet_name_round_trips():
    from estimating_tools.estimate_spec.build import sheet_qualified, split_sheet_ref

    for name in ("Crew Rates", "Bob's Sheet", "'", "a'b'c"):
        assert split_sheet_ref(sheet_qualified(name, "$A$1")) == (name, "$A$1")


def test_a_rendered_formula_against_an_apostrophe_sheet_is_valid():
    """Every address built from a sheet name -- the cross-sheet ref, the
    cover's per-item total and its sheet_ref -- carries the doubled quote."""
    from estimating_tools.estimate_spec.build import generate

    odd = "Bob's Shop Drawings"
    cfg = _probe_config(later_sheet=odd)
    cfg["items"][-1]["sheet"] = odd
    book = generate(cfg, load_pack(ALT_DIR))

    assert odd in book.sheetnames
    assert book["PROBE"]["B7"].value == "='Bob''s Shop Drawings'!$G$40"
    recap = [c.value for row in book["Recap"].iter_rows() for c in row
             if isinstance(c.value, str) and "Shop Drawings" in c.value]
    assert recap and all(f.startswith("='Bob''s Shop Drawings'!") for f in recap)


# --- fix round 1: the `exports:` template is the other broad catch -----------
def test_a_cross_sheet_ref_in_an_exports_template_is_not_swallowed():
    """layout_unit resolves a unit's own `exports:` templates behind a broad
    `except (KeyError, ValueError): pass` (a landmark dropped by a `when:` row
    simply does not publish). A cross-sheet ref written there must not vanish
    into it -- the export would go unpublished and the real error would surface
    later, somewhere else, as an unrelated unknown ref."""
    from estimating_tools.estimate_spec.build import CrossSheetRefError, generate

    pack = load_pack(ALT_DIR)          # load_pack is uncached: this copy is ours
    pack.units["crosslink"]["exports"]["probe_total"] = "{ref:Crew Rates!rate_sheet.table}"
    with pytest.raises(CrossSheetRefError) as exc_info:
        generate(_probe_config(), pack)
    msg = str(exc_info.value)
    assert "unit 'crosslink' on sheet 'PROBE' (exports:)" in msg
    assert "{ref:Crew Rates!rate_sheet.table}" in msg
    assert "publishes an address on its OWN sheet" in msg


# --- in_summary: a shared-cost sheet pushed into the items that pay for it ---
# The sheet carries cost the bid must pay once. Its TOTAL is referenced by the
# paying items (a cross-sheet {ref:}), so if the cover ALSO listed it, every
# dollar on it would reach the bid twice.
def _shared_cost_config():
    """The alt sample, with its last item marked as a shared-cost sheet."""
    cfg = _cfg()
    cfg["items"][-1]["in_summary"] = False
    return cfg


def test_the_cover_skips_a_sheet_marked_out_of_the_summary(alt):
    from estimating_tools.estimate_spec.build import generate

    shared = _cfg()["items"][-1]["sheet"]
    listed = generate(_shared_cost_config(), alt)["Recap"]
    rows = [c.value for row in listed.iter_rows() for c in row
            if isinstance(c.value, str) and shared in c.value]
    assert rows == [], f"{shared!r} is pushed into its paying items, not summed again"


def test_the_skipped_sheet_is_still_rendered_and_still_totals(alt):
    """Out of the SUMMARY, not out of the workbook: the paying items reach it
    with a cross-sheet ref, so it must exist and must publish its total."""
    from estimating_tools.estimate_spec.build import generate

    shared = _cfg()["items"][-1]["sheet"]
    book = generate(_shared_cost_config(), alt)
    assert shared in book.sheetnames
    plain = generate(_cfg(), alt)
    assert ([c.value for row in book[shared].iter_rows() for c in row]
            == [c.value for row in plain[shared].iter_rows() for c in row])


def test_in_summary_true_is_the_default_and_changes_nothing(alt):
    from estimating_tools.estimate_spec.build import generate

    cfg = _cfg()
    cfg["items"][-1]["in_summary"] = True
    assert (_cells_of(generate(cfg, alt)) == _cells_of(generate(_cfg(), alt)))


def _cells_of(wb):
    return {(ws.title, c.coordinate): c.value
            for ws in wb.worksheets for row in ws.iter_rows() for c in row}
