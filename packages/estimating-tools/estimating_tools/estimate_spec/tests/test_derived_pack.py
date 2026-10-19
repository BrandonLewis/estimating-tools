"""`derived_from:` -- a pack that keeps another pack's layout contract.

An org fork of fieldbook (its own name, trades, seed) declares
`derived_from: fieldbook`: it inherits fieldbook's `layout_map` when it
declares none, its workbook's `est.layout` says what it derives from, and
`readback_family` lets read-back know the workbook's grid is fieldbook's.
"""
from __future__ import annotations

import json

import pytest

from estimating_tools.estimate_spec import build as B
from estimating_tools.estimate_spec import format_pack as fp
from estimating_tools.estimate_spec.tests._fieldbook import (
    DERIVED_NAME, derived_pack_dir, example_spec, fieldbook_pack)
from estimating_tools.estimate_spec.tests._packs import LEGACY_DIR


def _layout(wb) -> dict:
    sheet, ref = wb.defined_names["est.layout"].attr_text.split("!")
    return json.loads(wb[sheet.strip("'")][ref.replace("$", "")].value)


def test_derived_from_is_loaded(tmp_path):
    pack = fp.load_pack(derived_pack_dir(tmp_path))
    assert pack.name == DERIVED_NAME
    assert pack.derived_from == "fieldbook"


def test_a_pack_that_declares_nothing_derives_from_nothing():
    assert fieldbook_pack().derived_from is None
    assert fp.load_pack(LEGACY_DIR).derived_from is None


@pytest.mark.parametrize("bad", ["", 1, ["fieldbook"], True])
def test_a_derived_from_that_is_not_a_name_is_refused(tmp_path, bad):
    root = derived_pack_dir(tmp_path, format_edit=lambda m: m.update(derived_from=bad))
    with pytest.raises(fp.FormatPackError, match="derived_from must be a non-empty string"):
        fp.load_pack(root)


def test_an_unknown_base_is_refused_naming_the_supported_ones(tmp_path):
    root = derived_pack_dir(tmp_path, derived_from="nosuchpack")
    with pytest.raises(fp.FormatPackError, match=r"derived_from 'nosuchpack'.*\['fieldbook'\]"):
        fp.load_pack(root)


def test_readback_family_is_the_base_or_the_packs_own_name(tmp_path):
    assert fp.readback_family(fp.load_pack(derived_pack_dir(tmp_path))) == "fieldbook"
    assert fp.readback_family(fieldbook_pack()) == "fieldbook"
    assert fp.readback_family(fp.load_pack(LEGACY_DIR)) == "legacy_demo"


def test_a_derived_pack_without_a_layout_map_inherits_the_bases(tmp_path):
    pack = fp.load_pack(derived_pack_dir(tmp_path))
    assert pack.layout_map == fieldbook_pack().layout_map


def test_a_derived_packs_own_layout_map_wins(tmp_path):
    def edit(meta):
        meta["layout_map"]["templates"]["item"]["description"] = "Our own template rows"
    pack = fp.load_pack(derived_pack_dir(tmp_path, keep_layout_map=True, format_edit=edit))
    assert pack.layout_map["templates"]["item"]["description"] == "Our own template rows"
    assert fieldbook_pack().layout_map["templates"]["item"]["description"] != \
        "Our own template rows"


def test_an_inherited_layout_map_is_the_packs_own_copy(tmp_path):
    pack = fp.load_pack(derived_pack_dir(tmp_path))
    pack.layout_map["tags"]["kind"] = "Z"
    assert fp.load_pack(derived_pack_dir(tmp_path)).layout_map["tags"]["kind"] == "Q"
    assert fieldbook_pack().layout_map["tags"]["kind"] == "Q"


def test_a_derived_pack_builds_a_navigator_workbook_stamped_with_its_own_name(tmp_path):
    pack = fp.load_pack(derived_pack_dir(tmp_path))
    wb = B.generate(example_spec(), pack)
    props = {p.name: p.value for p in wb.custom_doc_props.props}
    assert props["estimate_format"] == f"{DERIVED_NAME} 2 / fieldbook 2"
    assert "estimate_layout" in props
    assert "_templates" in wb.sheetnames
    assert {"est.index", "est.layout", "est.direct_total"} <= set(wb.defined_names)
    layout = _layout(wb)
    assert layout["pack"] == DERIVED_NAME
    assert layout["derived_from"] == "fieldbook"


def test_a_pack_that_derives_from_nothing_stamps_no_derived_from():
    assert "derived_from" not in _layout(B.generate(example_spec(), fieldbook_pack()))


# --- an inherited map must fit the pack that inherits it --------------------

@pytest.mark.parametrize("old, new", [
    ("cost_summary_total", "summary_total"),     # cells.direct_total's unit
    ("checks_head", "checks_top"),               # tables.checks' `after` anchor
    ("code_rows", "code_list"),                  # tables.codes' `span` anchor
])
def test_an_inherited_layout_map_that_does_not_fit_is_refused_at_load(tmp_path, old, new):
    root = derived_pack_dir(tmp_path, replace={old: new})
    with pytest.raises(fp.FormatPackError) as exc:
        fp.load_pack(root)
    message = str(exc.value)
    assert "layout_map inherited from fieldbook does not fit this pack" in message
    assert old in message
    assert "declare your own layout_map:" in message


def test_a_renamed_index_row_key_is_refused_at_load(tmp_path):
    root = derived_pack_dir(tmp_path, replace={'key: "cs[{it.item_no}]"': 'key: "row[{it.item_no}]"',
                                               "Cost Summary!cs[": "Cost Summary!row["})
    with pytest.raises(fp.FormatPackError, match=r"does not fit this pack.*cs\[\*\]\.bid"):
        fp.load_pack(root)


def test_a_declared_map_that_does_not_fit_is_a_format_pack_error_at_build(tmp_path, capsys,
                                                                       monkeypatch):
    from estimating_tools.estimate_spec import cli
    from estimating_tools.estimate_spec.tests._fieldbook import EXAMPLE
    monkeypatch.delenv("ESTIMATING_PROFILES", raising=False)

    def edit(meta):
        meta["layout_map"]["cells"]["direct_total"] = "Cost Summary!nothing.direct"
    root = derived_pack_dir(tmp_path, keep_layout_map=True, format_edit=edit)
    fp.load_pack(root)                       # its own map: not checked at load
    out = tmp_path / "o.xlsx"
    assert cli.main(["build", str(EXAMPLE), "-o", str(out), "--format", str(root)]) == 2
    err = capsys.readouterr().err
    assert "format pack error" in err and "invalid spec" not in err


# --- `--format fieldbook`: a base pack by name -------------------------------

def test_format_takes_a_base_packs_name(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)              # no ./fieldbook directory to find instead
    pack = fp.resolve_pack("fieldbook")
    assert pack.name == "fieldbook" and pack.chosen
    assert pack.root == fp.BUNDLED.resolve()


def test_a_directory_named_like_a_base_still_wins(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    mine = derived_pack_dir(tmp_path, name="mine")
    mine.rename(tmp_path / "fieldbook")
    assert fp.resolve_pack("fieldbook").name == "mine"


def test_build_with_format_fieldbook(monkeypatch, tmp_path):
    from estimating_tools.estimate_spec import cli
    from estimating_tools.estimate_spec.tests._fieldbook import EXAMPLE
    monkeypatch.chdir(tmp_path)
    out = tmp_path / "o.xlsx"
    assert cli.main(["build", str(EXAMPLE), "-o", str(out), "--format", "fieldbook"]) == 0
    from openpyxl import load_workbook
    props = {p.name: p.value for p in load_workbook(out).custom_doc_props.props}
    assert props["estimate_format"] == "fieldbook 2"


def test_an_extending_pack_inherits_the_family_and_its_own_layout_copy():
    from estimating_tools.estimate_spec.tests._fieldbook import DERIVED
    pack = fp.load_pack(DERIVED)
    base = fieldbook_pack()
    assert pack.derived_from == fp.readback_family(pack) == "fieldbook"
    assert pack.layout_map == base.layout_map and pack.layout_map is not base.layout_map
    wb = B.generate(example_spec(), pack)
    assert _layout(wb)["derived_from"] == "fieldbook"
