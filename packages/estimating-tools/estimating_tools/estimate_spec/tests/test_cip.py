"""cip.py — CIP concrete-structure buildup workbook generator (R7 Task 1).

Port of Work/Tools/structure_estimate.py. Fixture dicts below are transcribed
from real bid configs (cited per fixture); tests never read the Work tree at
runtime. Every derived cell must be a live Excel formula; openpyxl cannot
evaluate them, so assertions pin the emitted formula STRINGS and the yellow
input literals that feed them.
"""
from __future__ import annotations

import json
import zipfile

import openpyxl
import pytest

from estimating_tools.estimate_spec.cip import (
    CIP_BUILD_STAMP, DEFAULT_LABOR, DEFAULT_RATES, SAMPLE, build)

# Rates! rows are DEFAULT_RATES order, first label on row 2.
_RATE_ROW = {lab: i + 2 for i, (lab, _v) in enumerate(DEFAULT_RATES)}


def _find(ws, label):
    """First row whose column-A value equals label (exact)."""
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == label:
            return r
    raise AssertionError(f"label {label!r} not found on sheet {ws.title!r}")


def _col_a(ws):
    return [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]


def _build(config, tmp_path, name="out.xlsx"):
    out = tmp_path / name
    build(config, str(out))
    return openpyxl.load_workbook(str(out))


# =============================================================================
# Fixture configs (transcribed, load-bearing parts only)
# =============================================================================

# Source: the outfall fixture (values derived in-repo), modelled on an
#         outfall_structures.config.json — "S-280 Outfall ea (Sched E)":
# component mode, rebar_cost+rebar_vendor shorthand, per-phase days dict
# zeroing E/F, count: 2 (per-EACH sheet, bid pays 2 EA).
OUTFALL_CONFIG = {
    "title": "Outfall Fixture - Storm Drain Outfall Structures (CIP) Estimate",
    "basis_note": "transcribed fixture; see outfall_structures.config.json",
    "structures": [
        {
            "sheet": "S-280 Outfall ea (Sched E)",
            "title": "SACRAMENTO S-280 OUTFALL - Schedule E (per EACH; bid pays 2 EA)",
            "bid_item": "Sched E Item 12 - Outfall Structure (2 EA)",
            "note": "24in SD outfalls. Class A 4000psi 6-sack.",
            "count": 2,
            "components": [
                {"name": "Footing", "count": 1, "L": 7.83, "W": 6, "H": 0.67},
                {"name": "Headwall", "count": 1, "L": 5.84, "W": 0.67, "H": 5.84},
                {"name": "Wing walls", "count": 2, "L": 5.84, "W": 0.67, "H": 2.92},
            ],
            "rebar_cost": 400,
            "rebar_vendor": "Steel Sub (quoted)",
            "days": {"A": 1.0, "B": 0.25, "C": 1.5, "D": 0.5, "E": 0, "F": 0, "G": 0.5},
            "pickup_hr": 26, "tools_hr": 26, "skytrak_hr": 26, "pump_hr": 6,
            "short_loads": 1, "panel_hrs": 8,
            "fuel_pickup_hr": 26, "fuel_equip_hr": 26, "waterstop_lf": 0,
        },
    ],
}

# Source: the headwall fixture (values derived in-repo), modelled on a
#         headwalls_cutoffs.config.json — top-level `rates` override,
# full 5-phase labor_sections override, explicit `subs` (rebar allowance
# @ $280/CY PENDING QUOTE), "Elder Headwalls" structure.
HEADWALL_LABOR = [
    {"id": "A", "name": "[A] Layout / dowels (drill & epoxy) / fine grade",
     "fore": 1, "carp": 1, "lab": 2, "mason": 0, "days": 1.5},
    {"id": "B", "name": "[B] Forming (wall faces / pad edge forms)",
     "fore": 1, "carp": 2, "lab": 1, "mason": 0, "days": 3.0},
    {"id": "C", "name": "[C] Rebar, embeds & guardrail blockouts",
     "fore": 1, "carp": 1, "lab": 1, "mason": 0, "days": 1.5},
    {"id": "D", "name": "[D] Pour & finish",
     "fore": 1, "carp": 1, "lab": 2, "mason": 1, "days": 1.0},
    {"id": "E", "name": "[E] Strip, patch & clean-up",
     "fore": 1, "carp": 1, "lab": 1, "mason": 0, "days": 1.5},
]
HEADWALL_CONFIG = {
    "title": "Headwall Fixture - CIP Scope (Headwalls, Cutoff Walls) Estimate",
    "basis_note": "transcribed fixture; see headwalls_cutoffs.config.json",
    "rates": {"Concrete $/CY": 300, "Rebar $/CY": 0},
    "labor_sections": HEADWALL_LABOR,
    "structures": [
        {
            "sheet": "Elder Headwalls",
            "title": "ELDER CREEK CROSSING - CULVERT HEADWALLS (U/S + D/S)",
            "bid_item": "Elder - Culvert Headwalls (25 CY)",
            "components": [
                {"name": "Headwall U/S + D/S (1' thick)",
                 "count": 2, "L": 36.5, "W": 1.0, "H": 9.25},
            ],
            "days": {"A": 1.5, "B": 4.0, "C": 1.5, "D": 1.0, "E": 1.5},
            "pickup_hr": 76, "tools_hr": 76, "skytrak_hr": 76, "pump_hr": 8,
            "short_loads": 2, "panel_hrs": 16,
            "fuel_pickup_hr": 76, "fuel_equip_hr": 76,
            "subs": [
                {"name": "Rebar F&I + dowels (allowance)",
                 "vendor": "Rebar sub - PENDING QUOTE",
                 "qty": 25, "unit": "CY", "unit_price": 280, "kind": "rebar"},
            ],
        },
    ],
}

# Source: the box-culvert fixture (values derived in-repo), modelled on a
#         headwalls.config.json — "Culvert 1 (4)-12x5": NEGATIVE-count
# component prisms (box-opening deductions) and kind:"earthwork" subs.
BOX_CULVERT_COMPONENTS = [
    {"name": "Headwall wall (61.16 x 1.0 x 8.5)", "count": 1, "L": 61.16, "W": 1.0, "H": 8.5},
    {"name": "LESS (4) 12x5 box openings", "count": -4, "L": 12.0, "W": 1.0, "H": 5.0},
    {"name": "Headwall cutoff wall (61.16 x 1.0 x 4.0)", "count": 1, "L": 61.16, "W": 1.0, "H": 4.0},
    {"name": "Wingwall 1 footing (Sec B, 7'-9\" w)", "count": 1, "L": 14.25, "W": 7.75, "H": 1.16},
    {"name": "Wingwall 1 footing cutoff", "count": 1, "L": 14.25, "W": 1.5, "H": 3.0},
    {"name": "Wingwall 2 footing (Sec C, 4'-6\" w)", "count": 1, "L": 14.25, "W": 4.5, "H": 1.0},
    {"name": "Wingwall 2 footing cutoff", "count": 1, "L": 14.25, "W": 1.5, "H": 3.0},
    {"name": "Wingwall 1 stem (H 9.45)", "count": 1, "L": 14.25, "W": 1.0, "H": 9.45},
    {"name": "Wingwall 2 stem (H 8.0)", "count": 1, "L": 14.25, "W": 1.0, "H": 8.0},
]
BOX_CULVERT_CONFIG = {
    "title": "Box-Culvert Fixture - Box-Culvert Headwalls & Wingwalls (CIP)",
    "basis_note": "transcribed fixture; see headwalls.config.json",
    "labor_sections": [
        {"id": "A", "name": "[A] Layout / dowels / fine grade",
         "fore": 1, "carp": 1, "lab": 2, "mason": 0, "days": 2.0},
        {"id": "B", "name": "[B] Forming",
         "fore": 1, "carp": 2, "lab": 1, "mason": 0, "days": 5.0},
        {"id": "C", "name": "[C] Rebar, embeds & drains",
         "fore": 1, "carp": 1, "lab": 1, "mason": 0, "days": 2.0},
        {"id": "D", "name": "[D] Pour & finish",
         "fore": 1, "carp": 1, "lab": 2, "mason": 1, "days": 1.5},
        {"id": "E", "name": "[E] Strip, patch & clean-up",
         "fore": 1, "carp": 1, "lab": 1, "mason": 0, "days": 2.0},
    ],
    "structures": [
        {
            "sheet": "Culvert 1 (4)-12x5",
            "title": "CULVERT 1 - (4)-BARREL 12'x5' BOX - CIP HEADWALL + WINGWALLS",
            "bid_item": "STY - Culvert 1 Headwall & Wingwalls (CIP) ~40.5 CY",
            "components": BOX_CULVERT_COMPONENTS,
            "days": {"A": 2.0, "B": 5.0, "C": 2.0, "D": 1.5, "E": 2.0},
            "pickup_hr": 100, "tools_hr": 100, "skytrak_hr": 100, "pump_hr": 12,
            "short_loads": 3, "panel_hrs": 20,
            "fuel_pickup_hr": 100, "fuel_equip_hr": 100,
            "subs": [
                {"name": "Rebar F&I + dowels (allowance ~140 lb/CY)",
                 "vendor": "Rebar sub - PENDING QUOTE",
                 "qty": 41, "unit": "CY", "unit_price": 300, "kind": "rebar"},
                {"name": "Structure excavation (self-perform allowance)",
                 "vendor": "self-perform - VERIFY qty",
                 "qty": 200, "unit": "CY", "unit_price": 25, "kind": "earthwork"},
            ],
        },
    ],
}


# =============================================================================
# SAMPLE round-trip (closed-box mode)
# =============================================================================

@pytest.fixture(scope="module")
def sample_wb(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("cip_sample")
    return _build(SAMPLE, tmp, "sample.xlsx")


def test_sample_sheets(sample_wb):
    names = sample_wb.sheetnames
    assert "Rates" in names and "Summary" in names
    for s in SAMPLE["structures"]:
        assert s["sheet"] in names


def test_sample_rates_column(sample_wb):
    rs = sample_wb["Rates"]
    for i, (lab, default) in enumerate(DEFAULT_RATES):
        r = i + 2
        assert rs.cell(r, 1).value == lab
        assert rs.cell(r, 2).value == default


def test_sample_crew_and_phase_cost_formulas(sample_wb):
    ws = sample_wb[SAMPLE["structures"][0]["sheet"]]
    r = _find(ws, DEFAULT_LABOR[0][1])          # phase A row
    fore, carp, lab, mason, hrs = (
        _RATE_ROW["Labor FORE $/hr"], _RATE_ROW["Labor CARP $/hr"],
        _RATE_ROW["Labor LAB $/hr"], _RATE_ROW["Labor MASON $/hr"],
        _RATE_ROW["Hrs per day"])
    assert ws.cell(r, 6).value == (
        f"=B{r}*Rates!$B${fore}+C{r}*Rates!$B${carp}"
        f"+D{r}*Rates!$B${lab}+E{r}*Rates!$B${mason}")
    # phase cost = crew $/hr x Rates hrs/day x days cell
    assert ws.cell(r, 8).value == f"=F{r}*Rates!$B${hrs}*G{r}"
    # crew composition + days are yellow input literals
    assert ws.cell(r, 2).value == DEFAULT_LABOR[0][2]
    assert ws.cell(r, 7).value == DEFAULT_LABOR[0][6]


def test_sample_closed_box_wall_formula(sample_wb):
    ws = sample_wb[SAMPLE["structures"][0]["sheet"]]
    L = _find(ws, "Ext. Length")
    W = _find(ws, "Ext. Width")
    H = _find(ws, "Ext. Height")
    tw = _find(ws, "Wall thickness")
    tb = _find(ws, "Base thickness")
    td = _find(ws, "Deck thickness")
    walls = _find(ws, "Walls")
    assert ws.cell(walls, 2).value == (
        f"=2*(($B${L}-$B${tw})+($B${W}-$B${tw}))"
        f"*$B${tw}*($B${H}-$B${tb}-$B${td})")
    # and the waste rollup rides the Rates Waste-fraction cell
    cy = _find(ws, "TOTAL CONCRETE (w/ waste)")
    assert f"Rates!$B${_RATE_ROW['Waste fraction']}" in ws.cell(cy, 2).value


# =============================================================================
# Component mode (outfall fixture)
# =============================================================================

@pytest.fixture(scope="module")
def outfall_wb(tmp_path_factory):
    return _build(OUTFALL_CONFIG, tmp_path_factory.mktemp("cip_outfall"), "outfall.xlsx")


def test_outfall_component_cy_formulas(outfall_wb):
    ws = outfall_wb["S-280 Outfall ea (Sched E)"]
    for comp in OUTFALL_CONFIG["structures"][0]["components"]:
        r = _find(ws, comp["name"])
        assert ws.cell(r, 2).value == comp["count"]
        assert ws.cell(r, 3).value == comp["L"]
        assert ws.cell(r, 4).value == comp["W"]
        assert ws.cell(r, 5).value == comp["H"]
        assert ws.cell(r, 6).value == f"=B{r}*C{r}*D{r}*E{r}/27"


def test_outfall_waste_total_references_rates(outfall_wb):
    ws = outfall_wb["S-280 Outfall ea (Sched E)"]
    sub = _find(ws, "Neat concrete subtotal")
    first = _find(ws, "Footing")
    assert ws.cell(sub, 6).value == f"=SUM(F{first}:F{sub - 1})"
    cy = _find(ws, "TOTAL CONCRETE (w/ waste)")
    assert ws.cell(cy, 2).value == (
        f"=$F${sub}*(1+Rates!$B${_RATE_ROW['Waste fraction']})")


def test_outfall_days_override_zeroes_e_and_f(outfall_wb):
    ws = outfall_wb["S-280 Outfall ea (Sched E)"]
    # DEFAULT_LABOR phase names E and F; days dict zeroes both
    assert ws.cell(_find(ws, DEFAULT_LABOR[4][1]), 7).value == 0   # [E] deck form
    assert ws.cell(_find(ws, DEFAULT_LABOR[5][1]), 7).value == 0   # [F] pour deck
    assert ws.cell(_find(ws, DEFAULT_LABOR[0][1]), 7).value == 1.0  # [A] override


def test_outfall_rebar_shorthand_moves_rebar_to_subs(outfall_wb):
    ws = outfall_wb["S-280 Outfall ea (Sched E)"]
    labels = _col_a(ws)
    assert "SUBCONTRACTORS - SUB SPREAD" in labels
    r = _find(ws, "Rebar - furnish & install")
    assert ws.cell(r, 2).value == "Steel Sub (quoted)"
    assert ws.cell(r, 5).value == 400                       # amount as unit price
    assert ws.cell(r, 8).value == f"=C{r}*E{r}"
    # rebar does NOT appear in MATERIAL
    assert "Rebar (per $/CY basis)" not in labels


def test_outfall_summary_qty_and_extended(outfall_wb):
    sm = outfall_wb["Summary"]
    r = None
    for row in range(1, sm.max_row + 1):
        if sm.cell(row, 2).value == "S-280 Outfall ea (Sched E)":
            r = row
            break
    assert r is not None
    assert sm.cell(r, 6).value == 2                         # Qty (count)
    assert sm.cell(r, 7).value == f"=E{r}*F{r}"             # Extended = Each x Qty


# =============================================================================
# labor_sections override (headwall fixture)
# =============================================================================

@pytest.fixture(scope="module")
def headwall_wb(tmp_path_factory):
    return _build(HEADWALL_CONFIG, tmp_path_factory.mktemp("cip_headwall"), "headwall.xlsx")


def test_headwall_labor_sections_override(headwall_wb):
    ws = headwall_wb["Elder Headwalls"]
    labels = _col_a(ws)
    for sec in HEADWALL_LABOR:
        assert sec["name"] in labels
    # exactly the 5 overridden phase rows: none of the 7 defaults leak in
    for default in DEFAULT_LABOR:
        assert default[1] not in labels
    first = _find(ws, HEADWALL_LABOR[0]["name"])
    sub = _find(ws, "Labor subtotal")
    assert sub - first == len(HEADWALL_LABOR)
    assert ws.cell(sub, 8).value == f"=SUM(H{first}:H{sub - 1})"


def test_headwall_rates_override_lands_on_rates_sheet(headwall_wb):
    rs = headwall_wb["Rates"]
    assert rs.cell(_RATE_ROW["Concrete $/CY"], 2).value == 300
    assert rs.cell(_RATE_ROW["Rebar $/CY"], 2).value == 0


# =============================================================================
# Negative component counts (box-culvert fixture)
# =============================================================================

def test_box_culvert_negative_count_deduction(tmp_path):
    wb = _build(BOX_CULVERT_CONFIG, tmp_path, "box-culvert.xlsx")
    ws = wb["Culvert 1 (4)-12x5"]
    r = _find(ws, "LESS (4) 12x5 box openings")
    assert ws.cell(r, 2).value == -4                        # input literal
    assert ws.cell(r, 6).value == f"=B{r}*C{r}*D{r}*E{r}/27"  # formula unchanged
    first = _find(ws, BOX_CULVERT_COMPONENTS[0]["name"])
    sub = _find(ws, "Neat concrete subtotal")
    assert ws.cell(sub, 6).value == f"=SUM(F{first}:F{sub - 1})"
    # openpyxl cannot evaluate formulas; document the expected net here from
    # the fixture inputs: sum(count*L*W*H)/27 with the -4 opening deduction
    # nets ~40.5 CY (matches the bid_item note "~40.5 CY").
    net_cy = sum(c["count"] * c["L"] * c["W"] * c["H"]
                 for c in BOX_CULVERT_COMPONENTS) / 27
    assert net_cy == pytest.approx(40.5, abs=0.1)
    gross_cy = sum(c["count"] * c["L"] * c["W"] * c["H"]
                   for c in BOX_CULVERT_COMPONENTS if c["count"] > 0) / 27
    assert net_cy < gross_cy                                # deduction took effect
    # earthwork-kind subs land in the sub block too
    assert ws.cell(_find(ws, "Structure excavation (self-perform allowance)"),
                   8).value.startswith("=C")


# =============================================================================
# Sub Spread rollup sheet (headwall fixture)
# =============================================================================

def test_headwall_sub_spread_sheet(headwall_wb):
    ws = headwall_wb["Elder Headwalls"]
    r = _find(ws, "Rebar F&I + dowels (allowance)")
    assert ws.cell(r, 3).value == 25                        # qty input
    assert ws.cell(r, 5).value == 280                       # unit price input
    assert ws.cell(r, 8).value == f"=C{r}*E{r}"             # Amount formula
    ss = headwall_wb["Sub Spread"]
    # first data row cross-references the structure sheet's amount cell
    assert ss.cell(5, 1).value == "Rebar sub - PENDING QUOTE"
    assert ss.cell(5, 6).value == f"='Elder Headwalls'!$H${r}"
    assert ss.cell(5, 7).value == "=E5*F5"


def test_no_sub_spread_sheet_without_subs(sample_wb):
    assert "Sub Spread" not in sample_wb.sheetnames


def test_earthwork_only_subs_keep_rebar_in_material(tmp_path):
    # Deviation from the Work original: ANY subs list used to drop the
    # per-$/CY rebar line from MATERIAL — an earthwork-only subs list
    # silently un-priced the rebar. The port removes rebar from MATERIAL
    # only when it is actually subbed (rebar_cost or kind:"rebar").
    struct = {
        "sheet": "S1", "L": 9.5, "W": 9.5, "H": 11.0,
        "subs": [{"name": "Structure excavation", "vendor": "self-perform",
                  "qty": 100, "unit": "CY", "unit_price": 25,
                  "kind": "earthwork"}],
    }
    wb = _build({"structures": [struct]}, tmp_path, "earthwork.xlsx")
    labels = _col_a(wb["S1"])
    assert "Rebar (per $/CY basis)" in labels               # still self-performed
    assert "SUBCONTRACTORS - SUB SPREAD" in labels          # sub block still built


def test_rebar_kind_sub_removes_rebar_from_material(headwall_wb):
    # the headwall allowance shape: explicit kind:"rebar" sub, no rebar_cost
    assert "Rebar (per $/CY basis)" not in _col_a(headwall_wb["Elder Headwalls"])


# =============================================================================
# Validation — typos must refuse, not silently mis-bid
# =============================================================================

def test_kept_rebar_line_is_live_and_rate_pinned(tmp_path):
    # Review follow-up: the KEPT rebar line must stay live — qty rides the
    # structure's own total-CY cell and the rate the Rates sheet, so tuning
    # either reprices rebar without a rebuild.
    struct = {
        "sheet": "S1", "L": 9.5, "W": 9.5, "H": 11.0,
        "subs": [{"name": "Structure excavation", "vendor": "self-perform",
                  "qty": 100, "unit": "CY", "unit_price": 25,
                  "kind": "earthwork"}],
    }
    ws = _build({"structures": [struct]}, tmp_path, "kept.xlsx")["S1"]
    r = _find(ws, "Rebar (per $/CY basis)")
    assert str(ws.cell(r, 2).value).startswith("=$B$")       # live CY ref
    assert ws.cell(r, 4).value == "=Rates!$B$%d" % _RATE_ROW["Rebar $/CY"]
    assert ws.cell(r, 8).value == "=B%d*D%d" % (r, r)


def test_block_rate_refs_and_selfperform_composition(tmp_path):
    # Review follow-up: pin one EQUIPMENT rate ref (a swapped rate label in a
    # block items list is invisible to label-only tests) and the SELF-PERFORM
    # composition (four block subtotals, no more, no less).
    import re
    ws = _build({"structures": [{"sheet": "S1", "L": 9.5, "W": 9.5,
                                 "H": 11.0}]}, tmp_path, "blocks.xlsx")["S1"]
    r = _find(ws, "Pickup")
    assert ws.cell(r, 4).value == "=Rates!$B$%d" % _RATE_ROW["Pickup $/hr"]
    assert ws.cell(r, 8).value == "=B%d*D%d" % (r, r)
    sp = _find(ws, "SELF-PERFORM SUBTOTAL")
    assert re.fullmatch(r"=\$H\$\d+\+\$H\$\d+\+\$H\$\d+\+\$H\$\d+",
                        str(ws.cell(sp, 8).value))


def test_duplicate_sheet_names_raise(tmp_path):
    # openpyxl would silently rename the second sheet while Summary still
    # references the first — one structure counted twice, the other $0.
    config = {"structures": [
        {"sheet": "S1", "L": 9.5, "W": 9.5, "H": 11.0},
        {"sheet": "S1", "L": 4.0, "W": 4.0, "H": 5.0}]}
    with pytest.raises(ValueError, match="duplicate structure sheet"):
        build(config, str(tmp_path / "x.xlsx"))


def test_component_missing_key_raises(tmp_path):
    config = {"structures": [
        {"sheet": "S1", "components": [{"count": 1, "L": 2, "W": 2, "H": 2}]}]}
    with pytest.raises(ValueError, match="missing key"):
        build(config, str(tmp_path / "x.xlsx"))


def test_priceless_sub_raises(tmp_path):
    # A sub with neither unit_price nor amount prices at $0 — and a
    # kind:"rebar" one also removes the MATERIAL rebar line, silently
    # un-pricing the whole rebar scope.
    config = {"structures": [
        {"sheet": "S1", "L": 9.5, "W": 9.5, "H": 11.0,
         "subs": [{"name": "Rebar F&I", "vendor": "TBD", "qty": 3, "unit": "CY",
                   "kind": "rebar"}]}]}
    with pytest.raises(ValueError, match="unit_price.*or.*amount|'unit_price' or 'amount'"):
        build(config, str(tmp_path / "x.xlsx"))


def test_unknown_top_level_key_raises(tmp_path):
    config = dict(SAMPLE)
    config["labour_sections"] = []                          # typo of labor_sections
    with pytest.raises(ValueError, match="labour_sections"):
        build(config, str(tmp_path / "x.xlsx"))


def test_unknown_structure_key_raises(tmp_path):
    config = {"structures": [
        {"sheet": "S1", "L": 9.5, "W": 9.5, "H": 11.0, "rebar_vender": "X"}]}
    with pytest.raises(ValueError, match="rebar_vender"):
        build(config, str(tmp_path / "x.xlsx"))


def test_structure_without_geometry_raises(tmp_path):
    config = {"structures": [{"sheet": "S1", "title": "no geometry at all"}]}
    with pytest.raises(ValueError, match="S1"):
        build(config, str(tmp_path / "x.xlsx"))


def test_unknown_rate_label_raises(tmp_path):
    config = dict(SAMPLE)
    config["rates"] = {"Concrete $/YD": 300}                # typo of $/CY
    with pytest.raises(ValueError, match="Concrete \\$/YD"):
        build(config, str(tmp_path / "x.xlsx"))


# =============================================================================
# Determinism
# =============================================================================

def test_build_twice_byte_identical(tmp_path):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    build(OUTFALL_CONFIG, str(a))
    build(OUTFALL_CONFIG, str(b))
    assert a.read_bytes() == b.read_bytes()
    # zip member timestamps are pinned to the build stamp (not wall clock)
    with zipfile.ZipFile(str(a)) as z:
        stamps = {i.date_time for i in z.infolist()}
    assert stamps == {CIP_BUILD_STAMP.timetuple()[:6]}


# =============================================================================
# CLI (est-estimate cip-buildup)
# =============================================================================

def test_cli_init_then_build(tmp_path, capsys):
    from estimating_tools.estimate_spec.cli import main
    cfg = tmp_path / "sample.config.json"
    assert main(["cip-buildup", "--init", str(cfg)]) == 0
    config = json.loads(cfg.read_text(encoding="utf-8"))
    assert config == SAMPLE                                 # --init writes SAMPLE
    assert main(["cip-buildup", str(cfg)]) == 0
    out = capsys.readouterr().out
    assert "WROTE:" in out
    # output resolves relative to the config's directory via config["output"]
    assert (tmp_path / SAMPLE["output"]).exists()


def test_cli_output_flag(tmp_path):
    from estimating_tools.estimate_spec.cli import main
    cfg = tmp_path / "c.json"
    cfg.write_text(json.dumps(OUTFALL_CONFIG), encoding="utf-8")
    out = tmp_path / "custom.xlsx"
    assert main(["cip-buildup", str(cfg), "-o", str(out)]) == 0
    assert out.exists()


def test_cli_invalid_config_exits_nonzero(tmp_path, capsys):
    from estimating_tools.estimate_spec.cli import main
    cfg = tmp_path / "bad.json"
    cfg.write_text(json.dumps({"structures": [{"sheet": "S1"}]}), encoding="utf-8")
    assert main(["cip-buildup", str(cfg)]) != 0
    assert not (tmp_path / "bad.xlsx").exists()
