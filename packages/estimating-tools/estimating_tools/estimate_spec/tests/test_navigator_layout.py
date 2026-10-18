"""What `est-estimate build` writes for Estimate Navigator: defined names,
the `estimate_layout` marker, the layout JSON behind `est.layout` and the
hidden `_templates` sheet."""
from __future__ import annotations

import json
import re
import shutil

import pytest
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Font

from estimating_tools.estimate_spec import build as B
from estimating_tools.estimate_spec import format_pack as fp
from estimating_tools.estimate_spec import recipe as R
from estimating_tools.estimate_spec import layout_map as LM
from estimating_tools.estimate_spec.readback import layout as RL
from estimating_tools.estimate_spec.readback import load_workbook as rb_load
from estimating_tools.estimate_spec.readback import read_back
from estimating_tools.estimate_spec.tests._fieldbook import (
    FIELDBOOK_DIR, KIND, build, example_spec, fieldbook_pack, paste_template_rows, task_block)
from estimating_tools.estimate_spec.tests._packs import LEGACY_DIR, LEGACY_EXAMPLE

_PLAIN_REF = re.compile(r"^='((?:[^']|'')+)'!\$[A-Z]{1,3}\$[0-9]+$")


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    path = tmp_path_factory.mktemp("nav") / "built.xlsx"
    B.generate(example_spec(), None, spec_name="demo.estimate.yaml").save(path)
    return load_workbook(path)


def _dest(wb, name):
    ((sheet, ref),) = list(wb.defined_names[name].destinations)
    return sheet, ref


def _props(wb):
    return {p.name: p.value for p in wb.custom_doc_props.props}


def _layout(wb):
    """The layout JSON, read the way the add-in reads it: through `est.layout`."""
    sheet, ref = _dest(wb, "est.layout")
    assert sheet == LM.TEMPLATES_SHEET and ":" not in ref
    return json.loads(wb[sheet][ref.replace("$", "")].value)


def test_the_property_is_a_short_marker(wb):
    # Office.js truncates a custom property past 255 characters on desktop
    # Excel, so the property only says "this workbook has a layout".
    value = _props(wb)["estimate_layout"]
    assert value == str(LM.LAYOUT_VERSION) and len(value) <= 255


def test_the_layout_cell_sits_in_row_1_past_the_templates_content(wb):
    ws = wb[LM.TEMPLATES_SHEET]
    _, ref = _dest(wb, "est.layout")
    assert re.fullmatch(r"\$[A-Z]{1,3}\$1", ref)
    cell = ws[ref.replace("$", "")]
    used = [c.column for row in ws.iter_rows() for c in row
            if c.value is not None and c.coordinate != cell.coordinate]
    assert cell.column == max(used) + 2


def test_the_stamp_describes_the_layout(wb):
    layout = _layout(wb)
    assert layout["pack"] == "fieldbook" and layout["layout_version"] == LM.LAYOUT_VERSION
    assert layout["tables"] == ["checks", "codes", "index"]
    assert layout["tags"] == {"kind": "Q", "code_chain": "R"}
    assert layout["blocks"]["head"] == "TASK" and layout["templates"] is True
    assert layout["spec"] == "demo.estimate.yaml"
    assert "brand" not in layout


def test_the_index_lists_every_item_and_links_to_its_sheet(wb):
    sheet, ref = _dest(wb, "est.index.item_no")
    # The Cost Summary coerces a numeric item number ("1") to a number.
    items = [str(c.value) for (c,) in wb[sheet][ref.replace("$", "")]]
    assert items == [str(it["item_no"]) for it in example_spec()["items"]]
    sheet, ref = _dest(wb, "est.index.link")
    for (cell,) in wb[sheet][ref.replace("$", "")]:
        m = _PLAIN_REF.match(cell.value)
        assert m, cell.value
        assert m.group(1).replace("''", "'") in wb.sheetnames


def test_the_direct_total_is_the_cost_summary_total(wb):
    sheet, ref = _dest(wb, "est.direct_total")
    assert sheet == "Cost Summary" and wb[sheet][ref.replace("$", "")].value.startswith("=SUM(H")


def test_checks_and_codes_cover_their_tables(wb):
    sheet, ref = _dest(wb, "est.checks")
    rows = [[c.value for c in row] for row in wb[sheet][ref.replace("$", "")]]
    # One blank spacer row sits above each rule's group: the table spans
    # them, so a reader skips rows with no rule number.
    checks = [r for r in rows if r[0] is not None]
    assert len(checks) == 25 and len(rows) > len(checks)
    assert all(not any(r) for r in rows if r[0] is None)
    assert all(isinstance(r[5], str) and r[5].startswith("=") for r in checks)
    sheet, ref = _dest(wb, "est.codes.code")
    codes = [c.value for (c,) in wb[sheet][ref.replace("$", "")]]
    assert "02-100" in codes


def test_templates_is_a_hidden_last_sheet_with_blank_inputs(wb):
    assert wb.sheetnames[-1] == LM.TEMPLATES_SHEET
    ws = wb[LM.TEMPLATES_SHEET]
    assert ws.sheet_state == "hidden"
    for name, kind in (("est.tpl.crew", "CREW"), ("est.tpl.line", "LINE")):
        sheet, ref = _dest(wb, name)
        row = int(ref.split(":")[0].lstrip("$"))
        assert sheet == LM.TEMPLATES_SHEET and ws[f"{KIND}{row}"].value == kind
    sheet, ref = _dest(wb, "est.tpl.task")
    first, last = (int(p.lstrip("$")) for p in ref.split(":"))
    assert ws[f"{KIND}{first}"].value == "TASK" and ws[f"{KIND}{last}"].value == "SUBTOTAL"
    crew = int(_dest(wb, "est.tpl.crew")[1].split(":")[0].lstrip("$"))
    assert [ws[f"{c}{crew}"].value for c in "ADE"] == [None, None, None]
    assert [ws[f"{c}{first}"].value for c in "AB"] == [None, None]
    line = int(_dest(wb, "est.tpl.line")[1].split(":")[0].lstrip("$"))
    assert [ws[f"{c}{line}"].value for c in "ABCGHIJ"] == [None] * 7


def test_the_template_item_is_not_an_estimate_item(wb):
    sheet, ref = _dest(wb, "est.index.item_no")
    assert "TEMPLATE" not in [c.value for (c,) in wb[sheet][ref.replace("$", "")]]


def test_a_brand_rides_in_the_stamp(tmp_path):
    brand = {"name": "Example Co", "colors": {"light": {"accent": "#B4531A"}}}
    path = tmp_path / "b.xlsx"
    B.generate(example_spec(), None, brand=brand).save(path)
    assert _layout(load_workbook(path))["brand"] == brand


def test_cli_build_stamps_the_profile_brand_unless_told_not_to(tmp_path, monkeypatch):
    from estimating_tools.estimate_spec import cli
    from estimating_tools.estimate_spec.tests._fieldbook import EXAMPLE, FIELDBOOK_DIR
    prof = tmp_path / "prof"
    prof.mkdir()
    (prof / "profile.yaml").write_text(yaml.safe_dump(
        {"kind": "org", "name": "Example Paving",
         "brand": {"colors": {"light": {"accent": "#B4531A"}}}}), encoding="utf-8")
    monkeypatch.setenv("ESTIMATING_PROFILES", str(prof))
    out = tmp_path / "a.xlsx"
    assert cli.main(["build", str(EXAMPLE), "-o", str(out), "--format", str(FIELDBOOK_DIR)]) == 0
    assert _layout(load_workbook(out))["brand"] == {
        "name": "Example Paving", "colors": {"light": {"accent": "#B4531A"}}}
    out2 = tmp_path / "b.xlsx"
    assert cli.main(["build", str(EXAMPLE), "-o", str(out2), "--format", str(FIELDBOOK_DIR),
                     "--no-brand"]) == 0
    assert "brand" not in _layout(load_workbook(out2))


def test_no_stamp_means_no_navigator_anchors(tmp_path):
    path = tmp_path / "n.xlsx"
    B.generate(example_spec(), None, stamp=False).save(path)
    wb = load_workbook(path)
    assert "estimate_layout" not in _props(wb)
    assert LM.TEMPLATES_SHEET not in wb.sheetnames
    assert not [n for n in wb.defined_names if n.startswith("est.")]


def _pack_without_templates(tmp_path):
    root = tmp_path / "fieldbook-no-templates"
    shutil.copytree(FIELDBOOK_DIR, root)
    fmt = root / "format.yaml"
    doc = yaml.safe_load(fmt.read_text(encoding="utf-8"))
    del doc["layout_map"]["templates"]
    fmt.write_text(yaml.safe_dump(doc, sort_keys=False), encoding="utf-8")
    pack = fp.load_pack(root)
    assert pack.layout_map and "templates" not in pack.layout_map
    return pack


def test_a_layout_map_without_templates_still_gets_a_layout_cell(tmp_path):
    path = tmp_path / "nt.xlsx"
    B.generate(example_spec(), _pack_without_templates(tmp_path)).save(path)
    wb = load_workbook(path)
    assert wb.sheetnames[-1] == LM.TEMPLATES_SHEET
    ws = wb[LM.TEMPLATES_SHEET]
    assert ws.sheet_state == "hidden"
    _, ref = _dest(wb, "est.layout")
    assert [c.coordinate for row in ws.iter_rows() for c in row if c.value is not None]         == [ref.replace("$", "")]
    layout = _layout(wb)
    assert layout["templates"] is False and layout["tables"] == ["checks", "codes", "index"]
    assert not [n for n in wb.defined_names if n.startswith("est.tpl.")]
    result = read_back(example_spec(), rb_load(path))
    assert not result.blocked, result.findings.problems


def test_an_item_on_the_templates_sheet_name_is_refused():
    spec = example_spec()
    spec["items"][1]["sheet"] = "_Templates"
    with pytest.raises(R.RecipeError, match=r"item 2 \(_Templates\).*_templates"):
        B.generate(spec, None)


def test_a_table_ends_at_its_last_value_not_at_a_styled_empty_row():
    builder = B.EstimateBuilder(example_spec(), fieldbook_pack())
    wb = builder.run()
    ws = wb["Checks"]
    last = max(c.row for row in ws.iter_rows() for c in row if c.value is not None)
    ws[f"B{last + 40}"].font = Font(bold=True)        # styled, never typed into
    assert ws.max_row == last + 40
    B._write_layout(wb, builder)
    sheet, ref = _dest(wb, "est.checks")
    assert sheet == "Checks" and ref.endswith(f"${last}")


def test_a_pack_without_a_layout_map_writes_none_of_it(tmp_path):
    import yaml
    from estimating_tools.estimate_spec import format_pack as fp
    spec = yaml.safe_load(LEGACY_EXAMPLE.read_text(encoding="utf-8"))
    path = tmp_path / "legacy.xlsx"
    B.generate(spec, fp.load_pack(LEGACY_DIR)).save(path)
    wb = load_workbook(path)
    assert LM.TEMPLATES_SHEET not in wb.sheetnames
    assert not [n for n in wb.defined_names if n.startswith("est.")]
    assert "estimate_layout" not in {p.name for p in wb.custom_doc_props.props}


def test_read_back_names_the_same_templates_sheet():
    assert RL.TEMPLATES_SHEET == LM.TEMPLATES_SHEET


def test_template_clear_columns_are_read_backs_input_columns():
    clear = fieldbook_pack().layout_map["templates"]["clear"]
    assert clear["TASK"] == [RL.TASK_CODE, RL.TASK_DESC]
    assert clear["CREW"] == [RL.CREW_TRADE, RL.CREW_COUNT, RL.CREW_DAYS]


def test_an_unedited_workbook_with_templates_still_round_trips(tmp_path):
    path = tmp_path / "built.xlsx"
    build().save(path)
    result = read_back(example_spec(), rb_load(path))
    assert result.spec == example_spec() and not result.blocked


def _tpl_rows(wb, name):
    ((_, ref),) = list(wb.defined_names[name].destinations)
    first, last = (int(p.lstrip("$")) for p in ref.split(":"))
    return first, last


def test_a_crew_row_inserted_from_templates_reads_back(tmp_path):
    wb = build()
    sheet = "2 - 8IN MAIN"
    end = task_block(wb[sheet], "02-100")["END"][0]
    block = task_block(wb[sheet], "02-100")
    head, sub = block["TASK"][0], block["SUBTOTAL"][0]
    assert wb[sheet][f"K{sub}"].value == f"=SUM(K{head}:K{end})"
    paste_template_rows(wb, *_tpl_rows(wb, "est.tpl.crew"), sheet, end)
    ws = wb[sheet]
    # The pasted row chains to the row above it, prices its own count x days,
    # and the subtotal (pushed down one) grew to take it in.
    assert ws[f"R{end}"].value == f"=R{end - 1}"
    assert ws[f"F{end}"].value == f"=D{end}*E{end}*'Rates'!$B$2"
    # Excel's insert leaves a reference to a row above the insert point alone,
    # so the END row (pushed down one) still chains past the new row to the
    # row above it -- the same task code either way.
    assert ws[f"R{end + 1}"].value == f"=R{end - 1}"
    assert ws[f"K{sub + 1}"].value == f"=SUM(K{head}:K{end + 1})"
    ws[f"A{end}"], ws[f"D{end}"], ws[f"E{end}"] = "Laborer", 1, 3
    path = tmp_path / "crew.xlsx"
    wb.save(path)
    result = read_back(example_spec(), rb_load(path))
    assert not result.blocked, result.findings.problems
    crew = result.spec["items"][1]["labor"][0]["crew"]
    assert crew[-1] == {"trade": "Laborer", "count": 1, "days": 3}


def test_a_task_block_inserted_from_templates_reads_back(tmp_path):
    wb = build()
    sheet = "2 - 8IN MAIN"
    block = task_block(wb[sheet], "02-200")
    at = block["SUBTOTAL"][0] + 1
    first, last = _tpl_rows(wb, "est.tpl.task")
    paste_template_rows(wb, first, last, sheet, at)
    ws = wb[sheet]
    # The block keeps its internal offsets: its code chain starts at its own
    # TASK row and its subtotal sums exactly its own rows.
    assert ws[f"{KIND}{at}"].value == "TASK" and ws[f"{KIND}{at + 4}"].value == "SUBTOTAL"
    assert ws[f"R{at}"].value == f'=A{at}&""'
    assert ws[f"K{at + 4}"].value == f"=SUM(K{at}:K{at + 3})"
    ws[f"A{at}"], ws[f"B{at}"] = "02-300", "Pavement patch"
    crew = at + (int(_tpl_rows(wb, "est.tpl.crew")[0]) - first)
    ws[f"A{crew}"], ws[f"D{crew}"], ws[f"E{crew}"] = "Laborer", 2, 1
    path = tmp_path / "task.xlsx"
    wb.save(path)
    result = read_back(example_spec(), rb_load(path))
    assert not result.blocked, result.findings.problems
    labor = result.spec["items"][1]["labor"]
    assert [t["cost_code"] for t in labor] == ["02-100", "02-200", "02-300"]
    assert labor[-1]["crew"] == [{"trade": "Laborer", "count": 2, "days": 1}]
    assert not labor[-1].get("lines")          # the untouched prototype line reads as cleared
