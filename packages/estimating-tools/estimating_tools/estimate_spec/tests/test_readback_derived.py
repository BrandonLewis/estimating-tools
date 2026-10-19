"""Read-back of a workbook built with an org fork of fieldbook (`derived_from: fieldbook`).

Read-back resolves the pack the way build does (`--format`, the profile's
`estimating.format`, the bundled pack), accepts a workbook stamped with that
pack's name and version, refuses one another pack built, and baselines
against the pack it resolved -- never the shipped fieldbook.
"""
from __future__ import annotations

import copy

import pytest
import yaml
from openpyxl import load_workbook as open_xlsx

from estimating_tools.estimate_spec import build as B
from estimating_tools.estimate_spec import cli
from estimating_tools.estimate_spec import format_pack as fp
from estimating_tools.estimate_spec.readback import ReadBackRefused, load_workbook, read_back
from estimating_tools.estimate_spec.tests._fieldbook import (
    DERIVED_ITEM_HEAD, DERIVED_NAME, EXAMPLE, FIELDBOOK_DIR, build, derived_pack_dir,
    example_spec, rows_of_kind)
from estimating_tools.estimate_spec.tests._packs import LEGACY_DIR, legacy_example

MOB = "1 - MOBILIZATION"


@pytest.fixture(autouse=True)
def _no_profile(monkeypatch):
    """No org profile on the test machine decides which pack read-back uses."""
    monkeypatch.delenv("ESTIMATING_PROFILES", raising=False)


@pytest.fixture
def derived(tmp_path):
    return fp.load_pack(derived_pack_dir(tmp_path))


def _saved(wb, tmp_path, name="edited.xlsx"):
    path = tmp_path / name
    wb.save(path)
    return load_workbook(path)


def _superintendent_row(ws) -> int:
    (row,) = [r for r in rows_of_kind(ws, "CREW") if ws[f"A{r}"].value == "Superintendent"][:1]
    return row


def test_an_unedited_derived_workbook_reads_back_to_the_same_spec(derived, tmp_path):
    spec = example_spec()
    wb = B.generate(copy.deepcopy(spec), derived)
    # The fork's grid differs from fieldbook's on a sheet read-back compares
    # against its baseline: a fieldbook baseline would see column A move.
    assert wb["Cost Summary"]["A2"].value == DERIVED_ITEM_HEAD != build()["Cost Summary"]["A2"].value
    result = read_back(spec, _saved(wb, tmp_path), pack=derived)
    assert not result.changed and not result.blocked, result.findings.problems
    assert not [n for n in result.findings.notices if "Cost Summary" in n], result.findings.notices
    assert result.spec == spec


def test_a_fieldbook_baseline_would_have_noticed_the_forks_grid(derived, tmp_path, monkeypatch):
    """The check above is not vacuous: baselining the fork's workbook with
    the shipped fieldbook (the old behaviour) trips the Cost Summary notice."""
    spec = example_spec()
    wb = _saved(B.generate(copy.deepcopy(spec), derived), tmp_path)
    real = B.generate
    stock = fp.load_pack(FIELDBOOK_DIR)
    monkeypatch.setattr(B, "generate", lambda config, pack=None, **kw: real(config, stock, **kw))
    result = read_back(spec, wb, pack=derived)
    assert [n for n in result.findings.notices if n.startswith("Cost Summary: rows were")]


def test_a_crew_count_edited_in_a_derived_workbook_reads_back(derived, tmp_path):
    spec = example_spec()
    wb = B.generate(copy.deepcopy(spec), derived)
    ws = wb[MOB]
    ws[f"D{_superintendent_row(ws)}"] = 2
    result = read_back(spec, _saved(wb, tmp_path), pack=derived)
    assert not result.blocked, result.findings.problems
    assert result.changed
    (task,) = result.spec["items"][0]["labor"][:1]
    (supt,) = [c for c in task["crew"] if c["trade"] == "Superintendent"]
    assert supt["count"] == 2


def test_the_baseline_is_built_with_the_resolved_pack(derived, tmp_path, monkeypatch):
    spec = example_spec()
    wb = _saved(B.generate(copy.deepcopy(spec), derived), tmp_path)
    seen = []
    real = B.generate

    def spy(config, pack=None, **kw):
        seen.append(pack)
        return real(config, pack, **kw)

    monkeypatch.setattr(B, "generate", spy)
    read_back(spec, wb, pack=derived)
    assert seen and all(p is derived for p in seen)


def test_a_derived_workbook_read_against_the_bundled_pack_is_refused(derived, tmp_path):
    wb = _saved(B.generate(example_spec(), derived), tmp_path)
    with pytest.raises(ReadBackRefused) as exc:
        read_back(example_spec(), wb)
    message = str(exc.value)
    assert f"built in the {DERIVED_NAME!r} format" in message
    assert "'fieldbook'" in message and "--format" in message and "estimating.format" in message
    assert "--format fieldbook" not in message


def test_a_fieldbook_workbook_read_against_a_derived_pack_is_refused(derived, tmp_path):
    wb = _saved(B.generate(example_spec(), fp.load_pack(fp.BUNDLED)), tmp_path)
    with pytest.raises(ReadBackRefused,
                       match=rf"built in the 'fieldbook' format.*{DERIVED_NAME!r}") as exc:
        read_back(example_spec(), wb, pack=derived)
    # Advice that can be followed: the bundled pack by name. Never "point
    # the org profile at fieldbook" -- that would redirect every org build.
    assert "pass --format fieldbook" in str(exc.value)
    assert "estimating.format" not in str(exc.value)


def test_a_derived_workbook_from_another_version_is_refused(derived, tmp_path):
    wb = B.generate(example_spec(), derived)
    (prop,) = [p for p in wb.custom_doc_props.props if p.name == "estimate_format"]
    prop.value = f"{DERIVED_NAME} 7"
    with pytest.raises(ReadBackRefused,
                       match=rf"built with {DERIVED_NAME} 7; this read-back needs {DERIVED_NAME} 2"):
        read_back(example_spec(), _saved(wb, tmp_path), pack=derived)


def test_a_pack_outside_the_fieldbook_family_is_refused_by_family(tmp_path):
    legacy = fp.load_pack(LEGACY_DIR)
    wb = _saved(B.generate(legacy_example(), legacy), tmp_path)
    with pytest.raises(ReadBackRefused, match=r"'legacy_demo' family.*'fieldbook'"):
        read_back(legacy_example(), wb, pack=legacy)


# --- the CLI -----------------------------------------------------------------

@pytest.fixture
def derived_bid(tmp_path):
    """A spec, a derived pack directory, and the workbook built from both,
    with the mobilization superintendent's crew count edited to 2."""
    pack_dir = derived_pack_dir(tmp_path)
    spec = tmp_path / "bid.estimate.yaml"
    spec.write_text(EXAMPLE.read_text(encoding="utf-8"), encoding="utf-8")
    xlsx = tmp_path / "bid.xlsx"
    assert cli.main(["build", str(spec), "-o", str(xlsx), "--format", str(pack_dir)]) == 0
    wb = open_xlsx(xlsx)
    ws = wb[MOB]
    ws[f"D{_superintendent_row(ws)}"] = 2
    wb.save(xlsx)
    return spec, xlsx, pack_dir


def _supt_count(spec_path) -> int:
    spec = yaml.safe_load(spec_path.read_text(encoding="utf-8"))
    crew = spec["items"][0]["labor"][0]["crew"]
    return next(c["count"] for c in crew if c["trade"] == "Superintendent")


def test_cli_reads_back_a_derived_workbook_with_format(derived_bid, capsys):
    spec, xlsx, pack_dir = derived_bid
    assert _supt_count(spec) == 1
    assert cli.main(["read-back", str(spec), str(xlsx), "--format", str(pack_dir),
                     "--write"]) == 0
    assert _supt_count(spec) == 2
    assert "wrote 1 change(s)" in capsys.readouterr().out


def test_cli_without_format_or_profile_refuses_a_derived_workbook(derived_bid, capsys):
    spec, xlsx, _ = derived_bid
    before = spec.read_text(encoding="utf-8")
    assert cli.main(["read-back", str(spec), str(xlsx), "--write"]) == 2
    err = capsys.readouterr().err
    assert f"built in the {DERIVED_NAME!r} format" in err and "--format" in err
    assert spec.read_text(encoding="utf-8") == before


def test_cli_takes_the_pack_from_the_org_profile(derived_bid, tmp_path, monkeypatch, capsys):
    spec, xlsx, pack_dir = derived_bid
    profile = tmp_path / "org"
    profile.mkdir()
    (profile / "profile.yaml").write_text(
        f"estimating:\n  format: {pack_dir.as_posix()}\n", encoding="utf-8")
    monkeypatch.setenv("ESTIMATING_PROFILES", str(profile))
    assert cli.main(["read-back", str(spec), str(xlsx), "--write"]) == 0
    assert _supt_count(spec) == 2


def test_cli_reports_a_pack_it_cannot_load(derived_bid, tmp_path, capsys):
    spec, xlsx, _ = derived_bid
    assert cli.main(["read-back", str(spec), str(xlsx), "--format", str(tmp_path / "nope")]) == 2
    assert "format pack error" in capsys.readouterr().err


def test_cli_reads_a_stock_workbook_with_format_fieldbook_under_an_org_profile(
        tmp_path, monkeypatch, capsys):
    pack_dir = derived_pack_dir(tmp_path)
    profile = tmp_path / "org"
    profile.mkdir()
    (profile / "profile.yaml").write_text(
        f"estimating:\n  format: {pack_dir.as_posix()}\n", encoding="utf-8")
    spec = tmp_path / "bid.estimate.yaml"
    spec.write_text(EXAMPLE.read_text(encoding="utf-8"), encoding="utf-8")
    xlsx = tmp_path / "bid.xlsx"
    assert cli.main(["build", str(spec), "-o", str(xlsx), "--format", "fieldbook"]) == 0
    wb = open_xlsx(xlsx)
    ws = wb[MOB]
    ws[f"D{_superintendent_row(ws)}"] = 2
    wb.save(xlsx)
    monkeypatch.setenv("ESTIMATING_PROFILES", str(profile))
    assert cli.main(["read-back", str(spec), str(xlsx)]) == 2           # the profile's pack
    assert "pass --format fieldbook" in capsys.readouterr().err
    assert cli.main(["read-back", str(spec), str(xlsx), "--format", "fieldbook", "--write"]) == 0
    assert _supt_count(spec) == 2


def test_a_copied_derived_pack_refuses_a_stamp_without_the_base_version(derived):
    spec = example_spec()
    wb = B.generate(spec, derived)
    (stamp,) = [p for p in wb.custom_doc_props.props if p.name == "estimate_format"]
    assert stamp.value == f"{derived.name} {derived.version} / fieldbook 2"
    stamp.value = f"{derived.name} {derived.version}"
    with pytest.raises(ReadBackRefused, match="needs.*fieldbook 2.*rebuild"):
        read_back(spec, wb, pack=derived)
