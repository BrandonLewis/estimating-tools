"""Fieldbook phase 1: generic, pack-opted engine features."""
from __future__ import annotations

import openpyxl
import pytest

from estimating_tools.estimate_spec import build as B
from estimating_tools.estimate_spec import format_pack as fp
from estimating_tools.estimate_spec import recipe as R
from estimating_tools.estimate_spec.tests._packs import legacy_example, patched_pack


def _row_of(ws, col, value):
    return next(c.row for c in ws[col] if c.value == value)


# ---- forward references (characterization: no engine change) -------------

def test_a_cell_may_reference_an_export_laid_out_later_on_its_sheet(tmp_path):
    """Fieldbook's recap sits above the task blocks it totals."""
    def units(doc):
        doc["units"]["probe"] = {"rows": [
            {"cells": {"A": {"formula": "={ref:labor_total.labor_total}"}}}]}

    def recipes(doc):
        doc["recipes"]["setup"]["units"].insert(1, {"unit": "probe"})

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units, recipes=recipes))["SETUP"]
    formula = ws["A4"].value
    assert isinstance(formula, str), f"probe not at A4: {formula!r}"
    assert formula.startswith("=$F$")
    target_row = int(formula[len("=$F$"):])
    assert target_row > 4, f"expected a forward reference (row > 4), got row {target_row}"
    assert ws[f"E{target_row}"].value == "LABOR"


# ---- patched-pack helper itself --------------------------------------------

def test_patched_pack_breaks_anchor_sharing(tmp_path):
    """legacy_demo's recipes.yaml shares &header/*header; editing one
    recipe's unit must not silently edit another recipe's alias of it."""
    def recipes(doc):
        doc["recipes"]["setup"]["units"][0]["params"]["description"] = "MUTATED"

    pack = patched_pack(tmp_path, recipes=recipes)
    assert pack.recipes["recipes"]["setup"]["units"][0]["params"]["description"] == "MUTATED"
    assert pack.recipes["recipes"]["line_item"]["units"][0]["params"]["description"] != "MUTATED"


# ---- if_it ------------------------------------------------------------------

def test_if_it_renders_a_cell_only_for_matching_repeat_elements(tmp_path):
    def units(doc):
        row = doc["units"]["cost_block"]["rows"][1]          # the `repeat: p.lines` row
        row["cells"]["G"] = {"value": "HIT", "if_it": {"name": "Bedding rock"}}

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units))["2 - 8IN MAIN"]
    hits = {ws[f"A{c.row}"].value for c in ws["G"] if c.value == "HIT"}
    assert hits == {"Bedding rock"}


def test_if_it_compares_as_text(tmp_path):
    def units(doc):
        row = doc["units"]["cost_block"]["rows"][1]
        row["cells"]["G"] = {"value": "HIT", "if_it": {"unit_price": "31"}}   # spec holds 31

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units))["2 - 8IN MAIN"]
    assert [ws[f"A{c.row}"].value for c in ws["G"] if c.value == "HIT"] == ["Bedding rock"]


def test_if_it_on_a_non_repeat_row_fails_loudly(tmp_path):
    """`cost_block` row 0 is the single banner row (`cells`, no `repeat`):
    its `item` is always None, so `if_it` can never have a real element to
    compare against."""
    def units(doc):
        row = doc["units"]["cost_block"]["rows"][0]
        row["cells"]["A"]["if_it"] = {"name": "Bedding rock"}

    with pytest.raises(R.RecipeError, match=r"if_it needs a `repeat:` row over mappings"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


def test_if_it_with_a_string_value_fails_loudly(tmp_path):
    def units(doc):
        row = doc["units"]["cost_block"]["rows"][1]          # the `repeat: p.lines` row
        row["cells"]["G"] = {"value": "HIT", "if_it": "Bedding rock"}

    with pytest.raises(R.RecipeError, match=r"if_it must be a non-empty mapping"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


def test_if_it_with_an_empty_mapping_fails_loudly(tmp_path):
    def units(doc):
        row = doc["units"]["cost_block"]["rows"][1]
        row["cells"]["G"] = {"value": "HIT", "if_it": {}}

    with pytest.raises(R.RecipeError, match=r"if_it must be a non-empty mapping"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


# ---- locking and protection ----------------------------------------------

def _protected(tmp_path, password=""):
    def tokens(doc):
        doc["styles"]["price_input"]["locked"] = False

    def manifest(doc):
        doc["protect"] = {"sheets": True, "password": password}

    return patched_pack(tmp_path, tokens=tokens, format=manifest)


def test_protected_pack_locks_every_sheet_and_unlocks_only_unlocked_styles(tmp_path):
    wb = B.generate(legacy_example(), _protected(tmp_path))
    assert all(ws.protection.sheet for ws in wb.worksheets)
    ws = wb["2 - 8IN MAIN"]
    r = _row_of(ws, "A", "8-inch ductile iron pipe")
    assert ws[f"E{r}"].protection.locked is False      # price_input
    assert ws[f"F{r}"].protection.locked is True       # =C*E formula


def test_a_password_is_set_only_when_given(tmp_path):
    wb = B.generate(legacy_example(), _protected(tmp_path / "a", password="s3cret"))
    assert all(ws.protection.password for ws in wb.worksheets)
    wb = B.generate(legacy_example(), _protected(tmp_path / "b"))
    assert not any(ws.protection.password for ws in wb.worksheets)


@pytest.mark.parametrize("protect", [
    {"sheets": "yes"}, {"sheets": True, "password": 7}, {"sheets": True, "colour": "red"}, ["sheets"],
])
def test_a_malformed_protect_is_refused_at_load(tmp_path, protect):
    def manifest(doc):
        doc["protect"] = protect

    with pytest.raises(fp.FormatPackError, match="protect"):
        patched_pack(tmp_path, format=manifest)


def test_a_blank_optional_cell_still_locks_when_its_style_says_locked(tmp_path):
    """A line's price cell is optional and resolves empty (blank unit_price):
    write() never sets its value or runs the rest of `apply()`, but the
    style's `locked: False` still has to reach the cell, or a protected
    sheet would leave it permanently locked with no way to type a price in."""
    def units(doc):
        row = doc["units"]["cost_block"]["rows"][1]
        row["cells"]["E"]["optional"] = True

    def tokens(doc):
        doc["styles"]["price_input"]["locked"] = False

    def manifest(doc):
        doc["protect"] = {"sheets": True}

    pack = patched_pack(tmp_path, units=units, tokens=tokens, format=manifest)
    config = legacy_example()
    config["items"][1]["cost_details"][0]["lines"][1]["unit_price"] = ""   # Bedding rock

    wb = B.generate(config, pack)
    ws = wb["2 - 8IN MAIN"]
    r = _row_of(ws, "A", "Bedding rock")
    assert ws[f"E{r}"].value in (None, "")
    assert ws[f"E{r}"].protection.locked is False


def test_protected_sheets_allow_resize_but_block_other_protected_actions(tmp_path):
    wb = B.generate(legacy_example(), _protected(tmp_path))
    for ws in wb.worksheets:
        assert ws.protection.formatColumns is False
        assert ws.protection.formatRows is False
        assert ws.protection.insertRows is True    # untouched: openpyxl's blocked default


def test_locked_must_be_a_real_bool(tmp_path):
    def tokens(doc):
        doc["styles"]["price_input"]["locked"] = "false"

    with pytest.raises(R.RecipeError, match="price_input"):
        B.generate(legacy_example(), patched_pack(tmp_path, tokens=tokens))


def test_a_password_without_sheets_is_refused(tmp_path):
    def manifest(doc):
        doc["protect"] = {"password": "s3cret"}

    with pytest.raises(fp.FormatPackError, match="protect"):
        patched_pack(tmp_path, format=manifest)


# ---- manifest `strict:` -- schema enforcement on the build path -----------
#
# generate()/build() check cross-references (validate_cross_references) but
# never ran the pack's own composed schema -- a spec's refusals under
# scaffold.validate_spec (item_rules, defs.takeoff/spreads, etc.) were only
# ever checked there, so `est-estimate build` built right past them. `strict:
# true` in format.yaml opts a pack into the same jsonschema check, reusing
# validate_spec's try/except shape (a `schema violation: $[...]: <message>`
# ValueError), BEFORE anything is laid out. Absent (every pack before this),
# nothing changes -- proven by building the identical invalid spec against
# the unpatched pack and getting past schema entirely (whatever error it
# hits next, if any, is not this one).

def _strict(tmp_path):
    def manifest(doc):
        doc["strict"] = True

    return patched_pack(tmp_path, format=manifest)


def test_a_strict_pack_refuses_a_schema_invalid_spec_in_generate(tmp_path):
    config = legacy_example()
    config["project"] = ""                        # schema: project minLength 1
    with pytest.raises(ValueError, match=r"schema violation.*project"):
        B.generate(config, _strict(tmp_path))


def test_strict_absent_keeps_todays_behaviour(tmp_path):
    """The same schema-invalid spec, against the SAME pack minus `strict:`,
    is never schema-checked -- goldens built from unpatched packs are
    unaffected by this feature existing."""
    config = legacy_example()
    config["project"] = ""
    pack = patched_pack(tmp_path)                 # no strict: true
    assert pack.strict is False
    wb = B.generate(config, pack)                 # does not raise "schema violation"
    assert wb.sheetnames


@pytest.mark.parametrize("bad", ["yes", 1, None])
def test_a_malformed_manifest_strict_is_refused_at_load(tmp_path, bad):
    def manifest(doc):
        doc["strict"] = bad

    with pytest.raises(fp.FormatPackError, match="strict"):
        patched_pack(tmp_path, format=manifest)


# ---- conditional formats --------------------------------------------------

def _rules(ws):
    return {str(cf.sqref): cf.rules for cf in ws.conditional_formatting}


def test_a_row_conditional_is_written_once_per_placed_row(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "F", "formula": "=F{r}>10000", "style": "total"}]

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units))["2 - 8IN MAIN"]
    pipe = _row_of(ws, "A", "8-inch ductile iron pipe")
    rock = _row_of(ws, "A", "Bedding rock")
    rules = _rules(ws)
    assert rules[f"F{pipe}"][0].formula == [f"F{pipe}>10000"]
    assert rules[f"F{rock}"][0].formula == [f"F{rock}>10000"]
    assert rules[f"F{pipe}"][0].dxf.fill is not None      # `total` carries a fill


def test_a_conditional_may_span_a_row_range(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "A:F", "formula": "=$F{r}=0", "style": "total"}]

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units))["2 - 8IN MAIN"]
    pipe = _row_of(ws, "A", "8-inch ductile iron pipe")
    assert f"A{pipe}:F{pipe}" in _rules(ws)


def test_a_conditional_naming_an_unknown_style_stops_the_build(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "F", "formula": "=F{r}>0", "style": "no_such_style"}]

    with pytest.raises(ValueError, match="no_such_style"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


def test_a_conditional_on_a_single_non_repeat_row_lands_on_that_row(tmp_path):
    """The other `_place_row` call site -- `cost_block` row 0 is the single
    banner row (no `repeat:`)."""
    def units(doc):
        doc["units"]["cost_block"]["rows"][0]["conditional"] = [
            {"cells": "A", "formula": '=$A{r}<>""', "style": "total"}]

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units))["2 - 8IN MAIN"]
    banner = _row_of(ws, "A", "MATL")
    rules = _rules(ws)
    assert rules[f"A{banner}"][0].formula == [f'$A{banner}<>""']


def test_a_conditional_spanning_columns_rejects_an_unanchored_column_ref(tmp_path):
    """Over a multi-column range, Excel evaluates the formula relative to
    the range's top-left cell: `=F{r}=0` painted over A:F would check F for
    column A, G for column B, and so on. `=$F{r}=0` (already covered by
    test_a_conditional_may_span_a_row_range) is the correctly-anchored form
    that must still be accepted."""
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "A:F", "formula": "=F{r}=0", "style": "total"}]

    with pytest.raises(R.RecipeError, match="without a leading"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


# The anchoring check runs on the RESOLVED formula and must not mistake a
# function name, a scientific-notation literal, or a named range for an
# unanchored cell reference -- a false positive here blocks a legitimate
# formula, which is worse than a false negative.
@pytest.mark.parametrize("formula", [
    "=LOG10($F{r})>2",
    "=ATAN2($E{r},$F{r})>0",
    "=DAYS360($A{r},$B{r})>30",
    "=$F{r}>1E5",
    "=$F{r}>RATE2024",
    "=Checks!$B$3>0",
    "=SUM($A{r}:$F{r})>0",
    '=AND($G{r}<>"",$G{r}>10000)',
    '=$F{r}="F2"',
])
def test_a_multi_column_conditional_accepts_formulas_that_look_like_refs_but_are_not(tmp_path, formula):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "A:F", "formula": formula, "style": "total"}]

    B.generate(legacy_example(), patched_pack(tmp_path, units=units))   # must not raise


@pytest.mark.parametrize("formula", [
    "='Rates'!A2>0",     # quoted-sheet-prefixed but the column itself is bare
    "=F$5>0",            # row anchored, column is not -- the column is what matters
    "=SUM($A{r}:F{r})>0",  # second end of the range is unanchored
    "=f{r}>0",           # lowercase column still counts
])
def test_a_multi_column_conditional_rejects_unanchored_refs(tmp_path, formula):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "A:F", "formula": formula, "style": "total"}]

    with pytest.raises(R.RecipeError, match="without a leading"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


def test_a_multi_column_conditional_rejects_an_unanchored_ref_from_any_placeholder(tmp_path):
    """Not just `{r}` -- any placeholder (here `{row:head}`, a landmark) can
    resolve to an unanchored reference; the check runs on the RESOLVED
    formula, so it catches this regardless of which placeholder produced
    the final address."""
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "A:F", "formula": "=E{row:head}>0", "style": "total"}]

    with pytest.raises(R.RecipeError, match="without a leading"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


def test_a_conditional_cells_range_must_list_columns_left_to_right(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "F:A", "formula": "=$F{r}>0", "style": "total"}]

    with pytest.raises(R.RecipeError, match="before its second"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


def test_dxf_style_naming_an_unknown_font_color_stops_the_build(tmp_path):
    """A style used ONLY by the conditional (never by a normal cell), so the
    missing token is caught by `StyleFactory.dxf`, not by `apply()` on some
    unrelated cell that happens to share the style name."""
    def tokens(doc):
        doc["styles"]["cond_only"] = {"font_color": "no_such_color"}

    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "F", "formula": "=F{r}>0", "style": "cond_only"}]

    with pytest.raises(R.RecipeError, match="no_such_color"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units, tokens=tokens))


def test_dxf_style_naming_an_unknown_fill_stops_the_build(tmp_path):
    def tokens(doc):
        doc["styles"]["cond_only"] = {"fill": "no_such_fill"}

    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "F", "formula": "=F{r}>0", "style": "cond_only"}]

    with pytest.raises(R.RecipeError, match="no_such_fill"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units, tokens=tokens))


@pytest.mark.parametrize("conditional, match", [
    ({"cells": "F", "formula": "=F{r}>0", "style": "total"}, "must be a list"),
    ([{"cells": "F", "formula": "=F{r}>0"}], "exactly the keys"),               # missing style
    ([{"cells": "F", "formula": "=F{r}>0", "style": "total", "extra": 1}],
     "exactly the keys"),                                                       # extra key
    (["not-a-mapping"], "exactly the keys"),
    ([{"cells": "f", "formula": "=F{r}>0", "style": "total"}],
     "column or column range"),                                                # lowercase
    ([{"cells": "A1", "formula": "=F{r}>0", "style": "total"}],
     "column or column range"),                                                # digits in cells
    ([{"cells": "A:B:C", "formula": "=F{r}>0", "style": "total"}],
     "column or column range"),                                                # 3-way range
    ([{"cells": 5, "formula": "=F{r}>0", "style": "total"}],
     "column or column range"),                                                # not a string
    ([{"cells": "F", "formula": "F{r}>0", "style": "total"}],
     "starting with"),                                                          # no leading =
    ([{"cells": "F", "formula": 5, "style": "total"}], "starting with"),       # not a string
])
def test_a_malformed_conditional_is_refused_at_layout(tmp_path, conditional, match):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = conditional

    with pytest.raises(R.RecipeError, match=match):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


# ---- book sheets ----------------------------------------------------------

def _book_units(doc):
    doc["units"]["book_list"] = {"params": {"items": {"type": "list"}}, "rows": [
        {"repeat": "p.items", "cells": {
            "A": {"value": "{it.item_no}"},
            "B": {"formula": "={ref:{it.sheet}!job_total.grand_total}"}}}]}
    doc["units"]["book_note"] = {"params": {"text": {"type": "string"}}, "rows": [
        {"cells": {"A": {"value": "{p.text}"}}}]}


def _book_recipes(doc):
    del doc["recipes"]["cover"]
    doc["book"] = [
        {"sheet_name": "Front", "position": "before_items", "column_widths": {"A": 10, "B": 18},
         "units": [{"unit": "book_list", "params": {"items": "{book.items}"}}]},
        {"sheet_name": "Back", "position": "after_items",
         "units": [{"unit": "book_note", "params": {"text": "{config.project}"}}]},
    ]


def _book_pack(tmp_path):
    return patched_pack(tmp_path, units=_book_units, recipes=_book_recipes)


def test_book_sheets_bracket_the_item_sheets(tmp_path):
    legacy = B.generate(legacy_example(), patched_pack(tmp_path / "l"))
    wb = B.generate(legacy_example(), _book_pack(tmp_path / "b"))
    items = [s for s in legacy.sheetnames if s != "Bid Summary"]
    assert wb.sheetnames == ["Front", *items, "Back"]
    assert wb.active.title == "Front"


def test_a_book_sheet_reaches_every_item_total(tmp_path):
    legacy = B.generate(legacy_example(), patched_pack(tmp_path / "l"))["Bid Summary"]
    ws = B.generate(legacy_example(), _book_pack(tmp_path / "b"))["Front"]
    assert [ws[f"A{r}"].value for r in (1, 2, 3)] == [1, 2, 3]
    assert [ws[f"B{r}"].value for r in (1, 2, 3)] == [legacy[f"D{r}"].value for r in (5, 6, 7)]
    assert ws.column_dimensions["A"].width == 10


def test_a_book_sheet_skips_shared_cost_items(tmp_path):
    spec = legacy_example()
    spec["items"][2]["in_summary"] = False
    ws = B.generate(spec, _book_pack(tmp_path))["Front"]
    assert ws["A3"].value is None and ws["A2"].value == 2


def test_an_after_items_sheet_reads_the_config(tmp_path):
    spec = legacy_example()
    assert B.generate(spec, _book_pack(tmp_path))["Back"]["A1"].value == spec["project"]


@pytest.mark.parametrize("edit, why", [
    (lambda d: d.update(book=[{"sheet_name": "X", "position": "before_items", "units": []}]),
     "not both"),                                              # cover kept
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "middle", "units": []}])),
     "position"),
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items", "units": [],
                                "colour": "red"}])),
     "unknown keys"),
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items", "units": []},
                               {"sheet_name": "X", "position": "after_items", "units": []}])),
     "unique sheet_name"),
    (lambda d: (d["recipes"].pop("cover"), d.update(book=[])),
     "non-empty"),
])
def test_a_malformed_book_is_refused_at_load(tmp_path, edit, why):
    with pytest.raises(fp.FormatPackError, match=why):
        patched_pack(tmp_path, recipes=edit)


# ---- book validation: units/entries/sheet_name/column_widths --------------

@pytest.mark.parametrize("edit, why", [
    # units must be a list, not merely truthy
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items", "units": {}}])),
     "must be a list"),
    # each units[] entry must itself be a mapping
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items",
                                "units": ["not-a-mapping"]}])),
     "must be a mapping"),
    # sheet_name: non-empty string
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "", "position": "after_items", "units": []}])),
     "non-empty string"),
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": 5, "position": "after_items", "units": []}])),
     "non-empty string"),
    # sheet_name: Excel's 31-character limit
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X" * 32, "position": "after_items", "units": []}])),
     "longer than"),
    # sheet_name: none of Excel's reserved characters
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "Bad:Name", "position": "after_items", "units": []}])),
     "does not allow"),
    # column_widths must be a mapping
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items", "units": [],
                                "column_widths": ["A"]}])),
     "column_widths must be a mapping"),
    # column_widths keys must be column letters
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items", "units": [],
                                "column_widths": {"1": 10}}])),
     "column letter"),
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items", "units": [],
                                "column_widths": {"a": 10}}])),       # lowercase
     "column letter"),
    # column_widths keys must not reach past Excel's last column (XFD)
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items", "units": [],
                                "column_widths": {"ZZZ": 10}}])),
     "XFD"),
    # column_widths values must be numbers (bool excluded even though it is an int subclass)
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items", "units": [],
                                "column_widths": {"A": "wide"}}])),
     "must be a number"),
    (lambda d: (d["recipes"].pop("cover"),
                d.update(book=[{"sheet_name": "X", "position": "after_items", "units": [],
                                "column_widths": {"A": True}}])),
     "must be a number"),
])
def test_a_malformed_book_sheet_is_refused_at_load(tmp_path, edit, why):
    with pytest.raises(fp.FormatPackError, match=why):
        patched_pack(tmp_path, recipes=edit)


# ---- the paths fieldbook relies on -----------------------------------------

def test_moving_the_rate_sheet_onto_a_book_sheet_matches_legacy_vlookups(tmp_path):
    """The `rates` unit is anchored (its own sheet); fieldbook moves it off
    the item-recipe's `sheet: Rates` entry and onto an after_items book
    sheet instead. The crew-block VLOOKUPs on every item sheet have to keep
    resolving to the exact same 'Rates'!... range either way -- the global
    promotion mechanism (`globals.rate_lookup`, scope: workbook) does not
    care which layout path published `rates.table`."""
    def recipes(doc):
        setup_units = doc["recipes"]["setup"]["units"]
        rates_entry = next(u for u in setup_units if u.get("unit") == "rates")
        setup_units.remove(rates_entry)
        del doc["recipes"]["cover"]
        doc["book"] = [{"sheet_name": "Rates", "position": "after_items",
                        "units": [{"unit": "rates", "params": rates_entry["params"]}]}]

    legacy = B.generate(legacy_example(), patched_pack(tmp_path / "l"))
    moved = B.generate(legacy_example(), patched_pack(tmp_path / "m", recipes=recipes))
    item_sheets = [s for s in legacy.sheetnames if s not in ("Bid Summary", "Rates")]
    assert item_sheets, "fixture regressed: no item sheets to compare"
    found_a_vlookup = False
    for sheet in item_sheets:
        for row in legacy[sheet].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=VLOOKUP"):
                    found_a_vlookup = True
                    assert moved[sheet][cell.coordinate].value == cell.value
    assert found_a_vlookup, "fixture regressed: no VLOOKUP formula found to compare"
    assert moved["Rates"]["A2"].value == legacy["Rates"]["A2"].value


def test_an_item_sheet_can_ref_an_export_a_book_sheet_publishes(tmp_path):
    def units(doc):
        doc["units"]["book_echo"] = {
            "params": {"text": {"type": "string"}},
            "rows": [{"cells": {"A": {"value": "{p.text}"}}, "mark": "echo"}],
            "exports": {"echo": "$A${row:echo}"},
        }
        doc["units"]["probe"] = {"rows": [
            {"cells": {"A": {"formula": "={ref:Notes!book_echo.echo}"}}}]}

    def recipes(doc):
        doc["recipes"]["setup"]["units"].insert(0, {"unit": "probe"})
        del doc["recipes"]["cover"]
        doc["book"] = [{"sheet_name": "Notes", "position": "after_items",
                        "units": [{"unit": "book_echo", "params": {"text": "hello"}}]}]

    wb = B.generate(legacy_example(), patched_pack(tmp_path, units=units, recipes=recipes))
    assert wb["SETUP"]["A1"].value == "='Notes'!$A$1"
    assert wb["Notes"]["A1"].value == "hello"


# ---- book/item sheet-name collisions ---------------------------------------

def test_an_item_named_like_book0_is_refused(tmp_path):
    """book[0]'s sheet_name is pre-created as the workbook's placeholder
    active sheet before any item sheet is built -- an item sharing that
    name has to be caught explicitly (it isn't the generic already-exists
    check, since nothing has checked book names against items yet)."""
    spec = legacy_example()
    spec["items"][0]["sheet"] = "Front"              # _book_pack's book[0]
    with pytest.raises(R.RecipeError, match="collides with book sheet_name"):
        B.generate(spec, _book_pack(tmp_path))


def test_a_book_sheet_colliding_with_an_anchored_target_is_refused(tmp_path):
    """legacy_demo's own `setup` recipe anchors `rates` to `sheet: Rates`
    (untouched here) -- a book sheet also named 'Rates' collides with the
    sheet that anchored unit already created."""
    def units(doc):
        doc["units"]["book_note"] = {"params": {"text": {"type": "string"}}, "rows": [
            {"cells": {"A": {"value": "{p.text}"}}}]}

    def recipes(doc):
        del doc["recipes"]["cover"]
        doc["book"] = [{"sheet_name": "Rates", "position": "after_items",
                        "units": [{"unit": "book_note", "params": {"text": "x"}}]}]

    with pytest.raises(R.RecipeError, match="already exists in the workbook"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units, recipes=recipes))


def test_duplicate_item_sheet_names_are_refused(tmp_path):
    spec = legacy_example()
    spec["items"][1]["sheet"] = spec["items"][0]["sheet"]      # both "SETUP"
    with pytest.raises(R.RecipeError, match="already exists in the workbook"):
        B.generate(spec, patched_pack(tmp_path))


def test_an_item_named_like_the_legacy_cover_sheet_is_refused(tmp_path):
    """Not a book pack at all -- the generic already-exists check in
    build_item_sheet applies on the legacy `cover` path too, turning what
    used to be a silently-broken sheet (openpyxl would mint "Bid Summary1")
    into a named error."""
    spec = legacy_example()
    spec["items"][0]["sheet"] = "Bid Summary"                  # cover's sheet_name
    with pytest.raises(R.RecipeError, match="already exists in the workbook"):
        B.generate(spec, patched_pack(tmp_path))


# ---- "Sheet" is an ordinary name on the book path --------------------------

def test_an_item_may_be_named_sheet_on_the_book_path(tmp_path):
    """openpyxl's own default sheet ("Sheet") is removed before any item or
    book sheet is built on the book path (`_run_book`), so nothing reserves
    that name any more -- an item (or a book sheet) may use it freely."""
    spec = legacy_example()
    spec["items"][0]["sheet"] = "Sheet"
    wb = B.generate(spec, _book_pack(tmp_path))
    assert wb.sheetnames[0] == "Front"
    assert "Sheet" in wb.sheetnames
    assert wb.active.title == "Front"


# ---- sheet-name collisions are case-insensitive, like Excel's own ---------

def test_an_item_sheet_colliding_with_a_book_sheet_case_insensitively_is_refused(tmp_path):
    spec = legacy_example()
    spec["items"][0]["sheet"] = "front"                        # book[0] is "Front"
    with pytest.raises(R.RecipeError, match="collides with book sheet_name"):
        B.generate(spec, _book_pack(tmp_path))


def test_book_sheet_names_differing_only_by_case_are_refused_at_load(tmp_path):
    def recipes(doc):
        doc["recipes"].pop("cover")
        doc["book"] = [{"sheet_name": "Front", "position": "after_items", "units": []},
                       {"sheet_name": "FRONT", "position": "after_items", "units": []}]

    with pytest.raises(fp.FormatPackError, match="unique sheet_name"):
        patched_pack(tmp_path, recipes=recipes)


# ---- hidden columns and sheets ---------------------------------------------

def test_a_recipe_may_hide_columns(tmp_path):
    def recipes(doc):
        doc["recipes"]["line_item"]["hidden_columns"] = ["F"]

    wb = B.generate(legacy_example(), patched_pack(tmp_path, recipes=recipes))
    assert wb["2 - 8IN MAIN"].column_dimensions["F"].hidden is True
    assert not wb["SETUP"].column_dimensions["F"].hidden        # setup recipe untouched


def test_a_book_sheet_may_hide_columns_and_itself(tmp_path):
    def recipes(doc):
        _book_recipes(doc)
        doc["book"][1]["hidden"] = True
        doc["book"][0]["hidden_columns"] = ["B"]

    wb = B.generate(legacy_example(), patched_pack(tmp_path, units=_book_units, recipes=recipes))
    assert wb["Back"].sheet_state == "hidden"
    assert wb["Front"].sheet_state == "visible"
    assert wb["Front"].column_dimensions["B"].hidden is True


def test_a_hidden_book0_leaves_a_visible_sheet_active(tmp_path):
    """book[0] ("Front") is the sheet that would land at index 0 after the
    reorder; hiding it must not leave it the active sheet in memory -- Excel
    refuses to open a workbook whose active sheet is hidden."""
    def recipes(doc):
        _book_recipes(doc)
        doc["book"][0]["hidden"] = True

    wb = B.generate(legacy_example(), patched_pack(tmp_path, units=_book_units, recipes=recipes))
    assert wb["Front"].sheet_state == "hidden"
    assert wb.active is not None
    assert wb.active.sheet_state == "visible"
    assert wb.active.title != "Front"


@pytest.mark.parametrize("bad, match", [
    ({"hidden": "yes"}, "hidden must be true or false"),
    ({"hidden_columns": "B"}, "hidden_columns must be a list of column letters"),
    ({"hidden_columns": ["1"]}, "hidden_columns must be a list of column letters"),
])
def test_malformed_hiding_is_refused_at_load(tmp_path, bad, match):
    """`match=` pins each case to the hidden/hidden_columns message itself --
    without it, a case would also pass if `hidden`/`hidden_columns` reverted
    to an unknown book key (a different, less specific FormatPackError)."""
    def recipes(doc):
        _book_recipes(doc)
        doc["book"][0].update(bad)

    with pytest.raises(fp.FormatPackError, match=match):
        patched_pack(tmp_path, units=_book_units, recipes=recipes)


def test_a_malformed_recipe_level_hidden_columns_is_refused_at_load(tmp_path):
    def recipes(doc):
        doc["recipes"]["line_item"]["hidden_columns"] = "F"

    with pytest.raises(fp.FormatPackError, match=r"recipes\.line_item: hidden_columns"):
        patched_pack(tmp_path, recipes=recipes)


def test_hidden_columns_past_xfd_is_refused_at_load(tmp_path):
    def recipes(doc):
        _book_recipes(doc)
        doc["book"][0]["hidden_columns"] = ["ZZZ"]

    with pytest.raises(fp.FormatPackError, match="XFD"):
        patched_pack(tmp_path, units=_book_units, recipes=recipes)


# ---- dropdowns --------------------------------------------------------------

def _dv(ws):
    return {(dv.type, dv.formula1): sorted(str(r) for r in dv.sqref.ranges)
            for dv in ws.data_validations.dataValidation}


def _with_validation(validation, *, optional=False):
    def units(doc):
        cell = doc["units"]["cost_block"]["rows"][1]["cells"]["D"]   # the uom cell of each line
        cell["validation"] = validation
        if optional:
            cell["optional"] = True
    return units


def test_a_literal_list_becomes_a_dropdown_on_every_placed_cell(tmp_path):
    pack = patched_pack(tmp_path, units=_with_validation({"list": ["LF", "TON", "DAY"]}))
    ws = B.generate(legacy_example(), pack)["2 - 8IN MAIN"]
    rows = [_row_of(ws, "A", n) for n in ("8-inch ductile iron pipe", "Bedding rock", "Trench box")]
    assert _dv(ws)[("list", '"LF,TON,DAY"')] == sorted(f"D{r}" for r in rows)


def test_a_single_column_range_list_becomes_a_dropdown(tmp_path):
    """A literal `=range` (no placeholder involved) is the simple accepted
    shape: a single column, which Excel's List validation allows."""
    pack = patched_pack(tmp_path, units=_with_validation({"list": "='Rates'!$A$2:$A$6"}))
    ws = B.generate(legacy_example(), pack)["2 - 8IN MAIN"]
    rows = [_row_of(ws, "A", n) for n in ("8-inch ductile iron pipe", "Bedding rock", "Trench box")]
    assert _dv(ws)[("list", "'Rates'!$A$2:$A$6")] == sorted(f"D{r}" for r in rows)


def test_a_range_list_spanning_more_than_one_column_is_refused(tmp_path):
    """`{rate_lookup}` resolves to `'Rates'!$A$2:$B$6` -- 2 columns x 5 rows,
    a rectangle Excel's List validation refuses as a source (it wants a
    single cell, a single column, or a single row)."""
    pack = patched_pack(tmp_path, units=_with_validation({"list": "={rate_lookup}"}))
    with pytest.raises(R.RecipeError, match="single cell, a single column, or a single row"):
        B.generate(legacy_example(), pack)


@pytest.mark.parametrize("ref, spans", [
    ("'Codes'!$A:$A", False),                         # whole column, same column
    ("'Codes'!$A:$B", True),                          # whole column, different columns
    ("Codes!1:3", True),                              # whole row, different rows
    ("Codes!2:2", False),                             # whole row, same row
    ("CostCodes", False),                             # named range: fail-open
    ("'Codes'!CostCodes", False),                     # sheet-qualified named range: fail-open
    ("'Bob''s Codes'!$A$2:$A$9", False),              # escaped apostrophe sheet, single column
    ("'Bob''s Codes'!$A$2:$B$9", True),               # escaped apostrophe sheet, a rectangle
    ("'My Codes'!$A$2:$B$9", True),                   # quoted sheet (space), a rectangle
    ("'Codes'!$A$2:$H$2", False),                     # single row, many columns
    ("'Codes'!$A$2", False),                          # single cell
    ("$a$2:$b$9", True),                              # lowercase columns still count
    ("OFFSET(Codes!$A$1,0,0,5,2)", False),            # not a plain range: documented fail-open
])
def test_range_spans_rows_and_columns(ref, spans):
    assert B._range_spans_rows_and_columns(ref) is spans


def test_a_blank_optional_cell_still_gets_its_dropdown(tmp_path):
    spec = legacy_example()
    for line in spec["items"][1]["cost_details"][0]["lines"]:
        line.pop("uom", None)                                    # D resolves empty
    pack = patched_pack(tmp_path, units=_with_validation({"list": ["LF"]}, optional=True))
    ws = B.generate(spec, pack)["2 - 8IN MAIN"]
    r = _row_of(ws, "A", "8-inch ductile iron pipe")
    assert ws[f"D{r}"].value is None
    assert f"D{r}" in _dv(ws)[("list", '"LF"')]


def test_a_cell_skipped_by_if_it_gets_no_dropdown(tmp_path):
    """`if_it:` skips the placement before `_dropdown` ever runs -- only the
    one matching repeat element across the whole workbook gets the list."""
    def units(doc):
        row = doc["units"]["cost_block"]["rows"][1]
        row["cells"]["D"]["validation"] = {"list": ["LF", "TON", "DAY"]}
        row["cells"]["D"]["if_it"] = {"name": "Bedding rock"}

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units))["2 - 8IN MAIN"]
    rock = _row_of(ws, "A", "Bedding rock")
    assert _dv(ws)[("list", '"LF,TON,DAY"')] == [f"D{rock}"]


def test_a_validation_range_that_resolves_empty_stops_the_build(tmp_path):
    """A range's placeholder is free to resolve to nothing the same way any
    other cell's can -- but a bad list is never excused, even though the
    cell's own (unrelated) value might have been allowed to go blank."""
    def units(doc):
        doc["units"]["probe"] = {"rows": [
            {"cells": {"A": {"value": "x", "validation": {"list": "={p.rng}"}}}}]}

    def recipes(doc):
        doc["recipes"]["setup"]["units"].insert(0, {"unit": "probe", "params": {"rng": ""}})

    with pytest.raises(R.RecipeError, match="validation list"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units, recipes=recipes))


def test_a_bad_ref_in_a_validation_range_names_the_sheet_and_cell(tmp_path):
    pack = patched_pack(tmp_path, units=_with_validation({"list": "={ref:nope}"}))
    with pytest.raises(R.RecipeError, match=r"sheet '2 - 8IN MAIN' cell D\d+.*validation list"):
        B.generate(legacy_example(), pack)


def test_a_literal_list_dropdown_survives_save_and_reload(tmp_path):
    pack = patched_pack(tmp_path, units=_with_validation({"list": ["LF", "TON", "DAY"]}))
    wb = B.generate(legacy_example(), pack)
    out = tmp_path / "out.xlsx"
    wb.save(out)

    reloaded = openpyxl.load_workbook(out)
    ws = reloaded["2 - 8IN MAIN"]
    rows = [_row_of(ws, "A", n) for n in ("8-inch ductile iron pipe", "Bedding rock", "Trench box")]
    assert _dv(ws)[("list", '"LF,TON,DAY"')] == sorted(f"D{r}" for r in rows)


@pytest.mark.parametrize("bad", [
    "LF", {"list": []}, {"list": ["a,b"]}, {"list": "LF"}, {"choices": ["LF"]},
])
def test_a_malformed_validation_stops_the_build(tmp_path, bad):
    with pytest.raises(ValueError, match="validation"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=_with_validation(bad)))


def test_a_literal_list_over_255_chars_is_refused_at_layout(tmp_path):
    long_choice = "A" * 300
    with pytest.raises(R.RecipeError, match="255"):
        B.generate(legacy_example(),
                   patched_pack(tmp_path, units=_with_validation({"list": [long_choice]})))


def test_a_literal_choice_with_a_double_quote_is_refused_at_layout(tmp_path):
    with pytest.raises(R.RecipeError, match="double quotes"):
        B.generate(legacy_example(),
                   patched_pack(tmp_path, units=_with_validation({"list": ['L"F']})))


def test_a_bare_equals_range_is_refused_at_layout(tmp_path):
    with pytest.raises(R.RecipeError, match="validation"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=_with_validation({"list": "="})))


# ---- dropdown strictness ----------------------------------------------------

def _dvs(ws):
    return list(ws.data_validations.dataValidation)


def test_a_dropdown_rejects_typed_values_by_default(tmp_path):
    pack = patched_pack(tmp_path, units=_with_validation({"list": ["LF", "TON", "DAY"]}))
    (dv,) = _dvs(B.generate(legacy_example(), pack)["2 - 8IN MAIN"])
    assert dv.showErrorMessage is True
    assert dv.errorStyle == "stop"
    assert dv.errorTitle == "Not on the list"
    assert dv.error == "Pick a value from the list."


def test_a_non_strict_dropdown_allows_other_values(tmp_path):
    pack = patched_pack(tmp_path, units=_with_validation({"list": ["LF"], "strict": False}))
    (dv,) = _dvs(B.generate(legacy_example(), pack)["2 - 8IN MAIN"])
    assert not dv.showErrorMessage


def test_strict_and_non_strict_dropdowns_on_one_list_stay_separate(tmp_path):
    def units(doc):
        cells = doc["units"]["cost_block"]["rows"][1]["cells"]
        cells["A"]["validation"] = {"list": ["LF", "TON"]}
        cells["D"]["validation"] = {"list": ["LF", "TON"], "strict": False}

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units))["2 - 8IN MAIN"]
    by_strict = {bool(dv.showErrorMessage): sorted(str(r)[0] for r in dv.sqref.ranges)
                 for dv in _dvs(ws)}
    assert set(by_strict) == {True, False}
    assert set(by_strict[True]) == {"A"}
    assert set(by_strict[False]) == {"D"}


@pytest.mark.parametrize("bad", [
    {"list": ["LF"], "strict": "no"},
    {"list": ["LF"], "strict": 0},
    {"list": ["LF"], "strict": False, "extra": 1},
    {"strict": True},
])
def test_a_malformed_strict_is_refused(tmp_path, bad):
    with pytest.raises(R.RecipeError, match="validation"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=_with_validation(bad)))


# ---- dropdowns on protected sheets ------------------------------------------

def test_a_dropdown_on_a_locked_cell_of_a_protected_pack_is_refused(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["cells"]["D"]["validation"] = {"list": ["LF"]}

    def manifest(doc):
        doc["protect"] = {"sheets": True}

    with pytest.raises(R.RecipeError, match=r"cell D.*dropdown on a locked cell.*locked: false"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units, format=manifest))


def test_a_dropdown_on_an_unlocked_cell_of_a_protected_pack_is_allowed(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["cells"]["E"]["validation"] = {"list": ["1", "2"]}

    def tokens(doc):
        doc["styles"]["price_input"]["locked"] = False

    def manifest(doc):
        doc["protect"] = {"sheets": True}

    ws = B.generate(legacy_example(), patched_pack(
        tmp_path, units=units, tokens=tokens, format=manifest))["2 - 8IN MAIN"]
    assert _dvs(ws)


def test_a_cell_override_can_unlock_a_dropdown_on_a_protected_pack(tmp_path):
    def units(doc):
        cell = doc["units"]["cost_block"]["rows"][1]["cells"]["D"]
        cell["validation"] = {"list": ["LF"]}
        cell["locked"] = False

    def manifest(doc):
        doc["protect"] = {"sheets": True}

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units, format=manifest))[
        "2 - 8IN MAIN"]
    assert _dvs(ws)


# ---- conditional-format hardening --------------------------------------------

def test_a_bad_ref_in_a_conditional_names_its_site(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "F", "formula": "={ref:foo}>0", "style": "total"}]

    with pytest.raises(R.RecipeError,
                       match=r"sheet '2 - 8IN MAIN' row \d+ \(unit 'cost_block'.*conditional: .*foo"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


@pytest.mark.parametrize("conditional", [{}, "", None, 0])
def test_a_present_but_non_list_conditional_is_refused(tmp_path, conditional):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = conditional

    with pytest.raises(R.RecipeError, match="must be a list"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


def test_an_empty_conditional_list_is_a_no_op(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = []

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units))["2 - 8IN MAIN"]
    assert not _rules(ws)


@pytest.mark.parametrize("cells", ["ZZZ", "A:ZZZ", "XFE"])
def test_conditional_cells_past_xfd_are_refused(tmp_path, cells):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": cells, "formula": "=$F{r}>0", "style": "total"}]

    with pytest.raises(R.RecipeError, match="XFD"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


def test_conditional_cells_at_xfd_are_accepted(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "A:XFD", "formula": "=$F{r}>0", "style": "total"}]

    B.generate(legacy_example(), patched_pack(tmp_path, units=units))     # must not raise


def test_an_unanchored_ref_behind_an_escaped_apostrophe_sheet_is_refused(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "A:F", "formula": "='Bob''s'!F{r}>0", "style": "total"}]

    with pytest.raises(R.RecipeError, match="without a leading"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units))


def test_an_anchored_ref_behind_an_escaped_apostrophe_sheet_is_accepted(tmp_path):
    def units(doc):
        doc["units"]["cost_block"]["rows"][1]["conditional"] = [
            {"cells": "A:F", "formula": "='Bob''s'!$F${r}>0", "style": "total"}]

    B.generate(legacy_example(), patched_pack(tmp_path, units=units))     # must not raise


# ---- at least one visible sheet ----------------------------------------------

def test_a_workbook_with_every_sheet_hidden_is_refused(tmp_path):
    def recipes(doc):
        _book_recipes(doc)
        for sdef in doc["book"]:
            sdef["hidden"] = True

    spec = legacy_example()
    spec["items"] = []
    with pytest.raises(R.RecipeError, match="at least one sheet must be visible"):
        B.generate(spec, patched_pack(tmp_path, units=_book_units, recipes=recipes))


# ---- duplicate instance keys -----------------------------------------------

def test_a_duplicate_instance_key_on_one_sheet_is_refused(tmp_path):
    """Two `cost_block` instances publishing the same `cost[MATL]` key on
    "2 - 8IN MAIN" (two cost_details blocks both named "MATL") would
    otherwise overwrite each other's exports silently -- the second one
    wins, and whatever reads `cost[MATL]`'s exports (here, `cost_summary`)
    reads the wrong block's subtotal."""
    spec = legacy_example()
    spec["items"][1]["cost_details"].append(
        dict(spec["items"][1]["cost_details"][0]))  # another "MATL" block
    with pytest.raises(R.RecipeError, match=r"cost\[MATL\]"):
        B.generate(spec, patched_pack(tmp_path))


def test_an_exportless_unit_may_repeat_an_instance_key_on_one_sheet(tmp_path):
    """fieldbook's `code_rows` is laid out once per bid item on the Codes
    sheet, every call sharing the same (unit-name) key, with no `exports:` --
    nothing collides, so the duplicate-key guard must not fire for a unit
    that publishes no exports at all."""
    def units(doc):
        doc["units"]["no_export_probe"] = {
            "params": {"text": {"type": "string"}},
            "rows": [{"cells": {"A": {"value": "{p.text}"}}}],
        }

    def recipes(doc):
        doc["recipes"]["setup"]["units"].append(
            {"unit": "no_export_probe", "for_each": "item.labor",
             "params": {"text": "{it.slot}"}})

    spec = legacy_example()
    B.generate(spec, patched_pack(tmp_path, units=units, recipes=recipes))   # must not raise


# ---- ref_error: -------------------------------------------------------------

def _ref_error_probe(tmp_path, cell):
    def units(doc):
        doc["units"]["probe"] = {"rows": [{"cells": {"A": cell}}]}

    def recipes(doc):
        doc["recipes"]["setup"]["units"].insert(0, {"unit": "probe"})

    return B.generate(legacy_example(), patched_pack(tmp_path, units=units, recipes=recipes))


def test_ref_error_must_be_a_non_empty_string(tmp_path):
    with pytest.raises(R.RecipeError, match="ref_error must be a non-empty string"):
        _ref_error_probe(tmp_path, {"formula": "={ref:nope}", "ref_error": ""})


def test_a_non_string_ref_error_is_refused(tmp_path):
    with pytest.raises(R.RecipeError, match="ref_error must be a non-empty string"):
        _ref_error_probe(tmp_path, {"formula": "={ref:nope}", "ref_error": 5})


def test_ref_error_may_not_be_combined_with_optional(tmp_path):
    with pytest.raises(R.RecipeError, match="ref_error cannot be combined with optional"):
        _ref_error_probe(tmp_path, {"formula": "={ref:nope}", "ref_error": "friendlier",
                                     "optional": True})


def test_ref_error_renames_an_unresolved_ref_as_a_recipe_error(tmp_path):
    with pytest.raises(R.RecipeError, match="a much friendlier message"):
        _ref_error_probe(tmp_path, {"formula": "={ref:nope}", "ref_error": "a much friendlier message"})


def test_a_ref_error_template_that_itself_fails_to_resolve_raises_a_recipe_error(tmp_path):
    """ref_error is itself a template ({it}/{p.x}/etc resolve inside it) --
    if THAT fails to resolve, the build must still fail with a named
    RecipeError, not the bare KeyError/ValueError the engine used to leak."""
    with pytest.raises(R.RecipeError, match="ref_error itself failed to resolve"):
        _ref_error_probe(tmp_path, {"formula": "={ref:nope}", "ref_error": "bad {p.missing}"})


# ---- _coerce: a leading zero marks text, not a number ----------------------

@pytest.mark.parametrize("text,expected", [
    ("0", 0), ("10", 10), ("-5", -5), ("2A", "2A"),           # unchanged
    ("010", "010"), ("00", "00"), ("-010", "-010"), ("007", "007"),  # new: stays text
])
def test_coerce_keeps_a_leading_zero_padded_number_as_text(text, expected):
    assert B._coerce(text) == expected


# ---- freeze panes and print setup (fieldbook phase 3b) ---------------------

def test_a_recipe_may_freeze_panes_and_set_up_printing(tmp_path):
    def recipes(doc):
        doc["recipes"]["line_item"]["freeze"] = "A3"
        doc["recipes"]["line_item"]["print"] = {
            "orientation": "landscape", "fit_to_width": 1, "repeat_rows": "1:2"}

    wb = B.generate(legacy_example(), patched_pack(tmp_path, recipes=recipes))
    ws = wb["2 - 8IN MAIN"]
    assert ws.freeze_panes == "A3"
    assert ws.page_setup.orientation == "landscape"
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert (ws.page_setup.fitToWidth, ws.page_setup.fitToHeight) == (1, 0)
    assert ws.print_title_rows == "$1:$2"                  # openpyxl stores the rows absolute
    setup = wb["SETUP"]                                    # the setup recipe sets neither
    assert setup.freeze_panes is None
    assert setup.page_setup.orientation is None
    assert setup.print_title_rows is None


def test_a_book_sheet_may_freeze_panes_and_set_up_printing(tmp_path):
    def recipes(doc):
        _book_recipes(doc)
        doc["book"][0]["freeze"] = "B2"
        doc["book"][0]["print"] = {"orientation": "portrait", "fit_to_width": 1}

    wb = B.generate(legacy_example(), patched_pack(tmp_path, units=_book_units, recipes=recipes))
    front, back = wb["Front"], wb["Back"]
    assert front.freeze_panes == "B2"
    assert front.page_setup.orientation == "portrait"
    assert (front.page_setup.fitToWidth, front.page_setup.fitToHeight) == (1, 0)
    assert front.print_title_rows is None                  # no repeat_rows given
    assert back.freeze_panes is None and back.page_setup.orientation is None


def test_freeze_and_print_survive_save_and_reload(tmp_path):
    def recipes(doc):
        doc["recipes"]["line_item"]["freeze"] = "A2"
        doc["recipes"]["line_item"]["print"] = {
            "orientation": "landscape", "fit_to_width": 2, "repeat_rows": "1:1"}

    path = tmp_path / "out.xlsx"
    B.generate(legacy_example(), patched_pack(tmp_path, recipes=recipes)).save(path)
    ws = openpyxl.load_workbook(path)["2 - 8IN MAIN"]
    assert ws.freeze_panes == "A2"
    assert ws.page_setup.orientation == "landscape"
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert (ws.page_setup.fitToWidth, ws.page_setup.fitToHeight) == (2, 0)
    assert ws.print_title_rows == "$1:$1"


@pytest.mark.parametrize("bad, match", [
    ({"freeze": "a3"}, "freeze must be a cell reference"),
    ({"freeze": "A0"}, "freeze must be a cell reference"),
    ({"freeze": 3}, "freeze must be a cell reference"),
    ({"freeze": "ZZZ2"}, "freeze 'ZZZ2' is past Excel's last column, XFD"),
    ({"print": "landscape"}, "print must be a mapping"),
    ({"print": {}}, "print must be a mapping"),
    ({"print": {"paper": "A4"}}, r"print: unknown keys \['paper'\]"),
    ({"print": {"orientation": "sideways"}}, "print orientation must be portrait or landscape"),
    ({"print": {"fit_to_width": 0}}, "print fit_to_width must be a whole number of pages"),
    ({"print": {"fit_to_width": True}}, "print fit_to_width must be a whole number of pages"),
    ({"print": {"repeat_rows": "3"}}, "print repeat_rows must be"),
    ({"print": {"repeat_rows": "3:1"}}, "print repeat_rows must be"),
    ({"print": {"repeat_rows": "0:2"}}, "print repeat_rows must be"),
])
def test_malformed_freeze_or_print_is_refused_at_load(tmp_path, bad, match):
    def recipes(doc):
        doc["recipes"]["line_item"].update(bad)

    with pytest.raises(fp.FormatPackError, match=r"recipes\.line_item: " + match):
        patched_pack(tmp_path / "r", recipes=recipes)

    def book(doc):
        _book_recipes(doc)
        doc["book"][0].update(bad)

    with pytest.raises(fp.FormatPackError, match=r"book\[0\]: " + match):
        patched_pack(tmp_path / "b", units=_book_units, recipes=book)


# ---- concat / unique are checked at load (fieldbook phase 3b review) -------

@pytest.mark.parametrize("params, match", [
    ({"codes": {"concat": [], "sort": "code"}}, r"params\.codes: concat: unknown keys \['sort'\]"),
    ({"codes": {"concat": [], "unique": ""}}, r"params\.codes: concat: unique must be"),
    ({"codes": {"concat": [], "unique": 3}}, r"params\.codes: concat: unique must be"),
    ({"codes": {"unique": "code", "for_each": "config.items", "as": "{it}"}},
     r"params\.codes: unique: needs concat:"),
    ({"codes": {"concat": [{"concat": [], "uniq": "code"}]}},
     r"params\.codes\.concat\[0\]: concat: unknown keys \['uniq'\]"),
])
def test_a_malformed_concat_param_is_refused_at_load(tmp_path, params, match):
    def recipes(doc):
        units = doc["recipes"]["line_item"]["units"]
        units.append({"unit": units[1]["unit"], "key": "concat_probe", "params": params})

    with pytest.raises(fp.FormatPackError, match=r"recipes\.line_item\[\d+\]: " + match):
        patched_pack(tmp_path, recipes=recipes)


def test_a_well_formed_concat_param_loads(tmp_path):
    def recipes(doc):
        units = doc["recipes"]["line_item"]["units"]
        units.append({"unit": units[1]["unit"], "key": "concat_probe", "params": {
            "codes": {"unique": True, "concat": [["a"], {"concat": [], "unique": "code"}]}}})

    patched_pack(tmp_path, recipes=recipes)            # must not raise


def test_every_shipped_and_fixture_pack_passes_the_concat_check():
    for root in (fp.BUNDLED, fp.ENGINE_DIR / "tests" / "fixtures" / "packs" / "legacy_demo",
                 fp.ENGINE_DIR / "tests" / "fixtures" / "packs" / "alt"):
        fp.load_pack(root)


# ---- book sheets repeated per element (for_each) ---------------------------

def _pkg_units(doc):
    doc["units"]["pkg_echo"] = {
        "params": {"name": {"type": "string"}},
        "rows": [{"cells": {"A": {"value": "{p.name}"}}, "mark": "echo"}],
        "exports": {"echo": "$A${row:echo}"},
    }


def _pkg_recipes(probe_refs=(), **extra):
    def recipes(doc):
        del doc["recipes"]["cover"]
        doc["book"] = [{"sheet_name": "Pkg {sheet.name}", "position": "after_items",
                        "for_each": "config.x-packages",
                        "units": [{"unit": "pkg_echo", "key": "pkg[{sheet.name}]",
                                   "params": {"name": "{sheet.name}"}}], **extra}]
        if probe_refs:
            doc["recipes"]["setup"]["units"].insert(0, {"unit": "probe"})
    return recipes


def _probe_units(refs):
    def units(doc):
        _pkg_units(doc)
        doc["units"]["probe"] = {"rows": [
            {"cells": {"A": {"formula": "={ref:" + ref + "}"}}} for ref in refs]}
    return units


def _pkg_spec(*names):
    spec = legacy_example()
    spec["x-packages"] = [{"name": n} for n in names]
    return spec


def test_a_for_each_book_entry_lays_out_one_sheet_per_element(tmp_path):
    wb = B.generate(_pkg_spec("A", "B"),
                    patched_pack(tmp_path, units=_pkg_units, recipes=_pkg_recipes()))
    assert wb.sheetnames[-2:] == ["Pkg A", "Pkg B"]
    assert wb["Pkg A"]["A1"].value == "A"
    assert wb["Pkg B"]["A1"].value == "B"


def test_colliding_sheet_names_are_suffixed(tmp_path):
    stem = "S" * 27
    titles = [f"Pkg {stem}", f"Pkg {stem[:23]} (2)"]
    refs = [f"{t}!pkg[{stem}{c}].echo" for t, c in zip(titles, "XY")]
    pack = patched_pack(tmp_path, units=_probe_units(refs),
                        recipes=_pkg_recipes(probe_refs=refs))
    wb = B.generate(_pkg_spec(stem + "X", stem + "Y"), pack)
    assert wb.sheetnames[-2:] == titles
    assert all(len(t) <= 31 for t in titles)
    assert [wb["SETUP"][f"A{r}"].value for r in (1, 2)] == [f"='{t}'!$A$1" for t in titles]
    assert [wb[t]["A1"].value for t in titles] == [stem + "X", stem + "Y"]


def test_an_empty_for_each_list_lays_out_no_sheet(tmp_path):
    pack = patched_pack(tmp_path, units=_pkg_units, recipes=_pkg_recipes())
    legacy = B.generate(legacy_example(), patched_pack(tmp_path / "l"))
    wb = B.generate(_pkg_spec(), pack)
    assert wb.sheetnames == [s for s in legacy.sheetnames if s != "Bid Summary"]


def test_a_for_each_book_entry_honours_when_per_element(tmp_path):
    spec = _pkg_spec("A", "B")
    spec["x-packages"][0]["keep"] = True
    spec["x-packages"][1]["keep"] = False
    pack = patched_pack(tmp_path, units=_pkg_units, recipes=_pkg_recipes(when="sheet.keep"))
    wb = B.generate(spec, pack)
    assert "Pkg A" in wb.sheetnames and "Pkg B" not in wb.sheetnames


def test_an_expanded_book_title_colliding_with_an_item_sheet_is_refused(tmp_path):
    spec = _pkg_spec("A")
    spec["items"][0]["sheet"] = "pkg a"
    pack = patched_pack(tmp_path, units=_pkg_units, recipes=_pkg_recipes())
    with pytest.raises(R.RecipeError, match="collides with book sheet_name 'Pkg A'"):
        B.generate(spec, pack)


@pytest.mark.parametrize("bad, match", [
    ({"sheet_name": "Pkg"}, r"sheet_name must carry a \{sheet\.\.\.\} token"),
    ({"for_each": "items"}, r"for_each must be a config\. or book\. path"),
])
def test_a_malformed_for_each_book_entry_is_refused_at_load(tmp_path, bad, match):
    def recipes(doc):
        _pkg_recipes()(doc)
        doc["book"][0].update(bad)

    with pytest.raises(fp.FormatPackError, match=r"book\[0\]: " + match):
        patched_pack(tmp_path, units=_pkg_units, recipes=recipes)


def test_a_sheet_token_without_for_each_is_refused_at_load(tmp_path):
    def recipes(doc):
        _pkg_recipes()(doc)
        del doc["book"][0]["for_each"]

    with pytest.raises(fp.FormatPackError, match=r"book\[0\]: .*\{sheet\.\.\.\} .*for_each"):
        patched_pack(tmp_path, units=_pkg_units, recipes=recipes)


@pytest.mark.parametrize("name, taken, want", [
    ("Plain", set(), "Plain"),
    ("a[b]:c*d?e/f\\g", set(), "a-b--c-d-e-f-g"),
    ("X" * 40, set(), "X" * 31),
    ("Pkg A", {"pkg a"}, "Pkg A (2)"),
    ("Pkg A", {"Pkg A", "Pkg A (2)"}, "Pkg A (3)"),
    ("X" * 31, {"X" * 31}, "X" * 27 + " (2)"),
])
def test_sheet_title(name, taken, want):
    assert B.sheet_title(name, taken) == want


# ---- repeat_cols bands, column landmarks and per-row exports ---------------

def _band_units(band=None):
    """A `band` unit: a header row with one column pair per name (K/L, M/N,
    ...), and a `repeat:` row per line publishing `line[<key>].qty`."""
    def units(doc):
        repeat_cols = {"over": "p.names", "start": "K", "width": 2,
                       "cells": {"+0": {"value": "{v}"},
                                 "+1": {"formula": "=\"{c}|{c+1}|{c-0}\""}}}
        repeat_cols.update(band or {})
        doc["units"]["band"] = {
            "params": {"names": {"type": "list"}, "lines": {"type": "list"}},
            "rows": [
                {"cells": {"A": {"value": "HDR"},
                           "B": {"formula": "=INDEX(${col:vf}{r}:${col:vl+1}{r},1)"}},
                 "repeat_cols": repeat_cols,
                 "mark_cols_first": "vf", "mark_cols_last": "vl"},
                {"repeat": "p.lines",
                 "cells": {"A": {"value": "{it.key}"}, "B": {"formula": "={it.q}"}},
                 "export_each": {"key": "line[{it.key}]", "cells": {"qty": "$B${r}"}}},
            ],
            "exports": {"band_first": "${col:vf}$1"},
        }
        doc["units"]["probe"] = {"rows": [
            {"cells": {"A": {"formula": "={ref:SETUP!rows.line[a].qty}"},
                       "B": {"formula": "={ref:SETUP!rows.band_first}"}}}]}
    return units


def _band_recipes(names, lines=({"key": "a", "q": 1}, {"key": "b", "q": 2})):
    def recipes(doc):
        doc["recipes"]["setup"]["units"].insert(
            0, {"unit": "band", "key": "rows",
                "params": {"names": list(names), "lines": [dict(x) for x in lines]}})
        doc["recipes"]["line_item"]["units"].insert(0, {"unit": "probe"})
    return recipes


def test_repeat_cols_lays_one_column_pair_per_element(tmp_path):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][0]["cells"]["C"] = {"value": "{col:vl-2}|{col:vf+3}"}

    pack = patched_pack(tmp_path, units=units, recipes=_band_recipes(["X", "Y", "Z"]))
    ws = B.generate(legacy_example(), pack)["SETUP"]
    assert [ws[f"{c}1"].value for c in "KMO"] == ["X", "Y", "Z"]
    assert [ws[f"{c}1"].value for c in "LNP"] == ['="K|L|K"', '="M|N|M"', '="O|P|O"']
    assert ws["Q1"].value is None
    assert ws["B1"].value == "=INDEX($K1:$P1,1)"
    assert ws["C1"].value == "M|N"


def test_an_empty_repeat_cols_list_places_nothing_and_landmarks_read_start(tmp_path):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][0]["cells"]["C"] = {"value": "{col:vl}"}

    pack = patched_pack(tmp_path, units=units, recipes=_band_recipes([]))
    ws = B.generate(legacy_example(), pack)["SETUP"]
    assert all(ws.cell(1, c).value is None for c in range(11, 20))
    assert ws["B1"].value == "=INDEX($K1:$L1,1)"
    assert ws["C1"].value == "K"


def test_a_band_on_a_repeat_row_binds_its_as_name_beside_it(tmp_path):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][1]["repeat_cols"] = {
            "over": "p.vendors", "start": "D", "width": 1, "as": "q",
            "cells": {"+0": {"value": "{it.key}:{q.name}@{c}{r}"}}}

    def recipes(doc):
        _band_recipes(["X"])(doc)
        doc["recipes"]["setup"]["units"][0]["params"]["vendors"] = [
            {"name": "Acme"}, {"name": "Bolt"}]

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units, recipes=recipes))["SETUP"]
    assert [[ws[f"{c}{r}"].value for c in "DE"] for r in (2, 3)] == [
        ["a:Acme@D2", "a:Bolt@E2"], ["b:Acme@D3", "b:Bolt@E3"]]


def test_export_each_publishes_one_export_per_repeated_row(tmp_path):
    def units(doc):
        _band_units()(doc)

    pack = patched_pack(tmp_path, units=units, recipes=_band_recipes(["X"]))
    wb = B.generate(legacy_example(), pack)
    assert [wb["SETUP"][f"A{r}"].value for r in (2, 3)] == ["a", "b"]
    probe = wb["2 - 8IN MAIN"]
    assert probe["A1"].value == "=" + B.sheet_qualified("SETUP", "$B$2")
    assert probe["B1"].value == "=" + B.sheet_qualified("SETUP", "$K$1")


def test_export_each_key_naming_a_missing_field_is_refused(tmp_path):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][1]["export_each"]["key"] = "line[{it.nope}]"

    pack = patched_pack(tmp_path, units=units, recipes=_band_recipes(["X"]))
    with pytest.raises(R.RecipeError, match=r"export_each key .*nope"):
        B.generate(legacy_example(), pack)


def test_export_each_rows_sharing_a_key_are_refused(tmp_path):
    def units(doc):
        _band_units()(doc)

    pack = patched_pack(tmp_path, units=units, recipes=_band_recipes(
        ["X"], lines=({"key": "a", "q": 1}, {"key": "a", "q": 2})))
    with pytest.raises(R.RecipeError, match=r"line\[a\].*twice"):
        B.generate(legacy_example(), pack)


def test_a_unit_with_only_export_each_obeys_the_instance_key_guard(tmp_path):
    def units(doc):
        _band_units()(doc)
        del doc["units"]["band"]["exports"]

    def recipes(doc):
        _band_recipes(["X"])(doc)
        doc["recipes"]["setup"]["units"].insert(1, dict(doc["recipes"]["setup"]["units"][0]))

    with pytest.raises(R.RecipeError, match=r"instance key 'rows'.*twice"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units, recipes=recipes))


@pytest.mark.parametrize("band, match", [
    ({"over": None}, r"repeat_cols needs over: p\.NAME"),
    ({"over": "names"}, r"repeat_cols needs over: p\.NAME"),
    ({"start": "k1"}, r"repeat_cols start must be a column letter"),
    ({"width": 0}, r"repeat_cols width must be an integer >= 1"),
    ({"width": True}, r"repeat_cols width must be an integer >= 1"),
    ({"cells": {}}, r"repeat_cols cells must be a non-empty mapping"),
    ({"cells": {"+2": {"value": "x"}}}, r"repeat_cols cells key '\+2'"),
    ({"cells": {"1": {"value": "x"}}}, r"repeat_cols cells key '1'"),
    ({"cells": {"+0": "x"}}, r"repeat_cols cells\['\+0'\] must be a mapping"),
    ({"as": "it"}, r"repeat_cols as must be an identifier"),
    ({"as": "a b"}, r"repeat_cols as must be an identifier"),
    ({"bogus": 1}, r"repeat_cols: unknown keys \['bogus'\]"),
])
def test_a_malformed_repeat_cols_is_refused_at_load(tmp_path, band, match):
    def units(doc):
        _band_units(band)(doc)
        rc = doc["units"]["band"]["rows"][0]["repeat_cols"]
        for k, v in band.items():
            if v is None:
                rc.pop(k)

    with pytest.raises(fp.FormatPackError, match=r"units\.band\.rows\[0\]: " + match):
        patched_pack(tmp_path, units=units, recipes=_band_recipes(["X"]))


@pytest.mark.parametrize("edit, match", [
    (lambda rows: rows[0].pop("repeat_cols"), r"rows\[0\]: mark_cols_first needs repeat_cols"),
    (lambda rows: rows[0].update(export_each={"key": "k", "cells": {"q": "$B{r}"}}),
     r"rows\[0\]: export_each needs a `repeat:` row"),
    (lambda rows: rows[1]["export_each"].pop("key"), r"rows\[1\]: export_each needs key"),
    (lambda rows: rows[1]["export_each"].update(cells={}),
     r"rows\[1\]: export_each needs key"),
])
def test_malformed_column_marks_and_export_each_are_refused_at_load(tmp_path, edit, match):
    def units(doc):
        _band_units()(doc)
        edit(doc["units"]["band"]["rows"])

    with pytest.raises(fp.FormatPackError, match=r"units\.band\." + match):
        patched_pack(tmp_path, units=units, recipes=_band_recipes(["X"]))


# ---- cells past a band, per-element styles, comments, hidden columns -------

def test_a_cell_keyed_by_a_column_landmark_lands_past_the_band(tmp_path):
    def units(doc):
        _band_units()(doc)
        rows = doc["units"]["band"]["rows"]
        rows[0]["cells"]["{col:vl+2}"] = {"value": "AFTER"}
        rows[1]["cells"]["{col:vl+3}"] = {"formula": "={col:vl+2}{r}&\"!\""}

    pack = patched_pack(tmp_path, units=units, recipes=_band_recipes(["X", "Y"]))
    ws = B.generate(legacy_example(), pack)["SETUP"]
    assert ws["O1"].value == "AFTER"                      # vl = M, +2 = O
    assert [ws[f"P{r}"].value for r in (2, 3)] == ['=O2&"!"', '=O3&"!"']


def test_a_cell_keyed_by_an_unknown_landmark_is_refused(tmp_path):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][0]["cells"]["{col:nope+1}"] = {"value": "x"}

    with pytest.raises(R.RecipeError, match=r"cell key '\{col:nope\+1\}'.*nope"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units,
                                                  recipes=_band_recipes(["X"])))


def test_a_style_may_be_chosen_per_band_element(tmp_path):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][0]["repeat_cols"]["cells"]["+0"] = {"value": "{v.n}",
                                                                         "style": "{v.s}"}

    def recipes(doc):
        _band_recipes([])(doc)
        doc["recipes"]["setup"]["units"][0]["params"]["names"] = [
            {"s": "total", "n": "a"}, {"s": "price_input", "n": "b"}]

    ws = B.generate(legacy_example(), patched_pack(tmp_path, units=units, recipes=recipes))["SETUP"]
    assert ws["M1"].font.color.rgb.endswith("00B050")     # price_input's green
    assert ws["M1"].fill.fgColor.rgb != ws["K1"].fill.fgColor.rgb   # `total` carries a fill


def test_a_style_that_resolves_to_no_style_is_refused(tmp_path):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][0]["repeat_cols"]["cells"]["+0"]["style"] = "x_{v}"

    with pytest.raises(R.RecipeError, match=r"style 'x_A' is not in tokens.styles"):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units,
                                                  recipes=_band_recipes(["A"])))


def test_a_comment_is_written_only_when_it_resolves_to_text(tmp_path):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][1]["cells"]["A"]["comment"] = "{it.note}"

    pack = patched_pack(tmp_path, units=units, recipes=_band_recipes(
        ["X"], lines=({"key": "a", "q": 1, "note": "from a lump sum"},
                      {"key": "b", "q": 2, "note": ""})))
    ws = B.generate(legacy_example(), pack)["SETUP"]
    assert ws["A2"].comment.text == "from a lump sum"
    assert ws["A3"].comment is None


def test_hide_column_hides_the_cells_column(tmp_path):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][0]["cells"]["{col:vl+2}"] = {"value": "H", "hide_column": True}

    ws = B.generate(legacy_example(), patched_pack(
        tmp_path, units=units, recipes=_band_recipes(["X"])))["SETUP"]
    assert ws.column_dimensions["M"].hidden


@pytest.mark.parametrize("cell, match", [
    ({"value": "x", "hide_column": "yes"}, "hide_column must be true or false"),
    ({"value": "x", "comment": ""}, "comment must be a non-empty template"),
])
def test_malformed_comment_and_hide_column_are_refused(tmp_path, cell, match):
    def units(doc):
        _band_units()(doc)
        doc["units"]["band"]["rows"][0]["cells"]["C"] = cell

    with pytest.raises(R.RecipeError, match=match):
        B.generate(legacy_example(), patched_pack(tmp_path, units=units,
                                                  recipes=_band_recipes(["X"])))


def test_a_for_each_element_sees_its_position_as_it_index():
    scopes = {"config": {"xs": ["a", "b", "c"]}}
    got = R.resolve_value({"for_each": "config.xs", "as": {"v": "{it}", "i": {"from": "it_index"}}},
                          scopes)
    assert got == [{"v": "a", "i": 0}, {"v": "b", "i": 1}, {"v": "c", "i": 2}]


# ---- literal text -------------------------------------------------------------

def _with_note(note):
    spec = legacy_example()
    (item,) = [i for i in spec["items"] if i.get("sheet") == "2 - 8IN MAIN"]
    (line,) = [ln for blk in item["cost_details"] for ln in blk.get("lines", [])
               if ln["name"] == "Bedding rock"]
    line["note"] = note
    return spec


def _note_units(doc):
    row = doc["units"]["cost_block"]["rows"][1]          # the `repeat: p.lines` row
    row["cells"]["H"] = {"value": "{it.note}", "comment": "{it.note}",
                        "if_it": {"name": "Bedding rock"}}


def test_literal_text_is_never_resolved_as_placeholders(tmp_path):
    # Bound literally: a value's braces stay braces, even ones that look like
    # a placeholder the resolver knows ({r}) or one it does not ({monthly}).
    text = "per DOT index {monthly}, row {r}, {ref:nowhere}"
    spec = _with_note(B.literal(text))
    ws = B.generate(spec, patched_pack(tmp_path, units=_note_units))["2 - 8IN MAIN"]
    (cell,) = [c for c in ws["H"] if c.value == text]
    assert cell.comment.text == text
    assert ws[f"A{cell.row}"].value == "Bedding rock"


def test_the_same_text_unmarked_is_resolved(tmp_path):
    with pytest.raises(ValueError, match=r"unresolved placeholder \{monthly\}"):
        B.generate(_with_note("per DOT index {monthly}"), patched_pack(tmp_path, units=_note_units))
