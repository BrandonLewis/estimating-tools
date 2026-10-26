"""``simulate_delete_row`` rewrites references the way Excel does."""
from __future__ import annotations

from openpyxl import Workbook

from estimating_tools.estimate_spec.tests._fieldbook import simulate_delete_row


def _book():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    other = wb.create_sheet("T Two")
    return wb, ws, other


def test_refs_below_move_up_and_a_single_cell_ref_to_the_row_is_ref_error():
    wb, ws, _ = _book()
    ws["A1"] = "=A5+A6+A3"
    simulate_delete_row(ws, 5)
    assert ws["A1"].value == "=#REF!+A5+A3"


def test_a_range_shrinks_instead_of_becoming_ref_error():
    wb, ws, _ = _book()
    ws["A1"] = "=SUM(B3:B8)"        # row inside
    ws["A2"] = "=SUM(B3:B5)"        # row is the last endpoint
    ws["A3"] = "=SUM(B5:B8)"        # row is the first endpoint
    ws["A4"] = "=SUM(B5:B5)"        # the whole range is the row
    ws["A9"] = "=SUM(B6:B8)+SUM(B1:B2)"
    simulate_delete_row(ws, 5)
    assert ws["A1"].value == "=SUM(B3:B7)"
    assert ws["A2"].value == "=SUM(B3:B4)"
    assert ws["A3"].value == "=SUM(B5:B7)"
    assert ws["A4"].value == "=SUM(#REF!)"
    assert ws["A8"].value == "=SUM(B5:B7)+SUM(B1:B2)"


def test_other_sheet_refs_are_untouched_and_refs_into_this_sheet_follow():
    wb, ws, other = _book()
    ws["A1"] = "='T Two'!B9+'T Two'!B5+SUM('T Two'!C3:C8)"
    other["A1"] = "=S!B9+SUM(S!C3:C8)+B9"
    simulate_delete_row(ws, 5)
    assert ws["A1"].value == "='T Two'!B9+'T Two'!B5+SUM('T Two'!C3:C8)"
    assert other["A1"].value == "=S!B8+SUM(S!C3:C7)+B9"
