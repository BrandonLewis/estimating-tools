"""Build a HeavyCivil workbook from a get_estimate_workbook payload.

``build(doc, style=..., exported_at=..., host=...)`` returns an openpyxl
Workbook; ``save(wb, path)`` writes it byte-deterministically (the same
payload and stamp always give the same bytes). Nothing here talks to
BidSheet.ai: the skill fetches the payload over MCP and hands the JSON in.
"""
from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.packaging.custom import StringProperty
from openpyxl.workbook.defined_name import DefinedName

from estimating_tools.estimate_spec.build import sheet_qualified
from estimating_tools.heavycivil import FORMAT_NAME, FORMAT_VERSION, baseline, layout, model
from estimating_tools.heavycivil.columns import BF
from estimating_tools.heavycivil.sheets.bid_form import write_bid_form
from estimating_tools.heavycivil.sheets.checks import write_checks
from estimating_tools.heavycivil.sheets.common import last_column_letter, set_value
from estimating_tools.heavycivil.sheets.cost_report import write_cost_report
from estimating_tools.heavycivil.sheets.crews import write_crews
from estimating_tools.heavycivil.sheets.flat import write_flat
from estimating_tools.heavycivil.sheets.rates import write_rates
from estimating_tools.heavycivil.sheets.recap import write_recap
from estimating_tools.heavycivil.sheets.schedule import CAL_SHEET, write_calendar, write_schedule
from estimating_tools.heavycivil.sheets.shares import SHEET as SHARES_SHEET
from estimating_tools.heavycivil.sheets.shares import write_shares
from estimating_tools.xlsx_determinism import BUILD_STAMP, normalize_xlsx

WB_SHEET = "_wb"


def build(doc, *, style, exported_at, host) -> Workbook:
    data = model.parse(doc)
    # The payload as built (``merge_pages``' output when paged -- plain
    # ``json.loads`` data, no Decimal), for read-back. Encoded first: a value
    # JSON cannot carry (NaN in a field the model never reads) is refused
    # before any sheet is written.
    chunks, meta = baseline.encode(doc)
    plan = layout.plan(data, style)
    stamp = {"exported_at": exported_at, "host": host}

    book = Workbook()
    bid_ws = book.active
    bid_ws.title = plan.tabs["bid_form"]
    recap_ws = book.create_sheet(plan.tabs["recap"])
    cost_ws = book.create_sheet(plan.tabs["cost_report"])
    crews_ws = book.create_sheet(plan.tabs["crews"])
    rates_ws = book.create_sheet(plan.tabs["rates"])
    schedule_ws = book.create_sheet(plan.tabs["schedule"])
    checks_ws = book.create_sheet(plan.tabs["checks"])
    flat_ws = book.create_sheet(plan.tabs["flat"])
    shares_ws = book.create_sheet(SHARES_SHEET)
    cal_ws = book.create_sheet(CAL_SHEET)
    wb_ws = book.create_sheet(WB_SHEET)

    write_bid_form(bid_ws, plan, data, style)
    write_recap(recap_ws, plan, data, style)
    write_cost_report(cost_ws, plan, data, style)
    write_crews(crews_ws, plan, data, style)
    write_rates(rates_ws, plan, data, style)
    write_schedule(schedule_ws, plan, data, style)
    checks = write_checks(checks_ws, plan, data, style, stamp)
    write_flat(flat_ws, plan, data, style)
    write_shares(shares_ws, plan, data)
    write_calendar(cal_ws, plan)
    book.active = 0

    wb_ws.sheet_state = "hidden"
    wb_ws["A1"] = json.dumps(layout.layout_json(plan, data, style, meta), sort_keys=True,
                             separators=(",", ":"), default=str)
    for n, chunk in enumerate(chunks, start=1):
        # Pinned to text: a last chunk of base64 padding ("==") would
        # otherwise be stored as a formula.
        set_value(wb_ws[f"B{n}"], chunk)

    names = {
        "hc.cost_report": (plan.tabs["cost_report"],
                           (f"$A${layout.HEADER_ROW}:${last_column_letter('cost_report')}"
                            f"${max(plan.cost_report_last_row, layout.HEADER_ROW)}")),
        "hc.checks": (plan.tabs["checks"],
                      f"$A${layout.HEADER_ROW}:$I${max(checks['last'], layout.HEADER_ROW)}"),
        "hc.layout": (WB_SHEET, "$A$1"),
        "hc.bid_form": (plan.tabs["bid_form"],
                        f"$A${BF.HEADER_ROW}:${last_column_letter('bid_form')}"
                        f"${plan.bid_items_span[1]}"),
    }
    if plan.bid_total_rows:      # always, the implicit total included
        names["hc.bid_totals"] = (plan.tabs["bid_form"],
                                  f"$A${plan.bid_total_rows[0].row}:${BF.TOTAL}"
                                  f"${plan.bid_total_rows[-1].row}")
    if plan.schedule_rows:
        names["hc.schedule"] = (plan.tabs["schedule"],
                                f"$A${layout.HEADER_ROW}:${last_column_letter('schedule')}"
                                f"${max(plan.schedule_rows.values())}")
    for name, (sheet, ref) in sorted(names.items()):
        book.defined_names[name] = DefinedName(name, attr_text=sheet_qualified(sheet, ref))

    props = {"estimate_format": f"{FORMAT_NAME} {FORMAT_VERSION}",
             "estimate_id": str(data.estimate.id), "bidsheet_host": host,
             "exported_at": exported_at, "style_pack": f"{style.name} {style.version}",
             "source_sha": data.source_sha}
    for name, value in props.items():
        book.custom_doc_props.append(StringProperty(name=name, value=value))
    book.properties.created = BUILD_STAMP
    book.properties.modified = BUILD_STAMP
    return book


def _match_new_file_mode(tmp):
    """``mkstemp`` creates the temp file 0600; give it the mode a normally
    created file gets (0666 less the umask). The umask can only be read by
    setting it, so it is set and restored at once. On Windows chmod only
    toggles read-only, which this leaves off."""
    old = os.umask(0)
    os.umask(old)
    os.chmod(tmp, 0o666 & ~old)


def save(book, path) -> None:
    """Write atomically: a temp file beside ``path`` is saved and normalized,
    then moved over it. On any failure the temp file is removed and an
    existing ``path`` is left as it was."""
    path = Path(path)
    fd, tmp = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=".xlsx.tmp", dir=path.parent)
    os.close(fd)
    try:
        book.save(tmp)
        normalize_xlsx(tmp)
        _match_new_file_mode(tmp)
        os.replace(tmp, path)
    except BaseException:
        try:
            os.remove(tmp)
        except FileNotFoundError:
            pass
        raise
