"""Helpers for the read-back tests: build the fixture book, edit it the way
Excel would (insert, copy, cut, delete rows), save, read it back."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from estimating_tools.heavycivil import build, style
from estimating_tools.heavycivil.columns import CR, CW

FIXTURE = Path(__file__).parent / "fixtures" / "workbook.json"
CR_TAB, CW_TAB, BF_TAB, RT_TAB = "Cost Report", "Crews", "Bid Form", "Rates"


def fixture_doc():
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def linked_doc():
    """The fixture with resource 1 (CAT 336) linked to equipment rate 7."""
    doc = fixture_doc()
    doc["rates"]["equipment"].append({"id": 7, "name": "CAT 336", "unit": "HR",
                                      "rate": "185.00"})
    for res in doc["resources"]:
        if res["id"] == 1:
            res["equipment_rate_id"] = 7
    return doc


def build_book(doc, path):
    """Build ``doc`` and save it to ``path``; returns ``path``."""
    book = build.build(doc, style=style.load_style(), exported_at="2026-10-20T09:00:00Z",
                       host="www.bidsheet.ai")
    build.save(book, path)
    return Path(path)


def open_book(path):
    return load_workbook(path)


def save_book(wb, path):
    wb.save(path)
    return Path(path)


def find_row(ws, kind, ident=None, *, kind_col=CR.KIND, id_col=CR.ID, parent=None,
             parent_col=CR.PARENT, start=1):
    """The first row at or after ``start`` of ``kind`` (with id ``ident`` and/or
    parent ``parent`` when given)."""
    for r in range(start, ws.max_row + 1):
        if ws[f"{kind_col}{r}"].value != kind:
            continue
        if ident is not None and ws[f"{id_col}{r}"].value != ident:
            continue
        if parent is not None and ws[f"{parent_col}{r}"].value != parent:
            continue
        return r
    raise AssertionError(f"no {kind} row (id {ident}, parent {parent})")


def cr_row(ws, kind, ident=None, *, parent=None, start=1):
    return find_row(ws, kind, ident, parent=parent, start=start)


def cw_row(ws, kind, ident=None, *, parent=None):
    return find_row(ws, kind, ident, kind_col=CW.KIND, id_col=CW.ID, parent=parent,
                    parent_col=CW.CREW)


def block_end(ws, activity_id):
    """The END marker row of ``activity_id``'s first block."""
    return cr_row(ws, "END", parent=activity_id)


def copy_row(ws, src, dst):
    """Copy row ``src`` onto row ``dst`` the way Excel's copy does: values as
    they are, formulas translated by the row offset."""
    for col in range(1, ws.max_column + 1):
        s, d = ws.cell(row=src, column=col), ws.cell(row=dst, column=col)
        if s.data_type == "f" and isinstance(s.value, str):
            d.value = Translator(s.value, origin=s.coordinate).translate_formula(d.coordinate)
        else:
            d.value = s.value
            if isinstance(s.value, str):
                d.data_type = s.data_type


def insert_copies(ws, first, last, before):
    """Insert copies of rows ``first..last`` above row ``before`` (Excel's
    "Insert Copied Cells"). Returns the first inserted row. ``before`` may
    fall inside ``first..last``: Excel copies the range as it was, so each
    source row below the insert is read from where the insert pushed it."""
    n = last - first + 1
    ws.insert_rows(before, n)
    for k in range(n):
        src = first + k
        copy_row(ws, src + n if src >= before else src, before + k)
    return before


def cut_rows(ws, first, last, before):
    """Move rows ``first..last`` above row ``before`` (Excel's "Insert Cut
    Cells"). Returns the row the block now starts on."""
    n = last - first + 1
    insert_copies(ws, first, last, before)
    if first >= before:
        first += n
    ws.delete_rows(first, n)
    return before if before < first else before - n


def type_row(ws, before, values):
    """Insert a blank row above ``before`` and type ``{letter: value}`` in it."""
    ws.insert_rows(before)
    for letter, value in values.items():
        ws[f"{letter}{before}"] = value
    return before


def read(path):
    from estimating_tools.heavycivil.readback import read_book
    return read_book(path)


def ops_named(cs, name):
    return [op for op in cs.ops if op["op"] == name]
