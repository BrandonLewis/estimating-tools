"""Reading input cells, and overwritten formulas (spec §2 Overwritten formulas,
Case and codes)."""
from __future__ import annotations

import datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook

from estimating_tools.heavycivil.tests import _readback as H
from estimating_tools.heavycivil.columns import BF, CR, CW
from estimating_tools.heavycivil.readback import ReadRefused, baseline, fields, scan
from estimating_tools.heavycivil.readback.errors import Problems
from estimating_tools.heavycivil.readback.fields import Input, number_equal, read_value
from estimating_tools.heavycivil.readback.report import dec_str


@pytest.fixture
def path(tmp_path):
    return H.build_book(H.fixture_doc(), tmp_path / "book.xlsx")


def _edit(path, cells, sheet=H.CR_TAB):
    wb = H.open_book(path)
    ws = wb[sheet]
    for addr, value in cells(ws).items():
        ws[addr] = value
    H.save_book(wb, path)


def _problems(path):
    with pytest.raises(ReadRefused) as exc:
        H.read(path)
    return exc.value.problems


def test_an_overwritten_formula_is_refused_naming_each_cell(path):
    _edit(path, lambda ws: {f"{CR.LAB}{H.cr_row(ws, 'ACTIVITY', 2)}": 999,
                            f"{CR.TOTAL}{H.cr_row(ws, 'RESOURCE', 7)}": 5})
    problems = _problems(path)
    cells = {(p.row, p.column) for p in problems}
    wb = H.open_book(path)
    ws = wb[H.CR_TAB]
    assert (H.cr_row(ws, "ACTIVITY", 2), CR.LAB) in cells
    assert (H.cr_row(ws, "RESOURCE", 7), CR.TOTAL) in cells
    assert all("held a formula" in p.message for p in problems)
    assert any(p.message.startswith("Labor ") for p in problems)


def test_an_overwritten_crew_formula_is_refused(path):
    _edit(path, lambda ws: {f"{CW.RATE}{H.cw_row(ws, 'CREW', 1)}": 100}, sheet=H.CW_TAB)
    assert any(p.column == CW.RATE for p in _problems(path))


def test_the_activity_row_helper_is_checked_on_every_row_as_a_formula(path):
    """AJ is ``=IF(W{r}="ACTIVITY",ROW(),MAX(AJ$4:AJ{r-1}))`` on every Cost
    Report row, markers included. The check asks only that it is a formula
    (its text shifts when Excel inserts, deletes or copies rows), so an
    unedited book passes and a value over a marker's AJ is the one problem."""
    wb = H.open_book(path)
    ws = wb[H.CR_TAB]
    start = H.cr_row(ws, "START", parent=2)
    assert str(ws[f"{CR.ACT}{start}"].value).startswith('=IF(W')
    assert "MAX(AJ$4:" in ws[f"{CR.ACT}{start}"].value
    H.read(path)
    _edit(path, lambda ws: {f"{CR.ACT}{start}": 7})
    assert [(p.row, p.column) for p in _problems(path)] == [(start, CR.ACT)]


def test_a_cleared_formula_is_named_as_cleared(path):
    _edit(path, lambda ws: {f"{CR.LAB}{H.cr_row(ws, 'ACTIVITY', 2)}": None})
    [problem] = _problems(path)
    assert problem.column == CR.LAB
    assert "held a formula and was cleared" in problem.message


# -- an unedited book reads as no edit --------------------------------------------

def _scanned(path):
    book = baseline.load(path)
    problems = Problems()
    return book, scan.scan(book, problems), problems


def _base(book, line):
    m = book.model
    if line.kind in ("ITEM", "ITEM_BID"):
        return m.item(line.id)
    if line.kind == "ELEMENT_QTY":
        return next(q for e in m.elements if e.id == line.id
                    for q in e.quantities if q.key == line.key)
    if line.kind == "ACTIVITY":
        return m.activity(line.id)
    if line.kind == "STEP":
        return next(st for st in m.activity(line.parent).steps if st.id == line.id)
    if line.kind == "RESOURCE":
        return next(r for r in m.resources if r.id == line.id)
    if line.kind == "MEMBER":
        return next(mm for mm in m.crew(line.parent).members if mm.id == line.id)
    return None


def _assert_no_edit(path, touch=None):
    """Every input row equals its baseline. ``touch(book)`` edits the opened
    workbook in memory first (openpyxl writes floats at 16 significant
    digits, so Excel's 17-digit float noise does not survive a save)."""
    book, s, problems = _scanned(path)
    if touch is not None:
        touch(book)
    checked = 0
    for line in s.cost_report + s.crew_lines + s.bid_form:
        if (line.sheet, line.kind) not in fields.INPUTS:
            continue
        stepped = line.occ is not None and line.occ.has_steps_marker
        base = _base(book, line)
        assert fields.equal_to_base(book, line, line.kind, base, stepped=stepped), line
        assert fields.diff(book, line, line.kind, base, problems, stepped=stepped) == {}
        checked += 1
    assert checked and not problems.items


def test_an_unedited_book_is_no_edit(path):
    _assert_no_edit(path)


def test_a_blank_production_unit_in_the_baseline_is_no_edit(tmp_path):
    doc = H.fixture_doc()
    for a in doc["activities"]:
        a["production_unit"] = ""
    _assert_no_edit(H.build_book(doc, tmp_path / "book.xlsx"))


def test_float_noise_and_long_baseline_figures_are_no_edit(tmp_path):
    doc = H.fixture_doc()
    doc["items"][0]["estimated_unit_price"] = "0.3"
    doc["resources"][0]["unit_cost"] = "1234567.123456789012345"
    path = H.build_book(doc, tmp_path / "book.xlsx")

    def touch(book):
        ws = book.ws("cost_report")
        ws[f"{CR.UCOST}{H.cr_row(ws, 'ITEM', doc['items'][0]['id'])}"] = 0.1 + 0.2
    _assert_no_edit(path, touch)


def test_names_are_read_trimmed(tmp_path):
    doc = H.fixture_doc()
    doc["activities"][0]["name"] += " "
    path = H.build_book(doc, tmp_path / "book.xlsx")
    _assert_no_edit(path)
    wb = H.open_book(path)
    ws = wb[H.CR_TAB]
    r = H.cr_row(ws, "ACTIVITY", doc["activities"][0]["id"])
    ws[f"{CR.DESC}{r}"] = "  " + ws[f"{CR.DESC}{r}"].value + "  "
    H.save_book(wb, path)
    _assert_no_edit(path)


def test_the_read_cache_is_per_kind_and_steps(path):
    book, s, _ = _scanned(path)
    line = next(ln for ln in s.cost_report
                if ln.kind == "ACTIVITY" and not ln.occ.has_steps_marker)
    assert "duration_shifts" in fields.read_line(book, line, "ACTIVITY").values
    stepped = fields.read_line(book, line, "ACTIVITY", stepped=True)
    assert "duration_shifts" not in stepped.values


# -- percent points ---------------------------------------------------------------

PERCENT = "holds percent points -- type 15, not 15%"


def test_a_percent_typed_with_a_percent_sign_is_refused(path):
    wb = H.open_book(path)
    ws = wb[H.CR_TAB]
    r = H.cr_row(ws, "RESOURCE", 7)
    ws[f"{CR.BURD}{r}"] = 0.15
    ws[f"{CR.BURD}{r}"].number_format = "0%"
    bf = wb[H.BF_TAB]
    bf[f"{BF.VALUE}{BF.ROWS['markup_percentage']}"] = 0.1
    bf[f"{BF.VALUE}{BF.ROWS['markup_percentage']}"].number_format = "0.0%"
    H.save_book(wb, path)
    book, s, _ = _scanned(path)
    res = next(ln for ln in s.cost_report if ln.kind == "RESOURCE" and ln.id == 7)
    [problem] = fields.read_line(book, res, "RESOURCE").problems
    assert (problem.row, problem.column) == (r, CR.BURD)
    assert problem.message == f"Burden % {PERCENT}"
    est = next(ln for ln in s.bid_form if ln.name == "markup_percentage")
    [problem] = fields.read_estimate_input(book, est).problems
    assert problem.message == f"markup_percentage {PERCENT}"


def test_a_literal_percent_in_the_format_is_not_percent_points():
    cell = _cell(15)
    cell.number_format = '0.00"%"'
    assert read_value(Input("A", "f", "number"), cell) == (Decimal("15"), None)
    cell.number_format = "0%"
    assert read_value(Input("A", "f", "number"), cell) == (None, PERCENT)


# -- reading one cell -------------------------------------------------------------

def _cell(value):
    ws = Workbook().active
    ws["A1"] = value
    return ws["A1"]


@pytest.mark.parametrize("reader, value, expected", [
    ("number", " 2.50 ", Decimal("2.50")),
    ("number", 0.1, Decimal("0.1")),
    ("number", 30, Decimal("30")),
    ("number", 0.1 + 0.2, Decimal("0.3")),
    ("number", "$1,234.50", Decimal("1234.50")),
    ("number", " 1,000 ", Decimal("1000")),
    ("number", None, None),
    ("count", 2.0, 2),
    ("unit", "hours", "HOURS"),
    ("unit", None, None),
    ("step_unit", None, "DAYS"),
    ("time_unit", " ", ""),
    ("type", "LABOR", "labor"),
    ("basis", "Flat", "flat"),
    ("flag", "y", True),
    ("flag", "N", False),
    ("scope", "Excluded", "excluded"),
    ("name", "  ", None),
    ("name", " Excavate  ", "Excavate"),
    ("text", " CY ", "CY"),
    ("text", None, ""),
])
def test_read_value(reader, value, expected):
    assert read_value(Input("A", "f", reader), _cell(value)) == (expected, None)


@pytest.mark.parametrize("reader, value, message", [
    ("number", "abc", "'abc' is not a number"),
    ("number", True, "True is not a number"),
    ("number", "NaN", "'NaN' is not a number"),
    ("count", 1.5, "1.5 is not a whole number"),
    ("type", "labour", "'labour' is not one of labor, equipment, material, subcontractor, "
                       "other"),
    ("number", "=1+1", "this input holds a formula -- type a value instead"),
    ("number", datetime.datetime(2026, 1, 2), "a date is not a number"),
    ("number", datetime.time(10, 30), "a date is not a number"),
])
def test_read_value_refuses(reader, value, message):
    assert read_value(Input("A", "f", reader), _cell(value)) == (None, message)


def test_dec_str_and_number_equal():
    assert dec_str(Decimal("30.000")) == "30" and dec_str(Decimal("1E+1")) == "10"
    assert dec_str(Decimal("-0.00")) == "0"
    assert number_equal(Decimal("0.1"), Decimal("0.1000"))
    assert number_equal(None, None) and not number_equal(Decimal(1), None)
    assert number_equal(Decimal(repr(1000.01)), Decimal("1000.0100"))
    assert number_equal(Decimal(repr(0.1 + 0.2)), Decimal("0.3"))
    assert number_equal(Decimal("1234567.12345679"), Decimal("1234567.123456789012345"))
    assert not number_equal(Decimal("0.3"), Decimal("0.31"))


@pytest.mark.parametrize("sheet, cell", [
    (H.CR_TAB, lambda ws: f"{CR.LAB}{H.cr_row(ws, 'ACTIVITY', 2)}"),
    (H.CW_TAB, lambda ws: f"{CW.RATE}{H.cw_row(ws, 'CREW', 1)}"),
    (H.BF_TAB, lambda ws: "C7"),
], ids=["cost-report", "crews", "bid-form"])
def test_a_formula_that_lost_its_reference_is_refused(path, sheet, cell):
    """Excel writes #REF! into a formula whose row was deleted; the cell still
    holds a formula, so the formula check alone would pass it."""
    holder = {}

    def cells(ws):
        holder["addr"] = cell(ws)
        return {holder["addr"]: "=SUM(#REF!)*2"}
    _edit(path, cells, sheet=sheet)
    assert any("lost its reference (#REF!)" in p.message for p in _problems(path))


# -- export values a formula reads ------------------------------------------------

def _bf_row(ws, kind, ident):
    return H.find_row(ws, kind, ident, kind_col=BF.KIND, id_col=BF.ID)


@pytest.mark.parametrize("sheet, cell, value", [
    (H.CW_TAB, lambda ws: f"{CW.HPS}{H.cw_row(ws, 'CREW', 1)}", 8),
    (H.CW_TAB, lambda ws: f"{CW.CFACT}{H.cw_row(ws, 'CREW', 2)}", 1.25),
    (H.CW_TAB, lambda ws: f"{CW.HPS}{H.cw_row(ws, 'CREW', 2)}", None),
    (H.BF_TAB, lambda ws: f"{BF.QTY}{_bf_row(ws, 'ITEM_BID', 1)}", 12),
    (H.CR_TAB, lambda ws: f"{CR.UCOST}{H.cr_row(ws, 'ELEMENT_QTY', 1)}", 70),
], ids=["crew-hps", "crew-factor", "crew-hps-cleared", "bid-qty", "derived-qty"])
def test_an_edited_export_value_a_formula_reads_is_refused(path, sheet, cell, value):
    """The book shows these as export values, but its formulas read them: an
    edit would reprice the book while BidSheet.ai kept the figure. No apply op
    writes them, so the read is refused, naming the cell."""
    holder = {}

    def cells(ws):
        holder["addr"] = cell(ws)
        return {holder["addr"]: value}
    _edit(path, cells, sheet=sheet)
    [problem] = _problems(path)
    assert (problem.sheet, f"{problem.column}{problem.row}") == (sheet, holder["addr"])
    assert "read-only in the workbook -- change it in BidSheet.ai" in problem.message


def test_crew_export_values_with_float_noise_are_no_edit(tmp_path):
    doc = H.fixture_doc()
    doc["crews"][0]["hours_per_shift"] = "0.3"
    doc["crews"][1]["crew_factor"] = "1.123456789012345678"
    path = H.build_book(doc, tmp_path / "book.xlsx")
    wb = H.open_book(path)
    wb[H.CW_TAB][f"{CW.HPS}{H.cw_row(wb[H.CW_TAB], 'CREW', 1)}"] = 0.1 + 0.2
    H.save_book(wb, path)
    H.read(path)
