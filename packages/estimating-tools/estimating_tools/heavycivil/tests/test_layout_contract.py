"""Sub-project 3 reads edits back by these cells; pin them."""
from __future__ import annotations

from estimating_tools.heavycivil import build, layout, model, style
from estimating_tools.heavycivil.columns import BF, CR, CW


def test_hidden_id_and_v_match_the_payload(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = book["Cost Report"]
    for row in plan.cost_report:
        assert ws[f"{CR.KIND}{row.row}"].value == row.kind
        if row.kind in ("ITEM", "ACTIVITY", "RESOURCE", "STEP", "ELEMENT_QTY"):
            assert ws[f"{CR.ID}{row.row}"].value == row.id
            assert ws[f"{CR.V}{row.row}"].value == row.v
    crews = book["Crews"]
    for crew in wb.crews:
        r = plan.crew_rows[crew.id].row
        assert crews[f"{CW.ID}{r}"].value == crew.id and crews[f"{CW.V}{r}"].value == crew.v


def test_inputs_carry_the_input_fill(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    input_fill = "FF" + style.load_style().tokens["fills"]["input"]
    ws = book["Cost Report"]
    res = plan.resource_rows[wb.resources[0].id].row
    for col in (CR.QTY, CR.UCOST, CR.TYPE):
        assert ws[f"{col}{res}"].fill.fgColor.rgb == input_fill
    assert ws[f"{CR.TOTAL}{res}"].fill.fgColor.rgb != input_fill


def test_outline_levels_follow_the_plan(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = book["Cost Report"]
    for row in plan.cost_report:
        assert (ws.row_dimensions[row.row].outline_level or 0) == min(row.level, 7)


def test_end_rows_carry_kind_and_activity_but_no_id(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = book["Cost Report"]
    for a in wb.activities:
        r = plan.end_rows[a.id].row
        assert ws[f"{CR.KIND}{r}"].value == "END"
        assert ws[f"{CR.PARENT}{r}"].value == a.id
        assert ws[f"{CR.ID}{r}"].value is None and ws[f"{CR.V}{r}"].value is None
        assert "insert" in ws[f"{CR.DESC}{r}"].value


def test_start_rows_are_blank_short_markers(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = book["Cost Report"]
    for a in wb.activities:
        r = plan.start_rows[a.id].row
        assert ws[f"{CR.KIND}{r}"].value == "START"
        assert ws[f"{CR.PARENT}{r}"].value == a.id
        assert ws[f"{CR.ID}{r}"].value is None and ws[f"{CR.V}{r}"].value is None
        for col in (CR.QTY, CR.TYPE, CR.MH) + tuple(c for _, c in CR.SPLIT):
            assert ws[f"{col}{r}"].value is None
        assert ws.row_dimensions[r].height == 6


def test_activity_sums_reach_the_end_row(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = book["Cost Report"]
    for a in wb.activities:
        r = plan.activity_rows[a.id].row
        start, end = plan.start_rows[a.id].row, plan.end_rows[a.id].row
        assert f"({CR.EQP}{start}:{CR.EQP}{end})" in ws[f"{CR.EQP}{r}"].value


def _validations(book):
    out = {}
    for dv in book["Cost Report"].data_validations.dataValidation:
        out[dv.formula1] = (dv, {str(r) for r in dv.sqref.ranges})
    return out


def _covers(ranges, coord):
    from openpyxl.worksheet.cell_range import CellRange
    return any(CellRange(coord) <= CellRange(r) for r in ranges)


def test_list_validations_cover_the_coded_inputs(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    dvs = _validations(book)
    types = dvs['"labor,equipment,material,subcontractor,other"']
    basis, flag = dvs['"production,flat"'], dvs['"Y,N"']
    units = dvs['"DAYS,HOURS"']
    for dv, _ in (types, basis, flag, units):
        assert dv.type == "list" and dv.allow_blank
    assert len(book["Cost Report"].data_validations.dataValidation) == 4
    for res in wb.resources:
        r = plan.resource_rows[res.id].row
        assert _covers(types[1], f"{CR.TYPE}{r}")
        assert _covers(basis[1], f"{CR.BASIS}{r}")
        assert _covers(flag[1], f"{CR.BINC}{r}")
        assert _covers(units[1], f"{CR.PUNIT}{r}")
    for a in wb.activities:
        assert _covers(units[1], f"{CR.PUNIT}{plan.activity_rows[a.id].row}")
        for s in a.steps:
            assert _covers(units[1], f"{CR.UNIT}{plan.step_rows[s.id].row}")
    for a in wb.activities:
        for marker in (plan.start_rows[a.id].row, plan.end_rows[a.id].row):
            for dv, col in ((types, CR.TYPE), (basis, CR.BASIS), (flag, CR.BINC),
                            (units, CR.PUNIT)):
                assert _covers(dv[1], f"{col}{marker}")
    item = plan.item_rows[wb.items[0].id].row
    assert not _covers(types[1], f"{CR.TYPE}{item}")


def test_crew_header_hours_and_factor_read_as_export_values(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    tokens = style.load_style().tokens
    ws = book["Crews"]
    for crew in wb.crews:
        r = plan.crew_rows[crew.id].row
        for col in (CW.HPS, CW.CFACT):
            assert ws[f"{col}{r}"].fill.fgColor.rgb == "FF" + tokens["fills"]["export"]
            assert ws[f"{col}{r}"].font.i


def test_bid_form_estimate_rows_carry_their_input_marker(doc, book):
    """Read-back finds each estimate input and figure by its hidden name
    (T) on a marked row (J), not by its address."""
    wb = model.parse(doc)
    ws = book["Bid Form"]
    for name, r in BF.ROWS.items():
        assert ws[f"{BF.INPUT}{r}"].value == name
        expected = ("ESTIMATE" if r == BF.ROWS["markup_percentage"]
                    else "BID_INPUT" if name in BF.INPUTS else "BID_FIGURE")
        assert ws[f"{BF.KIND}{r}"].value == expected
    r0 = BF.ROWS["markup_percentage"]
    assert ws[f"{BF.ID}{r0}"].value == wb.estimate.id
    assert ws[f"{BF.V}{r0}"].value == wb.estimate.v
    assert ws.column_dimensions[BF.INPUT].hidden


def test_steps_rows_are_blank_short_markers(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = book["Cost Report"]
    for a in (a for a in wb.activities if a.steps):
        r = plan.steps_rows[a.id].row
        assert ws[f"{CR.KIND}{r}"].value == "STEPS"
        assert ws[f"{CR.PARENT}{r}"].value == a.id
        assert ws[f"{CR.ID}{r}"].value is None and ws[f"{CR.V}{r}"].value is None
        for col in (CR.DESC, CR.QTY, CR.UNIT, CR.MH) + tuple(c for _, c in CR.SPLIT):
            assert ws[f"{col}{r}"].value is None
        assert ws.row_dimensions[r].height == 6
        start = plan.start_rows[a.id].row
        assert f"SUMIF({CR.UNIT}{r}:{CR.UNIT}{start}," in \
            ws[f"{CR.SHIFTS}{plan.activity_rows[a.id].row}"].value


def test_step_unit_list_covers_the_whole_step_block(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    (units,) = [dv for dv in book["Cost Report"].data_validations.dataValidation
                if dv.formula1 == '"DAYS,HOURS"']
    ranges = {str(r) for r in units.sqref.ranges}
    for a in (a for a in wb.activities if a.steps):
        first, last = plan.step_span[a.id]
        # STEPS..last step; not START, whose D tops the resource block.
        for r in range(first, last):
            assert _covers(ranges, f"{CR.UNIT}{r}"), r
        assert not _covers(ranges, f"{CR.UNIT}{last}")


def test_crew_blocks_carry_their_markers(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = book["Crews"]
    for crew in wb.crews:
        head, end = plan.crew_span[crew.id]
        assert ws[f"{CW.KIND}{end}"].value == "CREW_END"
        assert ws[f"{CW.CREW}{end}"].value == crew.id
        assert ws[f"{CW.ID}{end}"].value is None
        assert "insert" in ws[f"{CW.NAME}{end}"].value
        assert ws[f"{CW.HEADCOUNT}{head}"].value == f"=SUM({CW.COUNT}{head}:{CW.COUNT}{end})"
        assert ws[f"{CW.RATE}{head}"].value == f"=SUM({CW.EXTENDED}{head}:{CW.EXTENDED}{end})"
    for note in (r for r in plan.crews if r.kind == "NOTE"):
        assert ws[f"{CW.KIND}{note.row}"].value == "NOTE"
        assert ws[f"{CW.CREW}{note.row}"].value == note.parent_id
        assert ws[f"{CW.ID}{note.row}"].value is None


def test_crew_head_and_markers_have_no_count_or_extended(doc, book):
    """The wider crew span sums C and Q, so CREW, NOTE and CREW_END rows
    must add nothing to them."""
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = book["Crews"]
    rows = [r.row for r in plan.crews if r.kind in ("CREW", "NOTE", "CREW_END")]
    assert rows
    for r in rows:
        assert ws[f"{CW.COUNT}{r}"].value is None, r
        assert ws[f"{CW.EXTENDED}{r}"].value is None, r


def test_every_data_row_carries_its_activity_row(doc, book):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = book["Cost Report"]
    assert ws.column_dimensions[CR.ACT].hidden
    for row in plan.cost_report:
        assert ws[f"{CR.ACT}{row.row}"].value.startswith(
            f'=IF({CR.KIND}{row.row}="ACTIVITY",ROW(),')
    for res in wb.resources:
        r = plan.resource_rows[res.id].row
        assert f"AJ{r}" in ws[f"{CR.H1}{r}"].value and f"AJ{r}" in ws[f"{CR.H2}{r}"].value


def test_activity_row_evaluates_to_the_owning_activity(doc, evaluated):
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    for res in wb.resources:
        r = plan.resource_rows[res.id].row
        assert evaluated[f"Cost Report!{CR.ACT}{r}"] == plan.activity_rows[res.activity_id].row


def test_resource_rows_carry_their_rate_link(doc):
    doc["resources"][0]["equipment_rate_id"] = 7
    doc["resources"][1].update(labor_rate_id=4, material_cost_id=9)
    wb = model.parse(doc)
    plan = layout.plan(wb, style.load_style())
    ws = build.build(doc, style=style.load_style(), exported_at="x", host="h")["Cost Report"]
    links = {res.id: ws[f"{CR.REF}{plan.resource_rows[res.id].row}"].value for res in wb.resources}
    assert links[wb.resources[0].id] == "equipment:7"
    assert links[wb.resources[1].id] == "labor:4"
    # The fixture links its "Extra" material line; the rest are unlinked.
    extra = next(r for r in wb.resources if r.name == "Extra")
    assert links[extra.id] == "material:1"
    assert all(v is None for rid, v in links.items()
               if rid not in (wb.resources[0].id, wb.resources[1].id, extra.id))
