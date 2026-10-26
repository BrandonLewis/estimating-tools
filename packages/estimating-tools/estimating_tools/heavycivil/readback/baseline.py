"""Open an edited HeavyCivil book and decode the baseline it carries.

The book carries the payload it was built from: ``_wb`` column B holds it
(zlib + base64, chunked) and the layout JSON in A1 names the chunk count and
the SHA-256 of the decoded JSON; ``heavycivil.baseline`` encodes it at build
and decodes it here. Read-back diffs the book against that, never against a
fresh fetch: a re-fetch cannot tell an estimator's edit from a change made in
BidSheet.ai since the export.
"""
from __future__ import annotations

import functools
import json
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from estimating_tools.heavycivil import FORMAT_NAME, layout, model
from estimating_tools.heavycivil import baseline as hc_baseline
from estimating_tools.heavycivil.build import WB_SHEET
from estimating_tools.heavycivil.readback.errors import REBUILD, refuse

#: The sheets read-back reads, by layout tab key.
READ_TABS = ("cost_report", "crews", "bid_form", "rates")
#: The sheets read-back walks by row kind; ``row_kinds`` must cover each.
SCANNED_TABS = ("cost_report", "crews", "bid_form")


@dataclass
class Index:
    """Lookups over the baseline built once per read: the payload's raw rows
    by id (for fields the model does not keep -- ``description``, ``trade``,
    ``pieces``, cost codes), each activity's children, depth and cost code,
    and each step's activity."""
    raw_activities: dict
    raw_resources: dict
    raw_steps: dict
    children: dict
    depth: dict
    activity_code: dict
    step_owner: dict

    @classmethod
    def of(cls, doc, data):
        from estimating_tools.heavycivil.readback.scan import norm_id

        raw_activities = {norm_id(r.get("id")): r for r in doc.get("activities") or ()}
        raw_steps = {norm_id(st.get("id")): st for r in raw_activities.values()
                     for st in r.get("steps") or ()}
        children = {}
        for a in data.activities:
            children.setdefault(a.parent_id, []).append(a)
        by_id = {a.id: a for a in data.activities}
        depth = {}
        for a in data.activities:
            d, node = 0, a
            while node.parent_id is not None:
                d, node = d + 1, by_id[node.parent_id]
            depth[a.id] = d
        return cls(
            raw_activities=raw_activities,
            raw_resources={norm_id(r.get("id")): r for r in doc.get("resources") or ()},
            raw_steps=raw_steps,
            children=children,
            depth=depth,
            activity_code={aid: norm_id((r.get("activity_code") or {}).get("id"))
                           for aid, r in raw_activities.items()},
            step_owner={st.id: a for a in data.activities for st in a.steps},
        )

    def subtree(self, activity_id) -> list:
        """Every activity under ``activity_id`` (not itself)."""
        out, frontier = [], [activity_id]
        while frontier:
            kids = [k for f in frontier for k in self.children.get(f, ())]
            out += kids
            frontier = [k.id for k in kids]
        return out


@dataclass
class Book:
    """An opened book: the openpyxl workbook (formulas, not cached values),
    its layout JSON, the decoded baseline payload (``doc``) and that payload
    parsed (``model``)."""
    path: Path
    wb: object
    layout: dict
    doc: dict
    model: model.Workbook
    kinds: dict = field(default_factory=dict)

    def tab(self, key):
        return self.layout["tabs"][key]

    @functools.cached_property
    def index(self) -> Index:
        return Index.of(self.doc, self.model)

    def ws(self, key):
        return self.wb[self.tab(key)]


def _row_kinds(layout_doc):
    """``{sheet key: {"kind_cell": letter, "kinds": {kind: identified_by}}}``
    from the layout's ``row_kinds``."""
    out = {}
    for entry in layout_doc.get("row_kinds") or ():
        sheet = out.setdefault(entry["sheet"], {"kind_cell": entry["kind_cell"], "kinds": {}})
        sheet["kinds"][entry["kind"]] = dict(entry["identified_by"])
    return out


def decode(ws, meta) -> dict:
    """The baseline payload from ``_wb`` column B (B1 down to the first empty
    cell), checked against ``meta`` (the layout JSON's ``baseline``) by the
    builder's own ``heavycivil.baseline.decode``."""
    if not isinstance(meta, dict):
        refuse(f"the book carries no baseline -- {REBUILD}", sheet=WB_SHEET)
    chunks, r = [], 1
    while ws.cell(row=r, column=2).value not in (None, ""):
        # Not only text: a chunk Excel turned into a number is handed on, so
        # ``heavycivil.baseline.decode`` names it ("chunk k ... is not text").
        chunks.append(ws.cell(row=r, column=2).value)
        r += 1
    try:
        return hc_baseline.decode(chunks, meta)
    except hc_baseline.BaselineError as exc:
        refuse(f"{exc} -- {REBUILD}", sheet=WB_SHEET)


def load(path) -> Book:
    """Open ``path`` and check it is a layout-``LAYOUT_VERSION`` HeavyCivil
    book with an intact baseline, or raise ReadRefused."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    path = Path(path)
    try:
        wb = load_workbook(path, data_only=False)
    except FileNotFoundError:
        refuse(f"no such file: {path}")
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        refuse(f"cannot read {path} as an .xlsx workbook: {exc}")
    if WB_SHEET not in wb.sheetnames:
        refuse(f"not a HeavyCivil workbook (no {WB_SHEET} sheet) -- {REBUILD}")
    try:
        layout_doc = json.loads(wb[WB_SHEET]["A1"].value or "")
    except (TypeError, ValueError):
        refuse(f"the layout in {WB_SHEET}!A1 is unreadable -- {REBUILD}", sheet=WB_SHEET)
    if not isinstance(layout_doc, dict) or layout_doc.get("format") != FORMAT_NAME:
        refuse(f"not a HeavyCivil workbook -- {REBUILD}", sheet=WB_SHEET)
    unreadable = f"the layout in {WB_SHEET}!A1 is unreadable -- {REBUILD}"
    version = layout_doc.get("layout_version")
    if isinstance(version, bool) or not isinstance(version, int):
        refuse(f"this book has an unrecognised layout version ({version!r}); read-back "
               f"reads layout {layout.LAYOUT_VERSION} -- {REBUILD}", sheet=WB_SHEET)
    if version != layout.LAYOUT_VERSION:
        refuse(f"this book has layout {version!r}; read-back reads layout "
               f"{layout.LAYOUT_VERSION} -- {REBUILD}", sheet=WB_SHEET)
    tabs = layout_doc.get("tabs")
    if not isinstance(tabs, dict) or not isinstance(layout_doc.get("row_kinds"), list):
        refuse(unreadable, sheet=WB_SHEET)
    try:
        kinds = _row_kinds(layout_doc)
    except (KeyError, TypeError, AttributeError, ValueError):
        refuse(unreadable, sheet=WB_SHEET)
    if any(key not in kinds for key in SCANNED_TABS):
        refuse(unreadable, sheet=WB_SHEET)
    doc = decode(wb[WB_SHEET], layout_doc.get("baseline"))
    try:
        data = model.parse(doc)
    except model.PayloadError as exc:
        refuse(f"the embedded baseline is not valid ({exc}) -- {REBUILD}", sheet=WB_SHEET)
    for key in READ_TABS:
        if tabs.get(key) not in wb.sheetnames:
            refuse(f"the {key.replace('_', ' ')} tab ({tabs.get(key)!r}) is missing or was "
                   f"renamed -- rename it back, or {REBUILD}")
    return Book(path=path, wb=wb, layout=layout_doc, doc=doc, model=data, kinds=kinds)
