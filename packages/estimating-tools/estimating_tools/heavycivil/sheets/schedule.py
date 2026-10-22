"""Schedule: one row per dated activity -- export start, live duration and
finish (the export finish moved by the change in whole working days) -- and
one column per week, shaded by conditional formatting between start and
finish. Read-only. The hidden ``_cal`` sheet lists the estimate's working
days for the finish lookup."""
from __future__ import annotations

from datetime import date

from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import column_index_from_string, get_column_letter

from estimating_tools.heavycivil import formulas as F
from estimating_tools.heavycivil.columns import SC
from estimating_tools.heavycivil.layout import FIRST_ROW, HEADER_ROW
from estimating_tools.heavycivil.sheets.common import export_multiplier, factory, put, setup_sheet

CAL_SHEET = "_cal"
WEEK_WIDTH = 4


def write_calendar(ws, plan):
    """``_cal``: the working days, ascending, from A1 (``MATCH`` type 1 needs
    them sorted)."""
    ws.sheet_state = "hidden"
    for n, day in enumerate(plan.calendar, start=1):
        ws[f"A{n}"] = day
        ws[f"A{n}"].number_format = "yyyy-mm-dd"


def write_schedule(ws, plan, wb, style):
    sf = factory(style)
    setup_sheet(ws, sf, style, "schedule", title="Schedule",
                subtitle="Starts as of export; finishes move with Shifts edits on the Cost Report")
    if not plan.schedule_rows:
        put(ws, sf, f"{SC.NAME}{FIRST_ROW}",
            "No schedule dates in this export -- set a schedule start in BidSheet.ai.", "note")
        return
    cost_tab, cal_last = plan.tabs["cost_report"], len(plan.calendar)
    first_week = column_index_from_string(SC.FIRST_WEEK)
    for n, monday in enumerate(plan.weeks):
        col = get_column_letter(first_week + n)
        put(ws, sf, f"{col}{HEADER_ROW}", monday, "calc_date")
        ws.column_dimensions[col].width = WEEK_WIDTH
    last_row = max(plan.schedule_rows.values())
    for activity_id, r in plan.schedule_rows.items():
        a = wb.activity(activity_id)
        act = plan.activity_rows[a.id].row
        put(ws, sf, f"{SC.CODE}{r}", a.code, "label")
        put(ws, sf, f"{SC.NAME}{r}", a.name, "label")
        put(ws, sf, f"{SC.ITEM}{r}", wb.item(a.item_id).item_number, "calc_text")
        put(ws, sf, f"{SC.START}{r}", date.fromisoformat(a.start_date), "export_date")
        put(ws, sf, f"{SC.DURATION}{r}", F.schedule_duration(r, cost_tab, act), "calc_qty")
        put(ws, sf, f"{SC.FINISH}{r}", F.schedule_finish(r, CAL_SHEET, cal_last), "calc_date")
        put(ws, None, f"{SC.KIND}{r}", "ACTIVITY")
        put(ws, None, f"{SC.ID}{r}", a.id)
        put(ws, None, f"{SC.XFINISH}{r}", date.fromisoformat(a.finish_date))
        ws[f"{SC.XFINISH}{r}"].number_format = "yyyy-mm-dd"
        # The payload's duration_days is per pass (Assembly.duration_days);
        # the export dates span it x the count chain (estimate_task_rollup).
        put(ws, None, f"{SC.XDURATION}{r}", a.duration_days or 0)
        put(ws, None, f"{SC.XMULT}{r}", export_multiplier(wb, a))
    if plan.weeks:
        last_week = get_column_letter(first_week + len(plan.weeks) - 1)
        area = f"{SC.FIRST_WEEK}{FIRST_ROW}:{last_week}{last_row}"
        font, fill = sf.dxf("schedule_bar")
        ws.conditional_formatting.add(area, FormulaRule(
            formula=[f"AND({SC.FIRST_WEEK}${HEADER_ROW}<=${SC.FINISH}{FIRST_ROW},"
                     f"{SC.FIRST_WEEK}${HEADER_ROW}+6>=${SC.START}{FIRST_ROW})"],
            font=font, fill=fill))
