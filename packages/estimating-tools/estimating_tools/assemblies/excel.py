"""Import an org's assembly library from a spreadsheet into the canonical model.

The pure builder (`build_assembly_library`) maps the org's columns to canonical fields
(identity default), groups rows by assembly, normalizes component_type, and validates
required fields — no openpyxl/MCP. Only `read_assembly_rows` reads the .xlsx (openpyxl).
"""
from __future__ import annotations

from estimating_tools.assemblies.load_model import (
    AssemblyLoadItem, CanonicalAssemblyLoad, LoadAssembly, LoadComponent,
)
from estimating_tools.assemblies.model import Assembly, CanonicalAssemblyLibrary, Component
from estimating_tools.assemblies.plan import VALID_COMPONENT_TYPES

_CANONICAL_FIELDS = (
    "assembly", "assembly_type", "cost_code", "assembly_description",
    "component_type", "component", "quantity", "unit", "unit_cost", "rate_ref", "trade",
    "component_description",
)
_REQUIRED = ("assembly", "component_type", "component")

# Estimate-context load (issue #201) adds the item binding + roll-up-contract
# columns on top of the library fields.
_LOAD_CANONICAL_FIELDS = _CANONICAL_FIELDS + (
    "estimate_item_id", "bid_item_code", "production_rate", "production_unit",
    "labor_cost_basis",
)
_LOAD_REQUIRED = ("estimate_item_id", "assembly", "component_type", "component")

_TYPE_ALIASES = {
    "lab": "labor", "labour": "labor",
    "mat": "material", "matl": "material",
    "equip": "equipment", "eqp": "equipment",
    "sub": "subcontractor",
}


def _effective_mapping(mapping):
    m = {f: f for f in _CANONICAL_FIELDS}      # identity default
    if mapping:
        m.update(mapping)
    return m


def _num(value, default):
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _int(value):
    """Parse an integer binding key. Blank -> None (the caller decides whether
    that's an error). A nonblank value that is not an integer is a HARD error:
    silently truncating 101.9 -> 101 would route the whole load to a different
    estimate item."""
    if value is None or value == "":
        return None
    try:
        f = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"estimate_item_id {value!r} is not a number")
    if not f.is_integer():
        raise ValueError(
            f"estimate_item_id {value!r} is not an integer (a fractional id would "
            "silently bind the load to the wrong estimate item)"
        )
    return int(f)


def _num_strict(value, default, field, assembly):
    """Like _num but a nonblank nonnumeric cell RAISES instead of falling back to
    the default. Blank still means "unset" (-> default). Used by the load builder
    so a typo can't silently become a phantom quantity (1.0) or a $0 unit_cost —
    the same class of trap plan_assembly_load refuses downstream (issue #201)."""
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        raise ValueError(
            f"{field} {value!r} in assembly {assembly!r} is not numeric "
            f"(leave the cell blank for the default, but a nonblank value must parse)"
        )


def _str(value):
    return str(value).strip() if value is not None else ""


def _normalize_component_type(value, assembly_name):
    key = _str(value).lower()
    norm = _TYPE_ALIASES.get(key, key)
    if norm not in VALID_COMPONENT_TYPES:
        raise ValueError(
            f"unrecognized component_type {value!r} in assembly {assembly_name!r} "
            f"(expected one of {sorted(VALID_COMPONENT_TYPES)} or an alias)"
        )
    return norm


def build_assembly_library(rows, mapping=None) -> CanonicalAssemblyLibrary:
    """Pure: spreadsheet rows (dicts keyed by the org's headers) + a column mapping
    -> a CanonicalAssemblyLibrary. No openpyxl/MCP."""
    if not rows:
        return CanonicalAssemblyLibrary(assemblies=[])
    m = _effective_mapping(mapping)
    present = set()
    for r in rows:
        present.update(r.keys())
    for field in _REQUIRED:
        if m[field] not in present:
            raise ValueError(
                f"required field {field!r} maps to column {m[field]!r}, not found in "
                f"the spreadsheet headers {sorted(present)}"
            )

    order: list[str] = []
    groups: dict[str, list[dict]] = {}
    for r in rows:
        name = _str(r.get(m["assembly"]))
        if not name:
            continue                            # blank-assembly / blank row -> skip
        if name not in groups:
            groups[name] = []
            order.append(name)
        groups[name].append(r)

    assemblies = []
    for name in order:
        group = groups[name]
        first = group[0]
        components = []
        for r in group:
            comp_name = _str(r.get(m["component"]))
            if not comp_name:
                raise ValueError(
                    f"required field 'component' is blank in a row of assembly "
                    f"{name!r} (a row with an assembly name needs a component name)"
                )
            components.append(Component(
                component_type=_normalize_component_type(r.get(m["component_type"]), name),
                name=comp_name,
                quantity=_num(r.get(m["quantity"]), 1.0),
                unit=_str(r.get(m["unit"])),
                unit_cost=_num(r.get(m["unit_cost"]), None),
                rate_ref=(_str(r.get(m["rate_ref"])) or None),
                trade=_str(r.get(m["trade"])),
                description=_str(r.get(m["component_description"])),
            ))
        assemblies.append(Assembly(
            name=name,
            assembly_type=_str(first.get(m["assembly_type"])),
            cost_code=(_str(first.get(m["cost_code"])) or None),
            description=_str(first.get(m["assembly_description"])),
            components=components,
        ))
    return CanonicalAssemblyLibrary(assemblies=assemblies)


def read_assembly_rows(path) -> list[dict]:
    """Read the first worksheet of an .xlsx into row-dicts keyed by header. The only
    openpyxl/file touch."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise ImportError(
            "reading .xlsx requires openpyxl: pip install 'estimating_tools[excel]'"
        ) from exc
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = None
        out: list[dict] = []
        for raw in ws.iter_rows(values_only=True):
            if raw is None or all(c is None for c in raw):
                continue
            if headers is None:
                headers = [_str(c) for c in raw]
                continue
            out.append({h: v for h, v in zip(headers, raw)})
        return out
    finally:
        wb.close()


def build_assembly_library_from_excel(path, mapping=None) -> CanonicalAssemblyLibrary:
    return build_assembly_library(read_assembly_rows(path), mapping)


# ---------- estimate-context load (issue #201) ----------

def build_assembly_load(rows, mapping=None) -> CanonicalAssemblyLoad:
    """Pure: spreadsheet/JSON row dicts (keyed by the org's headers) + a column
    mapping -> a CanonicalAssemblyLoad (issue #201). Rows are grouped by
    ``estimate_item_id``, then by assembly name within each item. The
    ``production_rate`` / ``production_unit`` / ``labor_cost_basis`` columns carry
    the roll-up contract; ``plan_assembly_load`` enforces it at push time (this
    builder only shapes — it does not judge $0 crews or phantom zeros). No
    openpyxl/MCP."""
    if not rows:
        return CanonicalAssemblyLoad(items=[])
    m = {f: f for f in _LOAD_CANONICAL_FIELDS}      # identity default
    if mapping:
        m.update(mapping)
    present = set()
    for r in rows:
        if not isinstance(r, dict):
            raise ValueError(
                f"each row must be an object/dict of {{header: value}}, got a "
                f"{type(r).__name__}"
            )
        present.update(r.keys())
    for field in _LOAD_REQUIRED:
        if m[field] not in present:
            raise ValueError(
                f"required field {field!r} maps to column {m[field]!r}, not found in "
                f"the row headers {sorted(present)}"
            )

    item_order: list[int] = []
    item_rows: dict[int, list[dict]] = {}
    for r in rows:
        name = _str(r.get(m["assembly"]))
        if not name:
            continue                            # blank-assembly / blank row -> skip
        eid = _int(r.get(m["estimate_item_id"]))
        if eid is None:
            raise ValueError(
                f"required field 'estimate_item_id' is blank/non-numeric in a row "
                f"of assembly {name!r} (every load row binds to an estimate item)"
            )
        if eid not in item_rows:
            item_rows[eid] = []
            item_order.append(eid)
        item_rows[eid].append(r)

    items: list[AssemblyLoadItem] = []
    for eid in item_order:
        rows_i = item_rows[eid]
        bid_code = _str(rows_i[0].get(m["bid_item_code"]))
        order: list[str] = []
        groups: dict[str, list[dict]] = {}
        for r in rows_i:
            name = _str(r.get(m["assembly"]))
            if name not in groups:
                groups[name] = []
                order.append(name)
            groups[name].append(r)

        assemblies: list[LoadAssembly] = []
        for name in order:
            group = groups[name]
            first = group[0]
            components: list[LoadComponent] = []
            for r in group:
                comp_name = _str(r.get(m["component"]))
                if not comp_name:
                    raise ValueError(
                        f"required field 'component' is blank in a row of assembly "
                        f"{name!r} (a row with an assembly name needs a component name)"
                    )
                components.append(LoadComponent(
                    component_type=_normalize_component_type(r.get(m["component_type"]), name),
                    name=comp_name,
                    quantity=_num_strict(r.get(m["quantity"]), 1.0, "quantity", name),
                    unit=_str(r.get(m["unit"])),
                    unit_cost=_num_strict(r.get(m["unit_cost"]), None, "unit_cost", name),
                    rate_ref=(_str(r.get(m["rate_ref"])) or None),
                    trade=_str(r.get(m["trade"])),
                    description=_str(r.get(m["component_description"])),
                    labor_cost_basis=(_str(r.get(m["labor_cost_basis"])).lower() or "production"),
                ))
            assemblies.append(LoadAssembly(
                name=name,
                assembly_type=_str(first.get(m["assembly_type"])),
                cost_code=(_str(first.get(m["cost_code"])) or None),
                description=_str(first.get(m["assembly_description"])),
                production_rate=_num_strict(first.get(m["production_rate"]), None,
                                            "production_rate", name),
                production_unit=(_str(first.get(m["production_unit"])) or "HOURS"),
                components=components,
            ))
        items.append(AssemblyLoadItem(
            estimate_item_id=eid, bid_item_code=bid_code, assemblies=assemblies,
        ))
    return CanonicalAssemblyLoad(items=items)


def build_assembly_load_from_excel(path, mapping=None) -> CanonicalAssemblyLoad:
    return build_assembly_load(read_assembly_rows(path), mapping)
