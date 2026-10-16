import openpyxl

from estimating_tools.importers.hcss import HcssHeavyBidImporter


def _hcss_workbook(path):
    """A minimal HeavyBid-style export: a header row with HCSS column names,
    a schedule-header row (description only), then item rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Bid Item", "Description", "Quantity", "Unit", "Unit Price", "Cost Code"])
    ws.append([None, "SCHEDULE A - ROADWAY", None, None, None, None])  # schedule header (desc only)
    ws.append(["1", "Mobilization", 1, "LS", 50000, "070-010"])
    ws.append(["2", "Clearing & Grubbing", 3, "AC", 4200, None])
    ws.append(["TOTAL", None, None, None, None, None])  # must be skipped
    wb.save(path)


def test_detect_true_for_hcss_shaped_xlsx(tmp_path):
    f = tmp_path / "export.xlsx"; _hcss_workbook(f)
    assert HcssHeavyBidImporter().detect(f) is True


def test_detect_false_for_unrelated_xlsx(tmp_path):
    f = tmp_path / "other.xlsx"
    wb = openpyxl.Workbook(); wb.active.append(["foo", "bar"]); wb.save(f)
    assert HcssHeavyBidImporter().detect(f) is False


def test_parse_maps_items_codes_and_schedule(tmp_path):
    f = tmp_path / "export.xlsx"; _hcss_workbook(f)
    est = HcssHeavyBidImporter().parse(f)
    # one schedule, two items (TOTAL row skipped)
    assert len(est.schedules) == 1
    sched = est.schedules[0]
    assert sched.name == "SCHEDULE A - ROADWAY"
    assert [i.item_number for i in sched.items] == ["1", "2"]
    first = sched.items[0]
    assert first.description == "Mobilization"
    assert first.item_code == "070-010"        # cost_code -> item_code
    assert first.quantity == 1 and first.unit == "LS" and first.unit_price == 50000
    assert sched.items[1].item_code is None     # blank cost code -> None
    assert est.project["source"] == "HCSS HeavyBid"


def test_parse_handles_dollar_and_comma_string_prices(tmp_path):
    f = tmp_path / "export.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Bid Item", "Description", "Quantity", "Unit", "Unit Price", "Cost Code"])
    ws.append(["1", "Mobilization", "1", "LS", "$50,000", "070-010"])
    wb.save(f)
    item = HcssHeavyBidImporter().parse(f).schedules[0].items[0]
    assert item.unit_price == 50000.0
    assert item.quantity == 1.0


def test_parse_skips_item_number_only_rows(tmp_path):
    f = tmp_path / "export.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Bid Item", "Description", "Quantity", "Unit", "Unit Price", "Cost Code"])
    ws.append(["1", "Mobilization", 1, "LS", 50000, "070-010"])
    ws.append(["2", None, None, None, None, None])  # spacer/continuation -> must be dropped
    wb.save(f)
    items = [i for s in HcssHeavyBidImporter().parse(f).schedules for i in s.items]
    assert [i.item_number for i in items] == ["1"]
