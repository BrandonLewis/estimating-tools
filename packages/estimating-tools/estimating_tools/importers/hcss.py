"""HCSS HeavyBid Excel-export adapter -> CanonicalEstimate.

Re-homes the parse logic from BidSheet.ai's
Projects/services/hcss_service.HeavyBidSpreadsheetImporter (Django-free).

Specifically re-homed from hcss_service.py:
- COLUMN_MAPPINGS dict (class attribute, ~line 466)
- HEADER_MARKERS list (class attribute, ~line 558)
- _find_header_row() module helper (~line 391)
- _resolve_columns() module helper (~line 417)
- _get_cell_str() module helper (~line 377)
- _get_cell_decimal() -> adapted as _get_cell_float() returning float|None
- _extract_schedule_label() from Projects/utils/schedule_detector._extract_label
- HeavyBidSpreadsheetImporter.parse() data-row loop (~line 614)
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from estimating_tools.importers.model import CanonicalEstimate, LineItem, Schedule

# ---------------------------------------------------------------------------
# Module helpers — re-homed verbatim from hcss_service.py (Django-free)
# ---------------------------------------------------------------------------


def _get_cell_str(ws, row, col, default=""):
    """Get cell value as stripped string.

    Re-homed from hcss_service._get_cell_str (~line 377).
    """
    val = ws.cell(row=row, column=col).value
    if val is None:
        return default
    return str(val).strip()


def _get_cell_float(ws, row, col):
    """Get cell value as float, returning None if not parseable.

    Re-homed from hcss_service._get_cell_decimal (~line 385), but returns
    float|None instead of Decimal|None (interchange format uses JSON numbers).
    """
    val = ws.cell(row=row, column=col).value
    if val is None:
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        s = str(val).strip().replace(",", "").replace("$", "")
        try:
            return float(s)
        except (TypeError, ValueError):
            return None


def _find_header_row(ws, marker_columns, max_rows=20):
    """Find the header row by searching for known column names.

    Re-homed verbatim from hcss_service._find_header_row (~line 391).

    Args:
        ws: Worksheet
        marker_columns: List of possible column name strings to look for
        max_rows: Max rows to search

    Returns:
        (header_row_number, {col_name_lower: col_index}) or (None, {})
    """
    markers_lower = {m.lower() for m in marker_columns}
    for row_num in range(1, min(ws.max_row + 1, max_rows + 1)):
        header_map = {}
        for col in range(1, 20):
            val = ws.cell(row=row_num, column=col).value
            if val:
                header_map[str(val).strip().lower()] = col

        if header_map.keys() & markers_lower:
            return row_num, header_map

    return None, {}


def _resolve_columns(header_map, col_mappings):
    """Resolve column names from header map using flexible name matching.

    Re-homed verbatim from hcss_service._resolve_columns (~line 417).

    Args:
        header_map: {lower_name: col_index}
        col_mappings: {target_name: [possible_variations]}

    Returns:
        {target_name: col_index}
    """
    resolved = {}
    for target, variations in col_mappings.items():
        for var in variations:
            if var.lower() in header_map:
                resolved[target] = header_map[var.lower()]
                break
    return resolved


def _extract_schedule_label(description):
    """Extract a short label from a schedule description.

    Re-homed from Projects/utils/schedule_detector._extract_label.
    Matches patterns like:
    - "SCHEDULE A - ROADWAY" -> "A"
    - "SCHEDULE 'B' - WATERLINE" -> "B"
    - "SCHEDULE \\"C\\"" -> "C"
    """
    patterns = [
        r"SCHEDULE\s+[‘\"]?([A-Z0-9]+)[‘\"]?",
        r"SCHEDULE\s+['‘’]([A-Z0-9]+)['‘’]",
    ]
    upper = description.upper()
    for pattern in patterns:
        match = re.search(pattern, upper)
        if match:
            return match.group(1)
    return ""


# ---------------------------------------------------------------------------
# Column mappings — re-homed verbatim from
# hcss_service.HeavyBidSpreadsheetImporter.COLUMN_MAPPINGS (~line 466)
# ---------------------------------------------------------------------------

COLUMN_MAPPINGS = {
    "item_number": [
        "bid item",
        "item number",
        "item #",
        "item no",
        "item",
        "biditem",
        "pay item",
        "pay item number",
    ],
    "description": [
        "description",
        "desc",
        "item description",
        "bid item description",
    ],
    "quantity": [
        "quantity",
        "qty",
        "est qty",
        "estimated quantity",
        "bid quantity",
        "bid qty",
    ],
    "unit": [
        "unit",
        "uom",
        "unit of measure",
        "bid unit",
    ],
    "unit_price": [
        "unit price",
        "price",
        "unit bid price",
        "bid price",
        "unit cost",
        "rate",
    ],
    "total": [
        "total",
        "amount",
        "extension",
        "extended",
        "total amount",
        "total price",
        "bid amount",
    ],
    "cost_code": [
        "cost code",
        "costcode",
        "cost cd",
        "cc",
    ],
    "activity_code": [
        "activity",
        "activity code",
        "act",
        "act code",
    ],
    "labor_cost": [
        "labor",
        "labor cost",
        "labor $",
        "labor amount",
    ],
    "equipment_cost": [
        "equipment",
        "equipment cost",
        "equip",
        "equip cost",
        "equipment $",
        "equip $",
    ],
    "material_cost": [
        "material",
        "material cost",
        "mat",
        "mat cost",
        "material $",
        "mat $",
    ],
    "subcontractor_cost": [
        "subcontractor",
        "sub",
        "sub cost",
        "subcontract",
        "subcontractor cost",
        "sub $",
    ],
}

# Re-homed verbatim from
# hcss_service.HeavyBidSpreadsheetImporter.HEADER_MARKERS (~line 558)
HEADER_MARKERS = [
    "bid item",
    "item number",
    "item #",
    "description",
    "quantity",
    "unit price",
    "cost code",
    "pay item",
]


# ---------------------------------------------------------------------------
# Importer class
# ---------------------------------------------------------------------------


class HcssHeavyBidImporter:
    name = "hcss-heavybid"

    def _open_and_locate(self, path: Path):
        """Return (ws, header_row, cols) or None if not an HCSS-shaped sheet."""
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception:
            return None
        ws = wb.active
        if ws is None:
            return None
        header_row, header_map = _find_header_row(ws, HEADER_MARKERS)
        if header_row is None:
            return None
        cols = _resolve_columns(header_map, COLUMN_MAPPINGS)
        if "item_number" not in cols:
            return None
        return ws, header_row, cols

    def detect(self, path: Path) -> bool:
        return self._open_and_locate(Path(path)) is not None

    def parse(self, path: Path) -> CanonicalEstimate:
        """Parse a HeavyBid Excel export into a CanonicalEstimate.

        Data-row loop re-homed from
        hcss_service.HeavyBidSpreadsheetImporter.parse (~line 614):
        - Schedule-header detection on description-only rows
        - TOTAL/SUBTOTAL row skip
        - cost_code -> item_code mapping (None if blank)
        """
        located = self._open_and_locate(Path(path))
        if located is None:
            raise ValueError(f"not an HCSS HeavyBid export: {path}")
        ws, header_row, cols = located

        # Ordered schedule accumulation; default bucket "Base Bid"
        order: list[str] = []
        buckets: dict[str, list[LineItem]] = {}
        current = None

        for row in range(header_row + 1, ws.max_row + 1):
            item_num = _get_cell_str(ws, row, cols["item_number"])
            if not item_num:
                # A description-only row may be a schedule/section header.
                desc = (
                    _get_cell_str(ws, row, cols["description"])
                    if "description" in cols
                    else ""
                )
                if desc and (_extract_schedule_label(desc) or "SCHEDULE" in desc.upper()):
                    current = desc
                    order.append(current)
                    buckets.setdefault(current, [])
                continue
            if item_num.upper() in ("TOTAL", "TOTAL:", "GRAND TOTAL", "SUBTOTAL"):
                continue

            code = _get_cell_str(ws, row, cols["cost_code"]) if "cost_code" in cols else ""
            breakdown = {}
            for field_name, key in (
                ("labor_cost", "labor"),
                ("equipment_cost", "equipment"),
                ("material_cost", "material"),
                ("subcontractor_cost", "subcontractor"),
            ):
                if field_name in cols:
                    v = _get_cell_float(ws, row, cols[field_name])
                    if v is not None:
                        breakdown[key] = v

            item = LineItem(
                item_number=item_num,
                description=(
                    _get_cell_str(ws, row, cols["description"])
                    if "description" in cols
                    else ""
                ),
                item_code=code or None,
                quantity=_get_cell_float(ws, row, cols["quantity"]) if "quantity" in cols else None,
                unit=_get_cell_str(ws, row, cols["unit"]) if "unit" in cols else "",
                unit_price=_get_cell_float(ws, row, cols["unit_price"]) if "unit_price" in cols else None,
                total=_get_cell_float(ws, row, cols["total"]) if "total" in cols else None,
                breakdown=breakdown or None,
            )
            # Guard re-homed from hcss_service.py line 695: skip rows that have
            # neither a description nor a quantity (spacer/continuation rows).
            if not (item.description or item.quantity is not None):
                continue
            bucket = current or "Base Bid"
            if bucket not in buckets:
                order.append(bucket)
                buckets[bucket] = []
            buckets[bucket].append(item)

        schedules = [Schedule(name=n, items=buckets[n]) for n in order if buckets.get(n)]
        if not schedules:
            schedules = [Schedule(name="Base Bid", items=[])]
        return CanonicalEstimate(
            project={"name": Path(path).stem, "source": "HCSS HeavyBid"},
            schedules=schedules,
        )
