"""Import an org's rate library from a spreadsheet into the canonical model.

The pure builder (`build_rate_library`) maps the org's columns to canonical fields
(identity default), routes each row by `rate_type` to a typed list, normalizes
numbers/booleans, and validates required fields — no openpyxl/MCP. Only `read_rate_rows`
reads the .xlsx (openpyxl)."""
from __future__ import annotations

from estimating_tools.rates.model import (
    CanonicalRateLibrary, EquipmentRate, LaborRate, MaterialCost, VALID_RATE_TYPES,
)

_CANONICAL_FIELDS = (
    "rate_type", "name", "rate", "unit", "region", "description", "burden_percentage",
    "fringe_hourly", "trade", "is_union", "union_local", "source", "includes_operator",
    "category", "supplier",
)
_REQUIRED = ("rate_type", "name", "rate")
_TRUTHY = {"yes", "y", "true", "1", "x"}
_TYPE_ALIASES = {"lab": "labor", "labour": "labor", "equip": "equipment",
                 "eqp": "equipment", "mat": "material", "matl": "material"}


def _effective_mapping(mapping):
    m = {f: f for f in _CANONICAL_FIELDS}      # identity default
    if mapping:
        m.update(mapping)
    return m


def _str(value):
    return str(value).strip() if value is not None else ""


def _truthy(value):
    return _str(value).lower() in _TRUTHY


def _num(value, default, label, name):
    text = _str(value)
    if not text:
        if default is None:
            raise ValueError(f"{label} is blank for rate {name!r} (a numeric value is required)")
        return default
    try:
        return float(text)
    except ValueError:
        raise ValueError(f"{label} {value!r} for rate {name!r} is not a number")


def _normalize_rate_type(value, name):
    key = _str(value).lower()
    norm = _TYPE_ALIASES.get(key, key)
    if norm not in VALID_RATE_TYPES:
        raise ValueError(
            f"unrecognized rate_type {value!r} for rate {name!r} "
            f"(expected one of {sorted(VALID_RATE_TYPES)} or an alias)"
        )
    return norm


def build_rate_library(rows, mapping=None) -> CanonicalRateLibrary:
    """Pure: spreadsheet rows (dicts keyed by the org's headers) + a column mapping ->
    a CanonicalRateLibrary (routed into labor/equipment/material). No openpyxl/MCP."""
    lib = CanonicalRateLibrary()
    if not rows:
        return lib
    m = _effective_mapping(mapping)
    present = set()
    for r in rows:
        present.update(r.keys())
    for field_name in _REQUIRED:
        if m[field_name] not in present:
            raise ValueError(
                f"required field {field_name!r} maps to column {m[field_name]!r}, not "
                f"found in the spreadsheet headers {sorted(present)}"
            )

    for r in rows:
        type_raw = _str(r.get(m["rate_type"]))
        if not type_raw:
            continue                            # blank rate_type / blank row -> skip
        name = _str(r.get(m["name"]))
        if not name:
            raise ValueError(
                f"required field 'name' is blank for a {type_raw!r} rate row "
                "(every rate needs a name)"
            )
        rate_type = _normalize_rate_type(type_raw, name)
        rate = _num(r.get(m["rate"]), None, "rate", name)

        if rate_type == "labor":
            lib.labor_rates.append(LaborRate(
                classification=name,
                base_hourly_rate=rate,
                burden_percentage=_num(r.get(m["burden_percentage"]), 35.0, "burden_percentage", name),
                fringe_hourly=_num(r.get(m["fringe_hourly"]), 0.0, "fringe_hourly", name),
                region=_str(r.get(m["region"])),
                trade=_str(r.get(m["trade"])),
                is_union=_truthy(r.get(m["is_union"])),
                union_local=_str(r.get(m["union_local"])),
                source=_str(r.get(m["source"])),
            ))
        elif rate_type == "equipment":
            lib.equipment_rates.append(EquipmentRate(
                equipment_type=name,
                hourly_rate=rate,
                description=_str(r.get(m["description"])),
                region=_str(r.get(m["region"])),
                includes_operator=_truthy(r.get(m["includes_operator"])),
                source=_str(r.get(m["source"])),
            ))
        else:  # material
            lib.material_costs.append(MaterialCost(
                material_name=name,
                unit_cost=rate,
                unit=_str(r.get(m["unit"])),
                category=_str(r.get(m["category"])),
                supplier=_str(r.get(m["supplier"])),
            ))
    return lib


def read_rate_rows(path) -> list[dict]:
    """Read the first worksheet of an .xlsx into row-dicts keyed by header. The only
    openpyxl/file touch. Values are the raw openpyxl cell values (an int stays int, a
    bool stays bool); build_rate_library coerces them via _str/_num."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise ImportError(
            "reading .xlsx requires openpyxl: pip install 'estimating_tools[excel]'"
        ) from exc
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = None
        out: list[dict] = []
        for raw in ws.iter_rows(values_only=True):
            if raw is None or all(c is None for c in raw):
                continue
            if headers is None:
                headers = [_str(c) for c in raw]
                continue
            out.append({h: v for h, v in zip(headers, raw)})
        return out
    finally:
        wb.close()


def build_rate_library_from_excel(path, mapping=None) -> CanonicalRateLibrary:
    return build_rate_library(read_rate_rows(path), mapping)
