import openpyxl
from estimating_tools.excel import excel_reader


def test_reads_back_written_cells(tmp_path):
    p = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Item"
    ws["B1"] = 42
    wb.save(p)

    # read_excel returns a dict; with has_header=False the row is a list of values.
    result = excel_reader.read_excel(str(p), has_header=False)

    assert result["success"] is True
    assert result["rows"] >= 1

    # First row should contain the written values.
    first_row = result["data"][0]
    assert "Item" in first_row
    assert 42 in first_row


def test_reads_back_with_header(tmp_path):
    p = tmp_path / "headed.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Name"
    ws["B1"] = "Quantity"
    ws["A2"] = "Concrete"
    ws["B2"] = 100
    wb.save(p)

    result = excel_reader.read_excel(str(p), has_header=True)

    assert result["success"] is True
    assert result["headers"] == ["Name", "Quantity"]
    assert result["data"][0]["Name"] == "Concrete"
    assert result["data"][0]["Quantity"] == 100
