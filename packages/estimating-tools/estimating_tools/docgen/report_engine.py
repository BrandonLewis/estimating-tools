"""Shared report-style helpers for project DOCX outputs.

Usage:
    from estimating_tools.docgen.report_engine import ReportStyle

    rs = ReportStyle("Sample Road Project", subtitle="Excavation and Backfill")
    rs.add_title_page()
    rs.add_heading("1. Section Heading", level=1)
    rs.add_para("Body paragraph.")
    rs.add_table(["Col A", "Col B"], [["a1", "b1"]])
    rs.add_callout("Title", ["bullet 1", "bullet 2"])
    rs.save("/path/to/Output.docx")

Style follows these report-style defaults:
- Times New Roman, body 10 pt dark gray, justified
- Headings dark blue (#1F3864), bold; H1 16, H2 13, H3 11
- Table headers dark-blue fill, white 9 pt bold
- Alternating light-gray (#F2F2F2) row shading
- Callout box: light blue (#EBF0F7) fill, dark-blue border
- 1.0" margins, portrait 8.5" x 11"
- Centered live page-number footer
"""
from __future__ import annotations
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from markdown_it import MarkdownIt


DARK_BLUE = RGBColor(0x1F, 0x38, 0x64)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY_HEX = "F2F2F2"
DARK_BLUE_HEX = "1F3864"
LIGHT_BLUE_HEX = "EBF0F7"


def _set_cell_shading(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _set_cell_borders(cell, color="1F3864", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), size)
        b.set(qn("w:color"), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


class ReportStyle:
    def __init__(self, project_name: str, *, subtitle: str = "", tagline: str = ""):
        self.project_name = project_name
        self.subtitle = subtitle
        self.tagline = tagline
        self.doc = Document()
        self._setup()

    def _setup(self):
        for section in self.doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)
            section.page_width = Inches(8.5)
            section.page_height = Inches(11.0)

        normal = self.doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(10)
        normal.font.color.rgb = DARK_GRAY

        for level, size in ((1, 16), (2, 13), (3, 11)):
            style = self.doc.styles[f"Heading {level}"]
            style.font.name = "Times New Roman"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = DARK_BLUE
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)

        self._add_page_number_footer()

    def _add_page_number_footer(self):
        section = self.doc.sections[0]
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.color.rgb = DARK_GRAY
        f1 = OxmlElement("w:fldChar")
        f1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        f2 = OxmlElement("w:fldChar")
        f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1)
        run._r.append(instr)
        run._r.append(f2)

    def add_title_page(self, info_rows=None):
        title_p = self.doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.space_before = Pt(120)
        title_p.paragraph_format.space_after = Pt(6)
        r = title_p.add_run(self.project_name.upper())
        r.font.name = "Times New Roman"
        r.font.size = Pt(26)
        r.bold = True
        r.font.color.rgb = DARK_BLUE

        if self.subtitle:
            sub_p = self.doc.add_paragraph()
            sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub_p.paragraph_format.space_after = Pt(4)
            r = sub_p.add_run(self.subtitle)
            r.font.name = "Times New Roman"
            r.font.size = Pt(20)
            r.bold = True
            r.font.color.rgb = DARK_BLUE

        if self.tagline:
            tag_p = self.doc.add_paragraph()
            tag_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tag_p.paragraph_format.space_after = Pt(20)
            r = tag_p.add_run(self.tagline)
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            r.italic = True
            r.font.color.rgb = DARK_BLUE

        divp = self.doc.add_paragraph()
        divp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rd = divp.add_run("____________________________")
        rd.font.name = "Times New Roman"
        rd.font.size = Pt(13)
        rd.font.color.rgb = DARK_BLUE

        if info_rows:
            t = self.doc.add_table(rows=len(info_rows), cols=2)
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for i, (k, v) in enumerate(info_rows):
                p1 = t.rows[i].cells[0].paragraphs[0]
                p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r1 = p1.add_run(k)
                r1.font.name = "Times New Roman"
                r1.font.size = Pt(10)
                r1.bold = True
                r1.font.color.rgb = DARK_BLUE
                p2 = t.rows[i].cells[1].paragraphs[0]
                r2 = p2.add_run(v)
                r2.font.name = "Times New Roman"
                r2.font.size = Pt(10)
                r2.font.color.rgb = DARK_GRAY

        self.doc.add_page_break()

    def add_heading(self, text: str, level: int = 1):
        h = self.doc.add_heading(level=level)
        r = h.add_run(text)
        r.font.name = "Times New Roman"
        return h

    def add_para(self, text: str = "", *, bold=False, italic=False, size=10,
                 color=DARK_GRAY, align=None, space_before=0, space_after=4,
                 justify=True):
        p = self.doc.add_paragraph()
        if align is not None:
            p.alignment = align
        elif justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            r = p.add_run(text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
            r.font.color.rgb = color
        return p

    def add_bullet(self, text: str, *, indent=0.25):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.space_after = Pt(2)
        for r in list(p.runs):
            r.text = ""
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.color.rgb = DARK_GRAY
        return p

    def add_table(self, headers, rows, *, col_widths=None):
        n_cols = len(headers)
        table = self.doc.add_table(rows=1, cols=n_cols)
        table.alignment = WD_ALIGN_PARAGRAPH.LEFT

        if col_widths:
            for i, w in enumerate(col_widths):
                for cell in table.columns[i].cells:
                    cell.width = Inches(w)

        hdr = table.rows[0].cells
        for i, htext in enumerate(headers):
            cell = hdr[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_shading(cell, DARK_BLUE_HEX)
            _set_cell_borders(cell)
            for p in cell.paragraphs:
                p.text = ""
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(htext)
                r.font.name = "Times New Roman"
                r.font.size = Pt(9)
                r.font.bold = True
                r.font.color.rgb = WHITE
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)

        for row_idx, row in enumerate(rows):
            rcells = table.add_row().cells
            for j, val in enumerate(row):
                cell = rcells[j]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _set_cell_borders(cell)
                if row_idx % 2 == 0:
                    _set_cell_shading(cell, LIGHT_GRAY_HEX)
                for p in cell.paragraphs:
                    p.text = ""
                    r = p.add_run(str(val))
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9)
                    r.font.color.rgb = DARK_GRAY
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
            if col_widths:
                for j, w in enumerate(col_widths):
                    rcells[j].width = Inches(w)

        self.add_para("", space_after=4)
        return table

    def add_callout(self, title: str, bullets):
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        _set_cell_shading(cell, LIGHT_BLUE_HEX)
        _set_cell_borders(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        p = cell.paragraphs[0]
        p.text = ""
        r = p.add_run(title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.bold = True
        r.font.color.rgb = DARK_BLUE
        p.paragraph_format.space_after = Pt(4)
        for b in bullets:
            bp = cell.add_paragraph()
            bp.paragraph_format.left_indent = Inches(0.15)
            bp.paragraph_format.space_after = Pt(2)
            br = bp.add_run("• " + b)
            br.font.name = "Times New Roman"
            br.font.size = Pt(9)
            br.font.color.rgb = DARK_GRAY
        self.add_para("", space_after=4)

    def add_block_quote(self, text: str, citation: str = ""):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.italic = True
        r.font.color.rgb = DARK_GRAY
        if citation:
            cp = self.doc.add_paragraph()
            cp.paragraph_format.left_indent = Inches(0.5)
            cp.paragraph_format.space_after = Pt(8)
            cr = cp.add_run("— " + citation)
            cr.font.name = "Times New Roman"
            cr.font.size = Pt(9)
            cr.italic = True
            cr.font.color.rgb = DARK_BLUE

    def save(self, path: Path | str):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        return path


# ---------------------------------------------------------------------------
# Engine B: markdown → ReportStyle adapter
# ---------------------------------------------------------------------------

_CALLOUT_TITLE_RE = re.compile(r"^\[!CALLOUT\]\s+(.+)$", re.IGNORECASE)


def _flatten_inline(token) -> str:
    """Collapse an inline token's children into a plain string (bold/italic stripped)."""
    parts: list[str] = []
    for child in token.children or []:
        if child.type == "text":
            parts.append(child.content)
        elif child.type == "softbreak":
            parts.append(" ")
        elif child.type == "code_inline":
            parts.append(child.content)
        elif child.type in (
            "strong_open", "strong_close",
            "em_open", "em_close",
            "s_open", "s_close",  # strikethrough: text preserved via child "text" tokens
        ):
            continue  # ReportStyle.add_para can't toggle inline runs; flatten for now.
    return "".join(parts).strip()


def _walk_blockquote(tokens, i, rs) -> int:
    """Handle a blockquote starting at tokens[i] (a `blockquote_open`).

    Two flavors:
        > [!CALLOUT] Title    → ReportStyle.add_callout(title, bullets)
        > body                → ReportStyle.add_block_quote(body)
    Returns the index of `blockquote_close`.
    """
    j = i + 1
    title: str | None = None
    bullets: list[str] = []
    body_lines: list[str] = []

    while j < len(tokens) and tokens[j].type != "blockquote_close":
        tok = tokens[j]
        if tok.type == "paragraph_open":
            inline = tokens[j + 1]
            text = _flatten_inline(inline)
            m = _CALLOUT_TITLE_RE.match(text)
            if m and title is None:
                title = m.group(1).strip()
            else:
                body_lines.append(text)
            j += 3  # paragraph_open, inline, paragraph_close
            continue
        if tok.type == "bullet_list_open":
            # Walk bullets inside the blockquote.
            k = j + 1
            while k < len(tokens) and tokens[k].type != "bullet_list_close":
                if tokens[k].type == "inline":
                    bullets.append(_flatten_inline(tokens[k]))
                k += 1
            j = k + 1
            continue
        j += 1

    if title is not None:
        rs.add_callout(title, bullets or body_lines)
    else:
        rs.add_block_quote(" ".join(body_lines))
        for b in bullets:
            rs.add_bullet(b)
    return j  # index of blockquote_close


def _walk_table(tokens, i, rs) -> int:
    """Convert a GFM table token sequence into ReportStyle.add_table()."""
    headers: list[str] = []
    rows: list[list[str]] = []
    j = i + 1
    in_header = False
    current: list[str] = []
    while j < len(tokens) and tokens[j].type != "table_close":
        tok = tokens[j]
        if tok.type == "thead_open":
            in_header = True
        elif tok.type == "thead_close":
            in_header = False
        elif tok.type == "tr_open":
            current = []
        elif tok.type == "tr_close":
            if in_header:
                headers = current
            else:
                rows.append(current)
        elif tok.type == "inline":
            current.append(_flatten_inline(tok))
        j += 1
    rs.add_table(headers, rows)
    return j  # index of table_close


def render_markdown(md_text: str, rs: ReportStyle) -> ReportStyle:
    """Walk a markdown document and emit it via ReportStyle on `rs`."""
    md = MarkdownIt("commonmark").enable("table").enable("strikethrough")
    tokens = md.parse(md_text)

    i = 0
    while i < len(tokens):
        tok = tokens[i]
        if tok.type == "heading_open":
            level = int(tok.tag[1])  # "h1" → 1
            inline = tokens[i + 1]
            rs.add_heading(_flatten_inline(inline), level=level)
            i += 3  # heading_open, inline, heading_close
            continue
        if tok.type == "paragraph_open":
            inline = tokens[i + 1]
            rs.add_para(_flatten_inline(inline))
            i += 3
            continue
        if tok.type == "bullet_list_open":
            j = i + 1
            while j < len(tokens) and tokens[j].type != "bullet_list_close":
                if tokens[j].type == "inline":
                    rs.add_bullet(_flatten_inline(tokens[j]))
                j += 1
            i = j + 1
            continue
        if tok.type == "blockquote_open":
            i = _walk_blockquote(tokens, i, rs) + 1
            continue
        if tok.type == "table_open":
            i = _walk_table(tokens, i, rs) + 1
            continue
        i += 1
    return rs
