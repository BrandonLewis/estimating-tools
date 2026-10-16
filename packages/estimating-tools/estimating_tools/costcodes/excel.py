"""Import an org's cost-code system from a spreadsheet into the canonical model.

The pure builder (`build_cost_code_system`) maps the org's columns to canonical fields
(identity default), derives cost types from the referenced cost_type codes, validates the
activity-code tree (unique codes, parent exists, no self-ref, no cycle), and returns the
activity codes in topological order (parents first) — no openpyxl/MCP. Only
`read_costcode_rows` reads the .xlsx (openpyxl)."""
from __future__ import annotations

from estimating_tools.costcodes.model import (
    ActivityCode, CanonicalCostCodeSystem, CostType,
)
from estimating_tools.costcodes.strategies import resolve_parents

_CANONICAL_FIELDS = (
    "code", "description", "category", "cost_type", "cost_type_name", "parent_code",
    "default_unit", "notes", "sort_order", "is_active", "structure_tags", "indent",
)
_REQUIRED = ("code",)
_FALSY = {"no", "false", "0", "n"}


def _effective_mapping(mapping):
    m = {f: f for f in _CANONICAL_FIELDS}      # identity default
    if mapping:
        m.update(mapping)
    return m


def _str(value):
    return str(value).strip() if value is not None else ""


def _int(value):
    text = _str(value)
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def _is_active(value):
    # falsy-override: blank -> True; an explicit falsy token -> False; else True
    return _str(value).lower() not in _FALSY


def _topo_sort(activity_codes):
    """Return activity_codes with every parent before its children. Roots and siblings keep
    input order. Assumes validation already ruled out missing parents / self-refs / cycles."""
    children: dict[str, list] = {a.code: [] for a in activity_codes}
    roots = []
    for a in activity_codes:
        if a.parent_code:
            children[a.parent_code].append(a)
        else:
            roots.append(a)
    ordered = []
    stack = list(reversed(roots))           # process roots in input order
    while stack:
        node = stack.pop()
        ordered.append(node)
        # push children in reverse so they pop in input order, after the parent
        for child in reversed(children[node.code]):
            stack.append(child)
    return ordered


def build_cost_code_system(rows, *, system_name, description="", version="", mapping=None,
                           parent_strategy="explicit_column", strategy_opts=None,
                           composition="composed", source_system="", edition="") -> CanonicalCostCodeSystem:
    """Pure: spreadsheet rows (dicts keyed by the org's headers) + a column mapping +
    a system name -> a CanonicalCostCodeSystem (derived cost types, validated + topo-sorted
    activity codes). Parents are resolved by the named strategy (default explicit_column).
    No openpyxl/MCP."""
    if not _str(system_name):
        raise ValueError("system_name is required (a non-blank name for the cost-code system)")
    system = CanonicalCostCodeSystem(
        name=_str(system_name), description=_str(description), version=_str(version),
        source_system=_str(source_system), edition=_str(edition),
        composition=_str(composition) or "composed",
    )
    if not rows:
        return system
    m = _effective_mapping(mapping)
    present = set()
    for r in rows:
        present.update(r.keys())
    for field_name in _REQUIRED:
        if m[field_name] not in present:
            raise ValueError(
                f"required field {field_name!r} maps to column {m[field_name]!r}, not "
                f"found in the spreadsheet headers {sorted(present)}"
            )

    parents = resolve_parents(parent_strategy, rows, m, **(strategy_opts or {}))

    activity_codes = []
    seen_codes = set()
    cost_type_order = []
    cost_type_names = {}
    for r in rows:
        code = _str(r.get(m["code"]))
        if not code:
            continue                            # blank code / blank row -> skip
        if code in seen_codes:
            raise ValueError(f"duplicate activity code {code!r} (codes must be unique)")
        seen_codes.add(code)
        ct_code = _str(r.get(m["cost_type"]))
        if ct_code and ct_code not in cost_type_names:
            cost_type_order.append(ct_code)
            cost_type_names[ct_code] = _str(r.get(m["cost_type_name"])) or ct_code
        raw_code = r.get(m["code"])
        tags_raw = _str(r.get(m["structure_tags"]))
        structure_tags = [t.strip() for t in tags_raw.split("|") if t.strip()]
        activity_codes.append(ActivityCode(
            code=code,
            source_code=str(raw_code) if raw_code is not None else "",
            description=_str(r.get(m["description"])),
            category=_str(r.get(m["category"])),
            cost_type_code=ct_code,
            parent_code=parents.get(code, ""),
            default_unit=_str(r.get(m["default_unit"])),
            structure_tags=structure_tags,
            notes=_str(r.get(m["notes"])),
            sort_order=_int(r.get(m["sort_order"])),
            is_active=_is_active(r.get(m["is_active"])),
        ))

    # validate the tree: self-ref, missing parent
    for a in activity_codes:
        if a.parent_code:
            if a.parent_code == a.code:
                raise ValueError(f"activity code {a.code!r} cannot be its own parent")
            if a.parent_code not in seen_codes:
                raise ValueError(
                    f"activity code {a.code!r} has parent_code {a.parent_code!r}, "
                    "which is not a code in the sheet"
                )
    # detect cycles — O(n) via a "safe" set of nodes whose chain is proven acyclic;
    # report only the codes in the cycle (not the prefix that leads into it).
    parent_of = {a.code: a.parent_code for a in activity_codes}
    safe = set()
    for start in parent_of:
        if start in safe:
            continue
        path = []
        index = {}
        node = start
        while node and node not in safe:
            if node in index:
                cycle = path[index[node]:]
                raise ValueError(f"cycle detected in activity-code parents: {cycle}")
            index[node] = len(path)
            path.append(node)
            node = parent_of.get(node, "")
        safe.update(path)

    system.cost_types = [
        CostType(code=c, name=cost_type_names[c], sort_order=i)
        for i, c in enumerate(cost_type_order)
    ]
    system.activity_codes = _topo_sort(activity_codes)
    _assign_levels(system.activity_codes)
    return system


def _assign_levels(ordered):
    """Set .level on each code (0 = root). ordered must be parents-first."""
    level_by_code = {}
    for a in ordered:
        level_by_code[a.code] = (
            0 if not a.parent_code else level_by_code.get(a.parent_code, 0) + 1
        )
        a.level = level_by_code[a.code]


def read_costcode_rows(path) -> list[dict]:
    """Read the first worksheet of an .xlsx into row-dicts keyed by header. The only
    openpyxl/file touch. Values are the raw openpyxl cell values; build_cost_code_system
    coerces them via _str/_int."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise ImportError(
            "reading .xlsx requires openpyxl: pip install 'estimating_tools[excel]'"
        ) from exc
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = None
        out: list[dict] = []
        for raw in ws.iter_rows(values_only=True):
            if raw is None or all(c is None for c in raw):
                continue
            if headers is None:
                headers = [_str(c) for c in raw]
                continue
            out.append({h: v for h, v in zip(headers, raw)})
        return out
    finally:
        wb.close()


def build_cost_code_system_from_excel(path, *, system_name, description="", version="", mapping=None,
                                      parent_strategy="explicit_column", strategy_opts=None,
                                      composition="composed", source_system="", edition="") -> CanonicalCostCodeSystem:
    return build_cost_code_system(read_costcode_rows(path), system_name=system_name,
                                  description=description, version=version, mapping=mapping,
                                  parent_strategy=parent_strategy, strategy_opts=strategy_opts,
                                  composition=composition, source_system=source_system,
                                  edition=edition)
