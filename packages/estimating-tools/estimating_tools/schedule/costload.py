"""Cost- & resource-load a CanonicalSchedule (5D spec).

The distribution itself (validate_costs, distribute_to_tasks,
distribute_to_task_codes, load_schedule_costs and their dataclasses) lives in
`bimforge.bim.costload`, so a consumer with only bimforge can write 5D; it is
re-exported here. This module adds the estimate side: deriving the cost record
from model quantities and the estimate workbook's item totals.
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Iterable, Mapping

from bimforge.bim.costload import (  # noqa: F401  (re-export)
    Coverage, LoadedSchedule, PeriodLoad, _bucketize, _d, _month_key,
    _week_key, _workdays, distribute_to_task_codes, distribute_to_tasks,
    load_schedule_costs, validate_costs,
)

# ===========================================================================
# model-QTO -> costs derivation (#205)
#
# The clean half of Work's build_5d.derive_costs, org-agnostic: the estimate
# workbook stays the single source of cost and the model the single source of
# quantity, so no money is re-keyed. The item's ONE governing pay code carries
# its whole burdened total; unit_price = total / modelled quantity. Output is
# the validate_costs shape below, which `est-bim estimate5d` embeds and
# cross-checks (its unit guard catches the $/CY-on-an-LF-item mistake).
# ===========================================================================

@dataclass(frozen=True)
class CodeSpec:
    """A pay code, the bid item it belongs to, and its modelled quantity/unit."""

    code: str
    item: str
    quantity: float
    unit: str = ""


def derive_costs(specs: Iterable[CodeSpec], governing: Mapping[str, str],
                 totals: Mapping[str, float]) -> tuple[dict, list[str]]:
    """Return ``(costs, unrated)`` from model quantities + workbook item totals.

    ``governing`` maps ``item -> the ONE code that carries it``; that code is
    priced ``cost = totals[item]``, ``unit_price = total / quantity``. Every
    other code is *unrated* (its cost rides in a governing code's rate). A code
    that governs several items gets the item-keyed shape ``unit_price``/``unit``
    of ``{item: …}`` with ``cost`` summed, so no item is silently overwritten.

    Fails loud — never a silent drop, the mistake the hand-kept tables made —
    when a ``governing`` mapping has no matching spec (the item would vanish and
    its codes fall to unrated), when a governing item is absent from ``totals``,
    or when a governed quantity is non-positive.
    """
    specs = list(specs)
    spec_of = {(s.item, s.code): s for s in specs}
    for item, code in governing.items():
        if (item, code) not in spec_of:
            raise ValueError(
                f"governing code {code!r} for item {item!r} has no matching "
                f"CodeSpec — cannot price it (the item would be dropped)")

    gov_items: dict[str, list[str]] = {}
    for item, code in governing.items():
        gov_items.setdefault(code, []).append(item)

    costs: dict[str, dict] = {}
    for code, items in gov_items.items():
        recs = []
        for item in items:
            if item not in totals:
                raise ValueError(
                    f"{code} governs {item!r}, which the estimate totals do "
                    f"not carry (found {sorted(totals)})")
            spec = spec_of[(item, code)]
            if spec.quantity <= 0:
                raise ValueError(
                    f"{code} has no modelled quantity to price for {item!r} "
                    f"(quantity={spec.quantity!r})")
            recs.append((item, float(totals[item]), spec.quantity, spec.unit))
        if len(recs) == 1:
            _item, total, qty, unit = recs[0]
            costs[code] = {"cost": total, "labor_hours": None,
                           "unit_price": total / qty, "unit": unit}
        else:                              # one code carrying several items
            costs[code] = {
                "cost": sum(t for _, t, _, _ in recs),
                "labor_hours": None,
                "unit_price": {it: t / q for it, t, q, _ in recs},
                "unit": {it: u for it, _, _, u in recs},
            }

    governing_codes = set(governing.values())
    unrated: list[str] = []
    seen: set[str] = set()
    for s in specs:
        if s.code not in governing_codes and s.code not in seen:
            seen.add(s.code)
            unrated.append(s.code)
    return costs, unrated


def tie_out(costs: Mapping[str, dict], specs: Iterable[CodeSpec],
            governing: Mapping[str, str],
            totals: Mapping[str, float]) -> tuple[float, float]:
    """Independent cross-check → ``(extended, priced)``.

    Sums over EVERY governing ``(item, code)``: ``extended`` re-derives
    ``unit_price * quantity`` (resolving the item-keyed rate for a plural code)
    and ``priced`` sums the item totals. Raises on any missing spec / cost /
    total rather than skipping it — otherwise a dropped item would fall out of
    both sides and the cross-check would pass while money is lost. The caller
    asserts the two tie within a cent.
    """
    spec_of = {(s.item, s.code): s for s in specs}
    extended = priced = 0.0
    for item, code in governing.items():
        if code not in costs:
            raise ValueError(
                f"tie_out: governing code {code!r} for {item!r} is not priced")
        if item not in totals:
            raise ValueError(
                f"tie_out: governing item {item!r} missing from totals")
        if (item, code) not in spec_of:
            raise ValueError(
                f"tie_out: no spec for governing {item!r}/{code!r}")
        up = costs[code]["unit_price"]
        rate = up[item] if isinstance(up, dict) else up
        extended += rate * spec_of[(item, code)].quantity
        priced += float(totals[item])
    return extended, priced


def item_totals_from_rows(rows: Iterable, *, prefix: str = "Item #"
                          ) -> dict[str, float]:
    """``{label: total}`` for ``(label, total)`` rows whose label starts with
    ``prefix`` and whose total is numeric.

    Walks the Summary rows rather than pinning cells: the per-item sheets grow a
    row whenever a note is added, so any fixed cell reference rots, but the
    ``label``/``total`` column pairing the Summary is built from is stable.
    """
    out: dict[str, float] = {}
    for label, total in rows:
        if (isinstance(label, str) and label.strip().startswith(prefix)
                and isinstance(total, (int, float))
                and not isinstance(total, bool)):
            out[label.strip()] = float(total)
    return out


def read_item_totals(path: str, *, sheet: str = "Summary",
                     label_col: str = "A", total_col: str = "E",
                     prefix: str = "Item #") -> dict[str, float]:
    """``{item: burdened total}`` from an estimate workbook's Summary sheet.

    Reads cached values (``data_only``): the workbook stores FORMULAS, so it
    must have been evaluated first (open in Excel, or
    ``calcbook.recalculate``). If ``prefix`` rows exist but none carries a
    cached number, the workbook was not evaluated — raise loudly rather than
    return an empty schedule.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb[sheet]
        rows = [(ws[f"{label_col}{r}"].value, ws[f"{total_col}{r}"].value)
                for r in range(1, ws.max_row + 1)]
    finally:
        wb.close()      # release the file handle (Windows would keep it locked)
    has_item_rows = any(isinstance(lbl, str) and lbl.strip().startswith(prefix)
                        for lbl, _ in rows)
    out = item_totals_from_rows(rows, prefix=prefix)
    if has_item_rows and not out:
        raise ValueError(
            f"{path}: found {prefix!r} rows but no cached totals — the "
            "workbook's formulas were not evaluated; recalc it first "
            "(calcbook.recalculate) or open it in Excel")
    return out
