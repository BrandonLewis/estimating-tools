"""Convert Gantt schedule workbooks into GanttWarrior `.gw.json`.

The input workbook is expected to have these sheets (created by the
`Gantt PROJ-001 SCH *.xlsx` schedule template):

- `Settings` — Project Start Date, Project Name/Number, Work Days/Week,
  Skip Holidays flag, and a list of holiday dates.
- `Schedule` — header row 4 with columns:
  `WBS | Task Name | Duration | Offset | Start Date | End Date |
   Predecessors | Crew | Manhours | Labor Cost | Total Cost | Progress`.

Offsets are whole work-day counts relative to the direct parent's start
(top-level rows are relative to the project start). Durations may be
fractional and are rounded up to the next whole work day. Weekends and
holidays are skipped when expanding `work_days`.

Output is written to `<stem>.gw.json` alongside the input by default; pass
`--output-dir` to redirect.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl


COLORS = ["red", "green", "blue", "yellow", "magenta", "cyan", "orange", "purple", "white"]
WEEKEND = {5, 6}  # Sat, Sun


@dataclass
class Settings:
    project_name: str
    project_number: str
    start_date: date
    work_days_per_week: int
    skip_holidays: bool
    holidays: set[date]

    @property
    def work_weekdays(self) -> set[int]:
        # 5-day week → Mon-Fri (weekdays 0-4). Anything else falls back to that
        # for now; the source workbooks only use 5 in practice.
        return {0, 1, 2, 3, 4}


@dataclass
class ScheduleRow:
    wbs: str
    name: str
    duration: float
    offset: int
    crew: str
    manhours: float
    progress: float


def _cell_value(ws, row: int, col: int):
    return ws.cell(row=row, column=col).value


def parse_settings(wb) -> Settings:
    ws = wb["Settings"]
    fields: dict[str, object] = {}
    holidays: set[date] = set()
    in_holidays = False
    for row in ws.iter_rows(values_only=True):
        label = row[0]
        if label == "HOLIDAYS":
            in_holidays = True
            continue
        if in_holidays:
            if isinstance(label, datetime):
                holidays.add(label.date())
            elif isinstance(label, date):
                holidays.add(label)
            continue
        if isinstance(label, str):
            fields[label.rstrip(":").strip()] = row[1]

    start = fields.get("Project Start Date")
    if isinstance(start, datetime):
        start_date = start.date()
    elif isinstance(start, date):
        start_date = start
    else:
        raise ValueError("Settings sheet missing 'Project Start Date'")

    skip = str(fields.get("Skip Holidays", "Yes")).strip().lower() in {"yes", "true", "1"}

    return Settings(
        project_name=str(fields.get("Project Name") or "").strip() or "Untitled Schedule",
        project_number=str(fields.get("Project Number") or "").strip(),
        start_date=start_date,
        work_days_per_week=int(fields.get("Work Days/Week") or 5),
        skip_holidays=skip,
        holidays=holidays,
    )


def parse_schedule(wb) -> list[ScheduleRow]:
    ws = wb["Schedule"]
    # Header is on row 4: WBS | Task Name | Duration | Offset | ...
    rows: list[ScheduleRow] = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        wbs = row[0]
        name = row[1]
        if not wbs or not name:
            continue
        duration = float(row[2] or 0)
        offset_raw = row[3]
        offset = int(offset_raw) if offset_raw is not None else 0
        crew = str(row[7] or "")
        manhours = float(row[8] or 0)
        progress = float(row[11] or 0)
        rows.append(
            ScheduleRow(
                wbs=str(wbs).strip(),
                name=str(name).strip(),
                duration=duration,
                offset=offset,
                crew=crew,
                manhours=manhours,
                progress=progress,
            )
        )
    return rows


def parent_wbs(wbs: str) -> str:
    if "." not in wbs:
        return ""
    return wbs.rsplit(".", 1)[0]


def _is_workday(d: date, work_weekdays: set[int], holidays: set[date], skip_holidays: bool) -> bool:
    if d.weekday() not in work_weekdays:
        return False
    if skip_holidays and d in holidays:
        return False
    return True


def shift_workdays(start: date, n: int, work_weekdays: set[int], holidays: set[date], skip_holidays: bool) -> date:
    """Return the work day reached by stepping `n` work days forward from `start`.

    If `start` is itself a work day, n=0 returns it. Non-work days are skipped
    over without consuming a step. Negative `n` is not supported.
    """
    if n < 0:
        raise ValueError("Negative work-day shifts are not supported")
    cur = start
    while not _is_workday(cur, work_weekdays, holidays, skip_holidays):
        cur += timedelta(days=1)
    steps = 0
    while steps < n:
        cur += timedelta(days=1)
        if _is_workday(cur, work_weekdays, holidays, skip_holidays):
            steps += 1
    return cur


def expand_workdays(start: date, count: int, work_weekdays: set[int], holidays: set[date], skip_holidays: bool) -> list[date]:
    """Return the next `count` consecutive work days starting at-or-after `start`."""
    if count <= 0:
        return []
    days: list[date] = []
    cur = start
    while len(days) < count:
        if _is_workday(cur, work_weekdays, holidays, skip_holidays):
            days.append(cur)
        cur += timedelta(days=1)
    return days


def build_tasks(rows: list[ScheduleRow], settings: Settings) -> list[dict]:
    by_wbs: dict[str, ScheduleRow] = {r.wbs: r for r in rows}
    starts: dict[str, date] = {}
    ids: dict[str, str] = {}

    # Stable color per top-level WBS group (e.g. A1a, A2a, ...)
    top_groups: list[str] = []
    for r in rows:
        top = r.wbs.split(".")[0]
        if top not in top_groups:
            top_groups.append(top)
    color_map = {g: COLORS[i % len(COLORS)] for i, g in enumerate(top_groups)}

    def task_start(wbs: str) -> date:
        if wbs in starts:
            return starts[wbs]
        row = by_wbs[wbs]
        parent = parent_wbs(wbs)
        base = task_start(parent) if parent else settings.start_date
        s = shift_workdays(
            base, row.offset, settings.work_weekdays, settings.holidays, settings.skip_holidays
        )
        starts[wbs] = s
        return s

    tasks: list[dict] = []
    for r in rows:
        ids[r.wbs] = uuid.uuid4().hex[:8]

    for r in rows:
        start = task_start(r.wbs)
        dur_days = max(1, math.ceil(r.duration))
        work_days = expand_workdays(
            start, dur_days, settings.work_weekdays, settings.holidays, settings.skip_holidays
        )
        end = work_days[-1]
        description_parts = []
        if r.crew:
            description_parts.append(f"Crew: {r.crew}")
        if r.manhours:
            description_parts.append(f"Manhours: {r.manhours:g}")
        if r.duration != dur_days:
            description_parts.append(f"Source duration: {r.duration:g} day(s)")
        description = "\n".join(description_parts)

        tasks.append(
            {
                "id": ids[r.wbs],
                "name": r.name,
                "wbs": r.wbs,
                "description": description,
                "start_date": start.isoformat(),
                "end_date": end.isoformat(),
                "duration_days": dur_days,
                "status": "not_started",
                "color": color_map[r.wbs.split(".")[0]],
                "progress": float(r.progress or 0),
                "dependencies": [],
                "assigned_to": r.crew,
                "priority": 0,
                "is_milestone": False,
                "parent_wbs": parent_wbs(r.wbs),
                "work_days": [d.isoformat() for d in work_days],
                "work_weekdays": sorted(settings.work_weekdays),
                "manually_edited": False,
                "time_blocks": [],
                "timezone": None,
            }
        )
    return tasks


def build_project(settings: Settings, tasks: list[dict], source_label: str) -> dict:
    name = source_label or settings.project_name
    description_bits = []
    if settings.project_number:
        description_bits.append(f"Project #: {settings.project_number}")
    if settings.project_name and settings.project_name != name:
        description_bits.append(f"Project Name: {settings.project_name}")
    return {
        "name": name,
        "description": " | ".join(description_bits),
        "tasks": tasks,
        "start_date": settings.start_date.isoformat(),
        "created_at": datetime.now().isoformat(),
        "default_work_weekdays": sorted(settings.work_weekdays),
        "holidays": sorted(d.isoformat() for d in settings.holidays),
        "timezone": "UTC",
        "calendar_granularity": 30,
    }


def _project_label(path: Path) -> str:
    # "Gantt PROJ-001 SCH A.xlsx" -> "PROJ-001 SCH A"
    stem = path.stem
    stem = re.sub(r"^(?:Gantt\s+)", "", stem, flags=re.IGNORECASE)
    return stem.strip()


def convert(xlsx_path: Path, output_dir: Path | None) -> Path:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    settings = parse_settings(wb)
    rows = parse_schedule(wb)
    tasks = build_tasks(rows, settings)
    project = build_project(settings, tasks, _project_label(xlsx_path))

    out_dir = output_dir or xlsx_path.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{xlsx_path.stem}.gw.json"
    out_path.write_text(json.dumps(project, indent=2))
    return out_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("xlsx", nargs="+", type=Path, help="Schedule workbook(s) to convert")
    parser.add_argument(
        "-d",
        "--output-dir",
        type=Path,
        default=None,
        help="Directory for the .gw.json files (default: alongside each input)",
    )
    args = parser.parse_args(argv)

    errors = 0
    for xlsx in args.xlsx:
        if not xlsx.exists():
            print(f"error: {xlsx} not found", file=sys.stderr)
            errors += 1
            continue
        try:
            out = convert(xlsx, args.output_dir)
        except Exception as e:
            print(f"error converting {xlsx}: {e}", file=sys.stderr)
            errors += 1
            continue
        print(f"wrote {out}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
