"""R8 sub-scope filter — scope.py + external_dependencies + CLI + Task Table.

Fixture: 3-level dotted WBS, 6 leaves under 3 top summaries, resources and
activity_codes populated, cross-scope dependencies (T2 depends on T1 inside
the pile scope and T5 outside it).
"""
import json
import zipfile

import pytest

from estimating_tools.schedule import cli, scope
from estimating_tools.schedule.importers import load_schedule
from bimforge.bim.schedule_model import (
    CanonicalSchedule, Dependency, ScheduleTask,
)


def _sched() -> CanonicalSchedule:
    return CanonicalSchedule(name="GC master", start_date="2026-09-01", tasks=[
        ScheduleTask(id="S1", name="Sitework", wbs="1"),
        ScheduleTask(id="S1a", name="Piles", wbs="1.1"),
        ScheduleTask(id="T1", name="Drive test pile", wbs="1.1.1",
                     start="2026-09-01", finish="2026-09-03",
                     duration_days=3.0,
                     resources=["Pile Crew"],
                     activity_codes=["Subcontractor: SUB-A"]),
        ScheduleTask(id="T2", name="Drive production piles", wbs="1.1.2",
                     start="2026-09-04", finish="2026-09-18",
                     duration_days=11.0,
                     resources=["Pile Crew"],
                     activity_codes=["Subcontractor: SUB-A"],
                     dependencies=[
                         Dependency(predecessor="T1"),
                         Dependency(predecessor="T5", type="SS", lag_days=2.0),
                     ]),
        ScheduleTask(id="S2", name="Concrete", wbs="2"),
        ScheduleTask(id="S2a", name="Footings", wbs="2.1"),
        ScheduleTask(id="T3", name="Form footings", wbs="2.1.1",
                     start="2026-09-21", finish="2026-09-25",
                     resources=["Concrete Crew"],
                     activity_codes=["Subcontractor: SUB-B"],
                     dependencies=[Dependency(predecessor="T2")]),
        ScheduleTask(id="T4", name="Pour footings", wbs="2.1.2",
                     start="2026-09-28", finish="2026-09-28",
                     is_milestone=True,
                     resources=["Concrete Crew"],
                     activity_codes=["Subcontractor: SUB-B"],
                     dependencies=[Dependency(predecessor="T3")]),
        ScheduleTask(id="S3", name="Paving", wbs="3"),
        ScheduleTask(id="S3a", name="Mainline", wbs="3.1"),
        ScheduleTask(id="T5", name="Grade subbase", wbs="3.1.1",
                     start="2026-09-02", finish="2026-09-10",
                     resources=["Paving Sub"],
                     activity_codes=["Phase: Early"]),
        ScheduleTask(id="T6", name="Pave mainline", wbs="3.1.2",
                     start="2026-09-11", finish="2026-09-18",
                     resources=["Paving Sub"],
                     activity_codes=["Phase: Early"],
                     dependencies=[Dependency(predecessor="T5")]),
    ])


# -- resource filter ---------------------------------------------------------

def test_resource_filter_keeps_matches_and_only_their_summaries():
    filtered, report = scope.filter_to_scope(_sched(), resource="pile")
    ids = [t.id for t in filtered.tasks]
    # 2 matched leaves + ONLY their ancestor summaries, document order
    assert ids == ["S1", "S1a", "T1", "T2"]
    assert "S2" not in ids and "S3" not in ids
    assert report.matched == 2
    assert report.summaries_kept == 2
    assert report.dropped_tasks == 8


def test_resource_substring_spans_multiple_names_unions_scope():
    # "crew" matches "Pile Crew" AND "Concrete Crew" — union, no ambiguity rule
    filtered, report = scope.filter_to_scope(_sched(), resource="crew")
    ids = [t.id for t in filtered.tasks]
    assert ids == ["S1", "S1a", "T1", "T2", "S2", "S2a", "T3", "T4"]
    assert report.matched == 4


def test_resource_zero_matches_raises_naming_needle():
    with pytest.raises(ValueError, match="nobody"):
        scope.filter_to_scope(_sched(), resource="nobody")


def test_exactly_one_selector_required():
    with pytest.raises(ValueError):
        scope.filter_to_scope(_sched())
    with pytest.raises(ValueError):
        scope.filter_to_scope(_sched(), resource="pile",
                              activity_code="SUB-A")


# -- dependency split --------------------------------------------------------

def test_dependency_split_internal_stays_external_moves():
    filtered, report = scope.filter_to_scope(_sched(), resource="pile")
    t2 = next(t for t in filtered.tasks if t.id == "T2")
    assert [d.predecessor for d in t2.dependencies] == ["T1"]
    assert [d.predecessor for d in t2.external_dependencies] == ["T5"]
    ext = t2.external_dependencies[0]
    assert ext.type == "SS" and ext.lag_days == 2.0
    assert report.internal_preds == 1
    assert report.external_preds == 1


def test_filter_does_not_mutate_input():
    sched = _sched()
    scope.filter_to_scope(sched, resource="pile")
    t2 = next(t for t in sched.tasks if t.id == "T2")
    assert [d.predecessor for d in t2.dependencies] == ["T1", "T5"]
    assert t2.external_dependencies == []


# -- activity-code exactly-one rule ------------------------------------------

def test_activity_code_ambiguous_needle_raises_listing_candidates():
    with pytest.raises(ValueError) as exc:
        scope.filter_to_scope(_sched(), activity_code="sub")
    assert "Subcontractor: SUB-A" in str(exc.value)
    assert "Subcontractor: SUB-B" in str(exc.value)


def test_activity_code_zero_matches_raises_naming_needle():
    with pytest.raises(ValueError, match="zzz"):
        scope.filter_to_scope(_sched(), activity_code="zzz")


def test_activity_code_exact_one_filters():
    filtered, report = scope.filter_to_scope(_sched(), activity_code="SUB-A")
    ids = [t.id for t in filtered.tasks]
    assert ids == ["S1", "S1a", "T1", "T2"]
    assert "Subcontractor: SUB-A" in report.resolved


# -- list helpers ------------------------------------------------------------

def test_list_resources_names_and_counts():
    assert scope.list_resources(_sched()) == [
        ("Concrete Crew", 2), ("Paving Sub", 2), ("Pile Crew", 2)]


def test_list_activity_codes_displays_and_counts():
    assert scope.list_activity_codes(_sched()) == [
        ("Phase: Early", 2), ("Subcontractor: SUB-A", 2),
        ("Subcontractor: SUB-B", 2)]


# -- validate() with external deps -------------------------------------------

def test_filtered_output_validates():
    filtered, _ = scope.filter_to_scope(_sched(), resource="pile")
    filtered.validate()  # external T5 pred must NOT be an unknown-target error


def test_external_dep_types_still_validated():
    sched = CanonicalSchedule(name="s", tasks=[
        ScheduleTask(id="a", name="a", wbs="1",
                     external_dependencies=[
                         Dependency(predecessor="ghost", type="XX")])])
    with pytest.raises(ValueError, match="XX"):
        sched.validate()


# -- JSON round-trip ---------------------------------------------------------

def test_round_trip_restores_external_dependencies():
    filtered, _ = scope.filter_to_scope(_sched(), resource="pile")
    back = CanonicalSchedule.from_json(filtered.to_json())
    t2 = next(t for t in back.tasks if t.id == "T2")
    assert [d.predecessor for d in t2.external_dependencies] == ["T5"]
    assert t2.external_dependencies[0].type == "SS"
    assert isinstance(t2.external_dependencies[0], Dependency)


def test_pre_r8_shape_round_trips_without_new_keys():
    doc = {"name": "S", "start_date": None,
           "work_weekdays": [0, 1, 2, 3, 4], "holidays": [],
           "tasks": [{"id": "t1", "name": "x", "wbs": "1",
                      "start": None, "finish": None, "duration_days": None,
                      "is_milestone": False, "dependencies": [],
                      "bid_item_codes": [],
                      "match_name_re": None, "match_ifc_class": None}]}
    out = CanonicalSchedule.from_dict(doc).to_dict()
    assert out == doc  # no external_dependencies/resources/activity_codes keys


# -- Task Table xlsx ---------------------------------------------------------

def _load_rows(path):
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_xlsx_headers_types_and_pred_rendering(tmp_path):
    filtered, _ = scope.filter_to_scope(_sched(), resource="pile")
    out = tmp_path / "scope.xlsx"
    scope.write_task_table(filtered, str(out))
    rows = _load_rows(out)
    assert rows[0] == ["ID", "Outline Level", "Task Name", "Type", "Duration",
                       "Start", "Finish", "Predecessors", "External Preds"]
    by_id = {r[0]: r for r in rows[1:]}
    assert by_id["S1"][3] == "Summary" and by_id["S1a"][3] == "Summary"
    assert by_id["T1"][3] == "Task"
    # indent by outline level
    assert by_id["T1"][1] == 3
    assert by_id["T1"][2].startswith("        ")
    # non-FS external pred renders id + type suffix + lag
    assert by_id["T2"][7] == "T1"
    assert by_id["T2"][8] == "T5SS+2d"


def test_xlsx_double_build_is_byte_equal(tmp_path):
    filtered, _ = scope.filter_to_scope(_sched(), resource="pile")
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    scope.write_task_table(filtered, str(a))
    scope.write_task_table(filtered, str(b))
    assert a.read_bytes() == b.read_bytes()
    # and zip member timestamps are pinned, not wall clock
    with zipfile.ZipFile(a) as z:
        stamps = {i.date_time for i in z.infolist()}
    assert len(stamps) == 1


# -- CLI ---------------------------------------------------------------------

def _sched_file(tmp_path):
    p = tmp_path / "gc.json"
    p.write_text(_sched().to_json(), encoding="utf-8")
    return str(p)


def test_cli_scope_resource_writes_loadable_json(tmp_path, capsys):
    src = _sched_file(tmp_path)
    out = tmp_path / "sub.json"
    assert cli.main(["scope", src, "--resource", "pile",
                     "-o", str(out)]) == 0
    err = capsys.readouterr().err
    assert "2" in err and "internal" in err and "external" in err
    back = load_schedule(str(out))
    assert [t.id for t in back.tasks] == ["S1", "S1a", "T1", "T2"]
    t2 = next(t for t in back.tasks if t.id == "T2")
    assert [d.predecessor for d in t2.external_dependencies] == ["T5"]


def test_cli_scope_xlsx_output(tmp_path):
    src = _sched_file(tmp_path)
    out = tmp_path / "sub.xlsx"
    assert cli.main(["scope", src, "--resource", "pile",
                     "--xlsx", str(out)]) == 0
    rows = _load_rows(out)
    assert rows[0][0] == "ID" and rows[0][8] == "External Preds"


def test_cli_list_resources(tmp_path, capsys):
    src = _sched_file(tmp_path)
    assert cli.main(["scope", src, "--list-resources"]) == 0
    out = capsys.readouterr().out
    assert "Pile Crew" in out and "2" in out


def test_cli_list_codes(tmp_path, capsys):
    src = _sched_file(tmp_path)
    assert cli.main(["scope", src, "--list-codes"]) == 0
    out = capsys.readouterr().out
    assert "Subcontractor: SUB-A" in out


def test_cli_both_selectors_exit_2(tmp_path):
    src = _sched_file(tmp_path)
    with pytest.raises(SystemExit) as exc:
        cli.main(["scope", src, "--resource", "pile",
                  "--activity-code", "SUB-A"])
    assert exc.value.code == 2


def test_cli_no_selector_no_list_exit_2(tmp_path):
    src = _sched_file(tmp_path)
    with pytest.raises(SystemExit) as exc:
        cli.main(["scope", src])
    assert exc.value.code == 2


def test_cli_no_match_is_clean_error(tmp_path, capsys):
    src = _sched_file(tmp_path)
    assert cli.main(["scope", src, "--resource", "nobody"]) == 1
    assert "error:" in capsys.readouterr().err.lower()


# -- review round ------------------------------------------------------------

def test_refilter_preserves_existing_external_dependencies():
    # An already-filtered schedule carries external deps of its own —
    # re-filtering must keep them (the GC ties are the module's point),
    # not rebuild the list from the internal partition alone.
    first, _ = scope.filter_to_scope(_sched(), resource="pile")
    t2 = next(t for t in first.tasks if t.id == "T2")
    assert [d.predecessor for d in t2.external_dependencies] == ["T5"]
    second, _ = scope.filter_to_scope(first, resource="pile")
    t2b = next(t for t in second.tasks if t.id == "T2")
    assert [d.predecessor for d in t2b.external_dependencies] == ["T5"]
    assert t2b.external_dependencies[0].type == "SS"
    assert t2b.external_dependencies[0].lag_days == 2.0


def test_ancestry_is_segment_based_not_string_prefix():
    # "1.1" must not claim ancestry over "1.10": a leaf at 1.10.1 keeps
    # summary 1.10 (and 1), never 1.1.
    sched = CanonicalSchedule(name="wide", tasks=[
        ScheduleTask(id="S1", name="One", wbs="1"),
        ScheduleTask(id="S11", name="One-one", wbs="1.1"),
        ScheduleTask(id="T-A", name="Other leaf", wbs="1.1.1",
                     resources=["Other Crew"]),
        ScheduleTask(id="S110", name="One-ten", wbs="1.10"),
        ScheduleTask(id="T-B", name="Target leaf", wbs="1.10.1",
                     resources=["Target Crew"]),
    ])
    filtered, _ = scope.filter_to_scope(sched, resource="target")
    assert [t.id for t in filtered.tasks] == ["S1", "S110", "T-B"]


def test_summary_row_dependency_is_counted_in_report():
    # Data and report must agree even when a SUMMARY carries a dependency.
    sched = _sched()
    s1a = next(t for t in sched.tasks if t.id == "S1a")
    s1a.dependencies = [Dependency(predecessor="T5")]
    filtered, report = scope.filter_to_scope(sched, resource="pile")
    kept = next(t for t in filtered.tasks if t.id == "S1a")
    assert [d.predecessor for d in kept.external_dependencies] == ["T5"]
    assert report.external_preds == 2          # T2's T5 + S1a's T5
    assert report.internal_preds == 1          # T2's T1


def test_list_helpers_count_leaves_only():
    sched = _sched()
    next(t for t in sched.tasks if t.id == "S1").resources = ["Ghost Crew"]
    names = [n for n, _c in scope.list_resources(sched)]
    assert "Ghost Crew" not in names           # listable == matchable


def test_duplicate_external_dep_in_both_lists_dedupes():
    # Copilot PR-128: a hand-merged input can carry the same tie in BOTH
    # dependencies and external_dependencies — the merge must not double
    # it (report and External Preds column would inflate).
    sched = _sched()
    t2 = next(t for t in sched.tasks if t.id == "T2")
    t2.external_dependencies = [
        Dependency(predecessor="T5", type="SS", lag_days=2.0)]  # dup of dep
    filtered, report = scope.filter_to_scope(sched, resource="pile")
    t2f = next(t for t in filtered.tasks if t.id == "T2")
    assert [d.predecessor for d in t2f.external_dependencies] == ["T5"]
    assert report.external_preds == 1
    # a DISTINCT tie (different lag) to the same predecessor is kept
    t2.external_dependencies.append(
        Dependency(predecessor="T5", type="FS", lag_days=0.0))
    filtered2, report2 = scope.filter_to_scope(sched, resource="pile")
    t2f2 = next(t for t in filtered2.tasks if t.id == "T2")
    assert len(t2f2.external_dependencies) == 2
    assert report2.external_preds == 2


def test_filter_keeps_the_calendars_kept_tasks_use():
    """#335: a kept task on a named calendar must still validate and embed."""
    from bimforge.bim.schedule_model import Calendar

    src = _sched()
    src.calendars = {
        "in-water": Calendar(blackouts=[{"from": "10-02", "to": "05-31"}]),
        "paving": Calendar(work_weekdays=[0, 1, 2, 3])}
    for t in src.tasks:
        if t.id in ("T1", "T2"):
            t.calendar = "in-water"
        if t.id in ("T5", "T6"):
            t.calendar = "paving"
    src.validate()
    filtered, _ = scope.filter_to_scope(src, resource="pile")
    filtered.validate()
    assert set(filtered.calendars) == {"in-water"}
    assert filtered.calendars["in-water"] == src.calendars["in-water"]
    assert filtered.calendars["in-water"] is not src.calendars["in-water"]
