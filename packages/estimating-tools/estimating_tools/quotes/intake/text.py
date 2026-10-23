"""Recover the text of each new document and decide its route: text or vision."""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path

from estimating_tools.quotes.intake.inventory import Document, Inventory

TEXT_CHARS_PER_PAGE = 200
PROBE_PAGES = 3
RENDER_DPI = 150
DOCX_METADATA_MARKER = "[docx file metadata -- not printed on the page]"
XLS_REASON = "xls is not readable here; save it as .xlsx"

_DOCX_PROPERTIES = ("author", "last_modified_by", "company", "title")
_TEXT_KINDS = {"docx", "xlsx", "csv", "txt", "md"}
_IMAGE_KINDS = {"png", "jpg", "jpeg"}


class UncachedFormulaError(ValueError):
    """A formula price cannot be transcribed without its saved value."""


@dataclass
class TextReport:
    text: int = 0
    vision: int = 0
    unreadable: int = 0
    lines: list[str] = field(default_factory=list)
    changed: bool = False   # the stage changed the inventory

    def summary(self) -> str:
        return f"{self.text} text, {self.vision} vision, {self.unreadable} unreadable"


def text_path(state_dir: Path, sha: str) -> Path:
    return Path(state_dir) / "text" / f"{sha}.txt"


def pages_dir(state_dir: Path, sha: str) -> Path:
    return Path(state_dir) / "pages" / sha


def read_pdf(path: Path) -> list[str]:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        return [page.extract_text() or "" for page in pdf.pages]


def decide_route(pages: list[str]) -> str:
    probe = pages[:PROBE_PAGES]
    if not probe:
        return "vision"
    mean = sum(len(p.strip()) for p in probe) / len(probe)
    return "text" if mean >= TEXT_CHARS_PER_PAGE else "vision"


def render_pdf(path: Path, out_dir: Path, dpi: int = RENDER_DPI) -> list[Path]:
    import pdfplumber

    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    with pdfplumber.open(path) as pdf:
        for n, page in enumerate(pdf.pages, start=1):
            target = out_dir / f"p{n:03d}.png"
            page.to_image(resolution=dpi).save(target)
            written.append(target)
    if not written:
        raise ValueError("the PDF has no pages")
    (out_dir / "manifest.json").write_text(json.dumps({"pages": len(written)}), encoding="utf-8")
    return written


def _docx_container_text(container) -> list[str]:
    parts = [para.text for para in container.paragraphs]
    for table in container.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return parts


def _docx_metadata(document) -> list[str]:
    props = getattr(document, "core_properties", None)
    if props is None:
        return []
    found = [f"{name}={value.strip()}"
             for name in _DOCX_PROPERTIES
             for value in [str(getattr(props, name, "") or "")]
             if value.strip()]
    return [DOCX_METADATA_MARKER, "  ".join(found)] if found else []


def _docx_parts(path: Path) -> tuple[str, str]:
    """The printed text (body, headers, footers) and the metadata block."""
    import docx

    document = docx.Document(str(path))
    lines = _docx_container_text(document)
    for section in document.sections:
        for attr in ("first_page_header", "header", "even_page_header",
                     "first_page_footer", "footer", "even_page_footer"):
            part = getattr(section, attr, None)
            if part is not None:
                lines.extend(t for t in _docx_container_text(part) if t.strip())
    return "\n".join(lines), "\n".join(_docx_metadata(document))


def read_docx(path: Path) -> str:
    body, meta = _docx_parts(path)
    return body + "\n" + meta if meta else body


def _xlsx_parts(path: Path) -> tuple[str, bool]:
    """The tab-joined text and whether any cell held a value."""
    import openpyxl

    book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    formulas = None
    has_content = False
    try:
        formulas = openpyxl.load_workbook(str(path), data_only=False, read_only=True)
        parts: list[str] = []
        for sheet in book.worksheets:
            parts.append(f"# {sheet.title}")
            for row, original in zip(sheet.iter_rows(values_only=True),
                                     formulas[sheet.title].iter_rows()):
                if any(cell.data_type == "f" and value is None
                       for cell, value in zip(original, row)):
                    raise UncachedFormulaError(
                        "no cached values in the workbook (open and save it in Excel)")
                cells = [str(c) for c in row if c is not None]
                if cells:
                    has_content = True
                    parts.append("\t".join(cells))
    finally:
        book.close()
        if formulas is not None:
            formulas.close()
    return "\n".join(parts), has_content


def read_xlsx(path: Path) -> str:
    return _xlsx_parts(path)[0]


def _xlsx_has_formulas(path: Path) -> bool:
    import openpyxl

    book = openpyxl.load_workbook(str(path), data_only=False, read_only=True)
    try:
        return any(isinstance(c, str) and c.startswith("=")
                   for sheet in book.worksheets
                   for row in sheet.iter_rows(values_only=True) for c in row)
    finally:
        book.close()


def read_text_file(path: Path) -> str:
    return Path(path).read_text(encoding="utf-8-sig", errors="replace")


_READERS = {"docx": read_docx, "xlsx": read_xlsx,
            "csv": read_text_file, "txt": read_text_file, "md": read_text_file}


def _write_text(state_dir: Path, sha: str, text: str) -> None:
    target = text_path(state_dir, sha)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(text, encoding="utf-8")


def _fail(doc: Document, reason: str) -> None:
    doc.route, doc.status, doc.reason = "unreadable", "error", reason


def _recover_one(doc: Document, state_dir: Path) -> None:
    path = Path(doc.paths[0])
    if doc.kind == "xls":
        _fail(doc, XLS_REASON)
    elif doc.kind in _IMAGE_KINDS:
        doc.route, doc.pages = "vision", 1
    elif doc.kind == "pdf":
        pages = read_pdf(path)
        route = decide_route(pages)
        if route == "text":
            body = "".join(f"\n\n--- page {n} ---\n\n{t}" for n, t in enumerate(pages, start=1))
            _write_text(state_dir, doc.sha256, body.lstrip("\n"))
        else:
            render_pdf(path, pages_dir(state_dir, doc.sha256))
        doc.route, doc.pages = route, len(pages)
    elif doc.kind in _TEXT_KINDS:
        if doc.kind == "docx":
            printed, meta = _docx_parts(path)
            text = printed + "\n" + meta if meta else printed
        elif doc.kind == "xlsx":
            try:
                text, has_content = _xlsx_parts(path)
            except UncachedFormulaError as exc:
                _fail(doc, str(exc))
                return
            printed = text if has_content else ""
        else:
            text = printed = _READERS[doc.kind](path)
        if not printed.strip():
            if doc.kind == "xlsx" and _xlsx_has_formulas(path):
                _fail(doc, "no cached values in the workbook (open and save it in Excel)")
            else:
                _fail(doc, "no text in the document")
        else:
            _write_text(state_dir, doc.sha256, text)
            doc.route = "text"


def recover_text(inv: Inventory, state_dir: Path) -> TextReport:
    report = TextReport()
    for doc in inv:
        if doc.status != "new" or doc.route is not None:
            continue
        inv.touch()
        report.changed = True
        try:
            _recover_one(doc, state_dir)
        except Exception as exc:  # noqa: BLE001 - one bad file must not stop the run
            _fail(doc, f"{type(exc).__name__}: {exc}")
        if doc.route == "text":
            report.text += 1
        elif doc.route == "vision":
            report.vision += 1
        elif doc.route == "unreadable":
            report.unreadable += 1
            report.lines.append(f"unreadable {doc.file_name}: {doc.reason}")
    return report


def image_paths(doc: Document, state_dir: Path) -> list[str] | None:
    if doc.kind in _IMAGE_KINDS:
        return [doc.paths[0]]
    folder = pages_dir(state_dir, doc.sha256)
    try:
        count = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))["pages"]
    except (OSError, ValueError, KeyError, TypeError):
        return None
    found = sorted(folder.glob("p*.png"))
    if len(found) != count:
        return None
    return [str(p) for p in found]
