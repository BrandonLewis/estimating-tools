"""Fixture builders shared by the quote-intake tests."""
from __future__ import annotations

import email.message
from pathlib import Path

import yaml

QUOTE_TEXT = (
    "DIRT MOVERS\n"
    "Quote 10/20/2026\n"
    "Project: Hartsook Creek Bridge\n"
    "Haul import and export 460 CY @ 16.50 = 7,590.00\n"
    "Prices are firm for 30 days from the quote date.\n"
    "Excludes permits, bonds, testing and traffic control.\n"
    "Payment terms net 30 on approved invoices.\n"
)


def make_pdf(path, pages):
    objs = []
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        objs.append((f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                     f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>").encode())
        lines = [] if text is None else text.split("\n")
        esc = lambda s: s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        body = ("BT /F1 11 Tf 72 740 Td 14 TL " + " ".join(f"({esc(l)}) Tj T*" for l in lines)
                + " ET").encode("latin-1")
        objs.append(b"<< /Length %d >>\nstream\n" % len(body) + body + b"\nendstream")
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    Path(path).write_bytes(bytes(out))
    return Path(path)


def make_docx(path, paragraphs, table=None, author=None):
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=max(len(r) for r in table))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = str(value)
    if author is not None:
        document.core_properties.author = author
    document.save(str(path))
    return Path(path)


def make_xlsx(path, sheets):
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(str(path))
    return Path(path)


def make_png(path, text=""):
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (200, 80), "white")
    if text:
        ImageDraw.Draw(image).text((8, 30), text, fill="black")
    image.save(str(path), "PNG")
    return Path(path)


def make_eml(path, *, subject, body_text, body_html, attachments):
    msg = email.message.EmailMessage()
    msg["Subject"] = subject
    msg["From"] = "sales@vendor.example"
    msg["To"] = "estimating@example.com"
    if body_text is not None:
        msg.set_content(body_text)
        if body_html is not None:
            msg.add_alternative(body_html, subtype="html")
    elif body_html is not None:
        msg.set_content(body_html, subtype="html")
    for att in attachments:
        att = Path(att)
        suffix = att.suffix.lower()
        maintype, subtype = {
            ".pdf": ("application", "pdf"), ".png": ("image", "png"),
            ".jpg": ("image", "jpeg"), ".jpeg": ("image", "jpeg"),
        }.get(suffix, ("application", "octet-stream"))
        msg.add_attachment(att.read_bytes(), maintype=maintype, subtype=subtype,
                           filename=att.name)
    Path(path).write_bytes(msg.as_bytes())
    return Path(path)


def spec_on_disk(tmp_path):
    from estimating_tools.estimate_spec.tests._fieldbook import packaged_spec

    spec = packaged_spec()
    spec["work_packages_quotes_dir"] = "quotes"
    spec["work_packages_inbox_dir"] = "inbox"
    spec_path = Path(tmp_path) / "demo.estimate.yaml"
    spec_path.write_text(yaml.safe_dump(spec, sort_keys=False), encoding="utf-8")
    (Path(tmp_path) / "quotes").mkdir(exist_ok=True)
    (Path(tmp_path) / "inbox").mkdir(exist_ok=True)
    return spec_path
