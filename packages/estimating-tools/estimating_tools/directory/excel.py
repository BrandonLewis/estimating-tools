"""Import an org's vendor/sub directory from a spreadsheet into the canonical model.

The pure builder (`build_directory`) maps the org's columns to canonical fields (identity
default), groups rows by organization, builds optional contacts, normalizes is_dbe/role,
and validates the required `organization` field — no openpyxl/MCP. Only
`read_directory_rows` reads the .xlsx (openpyxl)."""
from __future__ import annotations

from estimating_tools.directory.model import (
    CanonicalDirectory, DirectoryOrganization, DirectoryPerson, VALID_ROLES,
)

_CANONICAL_FIELDS = (
    "organization", "region", "license_number", "bidder_id", "org_phone", "org_email",
    "address", "is_dbe", "certification_number", "certifying_agency", "federal_id",
    "contact_email", "first_name", "last_name", "contact_phone", "job_title", "role",
)
_REQUIRED = ("organization",)
_TRUTHY = {"yes", "y", "true", "1", "x", "dbe"}


def _effective_mapping(mapping):
    m = {f: f for f in _CANONICAL_FIELDS}      # identity default
    if mapping:
        m.update(mapping)
    return m


def _str(value):
    return str(value).strip() if value is not None else ""


def _truthy(value):
    return _str(value).lower() in _TRUTHY


def _normalize_role(value, org_name):
    role = _str(value).lower()
    if role and role not in VALID_ROLES:
        raise ValueError(
            f"unrecognized role {value!r} for a contact of organization {org_name!r} "
            f"(expected one of {sorted(VALID_ROLES)} or blank)"
        )
    return role


def build_directory(rows, mapping=None) -> CanonicalDirectory:
    """Pure: spreadsheet rows (dicts keyed by the org's headers) + a column mapping ->
    a CanonicalDirectory. No openpyxl/MCP."""
    if not rows:
        return CanonicalDirectory(organizations=[])
    m = _effective_mapping(mapping)
    present = set()
    for r in rows:
        present.update(r.keys())
    for field_name in _REQUIRED:
        if m[field_name] not in present:
            raise ValueError(
                f"required field {field_name!r} maps to column {m[field_name]!r}, not "
                f"found in the spreadsheet headers {sorted(present)}"
            )

    order: list[str] = []
    groups: dict[str, list[dict]] = {}
    for r in rows:
        name = _str(r.get(m["organization"]))
        if not name:
            continue                            # blank-organization / blank row -> skip
        if name not in groups:
            groups[name] = []
            order.append(name)
        groups[name].append(r)

    organizations = []
    for name in order:
        group = groups[name]
        first = group[0]
        contacts = []
        seen_emails = set()
        for r in group:
            email = _str(r.get(m["contact_email"]))
            if not email:
                continue                        # org-only row -> no contact
            key = email.lower()
            if key in seen_emails:
                continue                        # dedup within the org, first-seen wins
            seen_emails.add(key)
            contacts.append(DirectoryPerson(
                email=email,
                first_name=_str(r.get(m["first_name"])),
                last_name=_str(r.get(m["last_name"])),
                phone_primary=_str(r.get(m["contact_phone"])),
                job_title=_str(r.get(m["job_title"])),
                role=_normalize_role(r.get(m["role"]), name),
            ))
        organizations.append(DirectoryOrganization(
            name=name,
            region=_str(first.get(m["region"])),
            license_number=_str(first.get(m["license_number"])),
            bidder_id=_str(first.get(m["bidder_id"])),
            phone=_str(first.get(m["org_phone"])),
            email=_str(first.get(m["org_email"])),
            address=_str(first.get(m["address"])),
            is_dbe=_truthy(first.get(m["is_dbe"])),
            certification_number=_str(first.get(m["certification_number"])),
            certifying_agency=_str(first.get(m["certifying_agency"])),
            federal_id=_str(first.get(m["federal_id"])),
            contacts=contacts,
        ))
    return CanonicalDirectory(organizations=organizations)


def read_directory_rows(path) -> list[dict]:
    """Read the first worksheet of an .xlsx into row-dicts keyed by header. The only
    openpyxl/file touch. Values are the raw openpyxl cell values (an int stays int, a
    bool stays bool, a date stays datetime); build_directory coerces them via _str."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise ImportError(
            "reading .xlsx requires openpyxl: pip install 'estimating_tools[excel]'"
        ) from exc
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = None
        out: list[dict] = []
        for raw in ws.iter_rows(values_only=True):
            if raw is None or all(c is None for c in raw):
                continue
            if headers is None:
                headers = [_str(c) for c in raw]
                continue
            out.append({h: v for h, v in zip(headers, raw)})
        return out
    finally:
        wb.close()


def build_directory_from_excel(path, mapping=None) -> CanonicalDirectory:
    return build_directory(read_directory_rows(path), mapping)
